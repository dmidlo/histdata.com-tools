"""Philadelphia Fed Survey of Professional Forecasters archive qualification.

The SPF is a quarterly forecast vintage, not an event-consensus substitute.
This module binds every 2000-present release PDF to the official release-date
ledger and to four current-quarter median forecasts from the point-in-time
official median workbooks.  The compact packaged manifest keeps hashes and
normalized values; callers may retain and replay the complete raw corpus.
"""

from __future__ import annotations

import base64
import hashlib
import json
import math
import re
from collections.abc import Mapping, Sequence
from dataclasses import dataclass
from datetime import date
from io import BytesIO
from pathlib import Path
from types import MappingProxyType
from typing import Any, cast
from urllib.parse import urljoin, urlsplit, urlunsplit

from openpyxl import load_workbook
from pypdf import PdfReader

from histdatacom.market_context.contracts import canonical_contract_json
from histdatacom.market_context.official_sources import (
    OfficialRawSnapshotV1,
    OfficialRequestMethod,
    OfficialSourceFormat,
    OfficialSourceRegistryV1,
    OfficialSourceRequestV1,
)
from histdatacom.market_context.us_backfill import (
    US_BACKFILL_START_DATE,
    UnitedStatesBackfillProfileV1,
    UnitedStatesProgramCoverageV1,
)
from histdatacom.runtime_contracts import JSONValue

SPF_INDEX_ENTRY_SCHEMA_VERSION = "histdatacom.spf-index-entry.v1"
SPF_RELEASE_INDEX_SCHEMA_VERSION = "histdatacom.spf-release-index.v1"
SPF_ARTIFACT_SCHEMA_VERSION = "histdatacom.spf-artifact.v1"
SPF_FORECAST_SCHEMA_VERSION = "histdatacom.spf-forecast.v1"
SPF_PUBLICATION_SCHEMA_VERSION = "histdatacom.spf-publication.v1"
SPF_ARCHIVE_MANIFEST_SCHEMA_VERSION = "histdatacom.spf-archive-manifest.v1"
SPF_SOURCE_KEY = "us.frb.philadelphia-spf"
SPF_PROGRAM_KEY = "us.philadelphia-fed.spf"
SPF_INDEX_URI = (
    "https://www.philadelphiafed.org/surveys-and-data/real-time-data-research/"
    "survey-of-professional-forecasters"
)
SPF_MEDIAN_FORECASTS_URI = (
    "https://www.philadelphiafed.org/surveys-and-data/real-time-data-research/"
    "median-forecasts"
)
SPF_ASSET_ROOT = (
    "https://www.philadelphiafed.org/-/media/FRBP/Assets/Surveys-And-Data/"
    "survey-of-professional-forecasters"
)
SPF_RELEASE_DATES_URI = f"{SPF_ASSET_ROOT}/spf-release-dates.txt"
SPF_MEDIAN_LEVEL_URI = f"{SPF_ASSET_ROOT}/historical-data/medianLevel.xlsx"
SPF_MEDIAN_GROWTH_URI = f"{SPF_ASSET_ROOT}/historical-data/medianGrowth.xlsx"
SPF_DOCUMENTATION_URI = f"{SPF_ASSET_ROOT}/spf-documentation.pdf"
SPF_FIRST_SURVEY_PERIOD = "2000-Q1"
SPF_LATEST_PACKAGED_SURVEY_PERIOD = "2026-Q3"
SPF_SHUTDOWN_DELAYED_PERIODS = (
    "2013-Q4",
    "2019-Q1",
    "2025-Q4",
    "2026-Q1",
)

MAX_SPF_RELEASES = 512
MAX_SPF_ARTIFACT_BYTES = 16 * 1024 * 1024
MAX_SPF_TOTAL_BYTES = 256 * 1024 * 1024
_SHA256_RE = re.compile(r"^[0-9a-f]{64}$")
_PERIOD_RE = re.compile(r"^(?P<year>\d{4})-Q(?P<quarter>[1-4])$")
_TITLE_RE = re.compile(
    r'"year":"(?P<year>\d{4})","url":"(?P<url>[^"]+)",'
    r'"title":"(?P<ordinal>First|Second|Third|Fourth) Quarter '
    r'(?P=year) Survey of Professional Forecasters","pdf":true'
)
_LEDGER_ROW_RE = re.compile(
    r"^\s*(?:(?P<year>\d{4})\s+)?Q(?P<quarter>[1-4])\s+"
    r"(?P<deadline>\d{1,2}/\d{1,2}/\d{2})(?:\*+)?\s+"
    r"(?P<release>\d{1,2}/\d{1,2}/\d{2})(?:\*+)?\s*$"
)
_ORDINAL_TO_QUARTER = {"First": 1, "Second": 2, "Third": 3, "Fourth": 4}


@dataclass(frozen=True, slots=True)
class _ForecastSpec:
    concept_key: str
    artifact_role: str
    worksheet: str
    column: str
    unit: str


_SPF_FORECAST_SPECS = (
    _ForecastSpec(
        "real-gdp-growth-current-quarter",
        "median-growth-workbook",
        "RGDP",
        "drgdp2",
        "annualized-percent-change",
    ),
    _ForecastSpec(
        "unemployment-rate-current-quarter",
        "median-level-workbook",
        "UNEMP",
        "UNEMP2",
        "percent",
    ),
    _ForecastSpec(
        "headline-cpi-inflation-current-quarter",
        "median-level-workbook",
        "CPI",
        "CPI2",
        "annualized-percent-change",
    ),
    _ForecastSpec(
        "three-month-treasury-bill-rate-current-quarter",
        "median-level-workbook",
        "TBILL",
        "TBILL2",
        "percent",
    ),
)
_FORECAST_BY_KEY = MappingProxyType(
    {item.concept_key: item for item in _SPF_FORECAST_SPECS}
)
_ROLE_TO_URI = MappingProxyType(
    {
        "archive-index": SPF_INDEX_URI,
        "median-forecasts-page": SPF_MEDIAN_FORECASTS_URI,
        "release-date-ledger": SPF_RELEASE_DATES_URI,
        "median-level-workbook": SPF_MEDIAN_LEVEL_URI,
        "median-growth-workbook": SPF_MEDIAN_GROWTH_URI,
        "documentation": SPF_DOCUMENTATION_URI,
    }
)
_ROLE_TO_FORMAT = MappingProxyType(
    {
        "archive-index": OfficialSourceFormat.HTML,
        "median-forecasts-page": OfficialSourceFormat.HTML,
        "release-date-ledger": OfficialSourceFormat.TEXT,
        "median-level-workbook": OfficialSourceFormat.XLSX,
        "median-growth-workbook": OfficialSourceFormat.XLSX,
        "documentation": OfficialSourceFormat.PDF,
    }
)


def _required_text(value: object, name: str) -> str:
    text = str(value).strip()
    if not text or len(text) > 4096:
        raise ValueError(f"{name} is invalid")
    return text


def _iso_date(value: object, name: str) -> str:
    text = _required_text(value, name)
    try:
        parsed = date.fromisoformat(text)
    except ValueError as exc:
        raise ValueError(f"{name} must be an ISO date") from exc
    if parsed.isoformat() != text:
        raise ValueError(f"{name} must be a canonical ISO date")
    return text


def _period(value: object, name: str) -> str:
    text = _required_text(value, name)
    if _PERIOD_RE.fullmatch(text) is None:
        raise ValueError(f"{name} must be a canonical year-quarter")
    return text


def _sha256(value: object, name: str) -> str:
    text = _required_text(value, name)
    if _SHA256_RE.fullmatch(text) is None:
        raise ValueError(f"{name} must be a lowercase SHA-256")
    return text


def _mapping(value: object, name: str) -> Mapping[str, Any]:
    if not isinstance(value, Mapping):
        raise TypeError(f"{name} must be a mapping")
    return cast(Mapping[str, Any], value)


def _sequence(value: object, name: str) -> Sequence[Any]:
    if not isinstance(value, Sequence) or isinstance(value, (str, bytes)):
        raise TypeError(f"{name} must be a sequence")
    return value


def _positive_int(value: object, name: str, maximum: int) -> int:
    if (
        not isinstance(value, int)
        or isinstance(value, bool)
        or not 1 <= value <= maximum
    ):
        raise ValueError(f"{name} is outside its bound")
    return value


def _finite(value: object, name: str) -> float:
    if isinstance(value, bool):
        raise TypeError(f"{name} must be numeric")
    try:
        result = float(cast(Any, value))
    except (TypeError, ValueError) as exc:
        raise ValueError(f"{name} must be numeric") from exc
    if not math.isfinite(result):
        raise ValueError(f"{name} must be finite")
    return result


def _stable_id(prefix: str, payload: Mapping[str, JSONValue]) -> str:
    digest = hashlib.sha256(
        canonical_contract_json(payload).encode("utf-8")
    ).hexdigest()
    return f"{prefix}:sha256:{digest}"


def _next_period(value: str) -> str:
    match = _PERIOD_RE.fullmatch(value)
    if match is None:
        raise ValueError("SPF period is invalid")
    year = int(match.group("year"))
    quarter = int(match.group("quarter"))
    if quarter == 4:
        return f"{year + 1:04d}-Q1"
    return f"{year:04d}-Q{quarter + 1}"


def _periods(first: str, last: str) -> tuple[str, ...]:
    result: list[str] = []
    current = first
    while current <= last:
        result.append(current)
        current = _next_period(current)
    return tuple(result)


def _year_quarter(value: str) -> tuple[int, int]:
    match = _PERIOD_RE.fullmatch(value)
    if match is None:
        raise ValueError("SPF period is invalid")
    return int(match.group("year")), int(match.group("quarter"))


def _normalize_uri(uri: str) -> str:
    parts = urlsplit(uri)
    return urlunsplit((parts.scheme, parts.netloc, parts.path, "", ""))


def _snapshot_artifact(
    snapshot: OfficialRawSnapshotV1, role: str
) -> PhiladelphiaFedSpfArtifactV1:
    return PhiladelphiaFedSpfArtifactV1(
        role=role,
        source_uri=snapshot.request.uri,
        source_format=snapshot.request.source_format,
        content_sha256=snapshot.content_sha256,
        content_length=len(snapshot.content),
    )


@dataclass(frozen=True, slots=True)
class PhiladelphiaFedSpfArtifactV1:
    """One content-addressed official SPF artifact."""

    role: str
    source_uri: str
    source_format: OfficialSourceFormat
    content_sha256: str
    content_length: int
    artifact_id: str = ""
    schema_version: str = SPF_ARTIFACT_SCHEMA_VERSION

    def __post_init__(self) -> None:
        if self.schema_version != SPF_ARTIFACT_SCHEMA_VERSION:
            raise ValueError("unsupported SPF artifact schema")
        role = _required_text(self.role, "role")
        if role not in {*_ROLE_TO_URI, "release-pdf"}:
            raise ValueError("SPF artifact role is invalid")
        uri = _required_text(self.source_uri, "source_uri")
        if not uri.startswith("https://www.philadelphiafed.org/"):
            raise ValueError("SPF artifact URI is invalid")
        if role in _ROLE_TO_URI and uri != _ROLE_TO_URI[role]:
            raise ValueError("SPF support-artifact URI differs")
        source_format = OfficialSourceFormat.from_value(self.source_format)
        if (
            role == "release-pdf"
            and source_format is not OfficialSourceFormat.PDF
        ):
            raise ValueError("SPF release artifact must be PDF")
        if (
            role in _ROLE_TO_FORMAT
            and source_format is not _ROLE_TO_FORMAT[role]
        ):
            raise ValueError("SPF support-artifact format differs")
        object.__setattr__(self, "role", role)
        object.__setattr__(self, "source_uri", uri)
        object.__setattr__(self, "source_format", source_format)
        object.__setattr__(
            self,
            "content_sha256",
            _sha256(self.content_sha256, "content_sha256"),
        )
        _positive_int(
            self.content_length, "content_length", MAX_SPF_ARTIFACT_BYTES
        )
        expected = _stable_id("spf-artifact", self.identity_payload())
        if self.artifact_id and self.artifact_id != expected:
            raise ValueError("SPF artifact identity differs")
        object.__setattr__(self, "artifact_id", expected)

    def identity_payload(self) -> dict[str, JSONValue]:
        return {
            "schema_version": self.schema_version,
            "role": self.role,
            "source_uri": self.source_uri,
            "source_format": self.source_format.value,
            "content_sha256": self.content_sha256,
            "content_length": self.content_length,
        }

    def to_dict(self) -> dict[str, JSONValue]:
        return {**self.identity_payload(), "artifact_id": self.artifact_id}

    @classmethod
    def from_dict(cls, data: Mapping[str, Any]) -> PhiladelphiaFedSpfArtifactV1:
        return cls(
            role=str(data.get("role", "")),
            source_uri=str(data.get("source_uri", "")),
            source_format=OfficialSourceFormat.from_value(
                str(data.get("source_format", ""))
            ),
            content_sha256=str(data.get("content_sha256", "")),
            content_length=cast(int, data.get("content_length")),
            artifact_id=str(data.get("artifact_id", "")),
            schema_version=str(data.get("schema_version", "")),
        )


@dataclass(frozen=True, slots=True)
class PhiladelphiaFedSpfIndexEntryV1:
    """One release-ledger row joined to its official archive PDF."""

    survey_period: str
    title: str
    true_deadline_date: str
    release_date: str
    release_pdf_uri: str
    shutdown_delayed: bool
    entry_id: str = ""
    schema_version: str = SPF_INDEX_ENTRY_SCHEMA_VERSION

    def __post_init__(self) -> None:
        if self.schema_version != SPF_INDEX_ENTRY_SCHEMA_VERSION:
            raise ValueError("unsupported SPF index-entry schema")
        period = _period(self.survey_period, "survey_period")
        if not isinstance(self.shutdown_delayed, bool):
            raise TypeError("shutdown_delayed must be boolean")
        year, quarter = _year_quarter(period)
        ordinals = ("First", "Second", "Third", "Fourth")
        expected_title = (
            f"{ordinals[quarter - 1]} Quarter {year} Survey of Professional "
            "Forecasters"
        )
        if self.title != expected_title:
            raise ValueError("SPF release title differs")
        deadline = _iso_date(self.true_deadline_date, "true_deadline_date")
        release = _iso_date(self.release_date, "release_date")
        if release < deadline:
            raise ValueError("SPF release predates its true deadline")
        uri = _required_text(self.release_pdf_uri, "release_pdf_uri")
        if not uri.startswith(
            f"{SPF_ASSET_ROOT}/{year}/"
        ) or not uri.lower().endswith(".pdf"):
            raise ValueError("SPF release PDF URI is invalid")
        if self.shutdown_delayed != (period in SPF_SHUTDOWN_DELAYED_PERIODS):
            raise ValueError("SPF shutdown-delay flag differs")
        object.__setattr__(self, "survey_period", period)
        object.__setattr__(self, "true_deadline_date", deadline)
        object.__setattr__(self, "release_date", release)
        object.__setattr__(self, "release_pdf_uri", uri)
        expected = _stable_id("spf-index-entry", self.identity_payload())
        if self.entry_id and self.entry_id != expected:
            raise ValueError("SPF index-entry identity differs")
        object.__setattr__(self, "entry_id", expected)

    def identity_payload(self) -> dict[str, JSONValue]:
        return {
            "schema_version": self.schema_version,
            "survey_period": self.survey_period,
            "title": self.title,
            "true_deadline_date": self.true_deadline_date,
            "release_date": self.release_date,
            "release_pdf_uri": self.release_pdf_uri,
            "shutdown_delayed": self.shutdown_delayed,
        }

    def to_dict(self) -> dict[str, JSONValue]:
        return {**self.identity_payload(), "entry_id": self.entry_id}

    @classmethod
    def from_dict(
        cls, data: Mapping[str, Any]
    ) -> PhiladelphiaFedSpfIndexEntryV1:
        return cls(
            survey_period=str(data.get("survey_period", "")),
            title=str(data.get("title", "")),
            true_deadline_date=str(data.get("true_deadline_date", "")),
            release_date=str(data.get("release_date", "")),
            release_pdf_uri=str(data.get("release_pdf_uri", "")),
            shutdown_delayed=cast(bool, data.get("shutdown_delayed")),
            entry_id=str(data.get("entry_id", "")),
            schema_version=str(data.get("schema_version", "")),
        )


@dataclass(frozen=True, slots=True)
class PhiladelphiaFedSpfReleaseIndexV1:
    """Joined official SPF archive and release-date ledger."""

    as_of_date: str
    archive_index: PhiladelphiaFedSpfArtifactV1
    release_date_ledger: PhiladelphiaFedSpfArtifactV1
    entries: tuple[PhiladelphiaFedSpfIndexEntryV1, ...]
    index_id: str = ""
    schema_version: str = SPF_RELEASE_INDEX_SCHEMA_VERSION

    def __post_init__(self) -> None:
        if self.schema_version != SPF_RELEASE_INDEX_SCHEMA_VERSION:
            raise ValueError("unsupported SPF release-index schema")
        as_of = _iso_date(self.as_of_date, "as_of_date")
        if self.archive_index.role != "archive-index":
            raise ValueError("SPF archive-index artifact differs")
        if self.release_date_ledger.role != "release-date-ledger":
            raise ValueError("SPF release-date artifact differs")
        entries = tuple(self.entries)
        if not entries or len(entries) > MAX_SPF_RELEASES:
            raise ValueError("SPF release count is outside its bound")
        periods = tuple(item.survey_period for item in entries)
        if periods != tuple(sorted(set(periods))):
            raise ValueError("SPF release periods must be unique and sorted")
        if periods != _periods(periods[0], periods[-1]):
            raise ValueError("SPF release periods are not contiguous")
        if any(item.release_date > as_of for item in entries):
            raise ValueError("SPF release follows the as-of date")
        object.__setattr__(self, "as_of_date", as_of)
        object.__setattr__(self, "entries", entries)
        expected = _stable_id("spf-release-index", self.identity_payload())
        if self.index_id and self.index_id != expected:
            raise ValueError("SPF release-index identity differs")
        object.__setattr__(self, "index_id", expected)

    @property
    def by_period(self) -> Mapping[str, PhiladelphiaFedSpfIndexEntryV1]:
        return MappingProxyType(
            {item.survey_period: item for item in self.entries}
        )

    def identity_payload(self) -> dict[str, JSONValue]:
        return {
            "schema_version": self.schema_version,
            "as_of_date": self.as_of_date,
            "archive_index": self.archive_index.to_dict(),
            "release_date_ledger": self.release_date_ledger.to_dict(),
            "entries": [item.to_dict() for item in self.entries],
        }

    def to_dict(self) -> dict[str, JSONValue]:
        return {**self.identity_payload(), "index_id": self.index_id}

    @classmethod
    def from_dict(
        cls, data: Mapping[str, Any]
    ) -> PhiladelphiaFedSpfReleaseIndexV1:
        return cls(
            as_of_date=str(data.get("as_of_date", "")),
            archive_index=PhiladelphiaFedSpfArtifactV1.from_dict(
                _mapping(data.get("archive_index"), "archive_index")
            ),
            release_date_ledger=PhiladelphiaFedSpfArtifactV1.from_dict(
                _mapping(data.get("release_date_ledger"), "release_date_ledger")
            ),
            entries=tuple(
                PhiladelphiaFedSpfIndexEntryV1.from_dict(
                    _mapping(item, "SPF index entry")
                )
                for item in _sequence(data.get("entries"), "entries")
            ),
            index_id=str(data.get("index_id", "")),
            schema_version=str(data.get("schema_version", "")),
        )


@dataclass(frozen=True, slots=True)
class PhiladelphiaFedSpfForecastV1:
    """One current-quarter median forecast from an official SPF workbook."""

    survey_period: str
    concept_key: str
    artifact_role: str
    worksheet: str
    column: str
    unit: str
    value: float
    lexical_value: str
    forecast_id: str = ""
    schema_version: str = SPF_FORECAST_SCHEMA_VERSION

    def __post_init__(self) -> None:
        if self.schema_version != SPF_FORECAST_SCHEMA_VERSION:
            raise ValueError("unsupported SPF forecast schema")
        period = _period(self.survey_period, "survey_period")
        concept = _required_text(self.concept_key, "concept_key")
        spec = _FORECAST_BY_KEY.get(concept)
        if spec is None:
            raise ValueError("SPF forecast concept is invalid")
        if (
            self.artifact_role,
            self.worksheet,
            self.column,
            self.unit,
        ) != (spec.artifact_role, spec.worksheet, spec.column, spec.unit):
            raise ValueError("SPF forecast source semantics differ")
        value = _finite(self.value, "value")
        lexical = _required_text(self.lexical_value, "lexical_value")
        if _finite(lexical, "lexical_value") != value:
            raise ValueError("SPF lexical and numeric values differ")
        object.__setattr__(self, "survey_period", period)
        object.__setattr__(self, "concept_key", concept)
        object.__setattr__(self, "value", value)
        expected = _stable_id("spf-forecast", self.identity_payload())
        if self.forecast_id and self.forecast_id != expected:
            raise ValueError("SPF forecast identity differs")
        object.__setattr__(self, "forecast_id", expected)

    def identity_payload(self) -> dict[str, JSONValue]:
        return {
            "schema_version": self.schema_version,
            "survey_period": self.survey_period,
            "concept_key": self.concept_key,
            "artifact_role": self.artifact_role,
            "worksheet": self.worksheet,
            "column": self.column,
            "unit": self.unit,
            "value": self.value,
            "lexical_value": self.lexical_value,
        }

    def to_dict(self) -> dict[str, JSONValue]:
        return {**self.identity_payload(), "forecast_id": self.forecast_id}

    @classmethod
    def from_dict(cls, data: Mapping[str, Any]) -> PhiladelphiaFedSpfForecastV1:
        return cls(
            survey_period=str(data.get("survey_period", "")),
            concept_key=str(data.get("concept_key", "")),
            artifact_role=str(data.get("artifact_role", "")),
            worksheet=str(data.get("worksheet", "")),
            column=str(data.get("column", "")),
            unit=str(data.get("unit", "")),
            value=_finite(data.get("value"), "value"),
            lexical_value=str(data.get("lexical_value", "")),
            forecast_id=str(data.get("forecast_id", "")),
            schema_version=str(data.get("schema_version", "")),
        )


@dataclass(frozen=True, slots=True)
class PhiladelphiaFedSpfPublicationV1:
    """One date-only SPF publication with four forecast measures."""

    survey_period: str
    true_deadline_date: str
    release_date: str
    shutdown_delayed: bool
    release_pdf: PhiladelphiaFedSpfArtifactV1
    forecasts: tuple[PhiladelphiaFedSpfForecastV1, ...]
    publication_id: str = ""
    schema_version: str = SPF_PUBLICATION_SCHEMA_VERSION

    def __post_init__(self) -> None:
        if self.schema_version != SPF_PUBLICATION_SCHEMA_VERSION:
            raise ValueError("unsupported SPF publication schema")
        period = _period(self.survey_period, "survey_period")
        if not isinstance(self.shutdown_delayed, bool):
            raise TypeError("shutdown_delayed must be boolean")
        deadline = _iso_date(self.true_deadline_date, "true_deadline_date")
        release = _iso_date(self.release_date, "release_date")
        if release < deadline:
            raise ValueError("SPF publication predates its deadline")
        if self.shutdown_delayed != (period in SPF_SHUTDOWN_DELAYED_PERIODS):
            raise ValueError("SPF publication shutdown flag differs")
        if self.release_pdf.role != "release-pdf":
            raise ValueError("SPF publication release artifact differs")
        forecasts = tuple(self.forecasts)
        if tuple(item.concept_key for item in forecasts) != tuple(
            item.concept_key for item in _SPF_FORECAST_SPECS
        ):
            raise ValueError("SPF publication forecast inventory differs")
        if any(item.survey_period != period for item in forecasts):
            raise ValueError("SPF publication contains another survey period")
        object.__setattr__(self, "survey_period", period)
        object.__setattr__(self, "true_deadline_date", deadline)
        object.__setattr__(self, "release_date", release)
        object.__setattr__(self, "forecasts", forecasts)
        expected = _stable_id("spf-publication", self.identity_payload())
        if self.publication_id and self.publication_id != expected:
            raise ValueError("SPF publication identity differs")
        object.__setattr__(self, "publication_id", expected)

    def identity_payload(self) -> dict[str, JSONValue]:
        return {
            "schema_version": self.schema_version,
            "survey_period": self.survey_period,
            "true_deadline_date": self.true_deadline_date,
            "release_date": self.release_date,
            "shutdown_delayed": self.shutdown_delayed,
            "release_pdf": self.release_pdf.to_dict(),
            "forecasts": [item.to_dict() for item in self.forecasts],
        }

    def to_dict(self) -> dict[str, JSONValue]:
        return {
            **self.identity_payload(),
            "publication_id": self.publication_id,
        }

    @classmethod
    def from_dict(
        cls, data: Mapping[str, Any]
    ) -> PhiladelphiaFedSpfPublicationV1:
        return cls(
            survey_period=str(data.get("survey_period", "")),
            true_deadline_date=str(data.get("true_deadline_date", "")),
            release_date=str(data.get("release_date", "")),
            shutdown_delayed=cast(bool, data.get("shutdown_delayed")),
            release_pdf=PhiladelphiaFedSpfArtifactV1.from_dict(
                _mapping(data.get("release_pdf"), "release_pdf")
            ),
            forecasts=tuple(
                PhiladelphiaFedSpfForecastV1.from_dict(
                    _mapping(item, "SPF forecast")
                )
                for item in _sequence(data.get("forecasts"), "forecasts")
            ),
            publication_id=str(data.get("publication_id", "")),
            schema_version=str(data.get("schema_version", "")),
        )


@dataclass(frozen=True, slots=True)
class PhiladelphiaFedSpfArchiveManifestV1:
    """Complete content-addressed SPF archive manifest."""

    registry_id: str
    profile_id: str
    release_index: PhiladelphiaFedSpfReleaseIndexV1
    support_artifacts: tuple[PhiladelphiaFedSpfArtifactV1, ...]
    publications: tuple[PhiladelphiaFedSpfPublicationV1, ...]
    raw_artifact_count: int
    total_content_bytes: int
    forecast_occurrence_count: int
    shutdown_delayed_count: int
    date_only_count: int
    manifest_id: str = ""
    schema_version: str = SPF_ARCHIVE_MANIFEST_SCHEMA_VERSION

    def __post_init__(self) -> None:
        if self.schema_version != SPF_ARCHIVE_MANIFEST_SCHEMA_VERSION:
            raise ValueError("unsupported SPF archive-manifest schema")
        registry_id = _required_text(self.registry_id, "registry_id")
        profile_id = _required_text(self.profile_id, "profile_id")
        if not registry_id.startswith("official-source-registry:sha256:"):
            raise ValueError("SPF registry identity is invalid")
        if not profile_id.startswith("us-backfill-profile:sha256:"):
            raise ValueError("SPF profile identity is invalid")
        support = tuple(self.support_artifacts)
        expected_roles = tuple(_ROLE_TO_URI)
        if tuple(item.role for item in support) != expected_roles:
            raise ValueError("SPF support-artifact inventory differs")
        if support[0] != self.release_index.archive_index:
            raise ValueError("SPF retained archive index differs")
        if support[2] != self.release_index.release_date_ledger:
            raise ValueError("SPF retained release ledger differs")
        publications = tuple(self.publications)
        if tuple(item.survey_period for item in publications) != tuple(
            item.survey_period for item in self.release_index.entries
        ):
            raise ValueError("SPF publication and index periods differ")
        if any(
            (
                publication.true_deadline_date,
                publication.release_date,
                publication.shutdown_delayed,
                publication.release_pdf.source_uri,
            )
            != (
                entry.true_deadline_date,
                entry.release_date,
                entry.shutdown_delayed,
                entry.release_pdf_uri,
            )
            for publication, entry in zip(
                publications, self.release_index.entries, strict=True
            )
        ):
            raise ValueError("SPF publication metadata differs from its index")
        artifacts = (*support, *(item.release_pdf for item in publications))
        if len({item.source_uri for item in artifacts}) != len(artifacts):
            raise ValueError("SPF artifact URI inventory is not unique")
        release_hashes = {
            publication.release_pdf.content_sha256
            for publication in publications
        }
        if len(release_hashes) != len(publications):
            raise ValueError("SPF release PDF hashes are not unique")
        expected_artifact_count = len(artifacts)
        expected_bytes = sum(item.content_length for item in artifacts)
        expected_forecasts = sum(len(item.forecasts) for item in publications)
        expected_delayed = sum(item.shutdown_delayed for item in publications)
        expected_dates = len(publications)
        expected_counts = (
            expected_artifact_count,
            expected_bytes,
            expected_forecasts,
            expected_delayed,
            expected_dates,
        )
        actual_counts = (
            self.raw_artifact_count,
            self.total_content_bytes,
            self.forecast_occurrence_count,
            self.shutdown_delayed_count,
            self.date_only_count,
        )
        if actual_counts != expected_counts:
            raise ValueError("SPF manifest summary counts differ")
        if expected_bytes > MAX_SPF_TOTAL_BYTES:
            raise ValueError("SPF retained corpus exceeds its byte bound")
        object.__setattr__(self, "registry_id", registry_id)
        object.__setattr__(self, "profile_id", profile_id)
        object.__setattr__(self, "support_artifacts", support)
        object.__setattr__(self, "publications", publications)
        expected = _stable_id("spf-archive-manifest", self.identity_payload())
        if self.manifest_id and self.manifest_id != expected:
            raise ValueError("SPF archive-manifest identity differs")
        object.__setattr__(self, "manifest_id", expected)

    @property
    def by_period(self) -> Mapping[str, PhiladelphiaFedSpfPublicationV1]:
        return MappingProxyType(
            {item.survey_period: item for item in self.publications}
        )

    def identity_payload(self) -> dict[str, JSONValue]:
        return {
            "schema_version": self.schema_version,
            "registry_id": self.registry_id,
            "profile_id": self.profile_id,
            "release_index": self.release_index.to_dict(),
            "support_artifacts": [
                item.to_dict() for item in self.support_artifacts
            ],
            "publications": [item.to_dict() for item in self.publications],
            "raw_artifact_count": self.raw_artifact_count,
            "total_content_bytes": self.total_content_bytes,
            "forecast_occurrence_count": self.forecast_occurrence_count,
            "shutdown_delayed_count": self.shutdown_delayed_count,
            "date_only_count": self.date_only_count,
        }

    def to_dict(self) -> dict[str, JSONValue]:
        return {**self.identity_payload(), "manifest_id": self.manifest_id}

    def to_json(self) -> str:
        return str(canonical_contract_json(self.to_dict()))

    @classmethod
    def from_dict(
        cls, data: Mapping[str, Any]
    ) -> PhiladelphiaFedSpfArchiveManifestV1:
        return cls(
            registry_id=str(data.get("registry_id", "")),
            profile_id=str(data.get("profile_id", "")),
            release_index=PhiladelphiaFedSpfReleaseIndexV1.from_dict(
                _mapping(data.get("release_index"), "release_index")
            ),
            support_artifacts=tuple(
                PhiladelphiaFedSpfArtifactV1.from_dict(
                    _mapping(item, "SPF support artifact")
                )
                for item in _sequence(
                    data.get("support_artifacts"), "support_artifacts"
                )
            ),
            publications=tuple(
                PhiladelphiaFedSpfPublicationV1.from_dict(
                    _mapping(item, "SPF publication")
                )
                for item in _sequence(data.get("publications"), "publications")
            ),
            raw_artifact_count=cast(int, data.get("raw_artifact_count")),
            total_content_bytes=cast(int, data.get("total_content_bytes")),
            forecast_occurrence_count=cast(
                int, data.get("forecast_occurrence_count")
            ),
            shutdown_delayed_count=cast(
                int, data.get("shutdown_delayed_count")
            ),
            date_only_count=cast(int, data.get("date_only_count")),
            manifest_id=str(data.get("manifest_id", "")),
            schema_version=str(data.get("schema_version", "")),
        )

    @classmethod
    def from_json(cls, text: str) -> PhiladelphiaFedSpfArchiveManifestV1:
        try:
            payload = json.loads(text)
        except json.JSONDecodeError as exc:
            raise ValueError("SPF manifest is invalid JSON") from exc
        return cls.from_dict(_mapping(payload, "SPF manifest"))


def _request(
    registry: OfficialSourceRegistryV1,
    uri: str,
    source_format: OfficialSourceFormat,
    *,
    as_of_date: str,
) -> OfficialSourceRequestV1:
    source = registry.source(SPF_SOURCE_KEY)
    return OfficialSourceRequestV1(
        source_key=source.source_key,
        source_id=source.source_id,
        method=OfficialRequestMethod.GET,
        uri=uri,
        source_format=source_format,
        parser_id=source.parser_id,
        parser_version=source.parser_version,
        window_start=US_BACKFILL_START_DATE,
        window_end=as_of_date,
    )


def build_philadelphia_fed_spf_support_requests(
    registry: OfficialSourceRegistryV1, *, as_of_date: str
) -> tuple[OfficialSourceRequestV1, ...]:
    """Plan the six official pages/files that define the SPF archive."""
    if not isinstance(registry, OfficialSourceRegistryV1):
        raise TypeError("SPF requests require a v1 registry")
    as_of = _iso_date(as_of_date, "as_of_date")
    return tuple(
        _request(
            registry,
            uri,
            _ROLE_TO_FORMAT[role],
            as_of_date=as_of,
        )
        for role, uri in _ROLE_TO_URI.items()
    )


def _validate_snapshot(
    snapshot: OfficialRawSnapshotV1,
    *,
    uri: str,
    source_format: OfficialSourceFormat,
) -> None:
    if not isinstance(snapshot, OfficialRawSnapshotV1):
        raise TypeError("SPF evidence requires a v1 snapshot")
    if (
        snapshot.request.source_key != SPF_SOURCE_KEY
        or snapshot.request.uri != uri
        or snapshot.request.source_format is not source_format
        or snapshot.status_code != 200
        or not snapshot.content
        or len(snapshot.content) > MAX_SPF_ARTIFACT_BYTES
    ):
        raise ValueError("SPF official snapshot differs")


def _parse_archive_entries(content: bytes) -> Mapping[str, tuple[str, str]]:
    try:
        source = content.decode("utf-8-sig")
    except UnicodeDecodeError as exc:
        raise ValueError("SPF archive page is not UTF-8") from exc
    result: dict[str, tuple[str, str]] = {}
    for match in _TITLE_RE.finditer(source):
        year = int(match.group("year"))
        quarter = _ORDINAL_TO_QUARTER[match.group("ordinal")]
        if year < 2000:
            continue
        period = f"{year:04d}-Q{quarter}"
        uri = _normalize_uri(urljoin(SPF_INDEX_URI, match.group("url")))
        title = match.group(0).split('"title":"', 1)[1].split('","pdf"', 1)[0]
        if period in result:
            raise ValueError("SPF archive page repeats a survey period")
        result[period] = (title, uri)
    if not result:
        raise ValueError("SPF archive page contains no 2000-present releases")
    return MappingProxyType(result)


def _parse_release_dates(content: bytes) -> Mapping[str, tuple[str, str]]:
    try:
        source = content.decode("utf-8-sig")
    except UnicodeDecodeError as exc:
        raise ValueError("SPF release-date ledger is not UTF-8") from exc
    if "Deadline and Release Dates for the Survey" not in source:
        raise ValueError("SPF release-date ledger title differs")
    current_year: int | None = None
    result: dict[str, tuple[str, str]] = {}
    for line in source.splitlines():
        match = _LEDGER_ROW_RE.fullmatch(line)
        if match is None:
            continue
        if match.group("year") is not None:
            current_year = int(cast(str, match.group("year")))
        if current_year is None or current_year < 2000:
            continue
        quarter = int(match.group("quarter"))
        period = f"{current_year:04d}-Q{quarter}"
        if period in result:
            raise ValueError("SPF release-date ledger repeats a survey period")
        values: list[str] = []
        for key in ("deadline", "release"):
            month, day, short_year = (
                int(item) for item in match.group(key).split("/")
            )
            year = 1900 + short_year if short_year >= 90 else 2000 + short_year
            values.append(date(year, month, day).isoformat())
        result[period] = (values[0], values[1])
    if not result:
        raise ValueError(
            "SPF release-date ledger contains no 2000-present rows"
        )
    return MappingProxyType(result)


def parse_philadelphia_fed_spf_release_index(
    archive_snapshot: OfficialRawSnapshotV1,
    ledger_snapshot: OfficialRawSnapshotV1,
    *,
    as_of_date: str,
) -> PhiladelphiaFedSpfReleaseIndexV1:
    """Join the official archive PDFs to true-deadline/release dates."""
    as_of = _iso_date(as_of_date, "as_of_date")
    _validate_snapshot(
        archive_snapshot,
        uri=SPF_INDEX_URI,
        source_format=OfficialSourceFormat.HTML,
    )
    _validate_snapshot(
        ledger_snapshot,
        uri=SPF_RELEASE_DATES_URI,
        source_format=OfficialSourceFormat.TEXT,
    )
    archive = _parse_archive_entries(archive_snapshot.content)
    ledger = _parse_release_dates(ledger_snapshot.content)
    periods = tuple(
        period
        for period in sorted(set(archive) & set(ledger))
        if ledger[period][1] <= as_of
    )
    if not periods:
        raise ValueError("SPF joined release index has no released surveys")
    if periods != _periods(SPF_FIRST_SURVEY_PERIOD, periods[-1]):
        raise ValueError("SPF joined release index is incomplete")
    if set(archive) != set(ledger):
        raise ValueError("SPF archive and release ledger inventories differ")
    entries = tuple(
        PhiladelphiaFedSpfIndexEntryV1(
            survey_period=period,
            title=archive[period][0],
            true_deadline_date=ledger[period][0],
            release_date=ledger[period][1],
            release_pdf_uri=archive[period][1],
            shutdown_delayed=period in SPF_SHUTDOWN_DELAYED_PERIODS,
        )
        for period in periods
    )
    return PhiladelphiaFedSpfReleaseIndexV1(
        as_of_date=as_of,
        archive_index=_snapshot_artifact(archive_snapshot, "archive-index"),
        release_date_ledger=_snapshot_artifact(
            ledger_snapshot, "release-date-ledger"
        ),
        entries=entries,
    )


def build_philadelphia_fed_spf_release_requests(
    registry: OfficialSourceRegistryV1,
    release_index: PhiladelphiaFedSpfReleaseIndexV1,
) -> tuple[OfficialSourceRequestV1, ...]:
    """Plan every occurrence-specific SPF release PDF request."""
    if not isinstance(registry, OfficialSourceRegistryV1):
        raise TypeError("SPF release requests require a v1 registry")
    if not isinstance(release_index, PhiladelphiaFedSpfReleaseIndexV1):
        raise TypeError("SPF release requests require a v1 release index")
    return tuple(
        _request(
            registry,
            item.release_pdf_uri,
            OfficialSourceFormat.PDF,
            as_of_date=release_index.as_of_date,
        )
        for item in release_index.entries
    )


def _workbook_forecasts(
    level: OfficialRawSnapshotV1,
    growth: OfficialRawSnapshotV1,
) -> Mapping[str, tuple[PhiladelphiaFedSpfForecastV1, ...]]:
    snapshots = {
        "median-level-workbook": level,
        "median-growth-workbook": growth,
    }
    values: dict[str, dict[str, PhiladelphiaFedSpfForecastV1]] = {}
    for spec in _SPF_FORECAST_SPECS:
        snapshot = snapshots[spec.artifact_role]
        workbook = load_workbook(
            BytesIO(snapshot.content), read_only=True, data_only=True
        )
        if spec.worksheet not in workbook.sheetnames:
            raise ValueError("SPF workbook omits a required worksheet")
        worksheet = workbook[spec.worksheet]
        rows = worksheet.iter_rows(values_only=True)
        try:
            headers = next(rows)
        except StopIteration as exc:
            raise ValueError("SPF workbook worksheet is empty") from exc
        names = tuple(str(item) if item is not None else "" for item in headers)
        if len(names) < 3 or names[:2] != ("YEAR", "QUARTER"):
            raise ValueError("SPF workbook worksheet header differs")
        try:
            column_index = tuple(name.upper() for name in names).index(
                spec.column.upper()
            )
        except ValueError as exc:
            raise ValueError(
                "SPF workbook omits a required forecast column"
            ) from exc
        for row in rows:
            if len(row) <= column_index or row[0] is None or row[1] is None:
                continue
            year = int(cast(float, row[0]))
            quarter = int(cast(float, row[1]))
            if year < 2000:
                continue
            raw_value = row[column_index]
            value = _finite(raw_value, f"{spec.worksheet}.{spec.column}")
            lexical = str(raw_value)
            period = f"{year:04d}-Q{quarter}"
            if spec.concept_key in values.setdefault(period, {}):
                raise ValueError("SPF workbook repeats a survey-period row")
            values[period][spec.concept_key] = PhiladelphiaFedSpfForecastV1(
                survey_period=period,
                concept_key=spec.concept_key,
                artifact_role=spec.artifact_role,
                worksheet=spec.worksheet,
                column=spec.column,
                unit=spec.unit,
                value=value,
                lexical_value=lexical,
            )
        workbook.close()
    expected_keys = tuple(item.concept_key for item in _SPF_FORECAST_SPECS)
    result: dict[str, tuple[PhiladelphiaFedSpfForecastV1, ...]] = {}
    for period, observations in values.items():
        if set(observations) != set(expected_keys):
            raise ValueError(
                "SPF workbook period has an incomplete forecast set"
            )
        result[period] = tuple(observations[key] for key in expected_keys)
    return MappingProxyType(result)


def _validate_pdf(snapshot: OfficialRawSnapshotV1, uri: str) -> None:
    _validate_snapshot(
        snapshot, uri=uri, source_format=OfficialSourceFormat.PDF
    )
    if not snapshot.content.startswith(b"%PDF-"):
        raise ValueError("SPF release PDF signature differs")
    try:
        pages = len(PdfReader(BytesIO(snapshot.content), strict=False).pages)
    except Exception as exc:
        raise ValueError("SPF release PDF cannot be read") from exc
    if pages < 1:
        raise ValueError("SPF release PDF contains no pages")


def build_philadelphia_fed_spf_archive_manifest(
    registry: OfficialSourceRegistryV1,
    profile: UnitedStatesBackfillProfileV1,
    release_index: PhiladelphiaFedSpfReleaseIndexV1,
    snapshots_by_uri: Mapping[str, OfficialRawSnapshotV1],
) -> PhiladelphiaFedSpfArchiveManifestV1:
    """Build a complete SPF manifest from one retained official corpus."""
    if not isinstance(registry, OfficialSourceRegistryV1):
        raise TypeError("SPF manifest requires a v1 registry")
    if not isinstance(profile, UnitedStatesBackfillProfileV1):
        raise TypeError("SPF manifest requires a v1 U.S. profile")
    if profile.registry_id != registry.registry_id:
        raise ValueError("SPF registry and profile identities differ")
    program = profile.by_key.get(SPF_PROGRAM_KEY)
    if program is None or program.source_key != SPF_SOURCE_KEY:
        raise ValueError("SPF program source differs")
    required_uris = {
        *_ROLE_TO_URI.values(),
        *(item.release_pdf_uri for item in release_index.entries),
    }
    if set(snapshots_by_uri) != required_uris:
        raise ValueError("SPF retained snapshot inventory differs")
    support: list[PhiladelphiaFedSpfArtifactV1] = []
    for role, uri in _ROLE_TO_URI.items():
        snapshot = snapshots_by_uri[uri]
        _validate_snapshot(
            snapshot, uri=uri, source_format=_ROLE_TO_FORMAT[role]
        )
        support.append(_snapshot_artifact(snapshot, role))
    if support[0] != release_index.archive_index or support[2] != (
        release_index.release_date_ledger
    ):
        raise ValueError("SPF index artifacts changed after parsing")
    forecasts = _workbook_forecasts(
        snapshots_by_uri[SPF_MEDIAN_LEVEL_URI],
        snapshots_by_uri[SPF_MEDIAN_GROWTH_URI],
    )
    publications: list[PhiladelphiaFedSpfPublicationV1] = []
    for entry in release_index.entries:
        if entry.survey_period not in forecasts:
            raise ValueError("SPF workbooks omit an indexed survey period")
        pdf = snapshots_by_uri[entry.release_pdf_uri]
        _validate_pdf(pdf, entry.release_pdf_uri)
        publications.append(
            PhiladelphiaFedSpfPublicationV1(
                survey_period=entry.survey_period,
                true_deadline_date=entry.true_deadline_date,
                release_date=entry.release_date,
                shutdown_delayed=entry.shutdown_delayed,
                release_pdf=_snapshot_artifact(pdf, "release-pdf"),
                forecasts=forecasts[entry.survey_period],
            )
        )
    artifacts = (*support, *(item.release_pdf for item in publications))
    return PhiladelphiaFedSpfArchiveManifestV1(
        registry_id=registry.registry_id,
        profile_id=profile.profile_id,
        release_index=release_index,
        support_artifacts=tuple(support),
        publications=tuple(publications),
        raw_artifact_count=len(artifacts),
        total_content_bytes=sum(item.content_length for item in artifacts),
        forecast_occurrence_count=sum(
            len(item.forecasts) for item in publications
        ),
        shutdown_delayed_count=sum(
            item.shutdown_delayed for item in publications
        ),
        date_only_count=len(publications),
    )


def replay_philadelphia_fed_spf_archive(
    registry: OfficialSourceRegistryV1,
    profile: UnitedStatesBackfillProfileV1,
    snapshots_by_uri: Mapping[str, OfficialRawSnapshotV1],
    expected: PhiladelphiaFedSpfArchiveManifestV1,
) -> PhiladelphiaFedSpfArchiveManifestV1:
    """Rebuild and compare the complete retained SPF corpus."""
    rebuilt_index = parse_philadelphia_fed_spf_release_index(
        snapshots_by_uri[SPF_INDEX_URI],
        snapshots_by_uri[SPF_RELEASE_DATES_URI],
        as_of_date=expected.release_index.as_of_date,
    )
    rebuilt = build_philadelphia_fed_spf_archive_manifest(
        registry, profile, rebuilt_index, snapshots_by_uri
    )
    if rebuilt != expected:
        raise ValueError("SPF retained-corpus replay differs")
    return rebuilt


def philadelphia_fed_spf_coverage_from_manifest(
    manifest: PhiladelphiaFedSpfArchiveManifestV1,
) -> UnitedStatesProgramCoverageV1:
    """Derive complete date-only forecast-vintage coverage for SPF."""
    if not isinstance(manifest, PhiladelphiaFedSpfArchiveManifestV1):
        raise TypeError("SPF coverage requires a v1 manifest")
    count = len(manifest.publications)
    artifacts = (
        *manifest.support_artifacts,
        *(item.release_pdf for item in manifest.publications),
    )
    return UnitedStatesProgramCoverageV1(
        program_key=SPF_PROGRAM_KEY,
        window_start_date=US_BACKFILL_START_DATE,
        window_end_date=manifest.release_index.as_of_date,
        expected_occurrence_count=count,
        schedule_count=count,
        initial_actual_count=count,
        previous_as_known_count=0,
        revision_count=0,
        exact_minute_count=0,
        forecast_count=manifest.forecast_occurrence_count,
        artifact_sha256s=tuple(item.content_sha256 for item in artifacts),
        gap_reasons=(),
        notes=(
            "Each occurrence is a date-only quarterly survey vintage, not a monthly event consensus.",
            "Column 2 is the official current-quarter nowcast horizon; column 1 is the known prior quarter.",
            "Real GDP growth, unemployment, headline CPI inflation, and the three-month bill rate remain separate concepts.",
            "The official ledger preserves true deadlines and four shutdown-delayed surveys.",
        ),
    )


def packaged_philadelphia_fed_spf_manifest_path() -> Path:
    return Path(__file__).with_name("assets") / "us_spf_archive_v1.json"


def packaged_philadelphia_fed_spf_ledger_path() -> Path:
    return (
        Path(__file__).with_name("assets") / "us_spf_release_dates_v1.txt.b64"
    )


def load_packaged_philadelphia_fed_spf_archive_manifest() -> (
    PhiladelphiaFedSpfArchiveManifestV1
):
    """Load and validate the packaged SPF archive manifest."""
    return PhiladelphiaFedSpfArchiveManifestV1.from_json(
        packaged_philadelphia_fed_spf_manifest_path().read_text(
            encoding="utf-8"
        )
    )


def load_packaged_philadelphia_fed_spf_release_dates() -> bytes:
    """Load the exact packaged official SPF release-date ledger."""
    try:
        content = base64.b64decode(
            packaged_philadelphia_fed_spf_ledger_path()
            .read_text(encoding="ascii")
            .strip(),
            validate=True,
        )
    except (OSError, UnicodeError, ValueError) as exc:
        raise ValueError("packaged SPF release-date ledger is invalid") from exc
    manifest = load_packaged_philadelphia_fed_spf_archive_manifest()
    artifact = manifest.release_index.release_date_ledger
    if (
        len(content) != artifact.content_length
        or hashlib.sha256(content).hexdigest() != artifact.content_sha256
    ):
        raise ValueError("packaged SPF release-date ledger differs")
    return content


__all__ = [
    "MAX_SPF_ARTIFACT_BYTES",
    "MAX_SPF_RELEASES",
    "MAX_SPF_TOTAL_BYTES",
    "SPF_ARCHIVE_MANIFEST_SCHEMA_VERSION",
    "SPF_ASSET_ROOT",
    "SPF_DOCUMENTATION_URI",
    "SPF_FIRST_SURVEY_PERIOD",
    "SPF_FORECAST_SCHEMA_VERSION",
    "SPF_INDEX_ENTRY_SCHEMA_VERSION",
    "SPF_INDEX_URI",
    "SPF_LATEST_PACKAGED_SURVEY_PERIOD",
    "SPF_MEDIAN_FORECASTS_URI",
    "SPF_MEDIAN_GROWTH_URI",
    "SPF_MEDIAN_LEVEL_URI",
    "SPF_PROGRAM_KEY",
    "SPF_PUBLICATION_SCHEMA_VERSION",
    "SPF_RELEASE_DATES_URI",
    "SPF_RELEASE_INDEX_SCHEMA_VERSION",
    "SPF_SHUTDOWN_DELAYED_PERIODS",
    "SPF_SOURCE_KEY",
    "PhiladelphiaFedSpfArchiveManifestV1",
    "PhiladelphiaFedSpfArtifactV1",
    "PhiladelphiaFedSpfForecastV1",
    "PhiladelphiaFedSpfIndexEntryV1",
    "PhiladelphiaFedSpfPublicationV1",
    "PhiladelphiaFedSpfReleaseIndexV1",
    "build_philadelphia_fed_spf_archive_manifest",
    "build_philadelphia_fed_spf_release_requests",
    "build_philadelphia_fed_spf_support_requests",
    "load_packaged_philadelphia_fed_spf_archive_manifest",
    "load_packaged_philadelphia_fed_spf_release_dates",
    "packaged_philadelphia_fed_spf_ledger_path",
    "packaged_philadelphia_fed_spf_manifest_path",
    "parse_philadelphia_fed_spf_release_index",
    "philadelphia_fed_spf_coverage_from_manifest",
    "replay_philadelphia_fed_spf_archive",
]
