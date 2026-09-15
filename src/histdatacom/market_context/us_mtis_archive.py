"""Occurrence-specific Census manufacturing-and-trade archive.

The official MTIS archive mixes text tables, legacy workbooks, modern XLSX
files, and PDF release documents.  This module keeps the three headline
concepts distinct, preserves the initially published value for every month,
and binds each revised-prior comparison to the preceding publication.
"""

from __future__ import annotations

import base64
import calendar
import hashlib
import html
import json
import math
import re
from collections.abc import Mapping, Sequence
from dataclasses import dataclass
from datetime import date, datetime
from html.parser import HTMLParser
from io import BytesIO
from pathlib import Path
from types import MappingProxyType
from typing import Any, cast
from urllib.parse import urljoin
from zoneinfo import ZoneInfo

import xlrd
from openpyxl import load_workbook
from pypdf import PdfReader

from histdatacom.market_context.contracts import canonical_contract_json
from histdatacom.market_context.official_sources import (
    OfficialRawSnapshotV1,
    OfficialRequestMethod,
    OfficialSourceFormat,
    OfficialSourceRegistryV1,
    OfficialSourceRequestV1,
)
from histdatacom.market_context.us_backfill import (
    US_BACKFILL_START_DATE,
    UnitedStatesBackfillProfileV1,
    UnitedStatesProgramCoverageV1,
)
from histdatacom.runtime_contracts import JSONValue

MTIS_INDEX_EVIDENCE_SCHEMA_VERSION = "histdatacom.mtis-index-evidence.v1"
MTIS_INDEX_ENTRY_SCHEMA_VERSION = "histdatacom.mtis-index-entry.v1"
MTIS_INDEX_SCHEMA_VERSION = "histdatacom.mtis-index.v1"
MTIS_ARTIFACT_SCHEMA_VERSION = "histdatacom.mtis-artifact.v1"
MTIS_MEASURE_SCHEMA_VERSION = "histdatacom.mtis-measure.v1"
MTIS_PUBLICATION_SCHEMA_VERSION = "histdatacom.mtis-publication.v1"
MTIS_ARCHIVE_MANIFEST_SCHEMA_VERSION = "histdatacom.mtis-archive-manifest.v1"
MTIS_INDEX_ENVELOPE_SCHEMA_VERSION = "histdatacom.mtis-index-envelope.v1"

MTIS_SOURCE_KEY = "us.census.mtis"
MTIS_PROGRAM_KEY = "us.census.manufacturing-trade-inventories"
MTIS_INDEX_URI = "https://www.census.gov/mtis/historic_releases.html"
MTIS_RETIRED_INDEX_URI = "https://www.census.gov/mtis/www/historical/index.html"
MTIS_DIRECTORY_URI = "https://www2.census.gov/mtis/historical/"
MTIS_BRIEFING_URI = "https://www.census.gov/economic-indicators/"
MTIS_RELEASE_SCHEDULE_URI = (
    "https://www.census.gov/economic-indicators/calendar-listview.html"
)
MTIS_FIRST_REFERENCE_PERIOD = "2000-01"
MTIS_PREDECESSOR_REFERENCE_PERIOD = "1999-12"
MTIS_LATEST_PACKAGED_REFERENCE_PERIOD = "2026-06"
MTIS_PREDECESSOR_URI = "https://www2.census.gov/mtis/historical/mtis9912.txt"
MTIS_MEASURE_KEYS = (
    "total-business-sales",
    "total-business-inventories",
    "total-business-inventories-sales-ratio",
)
MTIS_CORRECTED_PERIODS = ("2001-07", "2014-05", "2026-06")
MTIS_RELEASE_METADATA_FALLBACK_PERIODS = (
    "2001-04",
    "2003-04",
    "2009-11",
)
MTIS_DISRUPTION_PERIODS = (
    "2018-11",
    "2018-12",
    "2019-01",
    "2025-08",
    "2025-09",
    "2025-10",
    "2025-11",
    "2025-12",
    "2026-01",
)

MAX_MTIS_INDEX_BYTES = 2 * 1024 * 1024
MAX_MTIS_ARTIFACT_BYTES = 4 * 1024 * 1024
MAX_MTIS_RELEASES = 600
MAX_MTIS_TOTAL_BYTES = 256 * 1024 * 1024

_SHA256_RE = re.compile(r"^[0-9a-f]{64}$")
_YEAR_MONTH_RE = re.compile(r"^\d{4}-(?:0[1-9]|1[0-2])$")
_MONTH_NUMBERS = {
    name.lower(): number
    for number, name in enumerate(calendar.month_name)
    if name
}
_MONTH_PATTERN = "|".join(calendar.month_name[1:])
_DATE_RE = re.compile(
    rf"(?P<month>{_MONTH_PATTERN})\s+(?P<day>\d{{1,2}})\s*,?\s*"
    r"(?P<year>\d{4})",
    re.IGNORECASE,
)
_TIME_RE = re.compile(
    r"(?P<hour>8|10)\s*:\s*(?P<minute>30|00)\s*" r"(?:A\.?\s*M\.?)?",
    re.IGNORECASE,
)
_ZONE_RE = re.compile(r"\b(EST|EDT|ET)\b", re.IGNORECASE)
_ARTIFACT_RE = re.compile(
    r"mtis(?P<year>\d{2})(?P<month>\d{2})\." r"(?P<format>pdf|txt|xls|xlsx)$",
    re.IGNORECASE,
)


@dataclass(frozen=True, slots=True)
class _LegacyReleaseMetadata:
    release_date: str
    release_time: str
    reported_zone: str | None
    witness_uri: str
    source_header_lexical: str


_LEGACY_WITNESS_SPECS: tuple[tuple[str, str, str], ...] = (
    (
        "20000229040444",
        "http://www.census.gov/mtis/www/current.html",
        "2000-02-29T04:04:44Z",
    ),
    (
        "20000510124016",
        "http://www.census.gov/mtis/www/current.html",
        "2000-05-10T12:40:16Z",
    ),
    (
        "20000520033743",
        "http://www.census.gov/mtis/www/current.html",
        "2000-05-20T03:37:43Z",
    ),
    (
        "20000622191557",
        "http://www.census.gov/mtis/www/current.html",
        "2000-06-22T19:15:57Z",
    ),
    (
        "20000815053030",
        "http://www.census.gov/mtis/www/current.html",
        "2000-08-15T05:30:30Z",
    ),
    (
        "20001008090329",
        "http://www.census.gov/mtis/www/current.html",
        "2000-10-08T09:03:29Z",
    ),
    (
        "20001206223700",
        "http://www.census.gov/mtis/www/current.html",
        "2000-12-06T22:37:00Z",
    ),
)


def _legacy_replay_uri(timestamp: str, original_uri: str) -> str:
    return f"https://web.archive.org/web/{timestamp}id_/{original_uri}"


_LEGACY_WITNESS_URIS = tuple(
    _legacy_replay_uri(timestamp, original)
    for timestamp, original, _ in _LEGACY_WITNESS_SPECS
)


def _legacy_metadata() -> Mapping[str, _LegacyReleaseMetadata]:
    witness = _LEGACY_WITNESS_URIS
    rows = (
        (
            "2000-01",
            "2000-03-15",
            None,
            witness[0],
            "The Manufacturing and Trade Inventories and Sales Report for January is scheduled for release March 15, 2000 at 8:30 a.m.",
        ),
        (
            "2000-02",
            "2000-04-14",
            "EDT",
            witness[1],
            "FOR WIRE TRANSMISSION 8:30 A.M. EDT, Friday, April 14, 2000",
        ),
        (
            "2000-03",
            "2000-05-12",
            "EDT",
            witness[2],
            "FOR WIRE TRANSMISSION 8:30 A.M. EDT, Friday, May 12, 2000",
        ),
        (
            "2000-04",
            "2000-06-14",
            "EDT",
            witness[3],
            "FOR WIRE TRANSMISSION 8:30 A.M. EDT, Wednesday, June 14, 2000",
        ),
        (
            "2000-05",
            "2000-07-17",
            None,
            witness[3],
            "The Manufacturing and Trade Inventories and Sales Report for May is scheduled for release July 17, 2000 at 8:30 a.m.",
        ),
        (
            "2000-06",
            "2000-08-14",
            "EDT",
            witness[4],
            "FOR WIRE TRANSMISSION 8:30 A.M. EDT, Monday, August 14, 2000",
        ),
        (
            "2000-07",
            "2000-09-15",
            "EDT",
            witness[5],
            "FOR WIRE TRANSMISSION 8:30 A.M. EDT, Friday, September 15, 2000",
        ),
        (
            "2000-08",
            "2000-10-16",
            None,
            witness[5],
            "The Manufacturing and Trade Inventories and Sales Report for August is scheduled for release October 16, 2000 at 8:30 a.m.",
        ),
        (
            "2000-09",
            "2000-11-15",
            "EST",
            witness[6],
            "FOR WIRE TRANSMISSION 8:30 A.M. EST, Wednesday, November 15, 2000",
        ),
        (
            "2000-10",
            "2000-12-14",
            None,
            witness[6],
            "The Manufacturing and Trade Inventories and Sales Report for October is scheduled for release December 14, 2000 at 8:30 a.m.",
        ),
        (
            "2000-11",
            "2001-01-16",
            None,
            witness[6],
            "Scheduled release dates for 2001 are as follows: January 16",
        ),
        (
            "2000-12",
            "2001-02-14",
            None,
            witness[6],
            "Scheduled release dates for 2001 are as follows: January 16, February 14",
        ),
    )
    return MappingProxyType(
        {
            period: _LegacyReleaseMetadata(
                release_date=release_date,
                release_time="08:30",
                reported_zone=zone,
                witness_uri=uri,
                source_header_lexical=header,
            )
            for period, release_date, zone, uri, header in rows
        }
    )


_LEGACY_RELEASE_METADATA = _legacy_metadata()
_FALLBACK_METADATA: Mapping[str, tuple[str, str, str | None, str]] = (
    MappingProxyType(
        {
            "2001-04": ("2001-06-14", "08:30", None, "2001-03"),
            "2003-04": ("2003-06-12", "10:00", None, "2003-03"),
            "2009-11": ("2010-01-14", "10:00", "EST", "2009-10"),
        }
    )
)


def _required_text(value: object, name: str) -> str:
    result = str(value or "").strip()
    if not result or len(result.encode("utf-8")) > 32_768:
        raise ValueError(f"{name} is invalid")
    return result


def _optional_text(value: object) -> str | None:
    if value is None:
        return None
    result = str(value).strip()
    return result or None


def _mapping(value: object, name: str) -> Mapping[str, Any]:
    if not isinstance(value, Mapping):
        raise TypeError(f"{name} must be a mapping")
    return cast(Mapping[str, Any], value)


def _sequence(value: object, name: str) -> Sequence[Any]:
    if not isinstance(value, Sequence) or isinstance(value, (str, bytes)):
        raise TypeError(f"{name} must be a sequence")
    return value


def _sha256(value: object, name: str) -> str:
    result = _required_text(value, name).lower()
    if _SHA256_RE.fullmatch(result) is None:
        raise ValueError(f"{name} is not a SHA-256 digest")
    return result


def _year_month(value: object, name: str) -> str:
    result = _required_text(value, name)
    if _YEAR_MONTH_RE.fullmatch(result) is None:
        raise ValueError(f"{name} is not a year-month")
    return result


def _iso_date(value: object, name: str) -> str:
    result = _required_text(value, name)
    try:
        parsed = date.fromisoformat(result)
    except ValueError as exc:
        raise ValueError(f"{name} is not an ISO date") from exc
    if parsed.isoformat() != result:
        raise ValueError(f"{name} is not a canonical ISO date")
    return result


def _https_uri(value: object, name: str) -> str:
    result = _required_text(value, name)
    if not result.startswith("https://"):
        raise ValueError(f"{name} is not an HTTPS URI")
    return result


def _positive_int(value: object, name: str, maximum: int) -> int:
    if (
        not isinstance(value, int)
        or isinstance(value, bool)
        or not 1 <= value <= maximum
    ):
        raise ValueError(f"{name} is outside its bound")
    return value


def _stable_id(prefix: str, payload: Mapping[str, JSONValue]) -> str:
    digest = hashlib.sha256(
        canonical_contract_json(payload).encode("utf-8")
    ).hexdigest()
    return f"{prefix}:sha256:{digest}"


def _next_month(period: str) -> str:
    year, month = map(int, period.split("-"))
    month += 1
    if month == 13:
        year += 1
        month = 1
    return f"{year:04d}-{month:02d}"


def _previous_month(period: str) -> str:
    year, month = map(int, period.split("-"))
    month -= 1
    if month == 0:
        year -= 1
        month = 12
    return f"{year:04d}-{month:02d}"


def _month_end(period: str) -> str:
    year, month = map(int, period.split("-"))
    return f"{period}-{calendar.monthrange(year, month)[1]:02d}"


def _month_distance(left: str, right: str) -> int:
    left_year, left_month = map(int, left.split("-")[:2])
    right_year, right_month = map(int, right.split("-")[:2])
    return (right_year * 12 + right_month) - (left_year * 12 + left_month)


def _artifact_uri(period: str, extension: str) -> str:
    year, month = map(int, period.split("-"))
    return (
        f"{MTIS_DIRECTORY_URI}mtis{year % 100:02d}{month:02d}." f"{extension}"
    )


def _data_format(period: str) -> OfficialSourceFormat:
    if period <= "2012-12" or period == "2014-05":
        return OfficialSourceFormat.TEXT
    if period <= "2022-02":
        return OfficialSourceFormat.XLS
    return OfficialSourceFormat.XLSX


def _data_uri(period: str) -> str:
    extension = {
        OfficialSourceFormat.TEXT: "txt",
        OfficialSourceFormat.XLS: "xls",
        OfficialSourceFormat.XLSX: "xlsx",
    }[_data_format(period)]
    return _artifact_uri(period, extension)


def _pdf_uri(period: str) -> str | None:
    return _artifact_uri(period, "pdf") if period >= "2001-01" else None


class _AnchorParser(HTMLParser):
    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.anchors: list[tuple[str, str]] = []
        self._href: str | None = None
        self._text: list[str] = []

    def handle_starttag(
        self, tag: str, attrs: list[tuple[str, str | None]]
    ) -> None:
        if tag.lower() != "a":
            return
        values = {key.lower(): value or "" for key, value in attrs}
        self._href = values.get("href") or None
        self._text = []

    def handle_data(self, data: str) -> None:
        if self._href is not None:
            self._text.append(data)

    def handle_endtag(self, tag: str) -> None:
        if tag.lower() == "a" and self._href is not None:
            self.anchors.append(
                (self._href, " ".join(" ".join(self._text).split()))
            )
            self._href = None
            self._text = []


def _artifact_links(content: bytes, base_uri: str) -> set[str]:
    try:
        source = content.decode("utf-8-sig")
    except UnicodeDecodeError as exc:
        raise ValueError("MTIS HTML evidence is not UTF-8") from exc
    parser = _AnchorParser()
    parser.feed(source)
    return {
        urljoin(base_uri, html.unescape(href).strip())
        for href, _ in parser.anchors
        if _ARTIFACT_RE.search(html.unescape(href).strip()) is not None
    }


@dataclass(frozen=True, slots=True)
class CensusMTISIndexEvidenceV1:
    """One hash-bound official page or archived Census timing witness."""

    evidence_uri: str
    original_uri: str
    role: str
    captured_at_lexical: str | None
    content_sha256: str
    content_length: int
    evidence_id: str = ""
    schema_version: str = MTIS_INDEX_EVIDENCE_SCHEMA_VERSION

    def __post_init__(self) -> None:
        if self.schema_version != MTIS_INDEX_EVIDENCE_SCHEMA_VERSION:
            raise ValueError("unsupported MTIS index-evidence schema")
        evidence_uri = _https_uri(self.evidence_uri, "evidence_uri")
        original_uri = _required_text(self.original_uri, "original_uri")
        role = _required_text(self.role, "role")
        if role not in {
            "historical-index",
            "official-directory",
            "current-briefing",
            "archived-census-page",
        }:
            raise ValueError("unsupported MTIS evidence role")
        captured = _optional_text(self.captured_at_lexical)
        if role == "archived-census-page":
            if (
                evidence_uri not in _LEGACY_WITNESS_URIS
                or original_uri != "http://www.census.gov/mtis/www/current.html"
                or captured is None
            ):
                raise ValueError("MTIS legacy witness identity differs")
            datetime.fromisoformat(captured.replace("Z", "+00:00"))
        elif captured is not None or not original_uri.startswith("https://"):
            raise ValueError("MTIS official evidence metadata differs")
        sha = _sha256(self.content_sha256, "content_sha256")
        length = _positive_int(
            self.content_length, "content_length", MAX_MTIS_INDEX_BYTES
        )
        object.__setattr__(self, "evidence_uri", evidence_uri)
        object.__setattr__(self, "original_uri", original_uri)
        object.__setattr__(self, "role", role)
        object.__setattr__(self, "captured_at_lexical", captured)
        object.__setattr__(self, "content_sha256", sha)
        object.__setattr__(self, "content_length", length)
        expected = _stable_id("mtis-index-evidence", self.identity_payload())
        if self.evidence_id and self.evidence_id != expected:
            raise ValueError("MTIS index-evidence identity differs")
        object.__setattr__(self, "evidence_id", expected)

    def identity_payload(self) -> dict[str, JSONValue]:
        return {
            "schema_version": self.schema_version,
            "evidence_uri": self.evidence_uri,
            "original_uri": self.original_uri,
            "role": self.role,
            "captured_at_lexical": self.captured_at_lexical,
            "content_sha256": self.content_sha256,
            "content_length": self.content_length,
        }

    def to_dict(self) -> dict[str, JSONValue]:
        return {**self.identity_payload(), "evidence_id": self.evidence_id}

    @classmethod
    def from_dict(cls, data: Mapping[str, Any]) -> CensusMTISIndexEvidenceV1:
        return cls(
            evidence_uri=str(data.get("evidence_uri", "")),
            original_uri=str(data.get("original_uri", "")),
            role=str(data.get("role", "")),
            captured_at_lexical=_optional_text(data.get("captured_at_lexical")),
            content_sha256=str(data.get("content_sha256", "")),
            content_length=cast(int, data.get("content_length")),
            evidence_id=str(data.get("evidence_id", "")),
            schema_version=str(data.get("schema_version", "")),
        )


@dataclass(frozen=True, slots=True)
class CensusMTISReleaseIndexEntryV1:
    """One monthly MTIS release and any bounded archive correction."""

    reference_period: str
    listed_data_uri: str | None
    listed_pdf_uri: str | None
    data_uri: str
    pdf_uri: str | None
    data_format: OfficialSourceFormat
    correction_kind: str
    entry_id: str = ""
    schema_version: str = MTIS_INDEX_ENTRY_SCHEMA_VERSION

    def __post_init__(self) -> None:
        if self.schema_version != MTIS_INDEX_ENTRY_SCHEMA_VERSION:
            raise ValueError("unsupported MTIS index-entry schema")
        period = _year_month(self.reference_period, "reference_period")
        listed_data = _optional_text(self.listed_data_uri)
        listed_pdf = _optional_text(self.listed_pdf_uri)
        data_uri = _https_uri(self.data_uri, "data_uri")
        pdf_uri = _optional_text(self.pdf_uri)
        if listed_data is not None:
            listed_data = _https_uri(listed_data, "listed_data_uri")
        if listed_pdf is not None:
            listed_pdf = _https_uri(listed_pdf, "listed_pdf_uri")
        if pdf_uri is not None:
            pdf_uri = _https_uri(pdf_uri, "pdf_uri")
        data_format = OfficialSourceFormat.from_value(self.data_format)
        correction = _required_text(self.correction_kind, "correction_kind")
        expected_data = _data_uri(period)
        expected_pdf = _pdf_uri(period)
        expected_listed_data: str | None = expected_data
        expected_listed_pdf: str | None = expected_pdf
        expected_correction = "none"
        if period == "2001-07":
            expected_listed_pdf = None
            expected_correction = "unlisted-pdf"
        elif period == "2014-05":
            expected_listed_data = _artifact_uri(period, "xls")
            expected_correction = "broken-spreadsheet-text-fallback"
        elif period == "2026-06":
            expected_listed_data = None
            expected_listed_pdf = None
            expected_correction = "current-not-yet-indexed"
        if (
            data_uri,
            pdf_uri,
            data_format,
            listed_data,
            listed_pdf,
            correction,
        ) != (
            expected_data,
            expected_pdf,
            _data_format(period),
            expected_listed_data,
            expected_listed_pdf,
            expected_correction,
        ):
            raise ValueError("MTIS index-entry binding differs")
        object.__setattr__(self, "reference_period", period)
        object.__setattr__(self, "listed_data_uri", listed_data)
        object.__setattr__(self, "listed_pdf_uri", listed_pdf)
        object.__setattr__(self, "data_uri", data_uri)
        object.__setattr__(self, "pdf_uri", pdf_uri)
        object.__setattr__(self, "data_format", data_format)
        object.__setattr__(self, "correction_kind", correction)
        expected = _stable_id("mtis-index-entry", self.identity_payload())
        if self.entry_id and self.entry_id != expected:
            raise ValueError("MTIS index-entry identity differs")
        object.__setattr__(self, "entry_id", expected)

    @property
    def link_corrected(self) -> bool:
        return self.correction_kind != "none"

    def identity_payload(self) -> dict[str, JSONValue]:
        return {
            "schema_version": self.schema_version,
            "reference_period": self.reference_period,
            "listed_data_uri": self.listed_data_uri,
            "listed_pdf_uri": self.listed_pdf_uri,
            "data_uri": self.data_uri,
            "pdf_uri": self.pdf_uri,
            "data_format": self.data_format.value,
            "correction_kind": self.correction_kind,
        }

    def to_dict(self) -> dict[str, JSONValue]:
        return {**self.identity_payload(), "entry_id": self.entry_id}

    @classmethod
    def from_dict(
        cls, data: Mapping[str, Any]
    ) -> CensusMTISReleaseIndexEntryV1:
        return cls(
            reference_period=str(data.get("reference_period", "")),
            listed_data_uri=_optional_text(data.get("listed_data_uri")),
            listed_pdf_uri=_optional_text(data.get("listed_pdf_uri")),
            data_uri=str(data.get("data_uri", "")),
            pdf_uri=_optional_text(data.get("pdf_uri")),
            data_format=OfficialSourceFormat.from_value(
                str(data.get("data_format", ""))
            ),
            correction_kind=str(data.get("correction_kind", "")),
            entry_id=str(data.get("entry_id", "")),
            schema_version=str(data.get("schema_version", "")),
        )


@dataclass(frozen=True, slots=True)
class CensusMTISReleaseIndexV1:
    """Hash-bound enumeration of the complete selected MTIS corpus."""

    as_of_date: str
    evidence: tuple[CensusMTISIndexEvidenceV1, ...]
    releases: tuple[CensusMTISReleaseIndexEntryV1, ...]
    index_id: str = ""
    schema_version: str = MTIS_INDEX_SCHEMA_VERSION

    def __post_init__(self) -> None:
        if self.schema_version != MTIS_INDEX_SCHEMA_VERSION:
            raise ValueError("unsupported MTIS index schema")
        as_of = _iso_date(self.as_of_date, "as_of_date")
        evidence = tuple(self.evidence)
        roles = tuple(item.role for item in evidence)
        if (
            len(evidence) != 10
            or any(
                not isinstance(item, CensusMTISIndexEvidenceV1)
                for item in evidence
            )
            or roles.count("historical-index") != 1
            or roles.count("official-directory") != 1
            or roles.count("current-briefing") != 1
            or roles.count("archived-census-page") != 7
            or len({item.evidence_uri for item in evidence}) != len(evidence)
        ):
            raise ValueError("MTIS index evidence inventory differs")
        releases = tuple(self.releases)
        if (
            not releases
            or len(releases) > MAX_MTIS_RELEASES
            or any(
                not isinstance(item, CensusMTISReleaseIndexEntryV1)
                for item in releases
            )
        ):
            raise TypeError("MTIS index releases are invalid")
        periods = tuple(item.reference_period for item in releases)
        expected_periods: list[str] = []
        period = MTIS_FIRST_REFERENCE_PERIOD
        while period <= MTIS_LATEST_PACKAGED_REFERENCE_PERIOD:
            expected_periods.append(period)
            period = _next_month(period)
        if periods != tuple(expected_periods):
            raise ValueError("MTIS index monthly coverage differs")
        corrected = tuple(
            item.reference_period for item in releases if item.link_corrected
        )
        if corrected != MTIS_CORRECTED_PERIODS:
            raise ValueError("MTIS correction inventory differs")
        object.__setattr__(self, "as_of_date", as_of)
        object.__setattr__(self, "evidence", evidence)
        object.__setattr__(self, "releases", releases)
        expected = _stable_id("mtis-index", self.identity_payload())
        if self.index_id and self.index_id != expected:
            raise ValueError("MTIS index identity differs")
        object.__setattr__(self, "index_id", expected)

    @property
    def by_period(self) -> Mapping[str, CensusMTISReleaseIndexEntryV1]:
        return MappingProxyType(
            {item.reference_period: item for item in self.releases}
        )

    def identity_payload(self) -> dict[str, JSONValue]:
        return {
            "schema_version": self.schema_version,
            "as_of_date": self.as_of_date,
            "evidence": [item.to_dict() for item in self.evidence],
            "releases": [item.to_dict() for item in self.releases],
        }

    def to_dict(self) -> dict[str, JSONValue]:
        return {**self.identity_payload(), "index_id": self.index_id}

    @classmethod
    def from_dict(cls, data: Mapping[str, Any]) -> CensusMTISReleaseIndexV1:
        return cls(
            as_of_date=str(data.get("as_of_date", "")),
            evidence=tuple(
                CensusMTISIndexEvidenceV1.from_dict(
                    _mapping(item, "MTIS index evidence")
                )
                for item in _sequence(data.get("evidence"), "evidence")
            ),
            releases=tuple(
                CensusMTISReleaseIndexEntryV1.from_dict(
                    _mapping(item, "MTIS index entry")
                )
                for item in _sequence(data.get("releases"), "releases")
            ),
            index_id=str(data.get("index_id", "")),
            schema_version=str(data.get("schema_version", "")),
        )


@dataclass(frozen=True, slots=True)
class CensusMTISArtifactV1:
    """One immutable data table or release document."""

    reference_period: str
    artifact_uri: str
    source_format: OfficialSourceFormat
    role: str
    content_sha256: str
    content_length: int
    artifact_id: str = ""
    schema_version: str = MTIS_ARTIFACT_SCHEMA_VERSION

    def __post_init__(self) -> None:
        if self.schema_version != MTIS_ARTIFACT_SCHEMA_VERSION:
            raise ValueError("unsupported MTIS artifact schema")
        period = _year_month(self.reference_period, "reference_period")
        uri = _https_uri(self.artifact_uri, "artifact_uri")
        source_format = OfficialSourceFormat.from_value(self.source_format)
        role = _required_text(self.role, "role")
        if role == "data-table":
            if source_format not in {
                OfficialSourceFormat.TEXT,
                OfficialSourceFormat.XLS,
                OfficialSourceFormat.XLSX,
            } or uri != _data_uri(period):
                raise ValueError("MTIS data artifact binding differs")
        elif role == "release-document":
            if (
                source_format is not OfficialSourceFormat.PDF
                or uri != _pdf_uri(period)
            ):
                raise ValueError("MTIS PDF artifact binding differs")
        else:
            raise ValueError("unsupported MTIS artifact role")
        sha = _sha256(self.content_sha256, "content_sha256")
        length = _positive_int(
            self.content_length, "content_length", MAX_MTIS_ARTIFACT_BYTES
        )
        object.__setattr__(self, "reference_period", period)
        object.__setattr__(self, "artifact_uri", uri)
        object.__setattr__(self, "source_format", source_format)
        object.__setattr__(self, "role", role)
        object.__setattr__(self, "content_sha256", sha)
        object.__setattr__(self, "content_length", length)
        expected = _stable_id("mtis-artifact", self.identity_payload())
        if self.artifact_id and self.artifact_id != expected:
            raise ValueError("MTIS artifact identity differs")
        object.__setattr__(self, "artifact_id", expected)

    def identity_payload(self) -> dict[str, JSONValue]:
        return {
            "schema_version": self.schema_version,
            "reference_period": self.reference_period,
            "artifact_uri": self.artifact_uri,
            "source_format": self.source_format.value,
            "role": self.role,
            "content_sha256": self.content_sha256,
            "content_length": self.content_length,
        }

    def to_dict(self) -> dict[str, JSONValue]:
        return {**self.identity_payload(), "artifact_id": self.artifact_id}

    @classmethod
    def from_dict(cls, data: Mapping[str, Any]) -> CensusMTISArtifactV1:
        return cls(
            reference_period=str(data.get("reference_period", "")),
            artifact_uri=str(data.get("artifact_uri", "")),
            source_format=OfficialSourceFormat.from_value(
                str(data.get("source_format", ""))
            ),
            role=str(data.get("role", "")),
            content_sha256=str(data.get("content_sha256", "")),
            content_length=cast(int, data.get("content_length")),
            artifact_id=str(data.get("artifact_id", "")),
            schema_version=str(data.get("schema_version", "")),
        )


def _measure_value(
    value: object, lexical: object, measure_key: str, name: str
) -> tuple[int | float, str]:
    text = _required_text(lexical, f"{name}_lexical")
    if measure_key == MTIS_MEASURE_KEYS[2]:
        if re.fullmatch(r"\d\.\d{2}", text) is None:
            raise ValueError(f"{name}_lexical is not a two-decimal ratio")
        numeric: int | float = float(text)
        if not 0.1 <= numeric <= 10:
            raise ValueError(f"{name} ratio is outside its bound")
    else:
        if re.fullmatch(r"(?:[1-9]\d{0,2})(?:,\d{3})+", text) is None:
            raise ValueError(f"{name}_lexical is not a grouped level")
        numeric = int(text.replace(",", ""))
        if not 1 <= numeric <= 10_000_000:
            raise ValueError(f"{name} level is outside its bound")
    if (
        isinstance(value, bool)
        or not isinstance(value, (int, float))
        or not math.isfinite(float(value))
        or value != numeric
    ):
        raise ValueError(f"{name} lexical and numeric values differ")
    return numeric, text


@dataclass(frozen=True, slots=True)
class CensusMTISMeasureV1:
    """One concept's initial, prior-known, and revised-prior triplet."""

    reference_period: str
    measure_key: str
    unit: str
    actual_value: int | float
    actual_lexical: str
    previous_reference_period: str
    revised_previous_value: int | float
    revised_previous_lexical: str
    previous_as_known_value: int | float
    previous_as_known_lexical: str
    data_artifact_uri: str
    previous_artifact_uri: str
    comparison_basis: str = "prior-publication"
    measure_id: str = ""
    schema_version: str = MTIS_MEASURE_SCHEMA_VERSION

    def __post_init__(self) -> None:
        if self.schema_version != MTIS_MEASURE_SCHEMA_VERSION:
            raise ValueError("unsupported MTIS measure schema")
        period = _year_month(self.reference_period, "reference_period")
        measure_key = _required_text(self.measure_key, "measure_key")
        if measure_key not in MTIS_MEASURE_KEYS:
            raise ValueError("unsupported MTIS measure key")
        expected_unit = (
            "ratio"
            if measure_key == MTIS_MEASURE_KEYS[2]
            else "millions-of-us-dollars"
        )
        unit = _required_text(self.unit, "unit")
        if unit != expected_unit:
            raise ValueError("MTIS measure unit differs")
        actual, actual_lexical = _measure_value(
            self.actual_value, self.actual_lexical, measure_key, "actual"
        )
        revised, revised_lexical = _measure_value(
            self.revised_previous_value,
            self.revised_previous_lexical,
            measure_key,
            "revised_previous",
        )
        known, known_lexical = _measure_value(
            self.previous_as_known_value,
            self.previous_as_known_lexical,
            measure_key,
            "previous_as_known",
        )
        previous_period = _year_month(
            self.previous_reference_period, "previous_reference_period"
        )
        if previous_period != _previous_month(period):
            raise ValueError("MTIS previous reference period differs")
        data_uri = _https_uri(self.data_artifact_uri, "data_artifact_uri")
        previous_uri = _https_uri(
            self.previous_artifact_uri, "previous_artifact_uri"
        )
        if data_uri != _data_uri(period):
            raise ValueError("MTIS measure data-artifact binding differs")
        if previous_uri != _data_uri(previous_period):
            raise ValueError("MTIS measure prior-artifact binding differs")
        if self.comparison_basis != "prior-publication":
            raise ValueError("MTIS comparison basis differs")
        object.__setattr__(self, "reference_period", period)
        object.__setattr__(self, "measure_key", measure_key)
        object.__setattr__(self, "unit", unit)
        object.__setattr__(self, "actual_value", actual)
        object.__setattr__(self, "actual_lexical", actual_lexical)
        object.__setattr__(self, "previous_reference_period", previous_period)
        object.__setattr__(self, "revised_previous_value", revised)
        object.__setattr__(self, "revised_previous_lexical", revised_lexical)
        object.__setattr__(self, "previous_as_known_value", known)
        object.__setattr__(self, "previous_as_known_lexical", known_lexical)
        object.__setattr__(self, "data_artifact_uri", data_uri)
        object.__setattr__(self, "previous_artifact_uri", previous_uri)
        expected = _stable_id("mtis-measure", self.identity_payload())
        if self.measure_id and self.measure_id != expected:
            raise ValueError("MTIS measure identity differs")
        object.__setattr__(self, "measure_id", expected)

    @property
    def previous_was_revised(self) -> bool:
        return self.previous_as_known_value != self.revised_previous_value

    def identity_payload(self) -> dict[str, JSONValue]:
        return {
            "schema_version": self.schema_version,
            "reference_period": self.reference_period,
            "measure_key": self.measure_key,
            "unit": self.unit,
            "actual_value": self.actual_value,
            "actual_lexical": self.actual_lexical,
            "previous_reference_period": self.previous_reference_period,
            "revised_previous_value": self.revised_previous_value,
            "revised_previous_lexical": self.revised_previous_lexical,
            "previous_as_known_value": self.previous_as_known_value,
            "previous_as_known_lexical": self.previous_as_known_lexical,
            "data_artifact_uri": self.data_artifact_uri,
            "previous_artifact_uri": self.previous_artifact_uri,
            "comparison_basis": self.comparison_basis,
        }

    def to_dict(self) -> dict[str, JSONValue]:
        return {**self.identity_payload(), "measure_id": self.measure_id}

    @classmethod
    def from_dict(cls, data: Mapping[str, Any]) -> CensusMTISMeasureV1:
        return cls(
            reference_period=str(data.get("reference_period", "")),
            measure_key=str(data.get("measure_key", "")),
            unit=str(data.get("unit", "")),
            actual_value=cast(int | float, data.get("actual_value")),
            actual_lexical=str(data.get("actual_lexical", "")),
            previous_reference_period=str(
                data.get("previous_reference_period", "")
            ),
            revised_previous_value=cast(
                int | float, data.get("revised_previous_value")
            ),
            revised_previous_lexical=str(
                data.get("revised_previous_lexical", "")
            ),
            previous_as_known_value=cast(
                int | float, data.get("previous_as_known_value")
            ),
            previous_as_known_lexical=str(
                data.get("previous_as_known_lexical", "")
            ),
            data_artifact_uri=str(data.get("data_artifact_uri", "")),
            previous_artifact_uri=str(data.get("previous_artifact_uri", "")),
            comparison_basis=str(data.get("comparison_basis", "")),
            measure_id=str(data.get("measure_id", "")),
            schema_version=str(data.get("schema_version", "")),
        )


@dataclass(frozen=True, slots=True)
class CensusMTISPublicationV1:
    """One real monthly publication with three distinct concepts."""

    index_entry_id: str
    reference_period: str
    release_date: str
    released_lexical: str
    released_at_ns: int
    reported_zone: str | None
    zone_consistent: bool
    release_metadata_basis: str
    release_metadata_artifact_uri: str
    source_header_lexical: str
    parser_era: str
    federal_disruption: bool
    artifacts: tuple[CensusMTISArtifactV1, ...]
    measures: tuple[CensusMTISMeasureV1, ...]
    publication_id: str = ""
    schema_version: str = MTIS_PUBLICATION_SCHEMA_VERSION

    def __post_init__(self) -> None:
        if self.schema_version != MTIS_PUBLICATION_SCHEMA_VERSION:
            raise ValueError("unsupported MTIS publication schema")
        entry_id = _required_text(self.index_entry_id, "index_entry_id")
        if not entry_id.startswith("mtis-index-entry:sha256:"):
            raise ValueError("MTIS publication index identity differs")
        period = _year_month(self.reference_period, "reference_period")
        release_date = _iso_date(self.release_date, "release_date")
        released = _required_text(self.released_lexical, "released_lexical")
        try:
            local = datetime.fromisoformat(released).replace(
                tzinfo=ZoneInfo("America/New_York")
            )
        except ValueError as exc:
            raise ValueError("MTIS release timestamp is invalid") from exc
        expected_ns = int(local.timestamp() * 1_000_000_000)
        if (
            local.date().isoformat() != release_date
            or self.released_at_ns != expected_ns
        ):
            raise ValueError("MTIS release timestamp normalization differs")
        lag = _month_distance(period, release_date[:7])
        if lag not in {2, 3}:
            raise ValueError("MTIS release lag is outside its bound")
        disrupted = period in MTIS_DISRUPTION_PERIODS
        if self.federal_disruption is not disrupted or disrupted != (lag == 3):
            raise ValueError("MTIS disruption classification differs")
        reported_zone = _optional_text(self.reported_zone)
        if reported_zone not in {None, "EST", "EDT", "ET"}:
            raise ValueError("MTIS reported zone is invalid")
        consistent = reported_zone in {None, "ET", local.tzname()}
        if self.zone_consistent is not consistent:
            raise ValueError("MTIS release zone classification differs")
        basis = _required_text(
            self.release_metadata_basis, "release_metadata_basis"
        )
        expected_basis = (
            "archived-census-page"
            if period <= "2000-12"
            else (
                "previous-release-schedule"
                if period in MTIS_RELEASE_METADATA_FALLBACK_PERIODS
                else "native-pdf-header"
            )
        )
        if basis != expected_basis:
            raise ValueError("MTIS release-metadata basis differs")
        metadata_uri = _https_uri(
            self.release_metadata_artifact_uri,
            "release_metadata_artifact_uri",
        )
        if basis == "archived-census-page":
            if metadata_uri not in _LEGACY_WITNESS_URIS:
                raise ValueError("MTIS legacy timing witness differs")
        elif basis == "previous-release-schedule":
            prior = _FALLBACK_METADATA[period][3]
            if metadata_uri != _pdf_uri(prior):
                raise ValueError("MTIS prior schedule artifact differs")
        elif metadata_uri != _pdf_uri(period):
            raise ValueError("MTIS native PDF timing artifact differs")
        parser_era = _required_text(self.parser_era, "parser_era")
        expected_era = (
            "legacy-text"
            if period <= "2012-12" or period == "2014-05"
            else ("legacy-xls" if period <= "2022-02" else "modern-xlsx")
        )
        if parser_era != expected_era:
            raise ValueError("MTIS parser era differs")
        artifacts = tuple(self.artifacts)
        expected_formats = [_data_format(period)]
        if period >= "2001-01":
            expected_formats.append(OfficialSourceFormat.PDF)
        if tuple(item.source_format for item in artifacts) != tuple(
            expected_formats
        ) or any(item.reference_period != period for item in artifacts):
            raise ValueError("MTIS publication artifact binding differs")
        measures = tuple(self.measures)
        if tuple(
            item.measure_key for item in measures
        ) != MTIS_MEASURE_KEYS or any(
            item.reference_period != period for item in measures
        ):
            raise ValueError("MTIS publication concept inventory differs")
        if not isinstance(self.federal_disruption, bool):
            raise TypeError("federal_disruption must be boolean")
        object.__setattr__(self, "index_entry_id", entry_id)
        object.__setattr__(self, "reference_period", period)
        object.__setattr__(self, "release_date", release_date)
        object.__setattr__(self, "released_lexical", released)
        object.__setattr__(self, "reported_zone", reported_zone)
        object.__setattr__(self, "zone_consistent", consistent)
        object.__setattr__(self, "release_metadata_basis", basis)
        object.__setattr__(self, "release_metadata_artifact_uri", metadata_uri)
        object.__setattr__(
            self,
            "source_header_lexical",
            _required_text(self.source_header_lexical, "source_header_lexical"),
        )
        object.__setattr__(self, "parser_era", parser_era)
        object.__setattr__(self, "artifacts", artifacts)
        object.__setattr__(self, "measures", measures)
        expected = _stable_id("mtis-publication", self.identity_payload())
        if self.publication_id and self.publication_id != expected:
            raise ValueError("MTIS publication identity differs")
        object.__setattr__(self, "publication_id", expected)

    def identity_payload(self) -> dict[str, JSONValue]:
        return {
            "schema_version": self.schema_version,
            "index_entry_id": self.index_entry_id,
            "reference_period": self.reference_period,
            "release_date": self.release_date,
            "released_lexical": self.released_lexical,
            "released_at_ns": self.released_at_ns,
            "reported_zone": self.reported_zone,
            "zone_consistent": self.zone_consistent,
            "release_metadata_basis": self.release_metadata_basis,
            "release_metadata_artifact_uri": (
                self.release_metadata_artifact_uri
            ),
            "source_header_lexical": self.source_header_lexical,
            "parser_era": self.parser_era,
            "federal_disruption": self.federal_disruption,
            "artifacts": [item.to_dict() for item in self.artifacts],
            "measures": [item.to_dict() for item in self.measures],
        }

    def to_dict(self) -> dict[str, JSONValue]:
        return {
            **self.identity_payload(),
            "publication_id": self.publication_id,
        }

    @classmethod
    def from_dict(cls, data: Mapping[str, Any]) -> CensusMTISPublicationV1:
        return cls(
            index_entry_id=str(data.get("index_entry_id", "")),
            reference_period=str(data.get("reference_period", "")),
            release_date=str(data.get("release_date", "")),
            released_lexical=str(data.get("released_lexical", "")),
            released_at_ns=cast(int, data.get("released_at_ns")),
            reported_zone=_optional_text(data.get("reported_zone")),
            zone_consistent=cast(bool, data.get("zone_consistent")),
            release_metadata_basis=str(data.get("release_metadata_basis", "")),
            release_metadata_artifact_uri=str(
                data.get("release_metadata_artifact_uri", "")
            ),
            source_header_lexical=str(data.get("source_header_lexical", "")),
            parser_era=str(data.get("parser_era", "")),
            federal_disruption=cast(bool, data.get("federal_disruption")),
            artifacts=tuple(
                CensusMTISArtifactV1.from_dict(_mapping(item, "MTIS artifact"))
                for item in _sequence(data.get("artifacts"), "artifacts")
            ),
            measures=tuple(
                CensusMTISMeasureV1.from_dict(_mapping(item, "MTIS measure"))
                for item in _sequence(data.get("measures"), "measures")
            ),
            publication_id=str(data.get("publication_id", "")),
            schema_version=str(data.get("schema_version", "")),
        )


@dataclass(frozen=True, slots=True)
class CensusMTISArchiveManifestV1:
    """Compact replay receipt for all selected MTIS source artifacts."""

    registry_id: str
    profile_id: str
    release_index: CensusMTISReleaseIndexV1
    predecessor: CensusMTISArtifactV1
    predecessor_values: tuple[int | float, ...]
    predecessor_lexicals: tuple[str, ...]
    publications: tuple[CensusMTISPublicationV1, ...]
    raw_artifact_count: int
    unique_content_sha256_count: int
    total_content_bytes: int
    measure_occurrence_count: int
    previous_as_known_measure_count: int
    changed_revision_count: int
    corrected_index_entry_count: int
    legacy_text_publication_count: int
    legacy_xls_publication_count: int
    modern_xlsx_publication_count: int
    release_document_count: int
    archived_page_metadata_count: int
    previous_schedule_metadata_count: int
    native_pdf_metadata_count: int
    disruption_publication_count: int
    exact_minute_count: int
    zone_mismatch_count: int
    manifest_id: str = ""
    schema_version: str = MTIS_ARCHIVE_MANIFEST_SCHEMA_VERSION

    def __post_init__(self) -> None:
        if self.schema_version != MTIS_ARCHIVE_MANIFEST_SCHEMA_VERSION:
            raise ValueError("unsupported MTIS archive-manifest schema")
        registry_id = _required_text(self.registry_id, "registry_id")
        profile_id = _required_text(self.profile_id, "profile_id")
        if not registry_id.startswith("official-source-registry:sha256:"):
            raise ValueError("MTIS registry identity differs")
        if not profile_id.startswith("us-backfill-profile:sha256:"):
            raise ValueError("MTIS profile identity differs")
        if not isinstance(self.release_index, CensusMTISReleaseIndexV1):
            raise TypeError("MTIS manifest requires a v1 index")
        if (
            not isinstance(self.predecessor, CensusMTISArtifactV1)
            or self.predecessor.reference_period
            != MTIS_PREDECESSOR_REFERENCE_PERIOD
            or self.predecessor.artifact_uri != MTIS_PREDECESSOR_URI
            or self.predecessor.role != "data-table"
        ):
            raise ValueError("MTIS predecessor evidence differs")
        predecessor_values = tuple(self.predecessor_values)
        predecessor_lexicals = tuple(self.predecessor_lexicals)
        if len(predecessor_values) != 3 or len(predecessor_lexicals) != 3:
            raise ValueError("MTIS predecessor concept inventory differs")
        normalized_predecessor_values: list[int | float] = []
        normalized_predecessor_lexicals: list[str] = []
        for key, value, lexical in zip(
            MTIS_MEASURE_KEYS,
            predecessor_values,
            predecessor_lexicals,
        ):
            normalized, normalized_lexical = _measure_value(
                value, lexical, key, "predecessor"
            )
            normalized_predecessor_values.append(normalized)
            normalized_predecessor_lexicals.append(normalized_lexical)
        publications = tuple(self.publications)
        periods = tuple(item.reference_period for item in publications)
        if (
            len(publications) != len(self.release_index.releases)
            or any(
                not isinstance(item, CensusMTISPublicationV1)
                for item in publications
            )
            or periods
            != tuple(
                item.reference_period for item in self.release_index.releases
            )
            or tuple(item.released_at_ns for item in publications)
            != tuple(sorted({item.released_at_ns for item in publications}))
        ):
            raise ValueError("MTIS publication chronology differs")
        entries = self.release_index.by_period
        known: dict[str, tuple[int | float, str]] = {}
        for key, value in zip(MTIS_MEASURE_KEYS, normalized_predecessor_values):
            known[key] = (value, self.predecessor.artifact_uri)
        for publication in publications:
            entry = entries[publication.reference_period]
            if publication.index_entry_id != entry.entry_id:
                raise ValueError("MTIS publication index binding differs")
            expected_uris = {entry.data_uri}
            if entry.pdf_uri is not None:
                expected_uris.add(entry.pdf_uri)
            if {
                item.artifact_uri for item in publication.artifacts
            } != expected_uris:
                raise ValueError("MTIS publication artifact inventory differs")
            for measure in publication.measures:
                expected_known = known[measure.measure_key]
                if (
                    measure.previous_as_known_value != expected_known[0]
                    or measure.previous_artifact_uri != expected_known[1]
                ):
                    raise ValueError("MTIS prior lineage differs")
            for measure in publication.measures:
                known[measure.measure_key] = (
                    measure.actual_value,
                    measure.data_artifact_uri,
                )
        content_sha256s = [
            item.content_sha256 for item in self.release_index.evidence
        ]
        content_lengths = [
            item.content_length for item in self.release_index.evidence
        ]
        content_sha256s.append(self.predecessor.content_sha256)
        content_lengths.append(self.predecessor.content_length)
        for publication in publications:
            content_sha256s.extend(
                item.content_sha256 for item in publication.artifacts
            )
            content_lengths.extend(
                item.content_length for item in publication.artifacts
            )
        measures = tuple(
            measure
            for publication in publications
            for measure in publication.measures
        )
        observed = {
            "raw_artifact_count": len(content_sha256s),
            "unique_content_sha256_count": len(set(content_sha256s)),
            "total_content_bytes": sum(content_lengths),
            "measure_occurrence_count": len(measures),
            "previous_as_known_measure_count": len(measures),
            "changed_revision_count": sum(
                item.previous_was_revised for item in measures
            ),
            "corrected_index_entry_count": sum(
                item.link_corrected for item in self.release_index.releases
            ),
            "legacy_text_publication_count": sum(
                item.parser_era == "legacy-text" for item in publications
            ),
            "legacy_xls_publication_count": sum(
                item.parser_era == "legacy-xls" for item in publications
            ),
            "modern_xlsx_publication_count": sum(
                item.parser_era == "modern-xlsx" for item in publications
            ),
            "release_document_count": sum(
                artifact.role == "release-document"
                for publication in publications
                for artifact in publication.artifacts
            ),
            "archived_page_metadata_count": sum(
                item.release_metadata_basis == "archived-census-page"
                for item in publications
            ),
            "previous_schedule_metadata_count": sum(
                item.release_metadata_basis == "previous-release-schedule"
                for item in publications
            ),
            "native_pdf_metadata_count": sum(
                item.release_metadata_basis == "native-pdf-header"
                for item in publications
            ),
            "disruption_publication_count": sum(
                item.federal_disruption for item in publications
            ),
            "exact_minute_count": len(publications),
            "zone_mismatch_count": sum(
                not item.zone_consistent for item in publications
            ),
        }
        for name, expected in observed.items():
            actual = getattr(self, name)
            if (
                not isinstance(actual, int)
                or isinstance(actual, bool)
                or actual != expected
            ):
                raise ValueError(f"MTIS manifest {name} differs")
        if self.total_content_bytes > MAX_MTIS_TOTAL_BYTES:
            raise ValueError("MTIS corpus exceeds its total-byte bound")
        object.__setattr__(self, "registry_id", registry_id)
        object.__setattr__(self, "profile_id", profile_id)
        object.__setattr__(
            self, "predecessor_values", tuple(normalized_predecessor_values)
        )
        object.__setattr__(
            self,
            "predecessor_lexicals",
            tuple(normalized_predecessor_lexicals),
        )
        object.__setattr__(self, "publications", publications)
        expected_id = _stable_id(
            "mtis-archive-manifest", self.identity_payload()
        )
        if self.manifest_id and self.manifest_id != expected_id:
            raise ValueError("MTIS manifest identity differs")
        object.__setattr__(self, "manifest_id", expected_id)

    @property
    def by_period(self) -> Mapping[str, CensusMTISPublicationV1]:
        return MappingProxyType(
            {item.reference_period: item for item in self.publications}
        )

    def identity_payload(self) -> dict[str, JSONValue]:
        return {
            "schema_version": self.schema_version,
            "registry_id": self.registry_id,
            "profile_id": self.profile_id,
            "release_index": self.release_index.to_dict(),
            "predecessor": self.predecessor.to_dict(),
            "predecessor_values": list(self.predecessor_values),
            "predecessor_lexicals": list(self.predecessor_lexicals),
            "publications": [item.to_dict() for item in self.publications],
            "raw_artifact_count": self.raw_artifact_count,
            "unique_content_sha256_count": self.unique_content_sha256_count,
            "total_content_bytes": self.total_content_bytes,
            "measure_occurrence_count": self.measure_occurrence_count,
            "previous_as_known_measure_count": (
                self.previous_as_known_measure_count
            ),
            "changed_revision_count": self.changed_revision_count,
            "corrected_index_entry_count": self.corrected_index_entry_count,
            "legacy_text_publication_count": (
                self.legacy_text_publication_count
            ),
            "legacy_xls_publication_count": self.legacy_xls_publication_count,
            "modern_xlsx_publication_count": (
                self.modern_xlsx_publication_count
            ),
            "release_document_count": self.release_document_count,
            "archived_page_metadata_count": (self.archived_page_metadata_count),
            "previous_schedule_metadata_count": (
                self.previous_schedule_metadata_count
            ),
            "native_pdf_metadata_count": self.native_pdf_metadata_count,
            "disruption_publication_count": (self.disruption_publication_count),
            "exact_minute_count": self.exact_minute_count,
            "zone_mismatch_count": self.zone_mismatch_count,
        }

    def to_dict(self) -> dict[str, JSONValue]:
        return {**self.identity_payload(), "manifest_id": self.manifest_id}

    def to_json(self) -> str:
        return str(canonical_contract_json(self.to_dict()))

    @classmethod
    def from_dict(cls, data: Mapping[str, Any]) -> CensusMTISArchiveManifestV1:
        return cls(
            registry_id=str(data.get("registry_id", "")),
            profile_id=str(data.get("profile_id", "")),
            release_index=CensusMTISReleaseIndexV1.from_dict(
                _mapping(data.get("release_index"), "MTIS index")
            ),
            predecessor=CensusMTISArtifactV1.from_dict(
                _mapping(data.get("predecessor"), "MTIS predecessor")
            ),
            predecessor_values=tuple(
                cast(int | float, item)
                for item in _sequence(
                    data.get("predecessor_values"), "predecessor_values"
                )
            ),
            predecessor_lexicals=tuple(
                str(item)
                for item in _sequence(
                    data.get("predecessor_lexicals"),
                    "predecessor_lexicals",
                )
            ),
            publications=tuple(
                CensusMTISPublicationV1.from_dict(
                    _mapping(item, "MTIS publication")
                )
                for item in _sequence(data.get("publications"), "publications")
            ),
            raw_artifact_count=cast(int, data.get("raw_artifact_count")),
            unique_content_sha256_count=cast(
                int, data.get("unique_content_sha256_count")
            ),
            total_content_bytes=cast(int, data.get("total_content_bytes")),
            measure_occurrence_count=cast(
                int, data.get("measure_occurrence_count")
            ),
            previous_as_known_measure_count=cast(
                int, data.get("previous_as_known_measure_count")
            ),
            changed_revision_count=cast(
                int, data.get("changed_revision_count")
            ),
            corrected_index_entry_count=cast(
                int, data.get("corrected_index_entry_count")
            ),
            legacy_text_publication_count=cast(
                int, data.get("legacy_text_publication_count")
            ),
            legacy_xls_publication_count=cast(
                int, data.get("legacy_xls_publication_count")
            ),
            modern_xlsx_publication_count=cast(
                int, data.get("modern_xlsx_publication_count")
            ),
            release_document_count=cast(
                int, data.get("release_document_count")
            ),
            archived_page_metadata_count=cast(
                int, data.get("archived_page_metadata_count")
            ),
            previous_schedule_metadata_count=cast(
                int, data.get("previous_schedule_metadata_count")
            ),
            native_pdf_metadata_count=cast(
                int, data.get("native_pdf_metadata_count")
            ),
            disruption_publication_count=cast(
                int, data.get("disruption_publication_count")
            ),
            exact_minute_count=cast(int, data.get("exact_minute_count")),
            zone_mismatch_count=cast(int, data.get("zone_mismatch_count")),
            manifest_id=str(data.get("manifest_id", "")),
            schema_version=str(data.get("schema_version", "")),
        )

    @classmethod
    def from_json(cls, text: str) -> CensusMTISArchiveManifestV1:
        try:
            payload = json.loads(text)
        except json.JSONDecodeError as exc:
            raise ValueError("MTIS manifest is invalid JSON") from exc
        return cls.from_dict(_mapping(payload, "MTIS manifest"))


def build_census_mtis_index_requests(
    registry: OfficialSourceRegistryV1, *, as_of_date: str
) -> tuple[OfficialSourceRequestV1, ...]:
    """Plan the three official pages that enumerate the MTIS corpus."""
    if not isinstance(registry, OfficialSourceRegistryV1):
        raise TypeError("MTIS index requests require a v1 registry")
    as_of = _iso_date(as_of_date, "as_of_date")
    source = registry.source(MTIS_SOURCE_KEY)
    return tuple(
        OfficialSourceRequestV1(
            source_key=source.source_key,
            source_id=source.source_id,
            method=OfficialRequestMethod.GET,
            uri=uri,
            source_format=OfficialSourceFormat.HTML,
            parser_id=source.parser_id,
            parser_version=source.parser_version,
            window_start=US_BACKFILL_START_DATE,
            window_end=as_of,
            page_number=1,
        )
        for uri in (MTIS_INDEX_URI, MTIS_DIRECTORY_URI, MTIS_BRIEFING_URI)
    )


def _validate_index_snapshot(
    snapshot: OfficialRawSnapshotV1,
    *,
    uri: str,
    as_of_date: str,
) -> None:
    if not isinstance(snapshot, OfficialRawSnapshotV1):
        raise TypeError("MTIS index evidence requires a v1 snapshot")
    if (
        snapshot.request.source_key != MTIS_SOURCE_KEY
        or snapshot.request.uri != uri
        or snapshot.request.source_format is not OfficialSourceFormat.HTML
        or snapshot.request.window_start != US_BACKFILL_START_DATE
        or snapshot.request.window_end != as_of_date
        or snapshot.status_code != 200
        or not snapshot.content
        or len(snapshot.content) > MAX_MTIS_INDEX_BYTES
    ):
        raise ValueError("MTIS index snapshot differs")


def _index_evidence(
    snapshot: OfficialRawSnapshotV1, role: str
) -> CensusMTISIndexEvidenceV1:
    return CensusMTISIndexEvidenceV1(
        evidence_uri=snapshot.request.uri,
        original_uri=snapshot.request.uri,
        role=role,
        captured_at_lexical=None,
        content_sha256=snapshot.content_sha256,
        content_length=len(snapshot.content),
    )


def _legacy_evidence(
    legacy_witnesses_by_uri: Mapping[str, bytes],
) -> tuple[CensusMTISIndexEvidenceV1, ...]:
    if set(legacy_witnesses_by_uri) != set(_LEGACY_WITNESS_URIS):
        raise ValueError("MTIS legacy witness inventory differs")
    result: list[CensusMTISIndexEvidenceV1] = []
    for timestamp, original_uri, captured_at in _LEGACY_WITNESS_SPECS:
        uri = _legacy_replay_uri(timestamp, original_uri)
        content = legacy_witnesses_by_uri[uri]
        if not isinstance(content, bytes) or not content:
            raise TypeError("MTIS legacy witness must contain bytes")
        if len(content) > MAX_MTIS_INDEX_BYTES:
            raise ValueError("MTIS legacy witness exceeds its byte bound")
        source = content.decode("latin-1").lower()
        if "census" not in source or "inventories and sales" not in source:
            raise ValueError("MTIS legacy witness content differs")
        visible = html.unescape(re.sub(r"<[^>]+>", " ", source))
        visible = re.sub(r"\s+", " ", visible).strip()
        required_headers = tuple(
            item.source_header_lexical.lower()
            for item in _LEGACY_RELEASE_METADATA.values()
            if item.witness_uri == uri
        )
        if any(
            re.sub(r"\s+", " ", header).strip() not in visible
            for header in required_headers
        ):
            raise ValueError("MTIS legacy release-time evidence differs")
        result.append(
            CensusMTISIndexEvidenceV1(
                evidence_uri=uri,
                original_uri=original_uri,
                role="archived-census-page",
                captured_at_lexical=captured_at,
                content_sha256=hashlib.sha256(content).hexdigest(),
                content_length=len(content),
            )
        )
    return tuple(result)


def _validate_current_briefing(content: bytes) -> None:
    try:
        source = content.decode("utf-8-sig")
    except UnicodeDecodeError as exc:
        raise ValueError("MTIS current briefing is not UTF-8") from exc
    match = re.search(r"let\s+g_cidrOutput\s*=\s*(\{.*?\});", source)
    if match is None:
        raise ValueError("MTIS current briefing payload is absent")
    try:
        payload = json.loads(match.group(1))
    except json.JSONDecodeError as exc:
        raise ValueError("MTIS current briefing payload is invalid") from exc
    mtis = _mapping(_mapping(payload, "briefing").get("MTIS"), "MTIS briefing")
    expected = {
        "programName": "Manufacturing and Trade Inventories and Sales",
        "statPeriod": "June 2026",
        "relDate": "August 14th, 2026",
        "relTime": "10:00:00",
        "pdf_link": "/mtis/www/data/pdf/mtis_current.pdf",
        "xls_link": "/mtis/www/mtis_current.xlsx",
    }
    if any(mtis.get(key) != value for key, value in expected.items()):
        raise ValueError("MTIS current briefing facts differ")


def parse_census_mtis_release_index(
    snapshots_by_uri: Mapping[str, OfficialRawSnapshotV1],
    legacy_witnesses_by_uri: Mapping[str, bytes],
    *,
    as_of_date: str,
) -> CensusMTISReleaseIndexV1:
    """Parse the archive and close three explicitly bounded defects."""
    as_of = _iso_date(as_of_date, "as_of_date")
    required = {MTIS_INDEX_URI, MTIS_DIRECTORY_URI, MTIS_BRIEFING_URI}
    if set(snapshots_by_uri) != required:
        raise ValueError("MTIS official index-evidence inventory differs")
    for uri in required:
        _validate_index_snapshot(
            snapshots_by_uri[uri], uri=uri, as_of_date=as_of
        )
    historical = snapshots_by_uri[MTIS_INDEX_URI]
    directory = snapshots_by_uri[MTIS_DIRECTORY_URI]
    briefing = snapshots_by_uri[MTIS_BRIEFING_URI]
    historical_links = _artifact_links(historical.content, MTIS_INDEX_URI)
    directory_links = _artifact_links(directory.content, MTIS_DIRECTORY_URI)
    _validate_current_briefing(briefing.content)
    if as_of < "2026-08-14":
        raise ValueError("MTIS as-of date predates the selected current report")

    releases: list[CensusMTISReleaseIndexEntryV1] = []
    period = MTIS_FIRST_REFERENCE_PERIOD
    while period <= MTIS_LATEST_PACKAGED_REFERENCE_PERIOD:
        data_uri = _data_uri(period)
        pdf_uri = _pdf_uri(period)
        listed_data: str | None = data_uri
        listed_pdf: str | None = pdf_uri
        correction = "none"
        if period == "2001-07":
            listed_pdf = None
            correction = "unlisted-pdf"
            if pdf_uri not in directory_links:
                raise ValueError("MTIS July 2001 PDF is not directory-backed")
        elif period == "2014-05":
            listed_data = _artifact_uri(period, "xls")
            correction = "broken-spreadsheet-text-fallback"
            if data_uri not in directory_links:
                raise ValueError("MTIS May 2014 text fallback is not listed")
        elif period == "2026-06":
            listed_data = None
            listed_pdf = None
            correction = "current-not-yet-indexed"
        if listed_data is not None and listed_data not in historical_links:
            raise ValueError(f"MTIS index omits required data link {period}")
        if listed_pdf is not None and listed_pdf not in historical_links:
            raise ValueError(f"MTIS index omits required PDF link {period}")
        if period == "2026-06" and (
            data_uri in historical_links or pdf_uri in historical_links
        ):
            raise ValueError(
                "MTIS current-report correction is no longer needed"
            )
        releases.append(
            CensusMTISReleaseIndexEntryV1(
                reference_period=period,
                listed_data_uri=listed_data,
                listed_pdf_uri=listed_pdf,
                data_uri=data_uri,
                pdf_uri=pdf_uri,
                data_format=_data_format(period),
                correction_kind=correction,
            )
        )
        period = _next_month(period)

    evidence = (
        _index_evidence(historical, "historical-index"),
        _index_evidence(directory, "official-directory"),
        _index_evidence(briefing, "current-briefing"),
        *_legacy_evidence(legacy_witnesses_by_uri),
    )
    return CensusMTISReleaseIndexV1(
        as_of_date=as_of,
        evidence=evidence,
        releases=tuple(releases),
    )


def _release_request(
    registry: OfficialSourceRegistryV1,
    *,
    uri: str,
    period: str,
    source_format: OfficialSourceFormat,
) -> OfficialSourceRequestV1:
    source = registry.source(MTIS_SOURCE_KEY)
    return OfficialSourceRequestV1(
        source_key=source.source_key,
        source_id=source.source_id,
        method=OfficialRequestMethod.GET,
        uri=uri,
        source_format=source_format,
        parser_id=source.parser_id,
        parser_version=source.parser_version,
        window_start=f"{period}-01",
        window_end=_month_end(period),
    )


def build_census_mtis_archive_requests(
    registry: OfficialSourceRegistryV1,
    release_index: CensusMTISReleaseIndexV1,
) -> tuple[OfficialSourceRequestV1, ...]:
    """Plan the predecessor plus every data table and PDF request."""
    if not isinstance(registry, OfficialSourceRegistryV1):
        raise TypeError("MTIS requests require a v1 registry")
    if not isinstance(release_index, CensusMTISReleaseIndexV1):
        raise TypeError("MTIS requests require a v1 index")
    requests = [
        _release_request(
            registry,
            uri=MTIS_PREDECESSOR_URI,
            period=MTIS_PREDECESSOR_REFERENCE_PERIOD,
            source_format=OfficialSourceFormat.TEXT,
        )
    ]
    for entry in release_index.releases:
        requests.append(
            _release_request(
                registry,
                uri=entry.data_uri,
                period=entry.reference_period,
                source_format=entry.data_format,
            )
        )
        if entry.pdf_uri is not None:
            requests.append(
                _release_request(
                    registry,
                    uri=entry.pdf_uri,
                    period=entry.reference_period,
                    source_format=OfficialSourceFormat.PDF,
                )
            )
    if len({item.uri for item in requests}) != len(requests):
        raise ValueError("MTIS request plan repeats an artifact URI")
    return tuple(requests)


def _validate_release_snapshot(
    snapshot: OfficialRawSnapshotV1,
    *,
    uri: str,
    source_format: OfficialSourceFormat,
) -> None:
    if not isinstance(snapshot, OfficialRawSnapshotV1):
        raise TypeError("MTIS release requires a v1 snapshot")
    content = snapshot.content
    signature_ok = {
        OfficialSourceFormat.PDF: content.startswith(b"%PDF-"),
        OfficialSourceFormat.XLS: content.startswith(b"\xd0\xcf\x11\xe0"),
        OfficialSourceFormat.XLSX: content.startswith(b"PK"),
        OfficialSourceFormat.TEXT: not content.startswith((b"%PDF-", b"PK")),
    }.get(source_format, False)
    if (
        snapshot.request.source_key != MTIS_SOURCE_KEY
        or snapshot.request.uri != uri
        or snapshot.request.source_format is not source_format
        or snapshot.status_code != 200
        or not signature_ok
        or not content
        or len(content) > MAX_MTIS_ARTIFACT_BYTES
    ):
        raise ValueError("MTIS release snapshot differs")


def _decode_text(content: bytes) -> str:
    for encoding in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            return content.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise ValueError("MTIS text artifact cannot be decoded")


def _canonical_cell(value: int | float, measure_key: str) -> str:
    if measure_key == MTIS_MEASURE_KEYS[2]:
        return f"{float(value):.2f}"
    numeric = float(value)
    if not numeric.is_integer():
        raise ValueError("MTIS total-business level is not an integer")
    return f"{int(numeric):,}"


def _is_total_business(value: object) -> bool:
    normalized = re.sub(r"[^a-z]+", " ", str(value).lower()).strip()
    return normalized.startswith("total business")


def _text_table_cells(content: bytes, period: str) -> tuple[str, ...]:
    source = _decode_text(content)
    year, month = map(int, period.split("-"))
    data_month = re.search(
        rf"DATA\s+MONTH\s+0?{month}\s+{year}",
        source[:1000],
        re.IGNORECASE,
    )
    column_month = re.search(
        rf"\b{calendar.month_abbr[month]}\w*\.?\s+{year}\b",
        source[:1500],
        re.IGNORECASE,
    )
    if data_month is None and column_month is None:
        raise ValueError("MTIS text-table reference period differs")
    for line in source.splitlines():
        if re.match(r"^\s*TOTAL BUSINESS\s+", line, re.IGNORECASE):
            values = re.findall(r"(?<![A-Za-z])\d[\d,]*(?:\.\d+)?", line)
            if len(values) >= 9:
                return tuple(values[:9])
    raise ValueError("MTIS text table has no adjusted total-business row")


def _excel_table_cells(
    content: bytes,
    source_format: OfficialSourceFormat,
    period: str,
) -> tuple[str, ...]:
    rows: list[Sequence[object]] = []
    if source_format is OfficialSourceFormat.XLS:
        try:
            workbook = xlrd.open_workbook(file_contents=content)
            for sheet in workbook.sheets():
                rows.extend(
                    sheet.row_values(index) for index in range(sheet.nrows)
                )
        except Exception as exc:
            raise ValueError("MTIS XLS artifact is not readable") from exc
    elif source_format is OfficialSourceFormat.XLSX:
        try:
            workbook = load_workbook(
                BytesIO(content), read_only=True, data_only=True
            )
            try:
                for sheet in workbook.worksheets:
                    rows.extend(sheet.iter_rows(values_only=True))
            finally:
                workbook.close()
        except Exception as exc:
            raise ValueError("MTIS XLSX artifact is not readable") from exc
    else:
        raise ValueError("MTIS spreadsheet format differs")
    year, month = map(int, period.split("-"))
    heading = " ".join(str(item) for row in rows[:15] for item in row)
    if (
        re.search(
            rf"\b{calendar.month_abbr[month]}\w*\.?\s+{year}\b",
            heading,
            re.IGNORECASE,
        )
        is None
    ):
        raise ValueError("MTIS spreadsheet reference period differs")
    for row in rows:
        if not row or not _is_total_business(row[0]):
            continue
        values = [
            item
            for item in row[1:]
            if isinstance(item, (int, float)) and not isinstance(item, bool)
        ]
        if len(values) < 9:
            continue
        return tuple(
            _canonical_cell(value, MTIS_MEASURE_KEYS[index // 3])
            for index, value in enumerate(values[:9])
        )
    raise ValueError("MTIS spreadsheet has no adjusted total-business row")


def _source_values(
    snapshot: OfficialRawSnapshotV1, period: str
) -> tuple[tuple[int | float, str, int | float, str], ...]:
    if snapshot.request.source_format is OfficialSourceFormat.TEXT:
        cells = _text_table_cells(snapshot.content, period)
    else:
        cells = _excel_table_cells(
            snapshot.content, snapshot.request.source_format, period
        )
    pairs: list[tuple[int | float, str, int | float, str]] = []
    for key, current_index, revised_index in zip(
        MTIS_MEASURE_KEYS, (0, 3, 6), (1, 4, 7)
    ):
        current_lexical = (
            f"{float(cells[current_index].replace(',', '')):.2f}"
            if key == MTIS_MEASURE_KEYS[2]
            else f"{int(cells[current_index].replace(',', '')):,}"
        )
        revised_lexical = (
            f"{float(cells[revised_index].replace(',', '')):.2f}"
            if key == MTIS_MEASURE_KEYS[2]
            else f"{int(cells[revised_index].replace(',', '')):,}"
        )
        current, current_lexical = _measure_value(
            (
                float(current_lexical)
                if key == MTIS_MEASURE_KEYS[2]
                else int(current_lexical.replace(",", ""))
            ),
            current_lexical,
            key,
            "current",
        )
        revised, revised_lexical = _measure_value(
            (
                float(revised_lexical)
                if key == MTIS_MEASURE_KEYS[2]
                else int(revised_lexical.replace(",", ""))
            ),
            revised_lexical,
            key,
            "revised",
        )
        pairs.append((current, current_lexical, revised, revised_lexical))
    return tuple(pairs)


def _pdf_text(content: bytes, *, all_pages: bool = False) -> str:
    try:
        reader = PdfReader(BytesIO(content), strict=False)
        pages = reader.pages if all_pages else reader.pages[:2]
        source = "\n".join(page.extract_text() or "" for page in pages)
    except Exception as exc:
        raise ValueError("MTIS PDF artifact is not readable") from exc
    return re.sub(r"\s+", " ", source.replace("\xa0", " ")).strip()


def _native_release_metadata(
    content: bytes, period: str
) -> tuple[str, str, str | None, str]:
    source = _pdf_text(content)
    title = re.search(
        r"MANUFACTURING\s+AND\s+TRADE\s+INVENTORIES\s+AND\s+SALES",
        source,
        re.IGNORECASE,
    )
    prefix = source[: title.start()] if title is not None else source[:1200]
    date_match = _DATE_RE.search(prefix)
    time_match = _TIME_RE.search(prefix)
    if date_match is None or time_match is None:
        raise ValueError("MTIS PDF release header is not parseable")
    reference_heading = (
        source[title.end() : title.end() + 500] if title is not None else ""
    )
    year, month = map(int, period.split("-"))
    if (
        re.search(
            rf"\b{calendar.month_name[month]}\s+{year}",
            reference_heading,
            re.IGNORECASE,
        )
        is None
    ):
        raise ValueError("MTIS PDF reference period differs")
    start = max(0, min(date_match.start(), time_match.start()) - 30)
    finish = min(len(prefix), max(date_match.end(), time_match.end()) + 80)
    zone_match = _ZONE_RE.search(prefix[start:finish])
    release_date = (
        f"{int(date_match.group('year')):04d}-"
        f"{_MONTH_NUMBERS[date_match.group('month').lower()]:02d}-"
        f"{int(date_match.group('day')):02d}"
    )
    release_time = (
        f"{int(time_match.group('hour')):02d}:"
        f"{int(time_match.group('minute')):02d}"
    )
    zone = zone_match.group(1).upper() if zone_match is not None else None
    return release_date, release_time, zone, prefix[-700:].strip()


def _scheduled_release_metadata(
    content: bytes, period: str
) -> tuple[str, str, str | None, str]:
    release_date, release_time, zone, _ = _FALLBACK_METADATA[period]
    source = _pdf_text(content, all_pages=True)
    month_name = calendar.month_name[int(period[-2:])]
    date_value = date.fromisoformat(release_date)
    date_text = f"{calendar.month_name[date_value.month]} {date_value.day}, {date_value.year}"
    pattern = re.compile(
        rf"(?:Report|report)\s+for\s+{month_name}\s+is\s+scheduled\s+"
        rf"(?:for\s+release|to\s+be\s+released)\s+"
        rf"{re.escape(date_text)}\s+at\s+"
        rf"0?{int(release_time[:2])}:{release_time[3:]}\s+"
        rf"a\.?m\.?(?:\s+{zone})?",
        re.IGNORECASE,
    )
    match = pattern.search(source)
    if match is None:
        raise ValueError("MTIS previous-release schedule is not parseable")
    return release_date, release_time, zone, match.group(0)


def _release_metadata(
    period: str,
    snapshots_by_uri: Mapping[str, OfficialRawSnapshotV1],
) -> tuple[str, str, str | None, str, str, str]:
    legacy = _LEGACY_RELEASE_METADATA.get(period)
    if legacy is not None:
        return (
            legacy.release_date,
            legacy.release_time,
            legacy.reported_zone,
            legacy.source_header_lexical,
            "archived-census-page",
            legacy.witness_uri,
        )
    fallback = _FALLBACK_METADATA.get(period)
    if fallback is not None:
        prior_uri = cast(str, _pdf_uri(fallback[3]))
        release_date, release_time, zone, header = _scheduled_release_metadata(
            snapshots_by_uri[prior_uri].content, period
        )
        return (
            release_date,
            release_time,
            zone,
            header,
            "previous-release-schedule",
            prior_uri,
        )
    pdf_uri = cast(str, _pdf_uri(period))
    release_date, release_time, zone, header = _native_release_metadata(
        snapshots_by_uri[pdf_uri].content, period
    )
    return (
        release_date,
        release_time,
        zone,
        header,
        "native-pdf-header",
        pdf_uri,
    )


def _artifact(
    snapshot: OfficialRawSnapshotV1,
    *,
    period: str,
    role: str,
) -> CensusMTISArtifactV1:
    return CensusMTISArtifactV1(
        reference_period=period,
        artifact_uri=snapshot.request.uri,
        source_format=snapshot.request.source_format,
        role=role,
        content_sha256=snapshot.content_sha256,
        content_length=len(snapshot.content),
    )


def build_census_mtis_archive_manifest(
    registry: OfficialSourceRegistryV1,
    profile: UnitedStatesBackfillProfileV1,
    release_index: CensusMTISReleaseIndexV1,
    snapshots_by_uri: Mapping[str, OfficialRawSnapshotV1],
) -> CensusMTISArchiveManifestV1:
    """Reparse every table, timestamp, and three-concept vintage chain."""
    if not isinstance(registry, OfficialSourceRegistryV1):
        raise TypeError("MTIS manifest requires a v1 registry")
    if not isinstance(profile, UnitedStatesBackfillProfileV1):
        raise TypeError("MTIS manifest requires a v1 profile")
    if not isinstance(release_index, CensusMTISReleaseIndexV1):
        raise TypeError("MTIS manifest requires a v1 index")
    source = registry.source(MTIS_SOURCE_KEY)
    program = profile.by_key[MTIS_PROGRAM_KEY]
    if (
        profile.registry_id != registry.registry_id
        or program.source_key != MTIS_SOURCE_KEY
        or program.archive_uri != MTIS_INDEX_URI
        or source.archive_uri != MTIS_INDEX_URI
        or source.parser_id != "official.census-mtis.v1"
    ):
        raise ValueError("MTIS registry/profile binding differs")
    requests = build_census_mtis_archive_requests(registry, release_index)
    required_uris = {item.uri for item in requests}
    if set(snapshots_by_uri) != required_uris:
        missing = sorted(required_uris - set(snapshots_by_uri))
        unexpected = sorted(set(snapshots_by_uri) - required_uris)
        raise ValueError(
            "MTIS retained corpus differs: "
            f"missing={missing[:3]}, unexpected={unexpected[:3]}"
        )
    request_by_uri = {item.uri: item for item in requests}
    for uri, snapshot in snapshots_by_uri.items():
        request = request_by_uri[uri]
        _validate_release_snapshot(
            snapshot, uri=uri, source_format=request.source_format
        )

    predecessor_snapshot = snapshots_by_uri[MTIS_PREDECESSOR_URI]
    predecessor_rows = _source_values(
        predecessor_snapshot, MTIS_PREDECESSOR_REFERENCE_PERIOD
    )
    predecessor_values = tuple(item[0] for item in predecessor_rows)
    predecessor_lexicals = tuple(item[1] for item in predecessor_rows)
    predecessor = _artifact(
        predecessor_snapshot,
        period=MTIS_PREDECESSOR_REFERENCE_PERIOD,
        role="data-table",
    )
    known: dict[str, tuple[int | float, str, str]] = {
        key: (value, lexical, MTIS_PREDECESSOR_URI)
        for key, value, lexical in zip(
            MTIS_MEASURE_KEYS, predecessor_values, predecessor_lexicals
        )
    }

    publications: list[CensusMTISPublicationV1] = []
    for entry in release_index.releases:
        period = entry.reference_period
        data_snapshot = snapshots_by_uri[entry.data_uri]
        rows = _source_values(data_snapshot, period)
        artifacts = [_artifact(data_snapshot, period=period, role="data-table")]
        if entry.pdf_uri is not None:
            artifacts.append(
                _artifact(
                    snapshots_by_uri[entry.pdf_uri],
                    period=period,
                    role="release-document",
                )
            )
        measures: list[CensusMTISMeasureV1] = []
        for key, row in zip(MTIS_MEASURE_KEYS, rows):
            current, current_lexical, revised, revised_lexical = row
            prior_value, prior_lexical, prior_uri = known[key]
            measures.append(
                CensusMTISMeasureV1(
                    reference_period=period,
                    measure_key=key,
                    unit=(
                        "ratio"
                        if key == MTIS_MEASURE_KEYS[2]
                        else "millions-of-us-dollars"
                    ),
                    actual_value=current,
                    actual_lexical=current_lexical,
                    previous_reference_period=_previous_month(period),
                    revised_previous_value=revised,
                    revised_previous_lexical=revised_lexical,
                    previous_as_known_value=prior_value,
                    previous_as_known_lexical=prior_lexical,
                    data_artifact_uri=entry.data_uri,
                    previous_artifact_uri=prior_uri,
                )
            )
        for measure in measures:
            known[measure.measure_key] = (
                measure.actual_value,
                measure.actual_lexical,
                entry.data_uri,
            )
        (
            release_date,
            release_time,
            zone,
            source_header,
            metadata_basis,
            metadata_uri,
        ) = _release_metadata(period, snapshots_by_uri)
        released = f"{release_date}T{release_time}:00"
        local = datetime.fromisoformat(released).replace(
            tzinfo=ZoneInfo("America/New_York")
        )
        publications.append(
            CensusMTISPublicationV1(
                index_entry_id=entry.entry_id,
                reference_period=period,
                release_date=release_date,
                released_lexical=released,
                released_at_ns=int(local.timestamp() * 1_000_000_000),
                reported_zone=zone,
                zone_consistent=zone in {None, "ET", local.tzname()},
                release_metadata_basis=metadata_basis,
                release_metadata_artifact_uri=metadata_uri,
                source_header_lexical=source_header,
                parser_era=(
                    "legacy-text"
                    if entry.data_format is OfficialSourceFormat.TEXT
                    else (
                        "legacy-xls"
                        if entry.data_format is OfficialSourceFormat.XLS
                        else "modern-xlsx"
                    )
                ),
                federal_disruption=period in MTIS_DISRUPTION_PERIODS,
                artifacts=tuple(artifacts),
                measures=tuple(measures),
            )
        )

    content_sha256s = [item.content_sha256 for item in release_index.evidence]
    content_lengths = [item.content_length for item in release_index.evidence]
    content_sha256s.append(predecessor.content_sha256)
    content_lengths.append(predecessor.content_length)
    for publication in publications:
        content_sha256s.extend(
            item.content_sha256 for item in publication.artifacts
        )
        content_lengths.extend(
            item.content_length for item in publication.artifacts
        )
    all_measures = tuple(
        measure
        for publication in publications
        for measure in publication.measures
    )
    return CensusMTISArchiveManifestV1(
        registry_id=registry.registry_id,
        profile_id=profile.profile_id,
        release_index=release_index,
        predecessor=predecessor,
        predecessor_values=predecessor_values,
        predecessor_lexicals=predecessor_lexicals,
        publications=tuple(publications),
        raw_artifact_count=len(content_sha256s),
        unique_content_sha256_count=len(set(content_sha256s)),
        total_content_bytes=sum(content_lengths),
        measure_occurrence_count=len(all_measures),
        previous_as_known_measure_count=len(all_measures),
        changed_revision_count=sum(
            item.previous_was_revised for item in all_measures
        ),
        corrected_index_entry_count=sum(
            item.link_corrected for item in release_index.releases
        ),
        legacy_text_publication_count=sum(
            item.parser_era == "legacy-text" for item in publications
        ),
        legacy_xls_publication_count=sum(
            item.parser_era == "legacy-xls" for item in publications
        ),
        modern_xlsx_publication_count=sum(
            item.parser_era == "modern-xlsx" for item in publications
        ),
        release_document_count=sum(
            artifact.role == "release-document"
            for publication in publications
            for artifact in publication.artifacts
        ),
        archived_page_metadata_count=sum(
            item.release_metadata_basis == "archived-census-page"
            for item in publications
        ),
        previous_schedule_metadata_count=sum(
            item.release_metadata_basis == "previous-release-schedule"
            for item in publications
        ),
        native_pdf_metadata_count=sum(
            item.release_metadata_basis == "native-pdf-header"
            for item in publications
        ),
        disruption_publication_count=sum(
            item.federal_disruption for item in publications
        ),
        exact_minute_count=len(publications),
        zone_mismatch_count=sum(
            not item.zone_consistent for item in publications
        ),
    )


def replay_census_mtis_archive(
    registry: OfficialSourceRegistryV1,
    profile: UnitedStatesBackfillProfileV1,
    index_snapshots_by_uri: Mapping[str, OfficialRawSnapshotV1],
    legacy_witnesses_by_uri: Mapping[str, bytes],
    release_snapshots_by_uri: Mapping[str, OfficialRawSnapshotV1],
    expected: CensusMTISArchiveManifestV1,
) -> CensusMTISArchiveManifestV1:
    """Recompute and compare the complete retained MTIS corpus."""
    if not isinstance(expected, CensusMTISArchiveManifestV1):
        raise TypeError("MTIS replay requires a v1 manifest")
    release_index = parse_census_mtis_release_index(
        index_snapshots_by_uri,
        legacy_witnesses_by_uri,
        as_of_date=expected.release_index.as_of_date,
    )
    rebuilt = build_census_mtis_archive_manifest(
        registry,
        profile,
        release_index,
        release_snapshots_by_uri,
    )
    if rebuilt != expected:
        raise ValueError("MTIS retained-corpus replay differs")
    return rebuilt


def census_mtis_coverage_from_manifest(
    manifest: CensusMTISArchiveManifestV1,
) -> UnitedStatesProgramCoverageV1:
    """Derive complete publication-level coverage for the U.S. profile."""
    if not isinstance(manifest, CensusMTISArchiveManifestV1):
        raise TypeError("MTIS coverage requires a v1 manifest")
    publication_count = len(manifest.publications)
    return UnitedStatesProgramCoverageV1(
        program_key=MTIS_PROGRAM_KEY,
        window_start_date=US_BACKFILL_START_DATE,
        window_end_date=manifest.release_index.as_of_date,
        expected_occurrence_count=publication_count,
        schedule_count=publication_count,
        initial_actual_count=publication_count,
        previous_as_known_count=publication_count,
        revision_count=manifest.changed_revision_count,
        exact_minute_count=publication_count,
        forecast_count=0,
        artifact_sha256s=tuple(
            sorted(
                {
                    *(
                        item.content_sha256
                        for item in manifest.release_index.evidence
                    ),
                    manifest.predecessor.content_sha256,
                    *(
                        artifact.content_sha256
                        for publication in manifest.publications
                        for artifact in publication.artifacts
                    ),
                }
            )
        ),
        gap_reasons=(),
        notes=(
            "Each publication preserves sales, inventories, and their ratio as separate concepts.",
            "The May 2014 text table is the bounded fallback for one broken spreadsheet link.",
            "Seven archived Census pages support exact 2000 timing without becoming substitute producers.",
            "Three PDF timestamp gaps use only the preceding report's explicit next-release schedule.",
            "No historical event-level consensus is manufactured from another survey.",
        ),
    )


def packaged_census_mtis_archive_manifest_path() -> Path:
    """Return the packaged MTIS archive-manifest path."""
    return Path(__file__).with_name("assets") / "us_mtis_archive_v1.json"


def packaged_census_mtis_index_path() -> Path:
    """Return the packaged base64 MTIS index-evidence path."""
    return Path(__file__).with_name("assets") / "us_mtis_index_v1.json"


def load_packaged_census_mtis_archive_manifest() -> CensusMTISArchiveManifestV1:
    """Load and validate the packaged MTIS archive manifest."""
    return CensusMTISArchiveManifestV1.from_json(
        packaged_census_mtis_archive_manifest_path().read_text(encoding="utf-8")
    )


def load_packaged_census_mtis_index_evidence() -> (
    tuple[Mapping[str, OfficialRawSnapshotV1], Mapping[str, bytes]]
):
    """Restore exact official pages and archived Census timing witnesses."""
    try:
        payload = json.loads(
            packaged_census_mtis_index_path().read_text(encoding="utf-8")
        )
    except json.JSONDecodeError as exc:
        raise ValueError(
            "packaged MTIS index envelope is invalid JSON"
        ) from exc
    data = _mapping(payload, "MTIS index envelope")
    if data.get("schema_version") != MTIS_INDEX_ENVELOPE_SCHEMA_VERSION:
        raise ValueError("unsupported MTIS index-envelope schema")
    official_rows = _sequence(
        data.get("official_snapshots"), "official_snapshots"
    )
    if len(official_rows) != 3:
        raise ValueError("packaged MTIS official snapshot count differs")
    snapshots: dict[str, OfficialRawSnapshotV1] = {}
    for item in official_rows:
        row = _mapping(item, "MTIS official snapshot envelope")
        encoded = row.get("content_base64")
        if (
            not isinstance(encoded, str)
            or not encoded
            or len(encoded) > (MAX_MTIS_INDEX_BYTES * 4 // 3) + 4
        ):
            raise ValueError("MTIS official snapshot base64 is invalid")
        try:
            content = base64.b64decode(encoded, validate=True)
        except ValueError as exc:
            raise ValueError("MTIS official snapshot is not base64") from exc
        snapshot = OfficialRawSnapshotV1.restore(
            _mapping(row.get("snapshot"), "MTIS official snapshot"),
            content,
        )
        if snapshot.request.uri in snapshots:
            raise ValueError("packaged MTIS official snapshot repeats a URI")
        snapshots[snapshot.request.uri] = snapshot
    legacy_rows = _sequence(data.get("legacy_witnesses"), "legacy_witnesses")
    if len(legacy_rows) != 7:
        raise ValueError("packaged MTIS legacy witness count differs")
    witnesses: dict[str, bytes] = {}
    for item in legacy_rows:
        row = _mapping(item, "MTIS legacy witness envelope")
        uri = _https_uri(row.get("uri"), "legacy witness URI")
        encoded = row.get("content_base64")
        if (
            not isinstance(encoded, str)
            or not encoded
            or len(encoded) > (MAX_MTIS_INDEX_BYTES * 4 // 3) + 4
        ):
            raise ValueError("MTIS legacy witness base64 is invalid")
        if uri in witnesses:
            raise ValueError("packaged MTIS legacy witness repeats a URI")
        try:
            witnesses[uri] = base64.b64decode(encoded, validate=True)
        except ValueError as exc:
            raise ValueError("MTIS legacy witness is not base64") from exc
    if set(snapshots) != {
        MTIS_INDEX_URI,
        MTIS_DIRECTORY_URI,
        MTIS_BRIEFING_URI,
    } or set(witnesses) != set(_LEGACY_WITNESS_URIS):
        raise ValueError("packaged MTIS index evidence inventory differs")
    return MappingProxyType(snapshots), MappingProxyType(witnesses)


def load_packaged_census_mtis_index() -> CensusMTISReleaseIndexV1:
    """Replay the packaged evidence at the manifest's as-of date."""
    manifest = load_packaged_census_mtis_archive_manifest()
    snapshots, witnesses = load_packaged_census_mtis_index_evidence()
    return parse_census_mtis_release_index(
        snapshots,
        witnesses,
        as_of_date=manifest.release_index.as_of_date,
    )


__all__ = [
    "MTIS_ARCHIVE_MANIFEST_SCHEMA_VERSION",
    "MTIS_BRIEFING_URI",
    "MTIS_CORRECTED_PERIODS",
    "MTIS_DIRECTORY_URI",
    "MTIS_DISRUPTION_PERIODS",
    "MTIS_FIRST_REFERENCE_PERIOD",
    "MTIS_INDEX_ENVELOPE_SCHEMA_VERSION",
    "MTIS_INDEX_ENTRY_SCHEMA_VERSION",
    "MTIS_INDEX_EVIDENCE_SCHEMA_VERSION",
    "MTIS_INDEX_SCHEMA_VERSION",
    "MTIS_INDEX_URI",
    "MTIS_LATEST_PACKAGED_REFERENCE_PERIOD",
    "MTIS_MEASURE_KEYS",
    "MTIS_MEASURE_SCHEMA_VERSION",
    "MTIS_PREDECESSOR_REFERENCE_PERIOD",
    "MTIS_PREDECESSOR_URI",
    "MTIS_PROGRAM_KEY",
    "MTIS_PUBLICATION_SCHEMA_VERSION",
    "MTIS_RELEASE_METADATA_FALLBACK_PERIODS",
    "MTIS_RELEASE_SCHEDULE_URI",
    "MTIS_RETIRED_INDEX_URI",
    "MTIS_SOURCE_KEY",
    "CensusMTISArchiveManifestV1",
    "CensusMTISArtifactV1",
    "CensusMTISIndexEvidenceV1",
    "CensusMTISMeasureV1",
    "CensusMTISPublicationV1",
    "CensusMTISReleaseIndexEntryV1",
    "CensusMTISReleaseIndexV1",
    "build_census_mtis_archive_manifest",
    "build_census_mtis_archive_requests",
    "build_census_mtis_index_requests",
    "census_mtis_coverage_from_manifest",
    "load_packaged_census_mtis_archive_manifest",
    "load_packaged_census_mtis_index",
    "load_packaged_census_mtis_index_evidence",
    "packaged_census_mtis_archive_manifest_path",
    "packaged_census_mtis_index_path",
    "parse_census_mtis_release_index",
    "replay_census_mtis_archive",
]
