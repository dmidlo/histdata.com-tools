"""Reusable, fail-closed parsers for retained official-source snapshots.

The acquisition layer in :mod:`histdatacom.market_context.official_sources`
owns network policy and immutable response bytes.  This module turns those
bytes into protocol-neutral records without hiding parser identity, raw
provenance, source locators, schema drift, or reacquisition requirements.
"""

from __future__ import annotations

import csv
import hashlib
import io
import json
import math
import re
import zipfile
from collections.abc import Iterable, Mapping, Sequence
from dataclasses import dataclass, field
from datetime import date, datetime, time
from enum import Enum
from html.parser import HTMLParser
from types import MappingProxyType
from typing import Any, cast
from urllib.parse import urljoin
from xml.etree import ElementTree

import xlrd
from defusedxml.common import DefusedXmlException
from defusedxml.ElementTree import fromstring as safe_xml_fromstring
from openpyxl import load_workbook
from pypdf import PdfReader

from histdatacom.market_context.contracts import canonical_contract_json
from histdatacom.market_context.official_sources import (
    MAX_OFFICIAL_EVENTS,
    OfficialRawSnapshotV1,
    OfficialSourceEntryV1,
    OfficialSourceFormat,
    OfficialSourceParserV1,
    OfficialSourceRegistryV1,
    parse_official_snapshot,
)
from histdatacom.runtime_contracts import JSONValue

OFFICIAL_ADAPTER_RECORD_SCHEMA_VERSION = (
    "histdatacom.official-adapter-record.v1"
)
OFFICIAL_SPREADSHEET_LAYOUT_SCHEMA_VERSION = (
    "histdatacom.official-spreadsheet-layout.v1"
)
OFFICIAL_ADAPTER_QUALIFICATION_SCHEMA_VERSION = (
    "histdatacom.official-adapter-qualification.v1"
)
OFFICIAL_ADAPTER_AUDIT_SCHEMA_VERSION = "histdatacom.official-adapter-audit.v1"

MAX_OFFICIAL_RECORD_BYTES = 1024 * 1024
MAX_OFFICIAL_CELL_BYTES = 256 * 1024
MAX_OFFICIAL_COLUMNS = 16_384
MAX_OFFICIAL_SHEETS = 1_024
MAX_OFFICIAL_ARCHIVE_MEMBERS = 100_000
MAX_OFFICIAL_XML_DEPTH = 128

_KEY_RE = re.compile(r"^[a-z0-9][a-z0-9._:-]{0,255}$")
_SHA256_RE = re.compile(r"^[0-9a-f]{64}$")
_ICAL_LINE_RE = re.compile(r"^[A-Z0-9-]+(?:;[^:]*)?:")
_ARCHIVE_SUFFIXES = (
    ".7z",
    ".bz2",
    ".csv",
    ".gz",
    ".json",
    ".ods",
    ".pdf",
    ".tar",
    ".tgz",
    ".tsv",
    ".xls",
    ".xlsx",
    ".xml",
    ".zip",
)


class OfficialAdapterPack(str, Enum):
    """Stable families of protocol primitives."""

    SDMX = "sdmx"
    JSON_STAT = "json-stat"
    GENERIC_STRUCTURED = "generic-json-delimited"
    SPREADSHEET = "spreadsheet"
    HTML_RELEASE = "html-release"
    RELEASE_FEED = "rss-atom-ics"
    PDF = "pdf"
    STATIC_ARCHIVE = "static-archive"
    DATA_CATALOG = "data-catalog"


class OfficialAdapterRecordKind(str, Enum):
    """Protocol-neutral classifications emitted by built-in parsers."""

    JSON_RECORD = "json-record"
    DELIMITED_ROW = "delimited-row"
    SDMX_OBSERVATION = "sdmx-observation"
    SDMX_DATAFLOW = "sdmx-dataflow"
    SDMX_CODE = "sdmx-code"
    JSON_STAT_OBSERVATION = "json-stat-observation"
    SPREADSHEET_ROW = "spreadsheet-row"
    TEXT_DOCUMENT = "text-document"
    HTML_DOCUMENT = "html-document"
    HTML_TABLE_ROW = "html-table-row"
    HTML_LINK = "html-link"
    RELEASE_FEED_ITEM = "release-feed-item"
    PDF_PAGE = "pdf-page"
    ARCHIVE_MEMBER = "archive-member"
    ARCHIVE_LINK = "archive-link"
    DATASET_CATALOG_ITEM = "dataset-catalog-item"
    REACQUISITION_REQUIREMENT = "reacquisition-requirement"


class OfficialParserFailureCode(str, Enum):
    """Machine-readable reasons for an explicit parser refusal."""

    MALFORMED_DOCUMENT = "malformed-document"
    SCHEMA_DRIFT = "schema-drift"
    RESOURCE_LIMIT = "resource-limit"
    UNSUPPORTED_REPRESENTATION = "unsupported-representation"
    EXTRACTION_FAILED = "extraction-failed"


class OfficialParserError(ValueError):
    """Fail-closed parser error with stable protocol and drift context."""

    def __init__(
        self,
        code: OfficialParserFailureCode,
        message: str,
        *,
        parser_id: str,
        parser_version: str,
        source_key: str,
        locator: str = "$",
    ) -> None:
        self.code = code
        self.parser_id = parser_id
        self.parser_version = parser_version
        self.source_key = source_key
        self.locator = locator
        super().__init__(f"{code.value}: {message} [{source_key} {locator}]")


@dataclass(frozen=True, slots=True)
class OfficialSpreadsheetLayoutV1:
    """Replayable spreadsheet header selection and drift expectation."""

    sheet_name: str
    header_row: int
    expected_header_version: str | None = None
    schema_version: str = OFFICIAL_SPREADSHEET_LAYOUT_SCHEMA_VERSION

    def __post_init__(self) -> None:
        if self.schema_version != OFFICIAL_SPREADSHEET_LAYOUT_SCHEMA_VERSION:
            raise ValueError("unsupported official spreadsheet layout schema")
        name = str(self.sheet_name).strip()
        if not name or len(name) > 512:
            raise ValueError("spreadsheet sheet_name is invalid")
        object.__setattr__(self, "sheet_name", name)
        if (
            isinstance(self.header_row, bool)
            or not 1 <= self.header_row <= 1_000_000
        ):
            raise ValueError("spreadsheet header_row is invalid")
        if self.expected_header_version is not None:
            _digest(self.expected_header_version, "expected_header_version")

    @property
    def layout_id(self) -> str:
        return _stable_id("official-spreadsheet-layout", self.to_dict())

    def to_dict(self) -> dict[str, JSONValue]:
        return {
            "schema_version": self.schema_version,
            "sheet_name": self.sheet_name,
            "header_row": self.header_row,
            "expected_header_version": self.expected_header_version,
        }

    @classmethod
    def from_dict(cls, data: Mapping[str, Any]) -> OfficialSpreadsheetLayoutV1:
        """Restore and validate one serialized layout."""
        return cls(
            sheet_name=str(data.get("sheet_name", "")),
            header_row=_strict_int(data.get("header_row"), "header_row"),
            expected_header_version=_optional_string(
                data.get("expected_header_version")
            ),
            schema_version=str(data.get("schema_version", "")),
        )


@dataclass(frozen=True, slots=True)
class OfficialAdapterRecordV1:
    """One normalized record with an exact path back to retained raw bytes."""

    kind: OfficialAdapterRecordKind
    snapshot: OfficialRawSnapshotV1
    parser_id: str
    parser_version: str
    ordinal: int
    source_locator: str
    fields: Mapping[str, JSONValue]
    parser_configuration: Mapping[str, JSONValue] = field(default_factory=dict)
    schema_version: str = OFFICIAL_ADAPTER_RECORD_SCHEMA_VERSION

    def __post_init__(self) -> None:
        if self.schema_version != OFFICIAL_ADAPTER_RECORD_SCHEMA_VERSION:
            raise ValueError("unsupported official adapter record schema")
        if not isinstance(self.snapshot, OfficialRawSnapshotV1):
            raise TypeError("official adapter record requires a raw snapshot")
        parser_id = _key(self.parser_id, "parser_id")
        parser_version = _key(self.parser_version, "parser_version")
        if (
            parser_id != self.snapshot.request.parser_id
            or parser_version != self.snapshot.request.parser_version
        ):
            raise ValueError("adapter record parser differs from request")
        object.__setattr__(self, "parser_id", parser_id)
        object.__setattr__(self, "parser_version", parser_version)
        if (
            isinstance(self.ordinal, bool)
            or not 0 <= self.ordinal < MAX_OFFICIAL_EVENTS
        ):
            raise ValueError("official adapter record ordinal is invalid")
        locator = str(self.source_locator).strip()
        if not locator or len(locator.encode("utf-8")) > 16_384:
            raise ValueError("official adapter source locator is invalid")
        object.__setattr__(self, "source_locator", locator)
        fields = _json_mapping(self.fields, "fields")
        configuration = _json_mapping(
            self.parser_configuration, "parser_configuration"
        )
        object.__setattr__(self, "fields", MappingProxyType(fields))
        object.__setattr__(
            self, "parser_configuration", MappingProxyType(configuration)
        )
        if (
            len(canonical_contract_json(self.to_dict()).encode("utf-8"))
            > MAX_OFFICIAL_RECORD_BYTES
        ):
            raise ValueError("official adapter record exceeds byte bound")

    @property
    def record_id(self) -> str:
        return _stable_id("official-adapter-record", self.identity_payload())

    def identity_payload(self) -> dict[str, JSONValue]:
        return {
            "schema_version": self.schema_version,
            "kind": self.kind.value,
            "source_key": self.snapshot.request.source_key,
            "source_id": self.snapshot.request.source_id,
            "request_id": self.snapshot.request.request_id,
            "snapshot_id": self.snapshot.snapshot_id,
            "content_sha256": self.snapshot.content_sha256,
            "source_format": self.snapshot.request.source_format.value,
            "parser_id": self.parser_id,
            "parser_version": self.parser_version,
            "ordinal": self.ordinal,
            "source_locator": self.source_locator,
            "parser_configuration": dict(self.parser_configuration),
            "fields": dict(self.fields),
        }

    def to_dict(self) -> dict[str, JSONValue]:
        return {**self.identity_payload(), "record_id": self.record_id}

    @classmethod
    def restore(
        cls,
        data: Mapping[str, Any],
        snapshot: OfficialRawSnapshotV1,
    ) -> OfficialAdapterRecordV1:
        """Restore a record while re-verifying all raw provenance."""
        value = cls(
            kind=OfficialAdapterRecordKind(str(data.get("kind", ""))),
            snapshot=snapshot,
            parser_id=str(data.get("parser_id", "")),
            parser_version=str(data.get("parser_version", "")),
            ordinal=_strict_int(data.get("ordinal"), "ordinal"),
            source_locator=str(data.get("source_locator", "")),
            fields=_mapping(data.get("fields"), "fields"),
            parser_configuration=_mapping(
                data.get("parser_configuration"), "parser_configuration"
            ),
            schema_version=str(data.get("schema_version", "")),
        )
        if canonical_contract_json(value.to_dict()) != canonical_contract_json(
            dict(data)
        ):
            raise ValueError("restored official adapter record differs")
        return value


@dataclass(frozen=True, slots=True)
class OfficialAdapterQualificationV1:
    """Pinned first-party fixture expectation for one adapter pack."""

    pack: OfficialAdapterPack
    source_key: str
    source_format: OfficialSourceFormat
    parser_id: str
    parser_version: str
    fixture_uri: str
    content_sha256: str
    normalized_sha256: str
    expected_record_count: int
    schema_version: str = OFFICIAL_ADAPTER_QUALIFICATION_SCHEMA_VERSION

    def __post_init__(self) -> None:
        if self.schema_version != OFFICIAL_ADAPTER_QUALIFICATION_SCHEMA_VERSION:
            raise ValueError(
                "unsupported official adapter qualification schema"
            )
        object.__setattr__(
            self, "source_key", _key(self.source_key, "source_key")
        )
        object.__setattr__(self, "parser_id", _key(self.parser_id, "parser_id"))
        object.__setattr__(
            self, "parser_version", _key(self.parser_version, "parser_version")
        )
        uri = str(self.fixture_uri).strip()
        if not uri.startswith("https://"):
            raise ValueError("qualification fixture_uri must be HTTPS")
        object.__setattr__(self, "fixture_uri", uri)
        _digest(self.content_sha256, "content_sha256")
        _digest(self.normalized_sha256, "normalized_sha256")
        if (
            isinstance(self.expected_record_count, bool)
            or not 1 <= self.expected_record_count <= MAX_OFFICIAL_EVENTS
        ):
            raise ValueError("qualification expected_record_count is invalid")

    @property
    def qualification_id(self) -> str:
        return _stable_id("official-adapter-qualification", self.to_dict())

    def to_dict(self) -> dict[str, JSONValue]:
        return {
            "schema_version": self.schema_version,
            "pack": self.pack.value,
            "source_key": self.source_key,
            "source_format": self.source_format.value,
            "parser_id": self.parser_id,
            "parser_version": self.parser_version,
            "fixture_uri": self.fixture_uri,
            "content_sha256": self.content_sha256,
            "normalized_sha256": self.normalized_sha256,
            "expected_record_count": self.expected_record_count,
        }

    @classmethod
    def from_dict(
        cls, data: Mapping[str, Any]
    ) -> OfficialAdapterQualificationV1:
        """Restore and validate one pinned fixture qualification."""
        return cls(
            pack=OfficialAdapterPack(str(data.get("pack", ""))),
            source_key=str(data.get("source_key", "")),
            source_format=OfficialSourceFormat.from_value(
                str(data.get("source_format", ""))
            ),
            parser_id=str(data.get("parser_id", "")),
            parser_version=str(data.get("parser_version", "")),
            fixture_uri=str(data.get("fixture_uri", "")),
            content_sha256=str(data.get("content_sha256", "")),
            normalized_sha256=str(data.get("normalized_sha256", "")),
            expected_record_count=_strict_int(
                data.get("expected_record_count"), "expected_record_count"
            ),
            schema_version=str(data.get("schema_version", "")),
        )


@dataclass(frozen=True, slots=True)
class OfficialAdapterAuditV1:
    """Registry coverage and fixture-qualification audit result."""

    registry_id: str
    required_packs: tuple[OfficialAdapterPack, ...]
    available_packs: tuple[OfficialAdapterPack, ...]
    missing_parser_bindings: tuple[str, ...]
    missing_qualifications: tuple[OfficialAdapterPack, ...]
    schema_version: str = OFFICIAL_ADAPTER_AUDIT_SCHEMA_VERSION

    def __post_init__(self) -> None:
        if self.schema_version != OFFICIAL_ADAPTER_AUDIT_SCHEMA_VERSION:
            raise ValueError("unsupported official adapter audit schema")
        if not str(self.registry_id).startswith(
            "official-source-registry:sha256:"
        ):
            raise ValueError("official adapter audit registry_id is invalid")
        for name in ("required_packs", "available_packs"):
            values = tuple(getattr(self, name))
            if (
                tuple(sorted(set(values), key=lambda item: item.value))
                != values
            ):
                raise ValueError(f"{name} must be unique and sorted")
        missing = tuple(
            str(item).strip() for item in self.missing_parser_bindings
        )
        if any(not item for item in missing):
            raise ValueError("missing_parser_bindings contains an empty value")
        object.__setattr__(self, "missing_parser_bindings", missing)
        expected = tuple(
            sorted(
                set(self.missing_qualifications), key=lambda item: item.value
            )
        )
        if expected != self.missing_qualifications:
            raise ValueError("missing_qualifications must be unique and sorted")

    @property
    def complete(self) -> bool:
        return (
            not self.missing_parser_bindings and not self.missing_qualifications
        )

    def to_dict(self) -> dict[str, JSONValue]:
        return {
            "schema_version": self.schema_version,
            "registry_id": self.registry_id,
            "required_packs": [item.value for item in self.required_packs],
            "available_packs": [item.value for item in self.available_packs],
            "missing_parser_bindings": list(self.missing_parser_bindings),
            "missing_qualifications": [
                item.value for item in self.missing_qualifications
            ],
            "complete": self.complete,
        }

    @classmethod
    def from_dict(cls, data: Mapping[str, Any]) -> OfficialAdapterAuditV1:
        """Restore and verify the derived audit status."""
        value = cls(
            registry_id=str(data.get("registry_id", "")),
            required_packs=tuple(
                OfficialAdapterPack(str(item))
                for item in _sequence(
                    data.get("required_packs"), "required_packs"
                )
            ),
            available_packs=tuple(
                OfficialAdapterPack(str(item))
                for item in _sequence(
                    data.get("available_packs"), "available_packs"
                )
            ),
            missing_parser_bindings=tuple(
                str(item)
                for item in _sequence(
                    data.get("missing_parser_bindings"),
                    "missing_parser_bindings",
                )
            ),
            missing_qualifications=tuple(
                OfficialAdapterPack(str(item))
                for item in _sequence(
                    data.get("missing_qualifications"),
                    "missing_qualifications",
                )
            ),
            schema_version=str(data.get("schema_version", "")),
        )
        if data.get("complete") is not value.complete:
            raise ValueError("official adapter audit completion differs")
        return value


class _BaseOfficialParser:
    """Share identity-aware error and record construction."""

    parser_id = ""
    parser_version = "1"
    supported_formats: tuple[OfficialSourceFormat, ...] = ()
    pack: OfficialAdapterPack

    def parse(
        self, snapshot: OfficialRawSnapshotV1, *, max_events: int
    ) -> Sequence[Mapping[str, JSONValue]]:
        """Parse one snapshot in a concrete protocol implementation."""
        raise NotImplementedError

    def error(
        self,
        snapshot: OfficialRawSnapshotV1,
        code: OfficialParserFailureCode,
        message: str,
        locator: str = "$",
    ) -> OfficialParserError:
        return OfficialParserError(
            code,
            message,
            parser_id=self.parser_id,
            parser_version=self.parser_version,
            source_key=snapshot.request.source_key,
            locator=locator,
        )

    def records(
        self,
        snapshot: OfficialRawSnapshotV1,
        values: Iterable[
            tuple[OfficialAdapterRecordKind, str, Mapping[str, JSONValue]]
        ],
        *,
        max_events: int,
        configuration: Mapping[str, JSONValue] | None = None,
    ) -> tuple[Mapping[str, JSONValue], ...]:
        result: list[Mapping[str, JSONValue]] = []
        for ordinal, (kind, locator, fields) in enumerate(values):
            if ordinal >= max_events:
                raise self.error(
                    snapshot,
                    OfficialParserFailureCode.RESOURCE_LIMIT,
                    "normalized record count exceeds max_events",
                    locator,
                )
            record = OfficialAdapterRecordV1(
                kind=kind,
                snapshot=snapshot,
                parser_id=self.parser_id,
                parser_version=self.parser_version,
                ordinal=ordinal,
                source_locator=locator,
                fields=fields,
                parser_configuration=configuration or {},
            )
            result.append(record.to_dict())
        if not result:
            raise self.error(
                snapshot,
                OfficialParserFailureCode.SCHEMA_DRIFT,
                "document contains no recognized records",
            )
        return tuple(result)


class OfficialJsonParserV1(_BaseOfficialParser):
    """Generic JSON document/record parser with deterministic JSON Pointers."""

    parser_id = "official.json.v1"
    supported_formats = (OfficialSourceFormat.JSON, OfficialSourceFormat.CSV)
    pack = OfficialAdapterPack.GENERIC_STRUCTURED

    def parse(
        self, snapshot: OfficialRawSnapshotV1, *, max_events: int
    ) -> Sequence[Mapping[str, JSONValue]]:
        if snapshot.request.source_format is OfficialSourceFormat.CSV:
            return _parse_delimited(self, snapshot, max_events=max_events)
        return _parse_json(self, snapshot, max_events=max_events)


class OfficialCsvParserV1(_BaseOfficialParser):
    """Strict CSV/TSV parser, with HTML support for registered index routes."""

    parser_id = "official.csv.v1"
    supported_formats = (
        OfficialSourceFormat.CSV,
        OfficialSourceFormat.TSV,
        OfficialSourceFormat.HTML,
    )
    pack = OfficialAdapterPack.GENERIC_STRUCTURED

    def parse(
        self, snapshot: OfficialRawSnapshotV1, *, max_events: int
    ) -> Sequence[Mapping[str, JSONValue]]:
        if snapshot.request.source_format is OfficialSourceFormat.HTML:
            return _parse_html(self, snapshot, max_events=max_events)
        return _parse_delimited(self, snapshot, max_events=max_events)


class OfficialJsonStatParserV1(_BaseOfficialParser):
    """JSON-stat dataset decoder with category and dense-array checks."""

    parser_id = "official.json_stat.v1"
    supported_formats = (
        OfficialSourceFormat.JSON_STAT,
        OfficialSourceFormat.CSV,
    )
    pack = OfficialAdapterPack.JSON_STAT

    def parse(
        self, snapshot: OfficialRawSnapshotV1, *, max_events: int
    ) -> Sequence[Mapping[str, JSONValue]]:
        if snapshot.request.source_format is OfficialSourceFormat.CSV:
            return _parse_delimited(self, snapshot, max_events=max_events)
        return _parse_json_stat(self, snapshot, max_events=max_events)


class OfficialSdmx21ParserV1(_BaseOfficialParser):
    """SDMX 2.1 data, dataflow, and codelist parser."""

    parser_id = "official.sdmx_2.1.v1"
    supported_formats = (OfficialSourceFormat.SDMX_21,)
    pack = OfficialAdapterPack.SDMX

    def parse(
        self, snapshot: OfficialRawSnapshotV1, *, max_events: int
    ) -> Sequence[Mapping[str, JSONValue]]:
        return _parse_sdmx(self, snapshot, max_events=max_events)


class OfficialSdmx30ParserV1(_BaseOfficialParser):
    """SDMX 3.0 data/discovery parser with Eurostat JSON-stat support."""

    parser_id = "official.sdmx_3.0.v1"
    supported_formats = (
        OfficialSourceFormat.SDMX_30,
        OfficialSourceFormat.JSON_STAT,
    )
    pack = OfficialAdapterPack.SDMX

    def parse(
        self, snapshot: OfficialRawSnapshotV1, *, max_events: int
    ) -> Sequence[Mapping[str, JSONValue]]:
        if snapshot.request.source_format is OfficialSourceFormat.JSON_STAT:
            return _parse_json_stat(self, snapshot, max_events=max_events)
        return _parse_sdmx(self, snapshot, max_events=max_events)


class OfficialSpreadsheetParserV1(_BaseOfficialParser):
    """XLS/XLSX tabular extraction with replayable header selection."""

    parser_id = "official.xlsx.v1"
    supported_formats = (
        OfficialSourceFormat.XLS,
        OfficialSourceFormat.XLSX,
        OfficialSourceFormat.CSV,
    )
    pack = OfficialAdapterPack.SPREADSHEET

    def __init__(
        self, layouts: Sequence[OfficialSpreadsheetLayoutV1] = ()
    ) -> None:
        checked = tuple(layouts)
        if len({item.sheet_name for item in checked}) != len(checked):
            raise ValueError(
                "spreadsheet layouts contain duplicate sheet names"
            )
        self.layouts = checked

    def parse(
        self, snapshot: OfficialRawSnapshotV1, *, max_events: int
    ) -> Sequence[Mapping[str, JSONValue]]:
        if snapshot.request.source_format is OfficialSourceFormat.CSV:
            return _parse_delimited(self, snapshot, max_events=max_events)
        return _parse_spreadsheet(self, snapshot, max_events=max_events)


class OfficialHtmlParserV1(_BaseOfficialParser):
    """Release-index/press-release parser plus registered file delegates."""

    parser_id = "official.html.v1"
    supported_formats: tuple[OfficialSourceFormat, ...] = (
        OfficialSourceFormat.TEXT,
        OfficialSourceFormat.HTML,
        OfficialSourceFormat.CSV,
        OfficialSourceFormat.XLS,
        OfficialSourceFormat.XLSX,
        OfficialSourceFormat.PDF,
    )
    pack = OfficialAdapterPack.HTML_RELEASE

    def __init__(
        self, layouts: Sequence[OfficialSpreadsheetLayoutV1] = ()
    ) -> None:
        self.layouts = tuple(layouts)

    def parse(
        self, snapshot: OfficialRawSnapshotV1, *, max_events: int
    ) -> Sequence[Mapping[str, JSONValue]]:
        source_format = snapshot.request.source_format
        if source_format is OfficialSourceFormat.TEXT:
            return _parse_text(self, snapshot, max_events=max_events)
        if source_format is OfficialSourceFormat.HTML:
            return _parse_html(self, snapshot, max_events=max_events)
        if source_format in {
            OfficialSourceFormat.CSV,
            OfficialSourceFormat.TSV,
        }:
            return _parse_delimited(self, snapshot, max_events=max_events)
        if source_format in {
            OfficialSourceFormat.XLS,
            OfficialSourceFormat.XLSX,
        }:
            return _parse_spreadsheet(self, snapshot, max_events=max_events)
        return _parse_pdf(self, snapshot, max_events=max_events)


class OfficialCensusMartsParserV1(OfficialHtmlParserV1):
    """Parse Census MARTS archive envelopes with an exact registry identity."""

    parser_id = "official.census-marts.v1"
    supported_formats: tuple[OfficialSourceFormat, ...] = (
        OfficialSourceFormat.XLS,
        OfficialSourceFormat.XLSX,
        OfficialSourceFormat.HTML,
        OfficialSourceFormat.PDF,
    )


class OfficialCensusM3ParserV1(OfficialHtmlParserV1):
    """Parse Census M3 archive pages and PDFs with an exact identity."""

    parser_id = "official.census-m3.v1"
    supported_formats: tuple[OfficialSourceFormat, ...] = (
        OfficialSourceFormat.HTML,
        OfficialSourceFormat.PDF,
    )


class OfficialCensusNrcParserV1(OfficialHtmlParserV1):
    """Parse Census NRC archive pages and release documents."""

    parser_id = "official.census-nrc.v1"
    supported_formats: tuple[OfficialSourceFormat, ...] = (
        OfficialSourceFormat.HTML,
        OfficialSourceFormat.PDF,
        OfficialSourceFormat.TEXT,
    )


class OfficialCensusNrsParserV1(OfficialHtmlParserV1):
    """Parse Census NRS archive pages and release documents."""

    parser_id = "official.census-nrs.v1"
    supported_formats: tuple[OfficialSourceFormat, ...] = (
        OfficialSourceFormat.HTML,
        OfficialSourceFormat.PDF,
        OfficialSourceFormat.TEXT,
    )


class OfficialCensusFt900ParserV1(OfficialHtmlParserV1):
    """Parse Census/BEA FT-900 archive pages and monthly PDFs."""

    parser_id = "official.census-ft900.v1"
    supported_formats: tuple[OfficialSourceFormat, ...] = (
        OfficialSourceFormat.HTML,
        OfficialSourceFormat.PDF,
    )


class OfficialCensusMtisParserV1(OfficialHtmlParserV1):
    """Parse Census MTIS archive pages, tables, and release documents."""

    parser_id = "official.census-mtis.v1"
    supported_formats: tuple[OfficialSourceFormat, ...] = (
        OfficialSourceFormat.XLS,
        OfficialSourceFormat.XLSX,
        OfficialSourceFormat.TEXT,
        OfficialSourceFormat.HTML,
        OfficialSourceFormat.PDF,
    )


class OfficialPhiladelphiaFedSpfParserV1(OfficialHtmlParserV1):
    """Parse Philadelphia Fed SPF pages, workbooks, ledger, and PDFs."""

    parser_id = "official.philadelphia-spf.v1"
    supported_formats: tuple[OfficialSourceFormat, ...] = (
        OfficialSourceFormat.HTML,
        OfficialSourceFormat.TEXT,
        OfficialSourceFormat.XLSX,
        OfficialSourceFormat.PDF,
    )


class OfficialPhiladelphiaFedMbosParserV1(OfficialHtmlParserV1):
    """Parse Philadelphia Fed MBOS archives, pages, tables, and PDFs."""

    parser_id = "official.philadelphia-mbos.v1"
    supported_formats: tuple[OfficialSourceFormat, ...] = (
        OfficialSourceFormat.ARCHIVE,
        OfficialSourceFormat.HTML,
        OfficialSourceFormat.CSV,
        OfficialSourceFormat.XLS,
        OfficialSourceFormat.PDF,
    )

    def parse(
        self, snapshot: OfficialRawSnapshotV1, *, max_events: int
    ) -> Sequence[Mapping[str, JSONValue]]:
        if snapshot.request.source_format is OfficialSourceFormat.ARCHIVE:
            return _parse_archive(self, snapshot, max_events=max_events)
        return super().parse(snapshot, max_events=max_events)


class OfficialFederalReserveFomcParserV1(OfficialHtmlParserV1):
    """Bind Federal Reserve FOMC HTML and PDF artifacts exactly."""

    parser_id = "official.federal-reserve-fomc.v1"
    supported_formats: tuple[OfficialSourceFormat, ...] = (
        OfficialSourceFormat.HTML,
        OfficialSourceFormat.PDF,
    )


class OfficialFederalReserveH6ParserV1(OfficialHtmlParserV1):
    """Bind Federal Reserve H.6 HTML, JSON, and PDF artifacts exactly."""

    parser_id = "official.federal-reserve-h6.v1"
    supported_formats: tuple[OfficialSourceFormat, ...] = (
        OfficialSourceFormat.HTML,
        OfficialSourceFormat.JSON,
        OfficialSourceFormat.PDF,
        OfficialSourceFormat.TEXT,
    )

    def parse(
        self, snapshot: OfficialRawSnapshotV1, *, max_events: int
    ) -> Sequence[Mapping[str, JSONValue]]:
        if snapshot.request.source_format is OfficialSourceFormat.JSON:
            return _parse_json(self, snapshot, max_events=max_events)
        return super().parse(snapshot, max_events=max_events)


class OfficialDolEtaInitialClaimsParserV1(OfficialHtmlParserV1):
    """Bind ETA Initial Claims archive pages and release artifacts exactly."""

    parser_id = "official.dol-eta-initial-claims.v1"
    supported_formats: tuple[OfficialSourceFormat, ...] = (
        OfficialSourceFormat.HTML,
        OfficialSourceFormat.PDF,
    )

    def parse(
        self, snapshot: OfficialRawSnapshotV1, *, max_events: int
    ) -> Sequence[Mapping[str, JSONValue]]:
        if snapshot.request.source_format is OfficialSourceFormat.HTML:
            return _parse_html(
                self,
                snapshot,
                max_events=max_events,
                strict_tables=False,
            )
        return super().parse(snapshot, max_events=max_events)


class OfficialTreasuryMtsParserV1(OfficialHtmlParserV1):
    """Bind FiscalData MTS catalog, API, PDF, and release evidence."""

    parser_id = "official.treasury-mts.v1"
    supported_formats: tuple[OfficialSourceFormat, ...] = (
        OfficialSourceFormat.JSON,
        OfficialSourceFormat.HTML,
        OfficialSourceFormat.PDF,
    )

    def parse(
        self, snapshot: OfficialRawSnapshotV1, *, max_events: int
    ) -> Sequence[Mapping[str, JSONValue]]:
        if snapshot.request.source_format is OfficialSourceFormat.JSON:
            return _parse_json(self, snapshot, max_events=max_events)
        return super().parse(snapshot, max_events=max_events)


class OfficialReleaseFeedParserV1(_BaseOfficialParser):
    """RSS, Atom, and iCalendar schedule/release feed parser."""

    parser_id = "official.release_feed.v1"
    supported_formats = (
        OfficialSourceFormat.RSS,
        OfficialSourceFormat.ATOM,
        OfficialSourceFormat.ICS,
    )
    pack = OfficialAdapterPack.RELEASE_FEED

    def parse(
        self, snapshot: OfficialRawSnapshotV1, *, max_events: int
    ) -> Sequence[Mapping[str, JSONValue]]:
        if snapshot.request.source_format is OfficialSourceFormat.ICS:
            return _parse_ics(self, snapshot, max_events=max_events)
        return _parse_xml_feed(self, snapshot, max_events=max_events)


class OfficialPdfParserV1(_BaseOfficialParser):
    """Last-resort PDF text and conservative table-like row extraction."""

    parser_id = "official.pdf.v1"
    supported_formats = (OfficialSourceFormat.PDF,)
    pack = OfficialAdapterPack.PDF

    def parse(
        self, snapshot: OfficialRawSnapshotV1, *, max_events: int
    ) -> Sequence[Mapping[str, JSONValue]]:
        return _parse_pdf(self, snapshot, max_events=max_events)


class OfficialArchiveParserV1(_BaseOfficialParser):
    """Enumerate static HTML/JSON download indexes and archive members."""

    parser_id = "official.archive.v1"
    supported_formats = (
        OfficialSourceFormat.HTML,
        OfficialSourceFormat.JSON,
        OfficialSourceFormat.ARCHIVE,
    )
    pack = OfficialAdapterPack.STATIC_ARCHIVE

    def parse(
        self, snapshot: OfficialRawSnapshotV1, *, max_events: int
    ) -> Sequence[Mapping[str, JSONValue]]:
        return _parse_archive(self, snapshot, max_events=max_events)


class OfficialDataCatalogParserV1(_BaseOfficialParser):
    """Discover datasets in official JSON/DCAT-style catalogue feeds."""

    parser_id = "official.data_catalog.v1"
    supported_formats = (
        OfficialSourceFormat.DATA_CATALOG,
        OfficialSourceFormat.JSON,
        OfficialSourceFormat.RSS,
        OfficialSourceFormat.ATOM,
    )
    pack = OfficialAdapterPack.DATA_CATALOG

    def parse(
        self, snapshot: OfficialRawSnapshotV1, *, max_events: int
    ) -> Sequence[Mapping[str, JSONValue]]:
        return _parse_data_catalog(self, snapshot, max_events=max_events)


_BUILT_IN_PARSER_TYPES: tuple[type[_BaseOfficialParser], ...] = (
    OfficialJsonParserV1,
    OfficialCsvParserV1,
    OfficialJsonStatParserV1,
    OfficialSdmx21ParserV1,
    OfficialSdmx30ParserV1,
    OfficialSpreadsheetParserV1,
    OfficialHtmlParserV1,
    OfficialCensusMartsParserV1,
    OfficialCensusM3ParserV1,
    OfficialCensusNrcParserV1,
    OfficialCensusNrsParserV1,
    OfficialCensusFt900ParserV1,
    OfficialCensusMtisParserV1,
    OfficialDolEtaInitialClaimsParserV1,
    OfficialFederalReserveFomcParserV1,
    OfficialFederalReserveH6ParserV1,
    OfficialPhiladelphiaFedMbosParserV1,
    OfficialPhiladelphiaFedSpfParserV1,
    OfficialTreasuryMtsParserV1,
    OfficialReleaseFeedParserV1,
    OfficialPdfParserV1,
    OfficialArchiveParserV1,
    OfficialDataCatalogParserV1,
)


def built_in_official_source_parsers() -> Mapping[str, OfficialSourceParserV1]:
    """Return immutable fresh default parser instances keyed by parser ID."""
    values: dict[str, OfficialSourceParserV1] = {
        parser_type.parser_id: parser_type()
        for parser_type in _BUILT_IN_PARSER_TYPES
    }
    return MappingProxyType(values)


def resolve_official_source_parser(
    source: OfficialSourceEntryV1,
    *,
    parsers: Mapping[str, OfficialSourceParserV1] | None = None,
) -> OfficialSourceParserV1:
    """Resolve one exact registry binding without protocol fallback."""
    values = parsers or built_in_official_source_parsers()
    parser = values.get(source.parser_id)
    if parser is None:
        raise ValueError(f"no built-in parser for {source.parser_id}")
    if (
        parser.parser_id != source.parser_id
        or parser.parser_version != source.parser_version
    ):
        raise ValueError(
            "official parser identity differs from source registry"
        )
    supported = {
        OfficialSourceFormat.from_value(item)
        for item in parser.supported_formats
    }
    missing = set(source.formats) - supported
    if missing:
        rendered = ", ".join(sorted(item.value for item in missing))
        raise ValueError(
            f"official parser omits registered formats: {rendered}"
        )
    return parser


def parse_with_built_in_official_adapter(
    snapshot: OfficialRawSnapshotV1,
    source: OfficialSourceEntryV1,
    *,
    max_events: int,
    parsers: Mapping[str, OfficialSourceParserV1] | None = None,
) -> tuple[Mapping[str, JSONValue], ...]:
    """Resolve and run the exact built-in parser for a retained snapshot."""
    if source.source_key != snapshot.request.source_key:
        raise ValueError("official source differs from snapshot request")
    records: tuple[Mapping[str, JSONValue], ...] = parse_official_snapshot(
        snapshot,
        resolve_official_source_parser(source, parsers=parsers),
        max_events=max_events,
    )
    return records


def required_official_adapter_packs(
    registry: OfficialSourceRegistryV1,
) -> tuple[OfficialAdapterPack, ...]:
    """Derive packs required by current source formats and declared routes."""
    formats = {item for source in registry.sources for item in source.formats}
    result: set[OfficialAdapterPack] = set()
    mapping = {
        OfficialSourceFormat.SDMX_21: OfficialAdapterPack.SDMX,
        OfficialSourceFormat.SDMX_30: OfficialAdapterPack.SDMX,
        OfficialSourceFormat.JSON_STAT: OfficialAdapterPack.JSON_STAT,
        OfficialSourceFormat.JSON: OfficialAdapterPack.GENERIC_STRUCTURED,
        OfficialSourceFormat.CSV: OfficialAdapterPack.GENERIC_STRUCTURED,
        OfficialSourceFormat.TSV: OfficialAdapterPack.GENERIC_STRUCTURED,
        OfficialSourceFormat.XLS: OfficialAdapterPack.SPREADSHEET,
        OfficialSourceFormat.XLSX: OfficialAdapterPack.SPREADSHEET,
        OfficialSourceFormat.HTML: OfficialAdapterPack.HTML_RELEASE,
        OfficialSourceFormat.RSS: OfficialAdapterPack.RELEASE_FEED,
        OfficialSourceFormat.ATOM: OfficialAdapterPack.RELEASE_FEED,
        OfficialSourceFormat.ICS: OfficialAdapterPack.RELEASE_FEED,
        OfficialSourceFormat.TEXT: OfficialAdapterPack.HTML_RELEASE,
        OfficialSourceFormat.PDF: OfficialAdapterPack.PDF,
        OfficialSourceFormat.ARCHIVE: OfficialAdapterPack.STATIC_ARCHIVE,
        OfficialSourceFormat.DATA_CATALOG: OfficialAdapterPack.DATA_CATALOG,
    }
    result.update(mapping[item] for item in formats)
    if any(source.archive_uri is not None for source in registry.sources):
        result.add(OfficialAdapterPack.STATIC_ARCHIVE)
    return tuple(sorted(result, key=lambda item: item.value))


def audit_official_adapter_coverage(
    registry: OfficialSourceRegistryV1,
    qualifications: Sequence[OfficialAdapterQualificationV1] = (),
    *,
    parsers: Mapping[str, OfficialSourceParserV1] | None = None,
) -> OfficialAdapterAuditV1:
    """Audit exact registry/parser bindings and required fixture packs."""
    values = parsers or built_in_official_source_parsers()
    missing: list[str] = []
    for source in registry.sources:
        try:
            resolve_official_source_parser(source, parsers=values)
        except ValueError as exc:
            missing.append(f"{source.source_key}: {exc}")
    for qualification in qualifications:
        try:
            source = registry.source(qualification.source_key)
            parser = values[qualification.parser_id]
            if parser.parser_version != qualification.parser_version:
                raise ValueError("qualification parser version is unavailable")
            supported = {
                OfficialSourceFormat.from_value(item)
                for item in parser.supported_formats
            }
            if qualification.source_format not in supported:
                raise ValueError("qualification format is unsupported")
            if source.source_key != qualification.source_key:
                raise ValueError("qualification source is unavailable")
        except (KeyError, ValueError) as exc:
            missing.append(
                f"qualification {qualification.qualification_id}: {exc}"
            )
    required = required_official_adapter_packs(registry)
    qualified = {item.pack for item in qualifications}
    missing_qualifications = tuple(
        item for item in required if item not in qualified
    )
    available_set: set[OfficialAdapterPack] = set()
    for parser in values.values():
        candidate_pack = getattr(parser, "pack", None)
        if isinstance(candidate_pack, OfficialAdapterPack):
            available_set.add(candidate_pack)
    available = tuple(sorted(available_set, key=lambda item: item.value))
    return OfficialAdapterAuditV1(
        registry_id=registry.registry_id,
        required_packs=required,
        available_packs=available,
        missing_parser_bindings=tuple(sorted(missing)),
        missing_qualifications=missing_qualifications,
    )


def qualify_official_adapter_fixture(
    snapshot: OfficialRawSnapshotV1,
    source: OfficialSourceEntryV1,
    pack: OfficialAdapterPack,
    *,
    fixture_uri: str,
    max_events: int = 100_000,
    parser: OfficialSourceParserV1 | None = None,
    parsers: Mapping[str, OfficialSourceParserV1] | None = None,
) -> OfficialAdapterQualificationV1:
    """Parse a pinned first-party fixture and retain its exact expectations."""
    if source.source_key != snapshot.request.source_key:
        raise ValueError("official source differs from snapshot request")
    selected = parser or resolve_official_source_parser(source, parsers=parsers)
    records = parse_official_snapshot(snapshot, selected, max_events=max_events)
    digest = hashlib.sha256(
        canonical_contract_json(list(records)).encode("utf-8")
    ).hexdigest()
    return OfficialAdapterQualificationV1(
        pack=pack,
        source_key=source.source_key,
        source_format=snapshot.request.source_format,
        parser_id=snapshot.request.parser_id,
        parser_version=snapshot.request.parser_version,
        fixture_uri=fixture_uri,
        content_sha256=snapshot.content_sha256,
        normalized_sha256=digest,
        expected_record_count=len(records),
    )


def verify_official_adapter_fixture(
    snapshot: OfficialRawSnapshotV1,
    source: OfficialSourceEntryV1,
    qualification: OfficialAdapterQualificationV1,
    *,
    parser: OfficialSourceParserV1 | None = None,
    parsers: Mapping[str, OfficialSourceParserV1] | None = None,
) -> tuple[Mapping[str, JSONValue], ...]:
    """Replay a qualification and fail when raw or normalized bytes drift."""
    if (
        source.source_key != snapshot.request.source_key
        or qualification.source_key != source.source_key
        or qualification.source_format != snapshot.request.source_format
        or qualification.parser_id != snapshot.request.parser_id
        or qualification.parser_version != snapshot.request.parser_version
    ):
        raise ValueError("official qualification binding differs")
    if qualification.content_sha256 != snapshot.content_sha256:
        raise ValueError("official qualification raw fixture drifted")
    selected = parser or resolve_official_source_parser(source, parsers=parsers)
    records: tuple[Mapping[str, JSONValue], ...] = parse_official_snapshot(
        snapshot,
        selected,
        max_events=qualification.expected_record_count,
    )
    digest = hashlib.sha256(
        canonical_contract_json(list(records)).encode("utf-8")
    ).hexdigest()
    if (
        len(records) != qualification.expected_record_count
        or digest != qualification.normalized_sha256
    ):
        raise ValueError("official qualification normalized fixture drifted")
    return records


def _parse_json(
    parser: _BaseOfficialParser,
    snapshot: OfficialRawSnapshotV1,
    *,
    max_events: int,
) -> tuple[Mapping[str, JSONValue], ...]:
    payload = _load_json(parser, snapshot)
    record_lists: list[tuple[str, Sequence[Any]]] = []
    _find_json_record_lists(
        parser,
        snapshot,
        payload,
        pointer="$",
        depth=0,
        result=record_lists,
    )

    def values() -> (
        Iterable[tuple[OfficialAdapterRecordKind, str, Mapping[str, JSONValue]]]
    ):
        if record_lists:
            for pointer, items in record_lists:
                for index, item in enumerate(items):
                    yield (
                        OfficialAdapterRecordKind.JSON_RECORD,
                        f"{pointer}/{index}",
                        _json_record_fields(item),
                    )
            return
        yield (
            OfficialAdapterRecordKind.JSON_RECORD,
            "$",
            _json_record_fields(payload),
        )

    return parser.records(snapshot, values(), max_events=max_events)


def _find_json_record_lists(
    parser: _BaseOfficialParser,
    snapshot: OfficialRawSnapshotV1,
    value: JSONValue,
    *,
    pointer: str,
    depth: int,
    result: list[tuple[str, Sequence[Any]]],
) -> None:
    if depth > MAX_OFFICIAL_XML_DEPTH:
        raise parser.error(
            snapshot,
            OfficialParserFailureCode.RESOURCE_LIMIT,
            "JSON nesting exceeds depth bound",
            pointer,
        )
    if isinstance(value, list):
        if value and all(isinstance(item, Mapping) for item in value):
            result.append((pointer, value))
        for index, item in enumerate(value):
            _find_json_record_lists(
                parser,
                snapshot,
                item,
                pointer=f"{pointer}/{index}",
                depth=depth + 1,
                result=result,
            )
    elif isinstance(value, dict):
        for key, item in value.items():
            escaped = str(key).replace("~", "~0").replace("/", "~1")
            _find_json_record_lists(
                parser,
                snapshot,
                item,
                pointer=f"{pointer}/{escaped}",
                depth=depth + 1,
                result=result,
            )


def _json_record_fields(value: JSONValue) -> dict[str, JSONValue]:
    if isinstance(value, dict):
        return dict(value)
    return {"value": value}


def _parse_delimited(
    parser: _BaseOfficialParser,
    snapshot: OfficialRawSnapshotV1,
    *,
    max_events: int,
) -> tuple[Mapping[str, JSONValue], ...]:
    try:
        text = snapshot.content.decode("utf-8-sig")
    except UnicodeDecodeError as exc:
        raise parser.error(
            snapshot,
            OfficialParserFailureCode.MALFORMED_DOCUMENT,
            "delimited response is not UTF-8",
        ) from exc
    delimiter = (
        "\t"
        if snapshot.request.source_format is OfficialSourceFormat.TSV
        else ","
    )
    try:
        rows = list(
            csv.reader(io.StringIO(text), delimiter=delimiter, strict=True)
        )
    except csv.Error as exc:
        raise parser.error(
            snapshot,
            OfficialParserFailureCode.MALFORMED_DOCUMENT,
            f"invalid delimited response: {exc}",
        ) from exc
    if len(rows) < 2:
        raise parser.error(
            snapshot,
            OfficialParserFailureCode.SCHEMA_DRIFT,
            "delimited response requires a header and data row",
        )
    headers = _headers(parser, snapshot, rows[0], "$[0]")
    header_version = _header_version(headers)

    def values() -> (
        Iterable[tuple[OfficialAdapterRecordKind, str, Mapping[str, JSONValue]]]
    ):
        for index, row in enumerate(rows[1:], start=1):
            if not row or all(not cell.strip() for cell in row):
                continue
            if len(row) != len(headers):
                raise parser.error(
                    snapshot,
                    OfficialParserFailureCode.SCHEMA_DRIFT,
                    "delimited row width differs from header",
                    f"$[{index}]",
                )
            yield (
                OfficialAdapterRecordKind.DELIMITED_ROW,
                f"$[{index}]",
                {
                    "header_version": header_version,
                    "values": dict(zip(headers, row)),
                },
            )

    return parser.records(snapshot, values(), max_events=max_events)


def _parse_json_stat(
    parser: _BaseOfficialParser,
    snapshot: OfficialRawSnapshotV1,
    *,
    max_events: int,
) -> tuple[Mapping[str, JSONValue], ...]:
    payload = _load_json(parser, snapshot)
    if not isinstance(payload, dict):
        raise parser.error(
            snapshot,
            OfficialParserFailureCode.SCHEMA_DRIFT,
            "JSON-stat root must be an object",
        )
    datasets: list[tuple[str, Mapping[str, JSONValue]]] = []
    if payload.get("class") == "dataset":
        datasets.append(("$", payload))
    else:
        for key, value in payload.items():
            if isinstance(value, dict) and value.get("class") == "dataset":
                datasets.append((f"$/{_pointer(key)}", value))
    if not datasets:
        raise parser.error(
            snapshot,
            OfficialParserFailureCode.SCHEMA_DRIFT,
            "JSON-stat document contains no dataset",
        )

    def values() -> (
        Iterable[tuple[OfficialAdapterRecordKind, str, Mapping[str, JSONValue]]]
    ):
        for pointer, dataset in datasets:
            yield from _json_stat_dataset_records(
                parser, snapshot, pointer, dataset
            )

    return parser.records(snapshot, values(), max_events=max_events)


def _json_stat_dataset_records(
    parser: _BaseOfficialParser,
    snapshot: OfficialRawSnapshotV1,
    pointer: str,
    dataset: Mapping[str, JSONValue],
) -> Iterable[tuple[OfficialAdapterRecordKind, str, Mapping[str, JSONValue]]]:
    ids = dataset.get("id")
    sizes = dataset.get("size")
    dimensions = dataset.get("dimension")
    raw_values = dataset.get("value")
    if (
        not isinstance(ids, list)
        or not all(isinstance(item, str) and item for item in ids)
        or not isinstance(sizes, list)
        or not all(
            isinstance(item, int) and not isinstance(item, bool)
            for item in sizes
        )
        or not isinstance(dimensions, dict)
        or not isinstance(raw_values, (list, dict))
        or len(ids) != len(sizes)
    ):
        raise parser.error(
            snapshot,
            OfficialParserFailureCode.SCHEMA_DRIFT,
            "JSON-stat id, size, dimension, or value shape is invalid",
            pointer,
        )
    dimension_ids = [item for item in ids if isinstance(item, str)]
    dimension_sizes = [
        item
        for item in sizes
        if isinstance(item, int) and not isinstance(item, bool)
    ]
    if any(item <= 0 for item in dimension_sizes):
        raise parser.error(
            snapshot,
            OfficialParserFailureCode.SCHEMA_DRIFT,
            "JSON-stat dimensions must be non-empty",
            pointer,
        )
    total = math.prod(dimension_sizes)
    if total > MAX_OFFICIAL_EVENTS:
        raise parser.error(
            snapshot,
            OfficialParserFailureCode.RESOURCE_LIMIT,
            "JSON-stat dense shape exceeds global record bound",
            pointer,
        )
    categories = [
        _json_stat_categories(
            parser,
            snapshot,
            dimensions,
            name=name,
            size=size,
            pointer=pointer,
        )
        for name, size in zip(dimension_ids, dimension_sizes)
    ]
    dense_values: list[JSONValue] = [None] * total
    if isinstance(raw_values, list):
        if len(raw_values) != total:
            raise parser.error(
                snapshot,
                OfficialParserFailureCode.SCHEMA_DRIFT,
                "JSON-stat dense value count differs from size product",
                pointer,
            )
        dense_values = raw_values
    else:
        for raw_index, value in raw_values.items():
            try:
                index = int(raw_index)
            except ValueError as exc:
                raise parser.error(
                    snapshot,
                    OfficialParserFailureCode.SCHEMA_DRIFT,
                    "JSON-stat sparse value index is invalid",
                    pointer,
                ) from exc
            if not 0 <= index < total:
                raise parser.error(
                    snapshot,
                    OfficialParserFailureCode.SCHEMA_DRIFT,
                    "JSON-stat sparse value index is out of range",
                    pointer,
                )
            dense_values[index] = value
    statuses = dataset.get("status")
    for flat_index, value in enumerate(dense_values):
        offsets = _unravel_index(flat_index, dimension_sizes)
        coordinates: dict[str, JSONValue] = {
            name: categories[dimension_index][offset]
            for dimension_index, (name, offset) in enumerate(
                zip(dimension_ids, offsets)
            )
        }
        status: JSONValue = None
        if isinstance(statuses, list) and flat_index < len(statuses):
            status = statuses[flat_index]
        elif isinstance(statuses, dict):
            status = statuses.get(str(flat_index))
        yield (
            OfficialAdapterRecordKind.JSON_STAT_OBSERVATION,
            f"{pointer}/value/{flat_index}",
            {
                "dataset_label": dataset.get("label"),
                "coordinates": coordinates,
                "value": value,
                "status": status,
            },
        )


def _json_stat_categories(
    parser: _BaseOfficialParser,
    snapshot: OfficialRawSnapshotV1,
    dimensions: Mapping[str, JSONValue],
    *,
    name: str,
    size: int,
    pointer: str,
) -> list[str]:
    dimension = dimensions.get(name)
    if not isinstance(dimension, dict):
        raise parser.error(
            snapshot,
            OfficialParserFailureCode.SCHEMA_DRIFT,
            f"JSON-stat dimension {name} is absent",
            pointer,
        )
    category = dimension.get("category")
    index = category.get("index") if isinstance(category, dict) else None
    result: list[str]
    if isinstance(index, list) and all(isinstance(item, str) for item in index):
        result = cast(list[str], index)
    elif isinstance(index, dict):
        positions: list[tuple[int, str]] = []
        for key, position in index.items():
            if not isinstance(position, int) or isinstance(position, bool):
                raise parser.error(
                    snapshot,
                    OfficialParserFailureCode.SCHEMA_DRIFT,
                    "JSON-stat category position is invalid",
                    pointer,
                )
            positions.append((position, key))
        positions.sort()
        result = [key for _, key in positions]
    else:
        raise parser.error(
            snapshot,
            OfficialParserFailureCode.SCHEMA_DRIFT,
            "JSON-stat category index is invalid",
            pointer,
        )
    if len(result) != size or len(set(result)) != size:
        raise parser.error(
            snapshot,
            OfficialParserFailureCode.SCHEMA_DRIFT,
            "JSON-stat category count differs from declared size",
            pointer,
        )
    return result


def _unravel_index(flat_index: int, sizes: Sequence[int]) -> list[int]:
    offsets = [0] * len(sizes)
    remaining = flat_index
    for index in range(len(sizes) - 1, -1, -1):
        offsets[index] = remaining % sizes[index]
        remaining //= sizes[index]
    return offsets


def _parse_sdmx(
    parser: _BaseOfficialParser,
    snapshot: OfficialRawSnapshotV1,
    *,
    max_events: int,
) -> tuple[Mapping[str, JSONValue], ...]:
    stripped = snapshot.content.lstrip(b"\xef\xbb\xbf\x00\t\r\n ")
    if stripped.startswith((b"{", b"[")):
        values = _sdmx_json_records(parser, snapshot)
    elif stripped.startswith(b"<"):
        values = _sdmx_xml_records(parser, snapshot)
    else:
        values = _sdmx_csv_records(parser, snapshot)
    return parser.records(snapshot, values, max_events=max_events)


def _sdmx_csv_records(
    parser: _BaseOfficialParser,
    snapshot: OfficialRawSnapshotV1,
) -> Iterable[tuple[OfficialAdapterRecordKind, str, Mapping[str, JSONValue]]]:
    try:
        text = snapshot.content.decode("utf-8-sig")
        rows = list(csv.reader(io.StringIO(text), strict=True))
    except (UnicodeDecodeError, csv.Error) as exc:
        raise parser.error(
            snapshot,
            OfficialParserFailureCode.MALFORMED_DOCUMENT,
            "SDMX-CSV response is malformed",
        ) from exc
    if len(rows) < 2:
        raise parser.error(
            snapshot,
            OfficialParserFailureCode.SCHEMA_DRIFT,
            "SDMX-CSV requires a header and data row",
        )
    headers = _headers(parser, snapshot, rows[0], "$[0]")
    upper = {item.upper() for item in headers}
    if not ({"TIME_PERIOD", "OBS_VALUE"} <= upper):
        raise parser.error(
            snapshot,
            OfficialParserFailureCode.SCHEMA_DRIFT,
            "SDMX-CSV observation columns are absent",
            "$[0]",
        )
    header_version = _header_version(headers)
    for index, row in enumerate(rows[1:], start=1):
        if not row or all(not item.strip() for item in row):
            continue
        if len(row) != len(headers):
            raise parser.error(
                snapshot,
                OfficialParserFailureCode.SCHEMA_DRIFT,
                "SDMX-CSV row width differs from header",
                f"$[{index}]",
            )
        yield (
            OfficialAdapterRecordKind.SDMX_OBSERVATION,
            f"$[{index}]",
            {
                "header_version": header_version,
                "dimensions": dict(zip(headers, row)),
            },
        )


def _sdmx_json_records(
    parser: _BaseOfficialParser,
    snapshot: OfficialRawSnapshotV1,
) -> Iterable[tuple[OfficialAdapterRecordKind, str, Mapping[str, JSONValue]]]:
    payload = _load_json(parser, snapshot)
    if not isinstance(payload, dict):
        raise parser.error(
            snapshot,
            OfficialParserFailureCode.SCHEMA_DRIFT,
            "SDMX-JSON root must be an object",
        )
    data_sets = payload.get("dataSets")
    structure = payload.get("structure")
    if isinstance(data_sets, list) and isinstance(structure, dict):
        yield from _sdmx_json_data_records(
            parser, snapshot, data_sets, structure
        )
        return
    structures = payload.get("data")
    if isinstance(structures, dict):
        flows = structures.get("dataflows")
        codes = structures.get("codelists")
        found = False
        for kind, values in (("dataflow", flows), ("codelist", codes)):
            if not isinstance(values, list):
                continue
            for index, value in enumerate(values):
                if not isinstance(value, dict):
                    raise parser.error(
                        snapshot,
                        OfficialParserFailureCode.SCHEMA_DRIFT,
                        f"SDMX {kind} entry is not an object",
                        f"$/data/{kind}s/{index}",
                    )
                found = True
                yield (
                    (
                        OfficialAdapterRecordKind.SDMX_DATAFLOW
                        if kind == "dataflow"
                        else OfficialAdapterRecordKind.SDMX_CODE
                    ),
                    f"$/data/{kind}s/{index}",
                    dict(value),
                )
        if found:
            return
    raise parser.error(
        snapshot,
        OfficialParserFailureCode.SCHEMA_DRIFT,
        "unrecognized SDMX-JSON message shape",
    )


def _sdmx_json_data_records(
    parser: _BaseOfficialParser,
    snapshot: OfficialRawSnapshotV1,
    data_sets: Sequence[JSONValue],
    structure: Mapping[str, JSONValue],
) -> Iterable[tuple[OfficialAdapterRecordKind, str, Mapping[str, JSONValue]]]:
    dimensions = structure.get("dimensions")
    if not isinstance(dimensions, dict):
        raise parser.error(
            snapshot,
            OfficialParserFailureCode.SCHEMA_DRIFT,
            "SDMX-JSON dimensions are absent",
        )
    series_dimensions = _sdmx_json_dimensions(
        parser,
        snapshot,
        dimensions.get("series"),
        "$/structure/dimensions/series",
    )
    observation_dimensions = _sdmx_json_dimensions(
        parser,
        snapshot,
        dimensions.get("observation"),
        "$/structure/dimensions/observation",
    )
    for dataset_index, dataset in enumerate(data_sets):
        if not isinstance(dataset, dict):
            raise parser.error(
                snapshot,
                OfficialParserFailureCode.SCHEMA_DRIFT,
                "SDMX-JSON dataset is not an object",
                f"$/dataSets/{dataset_index}",
            )
        series = dataset.get("series")
        if not isinstance(series, dict):
            raise parser.error(
                snapshot,
                OfficialParserFailureCode.SCHEMA_DRIFT,
                "SDMX-JSON series mapping is absent",
                f"$/dataSets/{dataset_index}",
            )
        for series_key, series_value in series.items():
            if not isinstance(series_value, dict):
                raise parser.error(
                    snapshot,
                    OfficialParserFailureCode.SCHEMA_DRIFT,
                    "SDMX-JSON series value is not an object",
                    f"$/dataSets/{dataset_index}/series/{series_key}",
                )
            series_coordinates = _sdmx_json_key(
                parser,
                snapshot,
                series_key,
                series_dimensions,
                f"$/dataSets/{dataset_index}/series/{series_key}",
            )
            observations = series_value.get("observations")
            if not isinstance(observations, dict):
                raise parser.error(
                    snapshot,
                    OfficialParserFailureCode.SCHEMA_DRIFT,
                    "SDMX-JSON observations mapping is absent",
                    f"$/dataSets/{dataset_index}/series/{series_key}",
                )
            for observation_key, observation in observations.items():
                locator = (
                    f"$/dataSets/{dataset_index}/series/{series_key}/"
                    f"observations/{observation_key}"
                )
                observation_coordinates = _sdmx_json_key(
                    parser,
                    snapshot,
                    observation_key,
                    observation_dimensions,
                    locator,
                )
                if not isinstance(observation, list) or not observation:
                    raise parser.error(
                        snapshot,
                        OfficialParserFailureCode.SCHEMA_DRIFT,
                        "SDMX-JSON observation array is invalid",
                        locator,
                    )
                yield (
                    OfficialAdapterRecordKind.SDMX_OBSERVATION,
                    locator,
                    {
                        "dimensions": {
                            **series_coordinates,
                            **observation_coordinates,
                        },
                        "value": observation[0],
                        "attributes": observation[1:],
                    },
                )


def _sdmx_json_dimensions(
    parser: _BaseOfficialParser,
    snapshot: OfficialRawSnapshotV1,
    value: JSONValue,
    locator: str,
) -> list[tuple[str, list[str]]]:
    if not isinstance(value, list):
        raise parser.error(
            snapshot,
            OfficialParserFailureCode.SCHEMA_DRIFT,
            "SDMX-JSON dimension list is absent",
            locator,
        )
    result: list[tuple[str, list[str]]] = []
    for index, dimension in enumerate(value):
        if not isinstance(dimension, dict) or not isinstance(
            dimension.get("id"), str
        ):
            raise parser.error(
                snapshot,
                OfficialParserFailureCode.SCHEMA_DRIFT,
                "SDMX-JSON dimension definition is invalid",
                f"{locator}/{index}",
            )
        raw_values = dimension.get("values")
        if not isinstance(raw_values, list):
            raise parser.error(
                snapshot,
                OfficialParserFailureCode.SCHEMA_DRIFT,
                "SDMX-JSON dimension values are invalid",
                f"{locator}/{index}",
            )
        names: list[str] = []
        for value_index, item in enumerate(raw_values):
            if not isinstance(item, dict) or not isinstance(
                item.get("id"), str
            ):
                raise parser.error(
                    snapshot,
                    OfficialParserFailureCode.SCHEMA_DRIFT,
                    "SDMX-JSON dimension value is invalid",
                    f"{locator}/{index}/values/{value_index}",
                )
            names.append(cast(str, item["id"]))
        result.append((cast(str, dimension["id"]), names))
    return result


def _sdmx_json_key(
    parser: _BaseOfficialParser,
    snapshot: OfficialRawSnapshotV1,
    key: str,
    dimensions: Sequence[tuple[str, list[str]]],
    locator: str,
) -> dict[str, JSONValue]:
    try:
        indexes = [int(item) for item in key.split(":")]
    except ValueError as exc:
        raise parser.error(
            snapshot,
            OfficialParserFailureCode.SCHEMA_DRIFT,
            "SDMX-JSON dimension key is invalid",
            locator,
        ) from exc
    if len(indexes) != len(dimensions):
        raise parser.error(
            snapshot,
            OfficialParserFailureCode.SCHEMA_DRIFT,
            "SDMX-JSON dimension key width differs",
            locator,
        )
    result: dict[str, JSONValue] = {}
    for index, (name, values) in zip(indexes, dimensions):
        if not 0 <= index < len(values):
            raise parser.error(
                snapshot,
                OfficialParserFailureCode.SCHEMA_DRIFT,
                "SDMX-JSON dimension index is out of range",
                locator,
            )
        result[name] = values[index]
    return result


def _sdmx_xml_records(
    parser: _BaseOfficialParser,
    snapshot: OfficialRawSnapshotV1,
) -> Iterable[tuple[OfficialAdapterRecordKind, str, Mapping[str, JSONValue]]]:
    root = _load_xml(parser, snapshot)
    observation_count = 0
    series_items = (
        item for item in root.iter() if _local_name(item.tag) == "Series"
    )
    for series_count, series in enumerate(series_items, start=1):
        series_values = _sdmx_element_values(series, "SeriesKey")
        series_values.update(
            {str(key): str(value) for key, value in series.attrib.items()}
        )
        for observation in (
            item for item in series if _local_name(item.tag) == "Obs"
        ):
            observation_count += 1
            values = dict(series_values)
            values.update(
                {
                    str(key): str(value)
                    for key, value in observation.attrib.items()
                }
            )
            values.update(_sdmx_element_values(observation, "ObsDimension"))
            values.update(_sdmx_element_values(observation, "ObsValue"))
            values.update(_sdmx_element_values(observation, "Attributes"))
            json_dimensions = _json_text_mapping(values)
            yield (
                OfficialAdapterRecordKind.SDMX_OBSERVATION,
                f"/Series[{series_count}]/Obs[{observation_count}]",
                {"dimensions": json_dimensions},
            )
    if observation_count:
        return
    flow_count = 0
    for element in root.iter():
        if _local_name(element.tag) != "Dataflow":
            continue
        flow_count += 1
        yield (
            OfficialAdapterRecordKind.SDMX_DATAFLOW,
            f"/Dataflow[{flow_count}]",
            _xml_named_fields(element),
        )
    code_count = 0
    codelist_id: str | None = None
    for element in root.iter():
        name = _local_name(element.tag)
        if name == "Codelist":
            codelist_id = element.attrib.get("id")
        elif name == "Code":
            code_count += 1
            resolved_codelist_id = codelist_id or ""
            yield (
                OfficialAdapterRecordKind.SDMX_CODE,
                f"/Codelist[{resolved_codelist_id}]/Code[{code_count}]",
                {
                    "codelist_id": codelist_id,
                    **_xml_named_fields(element),
                },
            )
    if not flow_count and not code_count:
        raise parser.error(
            snapshot,
            OfficialParserFailureCode.SCHEMA_DRIFT,
            "unrecognized SDMX-XML message shape",
        )


def _sdmx_element_values(
    element: ElementTree.Element, child_name: str
) -> dict[str, str]:
    result: dict[str, str] = {}
    for child in element.iter():
        if _local_name(child.tag) != child_name:
            continue
        if "id" in child.attrib and "value" in child.attrib:
            result[child.attrib["id"]] = child.attrib["value"]
        elif "value" in child.attrib:
            result[child_name] = child.attrib["value"]
        for nested in child:
            if _local_name(nested.tag) == "Value":
                key = nested.attrib.get("id")
                value = nested.attrib.get("value")
                if key is not None and value is not None:
                    result[key] = value
    return result


def _xml_named_fields(element: ElementTree.Element) -> dict[str, JSONValue]:
    result: dict[str, JSONValue] = {
        str(key): str(value) for key, value in element.attrib.items()
    }
    names = [
        _normalized_text("".join(child.itertext()))
        for child in element
        if _local_name(child.tag) in {"Name", "Description"}
    ]
    names = [item for item in names if item]
    if names:
        result["names"] = _json_text_list(names)
    return result


def _parse_spreadsheet(
    parser: _BaseOfficialParser,
    snapshot: OfficialRawSnapshotV1,
    *,
    max_events: int,
) -> tuple[Mapping[str, JSONValue], ...]:
    layouts = tuple(
        cast(
            Sequence[OfficialSpreadsheetLayoutV1],
            getattr(parser, "layouts", ()),
        )
    )
    if snapshot.request.source_format is OfficialSourceFormat.XLSX:
        sheets = _xlsx_rows(parser, snapshot)
    elif snapshot.request.source_format is OfficialSourceFormat.XLS:
        sheets = _xls_rows(parser, snapshot)
    else:
        raise parser.error(
            snapshot,
            OfficialParserFailureCode.UNSUPPORTED_REPRESENTATION,
            "spreadsheet parser requires XLS or XLSX",
        )
    if len(sheets) > MAX_OFFICIAL_SHEETS:
        raise parser.error(
            snapshot,
            OfficialParserFailureCode.RESOURCE_LIMIT,
            "workbook sheet count exceeds bound",
        )
    by_sheet = {item.sheet_name: item for item in layouts}
    missing = set(by_sheet) - {name for name, _ in sheets}
    if missing:
        missing_text = ", ".join(sorted(missing))
        raise parser.error(
            snapshot,
            OfficialParserFailureCode.SCHEMA_DRIFT,
            f"configured workbook sheets are absent: {missing_text}",
        )

    def values() -> (
        Iterable[tuple[OfficialAdapterRecordKind, str, Mapping[str, JSONValue]]]
    ):
        for sheet_name, rows in sheets:
            layout = by_sheet.get(sheet_name)
            if layouts and layout is None:
                continue
            header_index = (
                layout.header_row - 1
                if layout is not None
                else _detect_header_row(rows)
            )
            if not 0 <= header_index < len(rows):
                raise parser.error(
                    snapshot,
                    OfficialParserFailureCode.SCHEMA_DRIFT,
                    "spreadsheet header row is absent",
                    f"{sheet_name}!R{header_index + 1}",
                )
            headers = _headers(
                parser,
                snapshot,
                rows[header_index],
                f"{sheet_name}!R{header_index + 1}",
            )
            header_version = _header_version(headers)
            if (
                layout is not None
                and layout.expected_header_version is not None
                and layout.expected_header_version != header_version
            ):
                raise parser.error(
                    snapshot,
                    OfficialParserFailureCode.SCHEMA_DRIFT,
                    "spreadsheet header version differs from expectation",
                    f"{sheet_name}!R{header_index + 1}",
                )
            for row_index, raw_row in enumerate(
                rows[header_index + 1 :], start=header_index + 2
            ):
                row = list(raw_row[: len(headers)])
                if len(row) < len(headers):
                    row.extend([None] * (len(headers) - len(row)))
                if all(item is None or item == "" for item in row):
                    continue
                yield (
                    OfficialAdapterRecordKind.SPREADSHEET_ROW,
                    f"{sheet_name}!R{row_index}",
                    {
                        "sheet_name": sheet_name,
                        "header_row": header_index + 1,
                        "header_version": header_version,
                        "values": dict(zip(headers, row)),
                    },
                )

    configuration: dict[str, JSONValue] = {
        "layouts": [
            {**item.to_dict(), "layout_id": item.layout_id} for item in layouts
        ],
        "header_selection": "configured" if layouts else "first-tabular-row-v1",
    }
    return parser.records(
        snapshot,
        values(),
        max_events=max_events,
        configuration=configuration,
    )


def _xlsx_rows(
    parser: _BaseOfficialParser, snapshot: OfficialRawSnapshotV1
) -> list[tuple[str, list[list[JSONValue]]]]:
    try:
        workbook = load_workbook(
            io.BytesIO(snapshot.content),
            read_only=True,
            data_only=False,
            keep_links=False,
        )
    except Exception as exc:  # pylint: disable=broad-exception-caught
        raise parser.error(
            snapshot,
            OfficialParserFailureCode.MALFORMED_DOCUMENT,
            f"XLSX workbook could not be opened: {exc}",
        ) from exc
    try:
        result: list[tuple[str, list[list[JSONValue]]]] = []
        for sheet in workbook.worksheets:
            rows: list[list[JSONValue]] = []
            for cells in sheet.iter_rows(values_only=True):
                if len(cells) > MAX_OFFICIAL_COLUMNS:
                    raise parser.error(
                        snapshot,
                        OfficialParserFailureCode.RESOURCE_LIMIT,
                        "spreadsheet column count exceeds bound",
                        f"{sheet.title}!R{len(rows) + 1}",
                    )
                rows.append([_spreadsheet_value(item) for item in cells])
            result.append((str(sheet.title), rows))
        return result
    finally:
        workbook.close()


def _xls_rows(
    parser: _BaseOfficialParser, snapshot: OfficialRawSnapshotV1
) -> list[tuple[str, list[list[JSONValue]]]]:
    try:
        workbook = xlrd.open_workbook(file_contents=snapshot.content)
    except Exception as exc:  # pylint: disable=broad-exception-caught
        raise parser.error(
            snapshot,
            OfficialParserFailureCode.MALFORMED_DOCUMENT,
            f"XLS workbook could not be opened: {exc}",
        ) from exc
    result: list[tuple[str, list[list[JSONValue]]]] = []
    for sheet in workbook.sheets():
        if sheet.ncols > MAX_OFFICIAL_COLUMNS:
            raise parser.error(
                snapshot,
                OfficialParserFailureCode.RESOURCE_LIMIT,
                "spreadsheet column count exceeds bound",
                sheet.name,
            )
        rows = [
            [_spreadsheet_value(item) for item in sheet.row_values(row)]
            for row in range(sheet.nrows)
        ]
        result.append((str(sheet.name), rows))
    return result


def _spreadsheet_value(value: Any) -> JSONValue:
    if value is None or isinstance(value, (str, bool, int)):
        result: JSONValue = value
    elif isinstance(value, float):
        if not math.isfinite(value):
            raise ValueError("spreadsheet cell is non-finite")
        result = value
    elif isinstance(value, (datetime, date, time)):
        result = value.isoformat()
    else:
        result = str(value)
    if (
        len(canonical_contract_json(result).encode("utf-8"))
        > MAX_OFFICIAL_CELL_BYTES
    ):
        raise ValueError("spreadsheet cell exceeds byte bound")
    return result


def _detect_header_row(rows: Sequence[Sequence[JSONValue]]) -> int:
    for index, row in enumerate(rows[:100]):
        names = [str(item).strip() if item is not None else "" for item in row]
        non_empty = [item for item in names if item]
        if len(non_empty) >= 2 and len(set(non_empty)) == len(non_empty):
            return index
    return 0


class _OfficialHtmlCollector(HTMLParser):
    """Collect semantic HTML components without executing content."""

    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.title_parts: list[str] = []
        self.document_parts: list[str] = []
        self.meta: dict[str, str] = {}
        self.links: list[dict[str, str]] = []
        self.tables: list[list[list[str]]] = []
        self.times: list[dict[str, str]] = []
        self.headings: list[str] = []
        self._title_depth = 0
        self._heading_depth = 0
        self._link: dict[str, str] | None = None
        self._link_parts: list[str] = []
        self._table: list[list[str]] | None = None
        self._row: list[str] | None = None
        self._cell_parts: list[str] | None = None
        self._heading_parts: list[str] = []
        self._time: dict[str, str] | None = None
        self._time_parts: list[str] = []

    def handle_starttag(
        self, tag: str, attrs: list[tuple[str, str | None]]
    ) -> None:
        values = {str(key).lower(): str(value or "") for key, value in attrs}
        tag = tag.lower()
        if tag == "title":
            self._title_depth += 1
        elif tag in {"h1", "h2", "h3", "h4", "h5", "h6"}:
            self._heading_depth += 1
            self._heading_parts = []
        elif tag == "meta":
            key = values.get("name") or values.get("property")
            if key and "content" in values:
                self.meta[key] = values["content"]
        elif tag == "a" and "href" in values:
            self._link = values
            self._link_parts = []
        elif tag == "table":
            self._table = []
        elif tag == "tr" and self._table is not None:
            self._row = []
        elif tag in {"td", "th"} and self._row is not None:
            self._cell_parts = []
        elif tag == "time":
            self._time = values
            self._time_parts = []

    def handle_endtag(self, tag: str) -> None:
        tag = tag.lower()
        if tag == "title" and self._title_depth:
            self._title_depth -= 1
        elif (
            tag in {"h1", "h2", "h3", "h4", "h5", "h6"} and self._heading_depth
        ):
            value = _normalized_text(" ".join(self._heading_parts))
            if value:
                self.headings.append(value)
            self._heading_depth -= 1
            self._heading_parts = []
        elif tag == "a" and self._link is not None:
            self._link["text"] = _normalized_text(" ".join(self._link_parts))
            self.links.append(self._link)
            self._link = None
            self._link_parts = []
        elif tag in {"td", "th"} and self._cell_parts is not None:
            if self._row is not None:
                self._row.append(_normalized_text(" ".join(self._cell_parts)))
            self._cell_parts = None
        elif tag == "tr" and self._row is not None:
            if self._table is not None and any(self._row):
                self._table.append(self._row)
            self._row = None
        elif tag == "table" and self._table is not None:
            if self._table:
                self.tables.append(self._table)
            self._table = None
        elif tag == "time" and self._time is not None:
            self._time["text"] = _normalized_text(" ".join(self._time_parts))
            self.times.append(self._time)
            self._time = None
            self._time_parts = []

    def handle_data(self, data: str) -> None:
        value = _normalized_text(data)
        if not value:
            return
        self.document_parts.append(value)
        if self._title_depth:
            self.title_parts.append(value)
        if self._heading_depth:
            self._heading_parts.append(value)
        if self._link is not None:
            self._link_parts.append(value)
        if self._cell_parts is not None:
            self._cell_parts.append(value)
        if self._time is not None:
            self._time_parts.append(value)


def _parse_text(
    parser: _BaseOfficialParser,
    snapshot: OfficialRawSnapshotV1,
    *,
    max_events: int,
) -> tuple[Mapping[str, JSONValue], ...]:
    """Retain one bounded UTF-8 official plain-text document as a record."""
    try:
        text = snapshot.content.decode("utf-8-sig")
    except UnicodeDecodeError as exc:
        raise parser.error(
            snapshot,
            OfficialParserFailureCode.MALFORMED_DOCUMENT,
            "text response is not UTF-8",
        ) from exc
    normalized = text.replace("\r\n", "\n").replace("\r", "\n").strip()
    if not normalized:
        raise parser.error(
            snapshot,
            OfficialParserFailureCode.SCHEMA_DRIFT,
            "text response is empty",
        )
    return parser.records(
        snapshot,
        (
            (
                OfficialAdapterRecordKind.TEXT_DOCUMENT,
                "/text",
                {
                    "line_count": normalized.count("\n") + 1,
                    "text": normalized,
                },
            ),
        ),
        max_events=max_events,
    )


def _parse_html(
    parser: _BaseOfficialParser,
    snapshot: OfficialRawSnapshotV1,
    *,
    max_events: int,
    strict_tables: bool = True,
) -> tuple[Mapping[str, JSONValue], ...]:
    try:
        text = snapshot.content.decode("utf-8-sig")
        collector = _OfficialHtmlCollector()
        collector.feed(text)
        collector.close()
    except (UnicodeDecodeError, ValueError) as exc:
        raise parser.error(
            snapshot,
            OfficialParserFailureCode.MALFORMED_DOCUMENT,
            f"HTML response is malformed: {exc}",
        ) from exc

    def values() -> (
        Iterable[tuple[OfficialAdapterRecordKind, str, Mapping[str, JSONValue]]]
    ):
        document_text = _normalized_text(" ".join(collector.document_parts))
        if document_text:
            headings = _json_text_list(collector.headings)
            times: list[JSONValue] = [
                _json_text_mapping(item) for item in collector.times
            ]
            metadata = _json_text_mapping(collector.meta)
            yield (
                OfficialAdapterRecordKind.HTML_DOCUMENT,
                "/html",
                {
                    "title": _normalized_text(" ".join(collector.title_parts)),
                    "headings": headings,
                    "times": times,
                    "metadata": metadata,
                    "text": document_text,
                },
            )
        for table_index, table in enumerate(collector.tables, start=1):
            if len(table) < 2:
                continue
            headers = _headers(
                parser,
                snapshot,
                table[0],
                f"/html/table[{table_index}]/tr[1]",
            )
            header_version = _header_version(headers)
            for row_index, row in enumerate(table[1:], start=2):
                if len(row) != len(headers):
                    if not strict_tables:
                        continue
                    raise parser.error(
                        snapshot,
                        OfficialParserFailureCode.SCHEMA_DRIFT,
                        "HTML table row width differs from header",
                        f"/html/table[{table_index}]/tr[{row_index}]",
                    )
                yield (
                    OfficialAdapterRecordKind.HTML_TABLE_ROW,
                    f"/html/table[{table_index}]/tr[{row_index}]",
                    {
                        "table_number": table_index,
                        "header_version": header_version,
                        "values": dict(zip(headers, row)),
                    },
                )
        base = snapshot.resolved_uri
        for link_index, link in enumerate(collector.links, start=1):
            href = link.get("href", "").strip()
            if not href:
                continue
            yield (
                OfficialAdapterRecordKind.HTML_LINK,
                f"/html/a[{link_index}]",
                {
                    "uri": urljoin(base, href),
                    "text": link.get("text", ""),
                    "media_type": link.get("type"),
                    "download": link.get("download"),
                },
            )

    return parser.records(snapshot, values(), max_events=max_events)


def _parse_xml_feed(
    parser: _BaseOfficialParser,
    snapshot: OfficialRawSnapshotV1,
    *,
    max_events: int,
) -> tuple[Mapping[str, JSONValue], ...]:
    root = _load_xml(parser, snapshot)
    expected = (
        "rss"
        if snapshot.request.source_format is OfficialSourceFormat.RSS
        else "feed"
    )
    if _local_name(root.tag).lower() != expected:
        raise parser.error(
            snapshot,
            OfficialParserFailureCode.SCHEMA_DRIFT,
            f"expected {expected} feed root",
        )
    item_name = "item" if expected == "rss" else "entry"
    items = [item for item in root.iter() if _local_name(item.tag) == item_name]

    def values() -> (
        Iterable[tuple[OfficialAdapterRecordKind, str, Mapping[str, JSONValue]]]
    ):
        for index, item in enumerate(items, start=1):
            fields: dict[str, JSONValue] = {}
            for child in item:
                name = _local_name(child.tag)
                if name == "link" and child.attrib.get("href"):
                    value = child.attrib["href"]
                else:
                    value = _normalized_text("".join(child.itertext()))
                if not value:
                    continue
                if name in fields:
                    previous = fields[name]
                    fields[name] = (
                        [*previous, value]
                        if isinstance(previous, list)
                        else [previous, value]
                    )
                else:
                    fields[name] = value
            if not fields.get("title"):
                raise parser.error(
                    snapshot,
                    OfficialParserFailureCode.SCHEMA_DRIFT,
                    "release feed item omits title",
                    f"/{expected}/{item_name}[{index}]",
                )
            yield (
                OfficialAdapterRecordKind.RELEASE_FEED_ITEM,
                f"/{expected}/{item_name}[{index}]",
                fields,
            )

    return parser.records(snapshot, values(), max_events=max_events)


def _parse_ics(
    parser: _BaseOfficialParser,
    snapshot: OfficialRawSnapshotV1,
    *,
    max_events: int,
) -> tuple[Mapping[str, JSONValue], ...]:
    try:
        text = snapshot.content.decode("utf-8-sig")
    except UnicodeDecodeError as exc:
        raise parser.error(
            snapshot,
            OfficialParserFailureCode.MALFORMED_DOCUMENT,
            "iCalendar response is not UTF-8",
        ) from exc
    physical = text.replace("\r\n", "\n").replace("\r", "\n").split("\n")
    lines: list[str] = []
    for line in physical:
        if line.startswith((" ", "\t")) and lines:
            lines[-1] += line[1:]
        else:
            lines.append(line)
    if not lines or lines[0].upper() != "BEGIN:VCALENDAR":
        raise parser.error(
            snapshot,
            OfficialParserFailureCode.SCHEMA_DRIFT,
            "iCalendar root is absent",
        )
    events: list[tuple[int, dict[str, JSONValue]]] = []
    current: dict[str, JSONValue] | None = None
    start_line = 0
    for line_number, line in enumerate(lines, start=1):
        upper = line.upper()
        if upper == "BEGIN:VEVENT":
            if current is not None:
                raise parser.error(
                    snapshot,
                    OfficialParserFailureCode.MALFORMED_DOCUMENT,
                    "nested iCalendar event",
                    f"line:{line_number}",
                )
            current = {}
            start_line = line_number
        elif upper == "END:VEVENT":
            if current is None:
                raise parser.error(
                    snapshot,
                    OfficialParserFailureCode.MALFORMED_DOCUMENT,
                    "unmatched iCalendar event terminator",
                    f"line:{line_number}",
                )
            events.append((start_line, current))
            current = None
        elif current is not None:
            if not _ICAL_LINE_RE.match(line):
                raise parser.error(
                    snapshot,
                    OfficialParserFailureCode.MALFORMED_DOCUMENT,
                    "invalid iCalendar property",
                    f"line:{line_number}",
                )
            left, value = line.split(":", 1)
            parts = left.split(";")
            name = parts[0].lower()
            parameters: dict[str, JSONValue] = {
                key.lower(): item
                for part in parts[1:]
                if "=" in part
                for key, item in (part.split("=", 1),)
            }
            decoded = (
                value.replace("\\n", "\n")
                .replace("\\,", ",")
                .replace("\\;", ";")
                .replace("\\\\", "\\")
            )
            parsed_field: JSONValue = (
                {"value": decoded, "parameters": parameters}
                if parameters
                else decoded
            )
            previous = current.get(name)
            current[name] = (
                [*previous, parsed_field]
                if isinstance(previous, list)
                else (
                    [previous, parsed_field]
                    if previous is not None
                    else parsed_field
                )
            )
    if current is not None:
        raise parser.error(
            snapshot,
            OfficialParserFailureCode.MALFORMED_DOCUMENT,
            "unterminated iCalendar event",
            f"line:{start_line}",
        )

    def values() -> (
        Iterable[tuple[OfficialAdapterRecordKind, str, Mapping[str, JSONValue]]]
    ):
        for line_number, fields in events:
            if "uid" not in fields or "dtstart" not in fields:
                raise parser.error(
                    snapshot,
                    OfficialParserFailureCode.SCHEMA_DRIFT,
                    "iCalendar event omits UID or DTSTART",
                    f"line:{line_number}",
                )
            yield (
                OfficialAdapterRecordKind.RELEASE_FEED_ITEM,
                f"line:{line_number}",
                fields,
            )

    return parser.records(snapshot, values(), max_events=max_events)


def _parse_pdf(
    parser: _BaseOfficialParser,
    snapshot: OfficialRawSnapshotV1,
    *,
    max_events: int,
) -> tuple[Mapping[str, JSONValue], ...]:
    try:
        reader = PdfReader(io.BytesIO(snapshot.content), strict=True)
    except Exception as exc:  # pylint: disable=broad-exception-caught
        raise parser.error(
            snapshot,
            OfficialParserFailureCode.MALFORMED_DOCUMENT,
            f"PDF could not be opened: {exc}",
        ) from exc
    if reader.is_encrypted:
        return parser.records(
            snapshot,
            (
                (
                    OfficialAdapterRecordKind.REACQUISITION_REQUIREMENT,
                    "/pdf",
                    {
                        "reason": "encrypted-pdf",
                        "requirement": (
                            "reacquire an unencrypted official artifact or "
                            "provide "
                            "an authorized decryption step before parsing"
                        ),
                    },
                ),
            ),
            max_events=max_events,
        )

    def values() -> (
        Iterable[tuple[OfficialAdapterRecordKind, str, Mapping[str, JSONValue]]]
    ):
        extracted = 0
        for page_number, page in enumerate(reader.pages, start=1):
            try:
                text = _normalized_text(page.extract_text() or "")
            except Exception as exc:  # pylint: disable=broad-exception-caught
                raise parser.error(
                    snapshot,
                    OfficialParserFailureCode.EXTRACTION_FAILED,
                    f"PDF page text extraction failed: {exc}",
                    f"/pdf/page[{page_number}]",
                ) from exc
            if not text:
                continue
            extracted += 1
            table_rows = [
                re.split(r"\s{2,}", line.strip())
                for line in (page.extract_text() or "").splitlines()
                if len(re.split(r"\s{2,}", line.strip())) >= 2
            ]
            yield (
                OfficialAdapterRecordKind.PDF_PAGE,
                f"/pdf/page[{page_number}]",
                {
                    "page_number": page_number,
                    "text": text,
                    "table_rows": _json_text_rows(table_rows),
                },
            )
        if not extracted:
            yield (
                OfficialAdapterRecordKind.REACQUISITION_REQUIREMENT,
                "/pdf",
                {
                    "reason": "image-only-or-empty-pdf",
                    "requirement": (
                        "reacquire a text-bearing official PDF or retain an "
                        "authorized OCR artifact bound to this content hash"
                    ),
                },
            )

    return parser.records(snapshot, values(), max_events=max_events)


def _parse_archive(
    parser: _BaseOfficialParser,
    snapshot: OfficialRawSnapshotV1,
    *,
    max_events: int,
) -> tuple[Mapping[str, JSONValue], ...]:
    source_format = snapshot.request.source_format
    if source_format is OfficialSourceFormat.HTML:
        links = _html_archive_links(parser, snapshot)
        return parser.records(snapshot, links, max_events=max_events)
    if source_format is OfficialSourceFormat.JSON:
        links = _json_archive_links(parser, snapshot)
        return parser.records(snapshot, links, max_events=max_events)
    try:
        archive = zipfile.ZipFile(io.BytesIO(snapshot.content))
    except (zipfile.BadZipFile, OSError) as exc:
        raise parser.error(
            snapshot,
            OfficialParserFailureCode.UNSUPPORTED_REPRESENTATION,
            "static archive is not an enumerable ZIP container",
        ) from exc
    with archive:
        infos = archive.infolist()
        if len(infos) > MAX_OFFICIAL_ARCHIVE_MEMBERS:
            raise parser.error(
                snapshot,
                OfficialParserFailureCode.RESOURCE_LIMIT,
                "archive member count exceeds bound",
            )

        def values() -> (
            Iterable[
                tuple[OfficialAdapterRecordKind, str, Mapping[str, JSONValue]]
            ]
        ):
            for index, info in enumerate(infos, start=1):
                yield (
                    OfficialAdapterRecordKind.ARCHIVE_MEMBER,
                    f"/zip/member[{index}]",
                    {
                        "name": info.filename,
                        "size_bytes": info.file_size,
                        "compressed_size_bytes": info.compress_size,
                        "crc32": f"{info.CRC:08x}",
                        "is_directory": info.is_dir(),
                    },
                )

        return parser.records(snapshot, values(), max_events=max_events)


def _html_archive_links(
    parser: _BaseOfficialParser,
    snapshot: OfficialRawSnapshotV1,
) -> Iterable[tuple[OfficialAdapterRecordKind, str, Mapping[str, JSONValue]]]:
    try:
        collector = _OfficialHtmlCollector()
        collector.feed(snapshot.content.decode("utf-8-sig"))
        collector.close()
    except (UnicodeDecodeError, ValueError) as exc:
        raise parser.error(
            snapshot,
            OfficialParserFailureCode.MALFORMED_DOCUMENT,
            "archive index HTML is malformed",
        ) from exc
    found = 0
    for index, link in enumerate(collector.links, start=1):
        href = link.get("href", "").strip()
        uri = urljoin(snapshot.resolved_uri, href)
        if not uri.lower().split("?", 1)[0].endswith(_ARCHIVE_SUFFIXES):
            continue
        found += 1
        yield (
            OfficialAdapterRecordKind.ARCHIVE_LINK,
            f"/html/a[{index}]",
            {"uri": uri, "text": link.get("text", "")},
        )
    if not found:
        raise parser.error(
            snapshot,
            OfficialParserFailureCode.SCHEMA_DRIFT,
            "archive index contains no recognized downloads",
        )


def _json_archive_links(
    parser: _BaseOfficialParser,
    snapshot: OfficialRawSnapshotV1,
) -> Iterable[tuple[OfficialAdapterRecordKind, str, Mapping[str, JSONValue]]]:
    payload = _load_json(parser, snapshot)
    found = 0
    for pointer, value in _walk_json(parser, snapshot, payload):
        if not isinstance(value, str):
            continue
        uri = urljoin(snapshot.resolved_uri, value)
        if not uri.lower().split("?", 1)[0].endswith(_ARCHIVE_SUFFIXES):
            continue
        found += 1
        yield (
            OfficialAdapterRecordKind.ARCHIVE_LINK,
            pointer,
            {"uri": uri},
        )
    if not found:
        raise parser.error(
            snapshot,
            OfficialParserFailureCode.SCHEMA_DRIFT,
            "archive JSON contains no recognized downloads",
        )


def _parse_data_catalog(
    parser: _BaseOfficialParser,
    snapshot: OfficialRawSnapshotV1,
    *,
    max_events: int,
) -> tuple[Mapping[str, JSONValue], ...]:
    stripped = snapshot.content.lstrip(b"\xef\xbb\xbf\x00\t\r\n ")
    if stripped.startswith((b"{", b"[")):
        values = _json_catalog_records(parser, snapshot)
    elif stripped.startswith(b"<"):
        values = _xml_catalog_records(parser, snapshot)
    else:
        raise parser.error(
            snapshot,
            OfficialParserFailureCode.UNSUPPORTED_REPRESENTATION,
            "data catalogue is neither JSON nor XML",
        )
    return parser.records(snapshot, values, max_events=max_events)


def _json_catalog_records(
    parser: _BaseOfficialParser,
    snapshot: OfficialRawSnapshotV1,
) -> Iterable[tuple[OfficialAdapterRecordKind, str, Mapping[str, JSONValue]]]:
    payload = _load_json(parser, snapshot)
    candidates: list[tuple[str, Mapping[str, JSONValue]]] = []
    for pointer, value in _walk_json(parser, snapshot, payload):
        if not isinstance(value, dict):
            continue
        lowered = {str(key).lower() for key in value}
        has_identity = bool(
            lowered & {"id", "identifier", "name", "title", "@id"}
        )
        has_catalog_shape = bool(
            lowered
            & {
                "distribution",
                "distributions",
                "resources",
                "downloadurl",
                "landingpage",
            }
        )
        if has_identity and has_catalog_shape:
            candidates.append((pointer, value))
    if not candidates:
        raise parser.error(
            snapshot,
            OfficialParserFailureCode.SCHEMA_DRIFT,
            "JSON catalogue contains no dataset-shaped records",
        )
    for pointer, candidate in candidates:
        yield (
            OfficialAdapterRecordKind.DATASET_CATALOG_ITEM,
            pointer,
            dict(candidate),
        )


def _xml_catalog_records(
    parser: _BaseOfficialParser,
    snapshot: OfficialRawSnapshotV1,
) -> Iterable[tuple[OfficialAdapterRecordKind, str, Mapping[str, JSONValue]]]:
    root = _load_xml(parser, snapshot)
    candidates = [
        item
        for item in root.iter()
        if _local_name(item.tag).lower() in {"dataset", "entry", "item"}
    ]
    if not candidates:
        raise parser.error(
            snapshot,
            OfficialParserFailureCode.SCHEMA_DRIFT,
            "XML catalogue contains no dataset-shaped records",
        )
    for index, item in enumerate(candidates, start=1):
        fields: dict[str, JSONValue] = {
            str(key): str(value) for key, value in item.attrib.items()
        }
        for child in item:
            name = _local_name(child.tag)
            value = child.attrib.get("href") or _normalized_text(
                "".join(child.itertext())
            )
            if value:
                fields[name] = value
        yield (
            OfficialAdapterRecordKind.DATASET_CATALOG_ITEM,
            f"/{_local_name(root.tag)}/{_local_name(item.tag)}[{index}]",
            fields,
        )


def _load_json(
    parser: _BaseOfficialParser, snapshot: OfficialRawSnapshotV1
) -> JSONValue:
    try:
        value = json.loads(snapshot.content.decode("utf-8-sig"))
        canonical_contract_json(value)
    except (
        UnicodeDecodeError,
        json.JSONDecodeError,
        TypeError,
        ValueError,
    ) as exc:
        raise parser.error(
            snapshot,
            OfficialParserFailureCode.MALFORMED_DOCUMENT,
            f"JSON response is malformed: {exc}",
        ) from exc
    return cast(JSONValue, value)


def _load_xml(
    parser: _BaseOfficialParser, snapshot: OfficialRawSnapshotV1
) -> ElementTree.Element:
    try:
        root = cast(ElementTree.Element, safe_xml_fromstring(snapshot.content))
    except (ElementTree.ParseError, DefusedXmlException) as exc:
        raise parser.error(
            snapshot,
            OfficialParserFailureCode.MALFORMED_DOCUMENT,
            f"XML response is malformed: {exc}",
        ) from exc
    _validate_xml_depth(parser, snapshot, root, depth=0)
    return root


def _validate_xml_depth(
    parser: _BaseOfficialParser,
    snapshot: OfficialRawSnapshotV1,
    element: ElementTree.Element,
    *,
    depth: int,
) -> None:
    if depth > MAX_OFFICIAL_XML_DEPTH:
        raise parser.error(
            snapshot,
            OfficialParserFailureCode.RESOURCE_LIMIT,
            "XML nesting exceeds depth bound",
        )
    for child in element:
        _validate_xml_depth(parser, snapshot, child, depth=depth + 1)


def _walk_json(
    parser: _BaseOfficialParser,
    snapshot: OfficialRawSnapshotV1,
    value: JSONValue,
) -> Iterable[tuple[str, JSONValue]]:
    stack: list[tuple[str, JSONValue, int]] = [("$", value, 0)]
    visited = 0
    while stack:
        pointer, item, depth = stack.pop()
        visited += 1
        if depth > MAX_OFFICIAL_XML_DEPTH:
            raise parser.error(
                snapshot,
                OfficialParserFailureCode.RESOURCE_LIMIT,
                "JSON nesting exceeds depth bound",
                pointer,
            )
        if visited > MAX_OFFICIAL_EVENTS * 16:
            raise parser.error(
                snapshot,
                OfficialParserFailureCode.RESOURCE_LIMIT,
                "JSON node count exceeds traversal bound",
                pointer,
            )
        yield pointer, item
        if isinstance(item, dict):
            children = [
                (f"{pointer}/{_pointer(key)}", child, depth + 1)
                for key, child in item.items()
            ]
            stack.extend(reversed(children))
        elif isinstance(item, list):
            children = [
                (f"{pointer}/{index}", child, depth + 1)
                for index, child in enumerate(item)
            ]
            stack.extend(reversed(children))


def _headers(
    parser: _BaseOfficialParser,
    snapshot: OfficialRawSnapshotV1,
    values: Sequence[JSONValue],
    locator: str,
) -> list[str]:
    if not values or len(values) > MAX_OFFICIAL_COLUMNS:
        raise parser.error(
            snapshot,
            OfficialParserFailureCode.RESOURCE_LIMIT,
            "header column count is invalid",
            locator,
        )
    result = [str(item).strip() if item is not None else "" for item in values]
    if any(not item for item in result) or len(set(result)) != len(result):
        raise parser.error(
            snapshot,
            OfficialParserFailureCode.SCHEMA_DRIFT,
            "headers must be non-empty and unique",
            locator,
        )
    return result


def _header_version(headers: Sequence[str]) -> str:
    return hashlib.sha256(
        canonical_contract_json(list(headers)).encode("utf-8")
    ).hexdigest()


def _json_text_list(values: Iterable[str]) -> list[JSONValue]:
    return list(values)


def _json_text_mapping(values: Mapping[str, str]) -> dict[str, JSONValue]:
    result: dict[str, JSONValue] = {}
    for key, value in values.items():
        result[key] = value
    return result


def _json_text_rows(values: Iterable[Iterable[str]]) -> list[JSONValue]:
    result: list[JSONValue] = []
    for row in values:
        result.append(_json_text_list(row))
    return result


def _json_mapping(
    value: Mapping[str, JSONValue], label: str
) -> dict[str, JSONValue]:
    if not isinstance(value, Mapping):
        raise TypeError(f"{label} must be a mapping")
    result = {str(key): item for key, item in value.items()}
    if len(result) != len(value) or any(not key for key in result):
        raise ValueError(f"{label} keys are invalid")
    try:
        canonical_contract_json(result)
    except (TypeError, ValueError) as exc:
        raise ValueError(f"{label} is not canonical JSON") from exc
    return result


def _mapping(value: Any, label: str) -> Mapping[str, JSONValue]:
    if not isinstance(value, Mapping):
        raise TypeError(f"{label} must be a mapping")
    return cast(Mapping[str, JSONValue], value)


def _sequence(value: Any, label: str) -> Sequence[Any]:
    if not isinstance(value, Sequence) or isinstance(value, (str, bytes)):
        raise TypeError(f"{label} must be a sequence")
    return value


def _strict_int(value: Any, label: str) -> int:
    if not isinstance(value, int) or isinstance(value, bool):
        raise TypeError(f"{label} must be an integer")
    return value


def _optional_string(value: Any) -> str | None:
    if value is None:
        return None
    if not isinstance(value, str):
        raise TypeError("optional value must be a string")
    return value


def _key(value: str, label: str) -> str:
    result = str(value).strip().lower()
    if not _KEY_RE.fullmatch(result):
        raise ValueError(f"{label} is invalid")
    return result


def _digest(value: str, label: str) -> str:
    result = str(value).strip().lower()
    if not _SHA256_RE.fullmatch(result):
        raise ValueError(f"{label} must be a SHA-256 digest")
    return result


def _stable_id(prefix: str, payload: Mapping[str, JSONValue]) -> str:
    digest = hashlib.sha256(
        canonical_contract_json(payload).encode("utf-8")
    ).hexdigest()
    return f"{prefix}:sha256:{digest}"


def _pointer(value: str) -> str:
    return str(value).replace("~", "~0").replace("/", "~1")


def _local_name(value: str) -> str:
    return value.rsplit("}", 1)[-1].rsplit(":", 1)[-1]


def _normalized_text(value: str) -> str:
    return " ".join(str(value).split())


__all__ = [
    "MAX_OFFICIAL_ARCHIVE_MEMBERS",
    "MAX_OFFICIAL_CELL_BYTES",
    "MAX_OFFICIAL_COLUMNS",
    "MAX_OFFICIAL_RECORD_BYTES",
    "MAX_OFFICIAL_SHEETS",
    "OFFICIAL_ADAPTER_AUDIT_SCHEMA_VERSION",
    "OFFICIAL_ADAPTER_QUALIFICATION_SCHEMA_VERSION",
    "OFFICIAL_ADAPTER_RECORD_SCHEMA_VERSION",
    "OFFICIAL_SPREADSHEET_LAYOUT_SCHEMA_VERSION",
    "OfficialAdapterAuditV1",
    "OfficialAdapterPack",
    "OfficialAdapterQualificationV1",
    "OfficialAdapterRecordKind",
    "OfficialAdapterRecordV1",
    "OfficialArchiveParserV1",
    "OfficialCensusFt900ParserV1",
    "OfficialCensusM3ParserV1",
    "OfficialCensusMartsParserV1",
    "OfficialCensusMtisParserV1",
    "OfficialCensusNrcParserV1",
    "OfficialCensusNrsParserV1",
    "OfficialCsvParserV1",
    "OfficialDataCatalogParserV1",
    "OfficialFederalReserveFomcParserV1",
    "OfficialFederalReserveH6ParserV1",
    "OfficialHtmlParserV1",
    "OfficialJsonParserV1",
    "OfficialJsonStatParserV1",
    "OfficialParserError",
    "OfficialParserFailureCode",
    "OfficialPdfParserV1",
    "OfficialPhiladelphiaFedMbosParserV1",
    "OfficialReleaseFeedParserV1",
    "OfficialSdmx21ParserV1",
    "OfficialSdmx30ParserV1",
    "OfficialSpreadsheetLayoutV1",
    "OfficialSpreadsheetParserV1",
    "OfficialTreasuryMtsParserV1",
    "audit_official_adapter_coverage",
    "built_in_official_source_parsers",
    "parse_with_built_in_official_adapter",
    "qualify_official_adapter_fixture",
    "required_official_adapter_packs",
    "resolve_official_source_parser",
    "verify_official_adapter_fixture",
]
