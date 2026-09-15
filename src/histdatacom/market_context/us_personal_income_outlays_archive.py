"""Occurrence-specific U.S. Personal Income and Outlays archive evidence.

The BEA archive spans multiple HTML layouts, two PDF-only publications, a
text time supplement, workbook-backed 2025 tables, split shutdown releases,
and a revision-only data update.  This module preserves those source facts
while reconstructing monthly personal-income, current-dollar PCE, headline
PCE-price, and core-PCE-price values and their previous-as-known lineage.
"""

from __future__ import annotations

import base64
import calendar
import hashlib
import html
import json
import math
import re
from collections.abc import Mapping, Sequence
from dataclasses import dataclass
from datetime import date, datetime, time
from enum import Enum
from io import BytesIO
from itertools import pairwise
from pathlib import Path
from types import MappingProxyType
from typing import Any, cast
from urllib.parse import urljoin, urlparse
from zoneinfo import ZoneInfo

from openpyxl import load_workbook
from pypdf import PdfReader

from histdatacom.market_context.contracts import canonical_contract_json
from histdatacom.market_context.economic_calendar import (
    EconomicReleaseStage,
    EconomicTimePrecision,
)
from histdatacom.market_context.official_sources import (
    OfficialRawSnapshotV1,
    OfficialRequestMethod,
    OfficialSourceFormat,
    OfficialSourceRegistryV1,
    OfficialSourceRequestV1,
    load_packaged_official_source_registry,
)
from histdatacom.market_context.us_backfill import (
    US_BACKFILL_START_DATE,
    UnitedStatesBackfillProfileV1,
    UnitedStatesProgramCoverageV1,
)
from histdatacom.runtime_contracts import JSONValue

BEA_PIO_INDEX_PAGE_SCHEMA_VERSION = "histdatacom.bea-pio-index-page.v1"
BEA_PIO_INDEX_ENTRY_SCHEMA_VERSION = "histdatacom.bea-pio-index-entry.v1"
BEA_PIO_INDEX_SCHEMA_VERSION = "histdatacom.bea-pio-index.v1"
BEA_PIO_MEASURE_VALUE_SCHEMA_VERSION = "histdatacom.bea-pio-measure-value.v1"
BEA_PIO_RELEASE_OBSERVATION_SCHEMA_VERSION = (
    "histdatacom.bea-pio-release-observation.v1"
)
BEA_PIO_ARCHIVE_MEASURE_SCHEMA_VERSION = (
    "histdatacom.bea-pio-archive-measure.v1"
)
BEA_PIO_ARCHIVE_PUBLICATION_SCHEMA_VERSION = (
    "histdatacom.bea-pio-archive-publication.v1"
)
BEA_PIO_ARCHIVE_MANIFEST_SCHEMA_VERSION = (
    "histdatacom.bea-pio-archive-manifest.v1"
)
BEA_PIO_INDEX_ENVELOPE_SCHEMA_VERSION = "histdatacom.bea-pio-index-envelope.v1"

BEA_PIO_SOURCE_KEY = "us.bea.personal-income-outlays"
BEA_PIO_PROGRAM_KEY = BEA_PIO_SOURCE_KEY
BEA_PIO_INDEX_URI = (
    "https://www.bea.gov/news/archive?"
    "created_1=All&field_related_product_target_id=476&title="
)
BEA_PIO_INDEX_PAGE_COUNT = 19
BEA_PIO_SUPPLEMENTAL_URI = (
    "https://www.bea.gov/news/2026/personal-income-and-outlays-july-2026"
)
BEA_PIO_SUPPLEMENTAL_DATE = "2026-08-26"
BEA_PIO_SUPPLEMENTAL_TITLE = "Personal Income and Outlays, July 2026"
BEA_PIO_PREDECESSOR_URI = (
    "https://www.bea.gov/sites/default/files/2018-03/bea99-38.pdf"
)
BEA_PIO_PREDECESSOR_DATE = "1999-12-23"
BEA_PIO_PREDECESSOR_REFERENCE_PERIOD = "1999-11"

BEA_PIO_PERSONAL_INCOME = "current-dollar-personal-income"
BEA_PIO_CURRENT_DOLLAR_PCE = "current-dollar-personal-consumption-expenditures"
BEA_PIO_HEADLINE_PRICE = "pce-price-index"
BEA_PIO_CORE_PRICE = "pce-price-index-excluding-food-and-energy"
BEA_PIO_MEASURE_KEYS = (
    BEA_PIO_PERSONAL_INCOME,
    BEA_PIO_CURRENT_DOLLAR_PCE,
    BEA_PIO_HEADLINE_PRICE,
    BEA_PIO_CORE_PRICE,
)

BEA_PIO_ARCHIVE_DATE_CORRECTIONS = MappingProxyType(
    {
        "https://www.bea.gov/news/2019/personal-income-and-outlays-december-2018-personal-income-january-2019": "2019-03-01"
    }
)

BEA_PIO_SUPPORT_ARTIFACTS = MappingProxyType(
    {
        "https://www.bea.gov/news/2000/personal-income-february-2000": (
            "https://www.bea.gov/sites/default/files/newsreleases/national/pi/2000/pdf/pi0200.pdf",
            OfficialSourceFormat.PDF,
        ),
        "https://www.bea.gov/news/2006/personal-income-and-outlays-october-2006": (
            "https://www.bea.gov/sites/default/files/newsreleases/national/pi/2006/pdf/pi1006.pdf",
            OfficialSourceFormat.PDF,
        ),
        "https://www.bea.gov/news/2008/personal-income-and-outlays-november-2008": (
            "https://www.bea.gov/sites/default/files/2020-03/pi1108.txt",
            OfficialSourceFormat.TEXT,
        ),
        "https://www.bea.gov/news/2025/personal-income-and-outlays-december-2024": (
            "https://www.bea.gov/sites/default/files/2025-01/pi1224.xlsx",
            OfficialSourceFormat.XLSX,
        ),
        "https://www.bea.gov/news/2025/personal-income-and-outlays-january-2025": (
            "https://www.bea.gov/sites/default/files/2025-02/pi0125_0.xlsx",
            OfficialSourceFormat.XLSX,
        ),
        "https://www.bea.gov/news/2025/personal-income-and-outlays-february-2025": (
            "https://www.bea.gov/sites/default/files/2025-03/pi0225.xlsx",
            OfficialSourceFormat.XLSX,
        ),
        "https://www.bea.gov/news/2025/personal-income-and-outlays-march-2025": (
            "https://www.bea.gov/sites/default/files/2025-04/pi0325.xlsx",
            OfficialSourceFormat.XLSX,
        ),
        "https://www.bea.gov/news/2025/personal-income-and-outlays-april-2025": (
            "https://www.bea.gov/sites/default/files/2025-05/pi0425.xlsx",
            OfficialSourceFormat.XLSX,
        ),
        "https://www.bea.gov/news/2025/personal-income-and-outlays-may-2025": (
            "https://www.bea.gov/sites/default/files/2025-06/pi0525.xlsx",
            OfficialSourceFormat.XLSX,
        ),
        "https://www.bea.gov/news/2025/personal-income-and-outlays-june-2025": (
            "https://www.bea.gov/sites/default/files/2025-07/pi0625.xlsx",
            OfficialSourceFormat.XLSX,
        ),
        "https://www.bea.gov/news/2025/personal-income-and-outlays-july-2025": (
            "https://www.bea.gov/sites/default/files/2025-08/pi0725.xlsx",
            OfficialSourceFormat.XLSX,
        ),
        "https://www.bea.gov/news/2025/personal-income-and-outlays-august-2025": (
            "https://www.bea.gov/sites/default/files/2025-09/pi0825.xlsx",
            OfficialSourceFormat.XLSX,
        ),
        "https://www.bea.gov/news/2025/personal-income-and-outlays-september-2025": (
            "https://www.bea.gov/sites/default/files/2025-12/pi0925.xlsx",
            OfficialSourceFormat.XLSX,
        ),
        "https://www.bea.gov/news/2025/personal-income-and-outlays-data-update-september-2025": (
            "https://www.bea.gov/sites/default/files/2025-12/pi0925-updated.xlsx",
            OfficialSourceFormat.XLSX,
        ),
    }
)

MAX_BEA_PIO_INDEX_PAGES = 64
MAX_BEA_PIO_RELEASES = 512
MAX_BEA_PIO_MEASURE_OCCURRENCES = 4096
MAX_BEA_PIO_INDEX_PAGE_BYTES = 4 * 1024 * 1024
MAX_BEA_PIO_RELEASE_BYTES = 32 * 1024 * 1024

_SHA256_RE = re.compile(r"^[0-9a-f]{64}$")
_YEAR_MONTH_RE = re.compile(r"^(?P<year>\d{4})-(?P<month>0[1-9]|1[0-2])$")
_MONTH_YEAR_RE = re.compile(
    r"\b(?P<month>January|February|March|April|May|June|July|August|"
    r"September|October|November|December)\s+(?P<year>\d{4})\b",
    re.IGNORECASE,
)
_LONG_DATE_RE = re.compile(
    r"^(?P<month>[A-Za-z]+)\s+(?P<day>\d{1,2})\s*,\s*(?P<year>\d{4})$"
)
_RELEASE_TIME_RE = re.compile(
    r"(?P<hour>8|10):(?P<minute>30|00)\s*"
    r"(?:A\.?M\.?|a\.?m\.?)\s*,?\s*(?P<zone>EST|EDT)\s*,\s*"
    r"(?:(?:Monday|Tuesday|Wednesday|Thursday|Friday)\s*,?\s*)?"
    r"(?P<date>[A-Za-z]+\s+\d{1,2}\s*,\s*\d{4})",
    re.IGNORECASE,
)
_ARCHIVE_ROW_RE = re.compile(r"(?is)<tr\b[^>]*>(?P<body>.*?)</tr>")
_ARCHIVE_TITLE_RE = re.compile(
    r"(?is)<td\b[^>]*class=[\"'][^\"']*views-field-title[^\"']*"
    r"[\"'][^>]*>.*?<a\b[^>]*href=[\"'](?P<uri>[^\"']+)"
    r"[\"'][^>]*>(?P<title>.*?)</a>"
)
_ARCHIVE_TIME_RE = re.compile(
    r"(?is)<td\b[^>]*class=[\"'][^\"']*views-field-created[^\"']*"
    r"[\"'][^>]*>.*?<time\b[^>]*datetime=[\"'](?P<stamp>[^\"']+)"
    r"[\"'][^>]*>"
)
_TOKEN_RE = re.compile(r"\.\.\.|-?(?:\d+\.\d+|\.\d+|\d+)")
_MONTH_NUMBERS = {
    name.lower(): number
    for number, name in enumerate(calendar.month_name)
    if name
}
_MONTH_NUMBERS.update(
    {name[:3].lower(): number for name, number in tuple(_MONTH_NUMBERS.items())}
)
_MONTH_NUMBERS["sept"] = 9
_UPDATE_TITLE_MARKERS = (
    "annual update",
    "historical revisions",
    "revised estimates",
    "data update",
    "revised january",
)
_PUBLICATION_KINDS = {
    "regular",
    "shutdown-split",
    "shutdown-catch-up",
    "shutdown-combined",
    "revision-only-update",
}


class BeaPioValueBasis(str, Enum):
    """Whether a monthly percentage is printed or level-derived."""

    SOURCE_PUBLISHED = "source-published"
    DERIVED_FROM_INDEX_LEVELS = "derived-from-index-levels"

    @classmethod
    def from_value(cls, value: str | BeaPioValueBasis) -> BeaPioValueBasis:
        if isinstance(value, cls):
            return value
        try:
            return cls(str(value))
        except ValueError as exc:
            raise ValueError("unsupported BEA PIO value basis") from exc


def _required_text(value: object, name: str) -> str:
    result = str(value).strip()
    if not result or len(result) > 8192:
        raise ValueError(f"{name} is invalid")
    return result


def _optional_text(value: object | None, name: str) -> str | None:
    if value is None:
        return None
    return _required_text(value, name)


def _iso_date(value: object, name: str) -> str:
    result = _required_text(value, name)
    try:
        parsed = date.fromisoformat(result)
    except ValueError as exc:
        raise ValueError(f"{name} must be an ISO date") from exc
    if parsed.isoformat() != result:
        raise ValueError(f"{name} must be a canonical ISO date")
    return result


def _year_month(value: object, name: str) -> str:
    result = _required_text(value, name)
    if _YEAR_MONTH_RE.fullmatch(result) is None:
        raise ValueError(f"{name} must be a canonical year-month")
    return result


def _previous_month(value: str) -> str:
    match = _YEAR_MONTH_RE.fullmatch(_year_month(value, "year_month"))
    assert match is not None
    year = int(match.group("year"))
    month = int(match.group("month"))
    if month == 1:
        return f"{year - 1:04d}-12"
    return f"{year:04d}-{month - 1:02d}"


def _next_month(value: str) -> str:
    match = _YEAR_MONTH_RE.fullmatch(_year_month(value, "year_month"))
    assert match is not None
    year = int(match.group("year"))
    month = int(match.group("month"))
    if month == 12:
        return f"{year + 1:04d}-01"
    return f"{year:04d}-{month + 1:02d}"


def _periods_ending(value: str, count: int) -> tuple[str, ...]:
    if not 1 <= count <= 32:
        raise ValueError("BEA PIO period count is outside its bound")
    result = [_year_month(value, "ending_period")]
    while len(result) < count:
        result.append(_previous_month(result[-1]))
    return tuple(reversed(result))


def _https_uri(value: object, name: str) -> str:
    result = _required_text(value, name)
    parsed = urlparse(result)
    if (
        parsed.scheme != "https"
        or parsed.hostname != "www.bea.gov"
        or parsed.username is not None
        or parsed.password is not None
        or parsed.fragment
    ):
        raise ValueError(f"{name} is not an official BEA HTTPS URI")
    return result


def _sha256(value: object, name: str) -> str:
    result = _required_text(value, name)
    if _SHA256_RE.fullmatch(result) is None:
        raise ValueError(f"{name} must be a lowercase SHA-256")
    return result


def _positive_int(value: object, name: str, maximum: int) -> int:
    if (
        not isinstance(value, int)
        or isinstance(value, bool)
        or not 1 <= value <= maximum
    ):
        raise ValueError(f"{name} is outside its bound")
    return value


def _nonnegative_int(value: object, name: str, maximum: int) -> int:
    if (
        not isinstance(value, int)
        or isinstance(value, bool)
        or not 0 <= value <= maximum
    ):
        raise ValueError(f"{name} is outside its bound")
    return value


def _finite(value: object, name: str) -> float:
    if isinstance(value, bool):
        raise TypeError(f"{name} must be numeric")
    try:
        result = float(cast(Any, value))
    except (TypeError, ValueError) as exc:
        raise ValueError(f"{name} must be numeric") from exc
    if not math.isfinite(result):
        raise ValueError(f"{name} must be finite")
    return result


def _numeric_lexical(value: object, name: str) -> tuple[str, float]:
    lexical = _required_text(value, name)
    if re.fullmatch(r"-?(?:0|[1-9]\d*)\.\d+", lexical) is None:
        raise ValueError(f"{name} must be a signed decimal lexical value")
    return lexical, _finite(lexical, name)


def _mapping(value: object, name: str) -> Mapping[str, Any]:
    if not isinstance(value, Mapping):
        raise TypeError(f"{name} must be a mapping")
    return cast(Mapping[str, Any], value)


def _sequence(value: object, name: str) -> Sequence[Any]:
    if not isinstance(value, Sequence) or isinstance(value, (str, bytes)):
        raise TypeError(f"{name} must be a sequence")
    return value


def _stable_id(prefix: str, payload: Mapping[str, JSONValue]) -> str:
    encoded = str(canonical_contract_json(payload)).encode("utf-8")
    return f"{prefix}:sha256:{hashlib.sha256(encoded).hexdigest()}"


def _visible_html(value: str) -> str:
    result = re.sub(r"(?i)<br\s*/?>", " ", value)
    result = re.sub(r"(?is)<script\b.*?</script>", " ", result)
    result = re.sub(r"(?is)<style\b.*?</style>", " ", result)
    result = re.sub(r"(?is)<[^>]+>", " ", result)
    return " ".join(html.unescape(result).replace("\xa0", " ").split())


def _decode_html(content: bytes) -> str:
    for encoding in ("utf-8-sig", "windows-1252"):
        try:
            return content.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise ValueError("BEA PIO HTML is neither UTF-8 nor Windows-1252")


def _pdf_text(content: bytes) -> str:
    try:
        reader = PdfReader(BytesIO(content))
        if not 1 <= len(reader.pages) <= 128:
            raise ValueError("BEA PIO PDF page count is outside its bound")
        result = "\n".join(page.extract_text() or "" for page in reader.pages)
    except Exception as exc:
        raise ValueError("BEA PIO PDF cannot be parsed") from exc
    if not result.strip():
        raise ValueError("BEA PIO PDF has no extractable text")
    return " ".join(result.split())


def _long_date(value: str) -> date:
    match = _LONG_DATE_RE.fullmatch(" ".join(value.split()))
    if match is None or match.group("month").lower() not in _MONTH_NUMBERS:
        raise ValueError("BEA PIO long date format changed")
    try:
        return date(
            int(match.group("year")),
            _MONTH_NUMBERS[match.group("month").lower()],
            int(match.group("day")),
        )
    except ValueError as exc:
        raise ValueError("BEA PIO long date is invalid") from exc


def _archive_page_uri(page_number: int) -> str:
    _nonnegative_int(page_number, "page_number", MAX_BEA_PIO_INDEX_PAGES - 1)
    if page_number == 0:
        return BEA_PIO_INDEX_URI
    return f"{BEA_PIO_INDEX_URI}&page={page_number}"


def _primary_period(uri: str, title: str) -> str:
    special = {
        "https://www.bea.gov/news/2019/personal-income-and-outlays-december-2018-personal-income-january-2019": "2019-01",
        "https://www.bea.gov/news/2019/personal-income-february-2019-personal-outlays-january-2019": "2019-02",
        "https://www.bea.gov/news/2019/personal-income-and-outlays-march-2019": "2019-03",
        "https://www.bea.gov/news/2026/personal-income-and-outlays-october-and-november-2025": "2025-11",
    }
    if uri in special:
        return special[uri]
    matches = tuple(_MONTH_YEAR_RE.finditer(title))
    if not matches:
        raise ValueError("BEA PIO title omits its reference month")
    match = matches[0]
    return (
        f"{int(match.group('year')):04d}-"
        f"{_MONTH_NUMBERS[match.group('month').lower()]:02d}"
    )


def _publication_kind(uri: str) -> str:
    if uri.endswith(
        (
            "personal-income-and-outlays-december-2018-personal-income-january-2019",
            "personal-income-february-2019-personal-outlays-january-2019",
        )
    ):
        return "shutdown-split"
    if uri.endswith("personal-income-and-outlays-march-2019"):
        return "shutdown-catch-up"
    if uri.endswith("personal-income-and-outlays-october-and-november-2025"):
        return "shutdown-combined"
    if uri.endswith("personal-income-and-outlays-data-update-september-2025"):
        return "revision-only-update"
    return "regular"


def _release_stage(uri: str) -> EconomicReleaseStage:
    if _publication_kind(uri) == "revision-only-update":
        return EconomicReleaseStage.REVISION
    return EconomicReleaseStage.INITIAL


@dataclass(frozen=True, slots=True)
class BeaPioReleaseIndexPageV1:
    """One exact page of the official BEA PIO archive inventory."""

    page_number: int
    source_uri: str
    content_sha256: str
    content_length: int
    schema_version: str = BEA_PIO_INDEX_PAGE_SCHEMA_VERSION

    def __post_init__(self) -> None:
        if self.schema_version != BEA_PIO_INDEX_PAGE_SCHEMA_VERSION:
            raise ValueError("unsupported BEA PIO index-page schema")
        page = _nonnegative_int(
            self.page_number, "page_number", MAX_BEA_PIO_INDEX_PAGES - 1
        )
        source = _https_uri(self.source_uri, "source_uri")
        if source != _archive_page_uri(page):
            raise ValueError("BEA PIO index-page URI differs from page number")
        object.__setattr__(self, "page_number", page)
        object.__setattr__(self, "source_uri", source)
        object.__setattr__(
            self,
            "content_sha256",
            _sha256(self.content_sha256, "content_sha256"),
        )
        _positive_int(
            self.content_length,
            "content_length",
            MAX_BEA_PIO_INDEX_PAGE_BYTES,
        )

    def to_dict(self) -> dict[str, JSONValue]:
        return {
            "schema_version": self.schema_version,
            "page_number": self.page_number,
            "source_uri": self.source_uri,
            "content_sha256": self.content_sha256,
            "content_length": self.content_length,
        }

    @classmethod
    def from_dict(cls, data: Mapping[str, Any]) -> BeaPioReleaseIndexPageV1:
        return cls(
            page_number=cast(int, data.get("page_number")),
            source_uri=str(data.get("source_uri", "")),
            content_sha256=str(data.get("content_sha256", "")),
            content_length=cast(int, data.get("content_length")),
            schema_version=str(data.get("schema_version", "")),
        )


@dataclass(frozen=True, slots=True)
class BeaPioReleaseIndexEntryV1:
    """One archive-listed or bounded supplemental PIO publication."""

    primary_reference_period: str
    release_stage: EconomicReleaseStage
    release_date: str
    archive_published_date: str
    archive_timestamp: str
    artifact_uri: str
    title: str
    publication_kind: str
    archive_listed: bool
    entry_id: str = ""
    schema_version: str = BEA_PIO_INDEX_ENTRY_SCHEMA_VERSION

    def __post_init__(self) -> None:
        if self.schema_version != BEA_PIO_INDEX_ENTRY_SCHEMA_VERSION:
            raise ValueError("unsupported BEA PIO index-entry schema")
        reference = _year_month(
            self.primary_reference_period, "primary_reference_period"
        )
        stage = EconomicReleaseStage.from_value(self.release_stage)
        if stage not in {
            EconomicReleaseStage.INITIAL,
            EconomicReleaseStage.REVISION,
        }:
            raise ValueError("BEA PIO release stage is invalid")
        released = _iso_date(self.release_date, "release_date")
        published = _iso_date(
            self.archive_published_date, "archive_published_date"
        )
        stamp = _required_text(self.archive_timestamp, "archive_timestamp")
        try:
            parsed_stamp = datetime.fromisoformat(stamp)
        except ValueError as exc:
            raise ValueError("BEA PIO archive timestamp is invalid") from exc
        if parsed_stamp.date().isoformat() != published:
            raise ValueError("BEA PIO archive timestamp date differs")
        artifact = _https_uri(self.artifact_uri, "artifact_uri")
        title = _required_text(self.title, "title")
        kind = _required_text(self.publication_kind, "publication_kind")
        if kind not in _PUBLICATION_KINDS or kind != _publication_kind(
            artifact
        ):
            raise ValueError("BEA PIO publication kind differs")
        if stage is not _release_stage(artifact):
            raise ValueError("BEA PIO release stage differs from publication")
        if reference != _primary_period(artifact, title):
            raise ValueError("BEA PIO primary period differs from title")
        expected_date = BEA_PIO_ARCHIVE_DATE_CORRECTIONS.get(
            artifact, published
        )
        if released != expected_date:
            raise ValueError(
                "BEA PIO release date differs from archive evidence"
            )
        if not isinstance(self.archive_listed, bool):
            raise TypeError("archive_listed must be boolean")
        if self.archive_listed is (artifact == BEA_PIO_SUPPLEMENTAL_URI):
            raise ValueError("BEA PIO archive-list status differs")
        object.__setattr__(self, "primary_reference_period", reference)
        object.__setattr__(self, "release_stage", stage)
        object.__setattr__(self, "release_date", released)
        object.__setattr__(self, "archive_published_date", published)
        object.__setattr__(self, "archive_timestamp", stamp)
        object.__setattr__(self, "artifact_uri", artifact)
        object.__setattr__(self, "title", title)
        object.__setattr__(self, "publication_kind", kind)
        expected = _stable_id("bea-pio-index-entry", self.identity_payload())
        if self.entry_id and self.entry_id != expected:
            raise ValueError("BEA PIO index-entry identity differs")
        object.__setattr__(self, "entry_id", expected)

    def identity_payload(self) -> dict[str, JSONValue]:
        return {
            "schema_version": self.schema_version,
            "primary_reference_period": self.primary_reference_period,
            "release_stage": self.release_stage.value,
            "release_date": self.release_date,
            "archive_published_date": self.archive_published_date,
            "archive_timestamp": self.archive_timestamp,
            "artifact_uri": self.artifact_uri,
            "title": self.title,
            "publication_kind": self.publication_kind,
            "archive_listed": self.archive_listed,
        }

    def to_dict(self) -> dict[str, JSONValue]:
        return {**self.identity_payload(), "entry_id": self.entry_id}

    @classmethod
    def from_dict(cls, data: Mapping[str, Any]) -> BeaPioReleaseIndexEntryV1:
        return cls(
            primary_reference_period=str(
                data.get("primary_reference_period", "")
            ),
            release_stage=EconomicReleaseStage.from_value(
                str(data.get("release_stage", ""))
            ),
            release_date=str(data.get("release_date", "")),
            archive_published_date=str(data.get("archive_published_date", "")),
            archive_timestamp=str(data.get("archive_timestamp", "")),
            artifact_uri=str(data.get("artifact_uri", "")),
            title=str(data.get("title", "")),
            publication_kind=str(data.get("publication_kind", "")),
            archive_listed=cast(bool, data.get("archive_listed")),
            entry_id=str(data.get("entry_id", "")),
            schema_version=str(data.get("schema_version", "")),
        )


@dataclass(frozen=True, slots=True)
class BeaPioReleaseIndexV1:
    """Content-addressed official PIO archive inventory."""

    pages: tuple[BeaPioReleaseIndexPageV1, ...]
    as_of_date: str
    releases: tuple[BeaPioReleaseIndexEntryV1, ...]
    index_id: str = ""
    schema_version: str = BEA_PIO_INDEX_SCHEMA_VERSION

    def __post_init__(self) -> None:
        if self.schema_version != BEA_PIO_INDEX_SCHEMA_VERSION:
            raise ValueError("unsupported BEA PIO release-index schema")
        pages = tuple(self.pages)
        if (
            len(pages) != BEA_PIO_INDEX_PAGE_COUNT
            or any(
                not isinstance(item, BeaPioReleaseIndexPageV1) for item in pages
            )
            or tuple(item.page_number for item in pages)
            != tuple(range(BEA_PIO_INDEX_PAGE_COUNT))
        ):
            raise ValueError("BEA PIO index pages are incomplete")
        as_of = _iso_date(self.as_of_date, "as_of_date")
        releases = tuple(self.releases)
        if (
            not releases
            or len(releases) > MAX_BEA_PIO_RELEASES
            or any(
                not isinstance(item, BeaPioReleaseIndexEntryV1)
                for item in releases
            )
        ):
            raise TypeError("BEA PIO index releases are invalid")
        if tuple(item.release_date for item in releases) != tuple(
            sorted(item.release_date for item in releases)
        ):
            raise ValueError("BEA PIO releases are not chronological")
        if len({item.artifact_uri for item in releases}) != len(releases):
            raise ValueError("BEA PIO index repeats an artifact URI")
        if any(
            not US_BACKFILL_START_DATE <= item.release_date <= as_of
            for item in releases
        ):
            raise ValueError("BEA PIO release lies outside the as-of window")
        revision_updates = tuple(
            item
            for item in releases
            if item.release_stage is EconomicReleaseStage.REVISION
        )
        if len(revision_updates) != 1 or revision_updates[0].artifact_uri != (
            "https://www.bea.gov/news/2025/"
            "personal-income-and-outlays-data-update-september-2025"
        ):
            raise ValueError("BEA PIO revision-only publication differs")
        object.__setattr__(self, "pages", pages)
        object.__setattr__(self, "as_of_date", as_of)
        object.__setattr__(self, "releases", releases)
        expected = _stable_id("bea-pio-release-index", self.identity_payload())
        if self.index_id and self.index_id != expected:
            raise ValueError("BEA PIO release-index identity differs")
        object.__setattr__(self, "index_id", expected)

    @property
    def by_uri(self) -> Mapping[str, BeaPioReleaseIndexEntryV1]:
        return MappingProxyType(
            {item.artifact_uri: item for item in self.releases}
        )

    def identity_payload(self) -> dict[str, JSONValue]:
        return {
            "schema_version": self.schema_version,
            "pages": [item.to_dict() for item in self.pages],
            "as_of_date": self.as_of_date,
            "releases": [item.to_dict() for item in self.releases],
        }

    def to_dict(self) -> dict[str, JSONValue]:
        return {**self.identity_payload(), "index_id": self.index_id}

    @classmethod
    def from_dict(cls, data: Mapping[str, Any]) -> BeaPioReleaseIndexV1:
        return cls(
            pages=tuple(
                BeaPioReleaseIndexPageV1.from_dict(
                    _mapping(item, "BEA PIO index page")
                )
                for item in _sequence(data.get("pages"), "pages")
            ),
            as_of_date=str(data.get("as_of_date", "")),
            releases=tuple(
                BeaPioReleaseIndexEntryV1.from_dict(
                    _mapping(item, "BEA PIO index entry")
                )
                for item in _sequence(data.get("releases"), "releases")
            ),
            index_id=str(data.get("index_id", "")),
            schema_version=str(data.get("schema_version", "")),
        )


@dataclass(frozen=True, slots=True)
class BeaPioMeasureValueV1:
    """One newly published or explicitly revised monthly PIO value."""

    reference_period: str
    measure_key: str
    release_stage: EconomicReleaseStage
    actual_value: float
    actual_lexical: str
    value_basis: BeaPioValueBasis
    value_locator: str
    value_artifact_uri: str
    value_content_sha256: str
    value_content_length: int
    reported_previous_value: float | None
    reported_previous_lexical: str | None
    previous_locator: str | None
    value_id: str = ""
    schema_version: str = BEA_PIO_MEASURE_VALUE_SCHEMA_VERSION

    def __post_init__(self) -> None:
        if self.schema_version != BEA_PIO_MEASURE_VALUE_SCHEMA_VERSION:
            raise ValueError("unsupported BEA PIO measure-value schema")
        reference = _year_month(self.reference_period, "reference_period")
        measure = _required_text(self.measure_key, "measure_key")
        if measure not in BEA_PIO_MEASURE_KEYS:
            raise ValueError("BEA PIO measure key is invalid")
        stage = EconomicReleaseStage.from_value(self.release_stage)
        if stage not in {
            EconomicReleaseStage.INITIAL,
            EconomicReleaseStage.REVISION,
        }:
            raise ValueError("BEA PIO measure release stage is invalid")
        actual_lexical, actual = _numeric_lexical(
            self.actual_lexical, "actual_lexical"
        )
        if not math.isclose(
            actual,
            _finite(self.actual_value, "actual_value"),
            rel_tol=0.0,
            abs_tol=1e-12,
        ):
            raise ValueError("BEA PIO actual lexical differs from actual_value")
        basis = BeaPioValueBasis.from_value(self.value_basis)
        if measure in {
            BEA_PIO_PERSONAL_INCOME,
            BEA_PIO_CURRENT_DOLLAR_PCE,
        } and (basis is not BeaPioValueBasis.SOURCE_PUBLISHED):
            raise ValueError("BEA PIO current-dollar measure cannot be derived")
        locator = _required_text(self.value_locator, "value_locator")
        value_uri = _https_uri(self.value_artifact_uri, "value_artifact_uri")
        value_sha = _sha256(self.value_content_sha256, "value_content_sha256")
        _positive_int(
            self.value_content_length,
            "value_content_length",
            MAX_BEA_PIO_RELEASE_BYTES,
        )
        previous_locator = _optional_text(
            self.previous_locator, "previous_locator"
        )
        if stage is EconomicReleaseStage.REVISION:
            if any(
                item is not None
                for item in (
                    self.reported_previous_value,
                    self.reported_previous_lexical,
                    previous_locator,
                )
            ):
                raise ValueError(
                    "revision-only PIO value cannot report a prior month"
                )
            previous_value = None
            previous_lexical = None
        else:
            if self.reported_previous_value is None:
                if (
                    self.reported_previous_lexical is not None
                    or previous_locator
                ):
                    raise ValueError(
                        "BEA PIO previous-value evidence is partial"
                    )
                previous_value = None
                previous_lexical = None
            else:
                if (
                    self.reported_previous_lexical is None
                    or previous_locator is None
                ):
                    raise ValueError(
                        "BEA PIO previous-value evidence is partial"
                    )
                previous_lexical, previous_value = _numeric_lexical(
                    self.reported_previous_lexical,
                    "reported_previous_lexical",
                )
                if not math.isclose(
                    previous_value,
                    _finite(
                        self.reported_previous_value,
                        "reported_previous_value",
                    ),
                    rel_tol=0.0,
                    abs_tol=1e-12,
                ):
                    raise ValueError("BEA PIO reported prior lexical differs")
        object.__setattr__(self, "reference_period", reference)
        object.__setattr__(self, "measure_key", measure)
        object.__setattr__(self, "release_stage", stage)
        object.__setattr__(self, "actual_value", actual)
        object.__setattr__(self, "actual_lexical", actual_lexical)
        object.__setattr__(self, "value_basis", basis)
        object.__setattr__(self, "value_locator", locator)
        object.__setattr__(self, "value_artifact_uri", value_uri)
        object.__setattr__(self, "value_content_sha256", value_sha)
        object.__setattr__(self, "reported_previous_value", previous_value)
        object.__setattr__(self, "reported_previous_lexical", previous_lexical)
        object.__setattr__(self, "previous_locator", previous_locator)
        expected = _stable_id("bea-pio-measure-value", self.identity_payload())
        if self.value_id and self.value_id != expected:
            raise ValueError("BEA PIO measure-value identity differs")
        object.__setattr__(self, "value_id", expected)

    def identity_payload(self) -> dict[str, JSONValue]:
        return {
            "schema_version": self.schema_version,
            "reference_period": self.reference_period,
            "measure_key": self.measure_key,
            "release_stage": self.release_stage.value,
            "actual_value": self.actual_value,
            "actual_lexical": self.actual_lexical,
            "value_basis": self.value_basis.value,
            "value_locator": self.value_locator,
            "value_artifact_uri": self.value_artifact_uri,
            "value_content_sha256": self.value_content_sha256,
            "value_content_length": self.value_content_length,
            "reported_previous_value": self.reported_previous_value,
            "reported_previous_lexical": self.reported_previous_lexical,
            "previous_locator": self.previous_locator,
        }

    def to_dict(self) -> dict[str, JSONValue]:
        return {**self.identity_payload(), "value_id": self.value_id}

    @classmethod
    def from_dict(cls, data: Mapping[str, Any]) -> BeaPioMeasureValueV1:
        return cls(
            reference_period=str(data.get("reference_period", "")),
            measure_key=str(data.get("measure_key", "")),
            release_stage=EconomicReleaseStage.from_value(
                str(data.get("release_stage", ""))
            ),
            actual_value=cast(float, data.get("actual_value")),
            actual_lexical=str(data.get("actual_lexical", "")),
            value_basis=BeaPioValueBasis.from_value(
                str(data.get("value_basis", ""))
            ),
            value_locator=str(data.get("value_locator", "")),
            value_artifact_uri=str(data.get("value_artifact_uri", "")),
            value_content_sha256=str(data.get("value_content_sha256", "")),
            value_content_length=cast(int, data.get("value_content_length")),
            reported_previous_value=cast(
                float | None, data.get("reported_previous_value")
            ),
            reported_previous_lexical=cast(
                str | None, data.get("reported_previous_lexical")
            ),
            previous_locator=cast(str | None, data.get("previous_locator")),
            value_id=str(data.get("value_id", "")),
            schema_version=str(data.get("schema_version", "")),
        )


@dataclass(frozen=True, slots=True)
class BeaPioReleaseObservationV1:
    """Parsed contents and time evidence for one official PIO publication."""

    index_entry_id: str
    release_date: str
    artifact_uri: str
    content_sha256: str
    content_length: int
    released_lexical: str
    time_precision: EconomicTimePrecision
    release_header_lexical: str
    reported_zone: str | None
    zone_consistent: bool | None
    release_time_locator: str
    source_era: str
    historical_update_notice: bool
    support_uri: str | None
    support_format: OfficialSourceFormat | None
    support_content_sha256: str | None
    support_content_length: int | None
    values: tuple[BeaPioMeasureValueV1, ...]
    observation_id: str = ""
    schema_version: str = BEA_PIO_RELEASE_OBSERVATION_SCHEMA_VERSION

    def __post_init__(self) -> None:
        if self.schema_version != BEA_PIO_RELEASE_OBSERVATION_SCHEMA_VERSION:
            raise ValueError("unsupported BEA PIO release-observation schema")
        index_id = _required_text(self.index_entry_id, "index_entry_id")
        if not index_id.startswith("bea-pio-index-entry:sha256:"):
            raise ValueError("BEA PIO observation index identity is invalid")
        release = _iso_date(self.release_date, "release_date")
        artifact = _https_uri(self.artifact_uri, "artifact_uri")
        content_sha = _sha256(self.content_sha256, "content_sha256")
        _positive_int(
            self.content_length,
            "content_length",
            MAX_BEA_PIO_RELEASE_BYTES,
        )
        precision = EconomicTimePrecision.from_value(self.time_precision)
        released = _required_text(self.released_lexical, "released_lexical")
        header = _required_text(
            self.release_header_lexical, "release_header_lexical"
        )
        zone = _optional_text(self.reported_zone, "reported_zone")
        if precision is EconomicTimePrecision.DATE_ONLY:
            if (
                released != release
                or zone is not None
                or self.zone_consistent is not None
            ):
                raise ValueError(
                    "date-only BEA PIO time evidence is inconsistent"
                )
        else:
            if precision is not EconomicTimePrecision.EXACT_MINUTE:
                raise ValueError(
                    "BEA PIO time precision must be date or exact-minute"
                )
            if (
                re.fullmatch(
                    rf"{re.escape(release)}T(?:08:30|10:00):00", released
                )
                is None
            ):
                raise ValueError("BEA PIO exact release time is invalid")
            if zone not in {"EST", "EDT"} or not isinstance(
                self.zone_consistent, bool
            ):
                raise ValueError("BEA PIO zone evidence is invalid")
        locator = _required_text(
            self.release_time_locator, "release_time_locator"
        )
        era = _required_text(self.source_era, "source_era")
        if era not in {
            "legacy-html",
            "preformatted-html",
            "semantic-table-html",
            "streamlined-html",
            "pdf-only",
            "xlsx-backed-html",
            "revision-only-xlsx",
        }:
            raise ValueError("BEA PIO source era is invalid")
        if not isinstance(self.historical_update_notice, bool):
            raise TypeError("historical_update_notice must be boolean")
        support_uri = _optional_text(self.support_uri, "support_uri")
        support_format = self.support_format
        support_sha = _optional_text(
            self.support_content_sha256, "support_content_sha256"
        )
        support_length = self.support_content_length
        supplied = (support_uri, support_format, support_sha, support_length)
        if any(item is not None for item in supplied):
            if any(item is None for item in supplied):
                raise ValueError("BEA PIO support evidence is partial")
            assert support_uri is not None
            assert support_format is not None
            assert support_sha is not None
            assert support_length is not None
            support_uri = _https_uri(support_uri, "support_uri")
            support_format = OfficialSourceFormat.from_value(support_format)
            support_sha = _sha256(support_sha, "support_content_sha256")
            _positive_int(
                support_length,
                "support_content_length",
                MAX_BEA_PIO_RELEASE_BYTES,
            )
            if BEA_PIO_SUPPORT_ARTIFACTS.get(artifact) != (
                support_uri,
                support_format,
            ):
                raise ValueError(
                    "BEA PIO support artifact differs from registry"
                )
        elif artifact in BEA_PIO_SUPPORT_ARTIFACTS:
            raise ValueError("BEA PIO observation omits required support")
        values = tuple(self.values)
        if (
            not values
            or len(values) > 32
            or any(
                not isinstance(item, BeaPioMeasureValueV1) for item in values
            )
        ):
            raise TypeError("BEA PIO observation values are invalid")
        keys = tuple(
            (item.reference_period, item.measure_key) for item in values
        )
        if keys != tuple(sorted(keys)) or len(set(keys)) != len(keys):
            raise ValueError(
                "BEA PIO observation values are not uniquely sorted"
            )
        allowed_value_uris = {artifact}
        if support_uri is not None:
            allowed_value_uris.add(support_uri)
        if any(
            item.value_artifact_uri not in allowed_value_uris for item in values
        ):
            raise ValueError(
                "BEA PIO value evidence is outside its publication"
            )
        object.__setattr__(self, "index_entry_id", index_id)
        object.__setattr__(self, "release_date", release)
        object.__setattr__(self, "artifact_uri", artifact)
        object.__setattr__(self, "content_sha256", content_sha)
        object.__setattr__(self, "released_lexical", released)
        object.__setattr__(self, "time_precision", precision)
        object.__setattr__(self, "release_header_lexical", header)
        object.__setattr__(self, "reported_zone", zone)
        object.__setattr__(self, "release_time_locator", locator)
        object.__setattr__(self, "source_era", era)
        object.__setattr__(self, "support_uri", support_uri)
        object.__setattr__(self, "support_format", support_format)
        object.__setattr__(self, "support_content_sha256", support_sha)
        object.__setattr__(self, "support_content_length", support_length)
        object.__setattr__(self, "values", values)
        expected = _stable_id(
            "bea-pio-release-observation", self.identity_payload()
        )
        if self.observation_id and self.observation_id != expected:
            raise ValueError("BEA PIO release-observation identity differs")
        object.__setattr__(self, "observation_id", expected)

    @property
    def by_key(self) -> Mapping[tuple[str, str], BeaPioMeasureValueV1]:
        return MappingProxyType(
            {
                (item.reference_period, item.measure_key): item
                for item in self.values
            }
        )

    def identity_payload(self) -> dict[str, JSONValue]:
        return {
            "schema_version": self.schema_version,
            "index_entry_id": self.index_entry_id,
            "release_date": self.release_date,
            "artifact_uri": self.artifact_uri,
            "content_sha256": self.content_sha256,
            "content_length": self.content_length,
            "released_lexical": self.released_lexical,
            "time_precision": self.time_precision.value,
            "release_header_lexical": self.release_header_lexical,
            "reported_zone": self.reported_zone,
            "zone_consistent": self.zone_consistent,
            "release_time_locator": self.release_time_locator,
            "source_era": self.source_era,
            "historical_update_notice": self.historical_update_notice,
            "support_uri": self.support_uri,
            "support_format": (
                None
                if self.support_format is None
                else self.support_format.value
            ),
            "support_content_sha256": self.support_content_sha256,
            "support_content_length": self.support_content_length,
            "values": [item.to_dict() for item in self.values],
        }

    def to_dict(self) -> dict[str, JSONValue]:
        return {
            **self.identity_payload(),
            "observation_id": self.observation_id,
        }

    @classmethod
    def from_dict(cls, data: Mapping[str, Any]) -> BeaPioReleaseObservationV1:
        support_format_value = data.get("support_format")
        return cls(
            index_entry_id=str(data.get("index_entry_id", "")),
            release_date=str(data.get("release_date", "")),
            artifact_uri=str(data.get("artifact_uri", "")),
            content_sha256=str(data.get("content_sha256", "")),
            content_length=cast(int, data.get("content_length")),
            released_lexical=str(data.get("released_lexical", "")),
            time_precision=EconomicTimePrecision.from_value(
                str(data.get("time_precision", ""))
            ),
            release_header_lexical=str(data.get("release_header_lexical", "")),
            reported_zone=cast(str | None, data.get("reported_zone")),
            zone_consistent=cast(bool | None, data.get("zone_consistent")),
            release_time_locator=str(data.get("release_time_locator", "")),
            source_era=str(data.get("source_era", "")),
            historical_update_notice=cast(
                bool, data.get("historical_update_notice")
            ),
            support_uri=cast(str | None, data.get("support_uri")),
            support_format=(
                None
                if support_format_value is None
                else OfficialSourceFormat.from_value(str(support_format_value))
            ),
            support_content_sha256=cast(
                str | None, data.get("support_content_sha256")
            ),
            support_content_length=cast(
                int | None, data.get("support_content_length")
            ),
            values=tuple(
                BeaPioMeasureValueV1.from_dict(
                    _mapping(item, "BEA PIO measure value")
                )
                for item in _sequence(data.get("values"), "values")
            ),
            observation_id=str(data.get("observation_id", "")),
            schema_version=str(data.get("schema_version", "")),
        )


@dataclass(frozen=True, slots=True)
class BeaPioArchiveMeasureV1:
    """One PIO value with its ex-ante comparison lineage."""

    value: BeaPioMeasureValueV1
    comparison_reference_period: str
    comparison_basis: str
    previous_artifact_uri: str
    previous_content_sha256: str
    previous_content_length: int
    previous_as_known_value: float
    previous_as_known_lexical: str
    revised_previous_value: float
    revised_previous_lexical: str
    revision_comparable: bool
    normalized_sha256: str
    measure_id: str = ""
    schema_version: str = BEA_PIO_ARCHIVE_MEASURE_SCHEMA_VERSION

    def __post_init__(self) -> None:
        if self.schema_version != BEA_PIO_ARCHIVE_MEASURE_SCHEMA_VERSION:
            raise ValueError("unsupported BEA PIO archive-measure schema")
        if not isinstance(self.value, BeaPioMeasureValueV1):
            raise TypeError("BEA PIO archive measure requires a parsed value")
        comparison = _year_month(
            self.comparison_reference_period, "comparison_reference_period"
        )
        expected_comparison = (
            self.value.reference_period
            if self.value.release_stage is EconomicReleaseStage.REVISION
            else _previous_month(self.value.reference_period)
        )
        if comparison != expected_comparison:
            raise ValueError("BEA PIO comparison period differs from stage")
        basis = _required_text(self.comparison_basis, "comparison_basis")
        if basis not in {
            "prior-publication",
            "same-publication",
            "prior-not-restated",
            "basis-transition",
            "revision-publication",
        }:
            raise ValueError("BEA PIO comparison basis is invalid")
        previous_uri = _https_uri(
            self.previous_artifact_uri, "previous_artifact_uri"
        )
        previous_sha = _sha256(
            self.previous_content_sha256, "previous_content_sha256"
        )
        _positive_int(
            self.previous_content_length,
            "previous_content_length",
            MAX_BEA_PIO_RELEASE_BYTES,
        )
        previous_lexical, previous = _numeric_lexical(
            self.previous_as_known_lexical, "previous_as_known_lexical"
        )
        revised_lexical, revised = _numeric_lexical(
            self.revised_previous_lexical, "revised_previous_lexical"
        )
        if not math.isclose(
            previous,
            _finite(self.previous_as_known_value, "previous_as_known_value"),
            rel_tol=0.0,
            abs_tol=1e-12,
        ) or not math.isclose(
            revised,
            _finite(self.revised_previous_value, "revised_previous_value"),
            rel_tol=0.0,
            abs_tol=1e-12,
        ):
            raise ValueError("BEA PIO comparison lexical values differ")
        if not isinstance(self.revision_comparable, bool):
            raise TypeError("revision_comparable must be boolean")
        expected_comparable = basis in {
            "prior-publication",
            "revision-publication",
        }
        if self.revision_comparable is not expected_comparable:
            raise ValueError(
                "BEA PIO revision comparability differs from basis"
            )
        if self.value.release_stage is EconomicReleaseStage.REVISION and not (
            basis == "revision-publication"
            and math.isclose(
                revised,
                self.value.actual_value,
                rel_tol=0.0,
                abs_tol=1e-12,
            )
        ):
            raise ValueError("BEA PIO revision-only value differs")
        normalized = _sha256(self.normalized_sha256, "normalized_sha256")
        expected_normalized = _normalized_sha256(
            self.value,
            comparison,
            basis,
            previous_lexical,
            revised_lexical,
        )
        if normalized != expected_normalized:
            raise ValueError("BEA PIO normalized digest differs")
        object.__setattr__(self, "comparison_reference_period", comparison)
        object.__setattr__(self, "comparison_basis", basis)
        object.__setattr__(self, "previous_artifact_uri", previous_uri)
        object.__setattr__(self, "previous_content_sha256", previous_sha)
        object.__setattr__(self, "previous_as_known_value", previous)
        object.__setattr__(self, "previous_as_known_lexical", previous_lexical)
        object.__setattr__(self, "revised_previous_value", revised)
        object.__setattr__(self, "revised_previous_lexical", revised_lexical)
        object.__setattr__(self, "normalized_sha256", normalized)
        expected = _stable_id(
            "bea-pio-archive-measure", self.identity_payload()
        )
        if self.measure_id and self.measure_id != expected:
            raise ValueError("BEA PIO archive-measure identity differs")
        object.__setattr__(self, "measure_id", expected)

    @property
    def previous_was_revised(self) -> bool:
        return self.revision_comparable and not math.isclose(
            self.previous_as_known_value,
            self.revised_previous_value,
            rel_tol=0.0,
            abs_tol=1e-12,
        )

    def identity_payload(self) -> dict[str, JSONValue]:
        return {
            "schema_version": self.schema_version,
            "value": self.value.to_dict(),
            "comparison_reference_period": self.comparison_reference_period,
            "comparison_basis": self.comparison_basis,
            "previous_artifact_uri": self.previous_artifact_uri,
            "previous_content_sha256": self.previous_content_sha256,
            "previous_content_length": self.previous_content_length,
            "previous_as_known_value": self.previous_as_known_value,
            "previous_as_known_lexical": self.previous_as_known_lexical,
            "revised_previous_value": self.revised_previous_value,
            "revised_previous_lexical": self.revised_previous_lexical,
            "revision_comparable": self.revision_comparable,
            "normalized_sha256": self.normalized_sha256,
        }

    def to_dict(self) -> dict[str, JSONValue]:
        return {**self.identity_payload(), "measure_id": self.measure_id}

    @classmethod
    def from_dict(cls, data: Mapping[str, Any]) -> BeaPioArchiveMeasureV1:
        return cls(
            value=BeaPioMeasureValueV1.from_dict(
                _mapping(data.get("value"), "BEA PIO value")
            ),
            comparison_reference_period=str(
                data.get("comparison_reference_period", "")
            ),
            comparison_basis=str(data.get("comparison_basis", "")),
            previous_artifact_uri=str(data.get("previous_artifact_uri", "")),
            previous_content_sha256=str(
                data.get("previous_content_sha256", "")
            ),
            previous_content_length=cast(
                int, data.get("previous_content_length")
            ),
            previous_as_known_value=cast(
                float, data.get("previous_as_known_value")
            ),
            previous_as_known_lexical=str(
                data.get("previous_as_known_lexical", "")
            ),
            revised_previous_value=cast(
                float, data.get("revised_previous_value")
            ),
            revised_previous_lexical=str(
                data.get("revised_previous_lexical", "")
            ),
            revision_comparable=cast(bool, data.get("revision_comparable")),
            normalized_sha256=str(data.get("normalized_sha256", "")),
            measure_id=str(data.get("measure_id", "")),
            schema_version=str(data.get("schema_version", "")),
        )


@dataclass(frozen=True, slots=True)
class BeaPioArchivePublicationV1:
    """One official publication and all newly issued/revised measure values."""

    observation: BeaPioReleaseObservationV1
    measures: tuple[BeaPioArchiveMeasureV1, ...]
    publication_id: str = ""
    schema_version: str = BEA_PIO_ARCHIVE_PUBLICATION_SCHEMA_VERSION

    def __post_init__(self) -> None:
        if self.schema_version != BEA_PIO_ARCHIVE_PUBLICATION_SCHEMA_VERSION:
            raise ValueError("unsupported BEA PIO archive-publication schema")
        if not isinstance(self.observation, BeaPioReleaseObservationV1):
            raise TypeError("BEA PIO publication requires an observation")
        measures = tuple(self.measures)
        if (
            not measures
            or any(
                not isinstance(item, BeaPioArchiveMeasureV1)
                for item in measures
            )
            or tuple(item.value.value_id for item in measures)
            != tuple(item.value_id for item in self.observation.values)
        ):
            raise ValueError("BEA PIO publication measure package differs")
        object.__setattr__(self, "measures", measures)
        expected = _stable_id(
            "bea-pio-archive-publication", self.identity_payload()
        )
        if self.publication_id and self.publication_id != expected:
            raise ValueError("BEA PIO archive-publication identity differs")
        object.__setattr__(self, "publication_id", expected)

    def identity_payload(self) -> dict[str, JSONValue]:
        return {
            "schema_version": self.schema_version,
            "observation": self.observation.to_dict(),
            "measures": [item.to_dict() for item in self.measures],
        }

    def to_dict(self) -> dict[str, JSONValue]:
        return {
            **self.identity_payload(),
            "publication_id": self.publication_id,
        }

    @classmethod
    def from_dict(cls, data: Mapping[str, Any]) -> BeaPioArchivePublicationV1:
        return cls(
            observation=BeaPioReleaseObservationV1.from_dict(
                _mapping(data.get("observation"), "BEA PIO observation")
            ),
            measures=tuple(
                BeaPioArchiveMeasureV1.from_dict(
                    _mapping(item, "BEA PIO archive measure")
                )
                for item in _sequence(data.get("measures"), "measures")
            ),
            publication_id=str(data.get("publication_id", "")),
            schema_version=str(data.get("schema_version", "")),
        )


@dataclass(frozen=True, slots=True)
class BeaPioArchiveManifestV1:
    """Compact deterministic receipt for the full PIO publication corpus."""

    registry_id: str
    profile_id: str
    release_index: BeaPioReleaseIndexV1
    predecessor_uri: str
    predecessor_content_sha256: str
    predecessor_content_length: int
    publications: tuple[BeaPioArchivePublicationV1, ...]
    raw_artifact_count: int
    total_content_bytes: int
    initial_publication_count: int
    revision_publication_count: int
    measure_occurrence_count: int
    personal_income_initial_count: int
    current_dollar_pce_initial_count: int
    headline_price_initial_count: int
    core_price_initial_count: int
    comparable_revision_count: int
    personal_income_revision_count: int
    current_dollar_pce_revision_count: int
    headline_price_revision_count: int
    core_price_revision_count: int
    derived_price_occurrence_count: int
    historical_update_notice_count: int
    shutdown_publication_count: int
    support_artifact_count: int
    exact_minute_count: int
    date_only_count: int
    zone_mismatch_count: int
    archive_date_mismatch_count: int
    legacy_html_count: int
    preformatted_html_count: int
    semantic_table_html_count: int
    streamlined_html_count: int
    pdf_only_count: int
    xlsx_backed_html_count: int
    manifest_id: str = ""
    schema_version: str = BEA_PIO_ARCHIVE_MANIFEST_SCHEMA_VERSION

    def __post_init__(self) -> None:
        if self.schema_version != BEA_PIO_ARCHIVE_MANIFEST_SCHEMA_VERSION:
            raise ValueError("unsupported BEA PIO archive-manifest schema")
        registry = _required_text(self.registry_id, "registry_id")
        profile = _required_text(self.profile_id, "profile_id")
        if not registry.startswith("official-source-registry:sha256:"):
            raise ValueError("BEA PIO manifest registry identity is invalid")
        if not profile.startswith("us-backfill-profile:sha256:"):
            raise ValueError("BEA PIO manifest profile identity is invalid")
        if not isinstance(self.release_index, BeaPioReleaseIndexV1):
            raise TypeError("BEA PIO manifest requires a v1 release index")
        predecessor_uri = _https_uri(self.predecessor_uri, "predecessor_uri")
        if predecessor_uri != BEA_PIO_PREDECESSOR_URI:
            raise ValueError("BEA PIO predecessor URI differs")
        predecessor_sha = _sha256(
            self.predecessor_content_sha256,
            "predecessor_content_sha256",
        )
        _positive_int(
            self.predecessor_content_length,
            "predecessor_content_length",
            MAX_BEA_PIO_RELEASE_BYTES,
        )
        publications = tuple(self.publications)
        if len(publications) != len(self.release_index.releases) or any(
            not isinstance(item, BeaPioArchivePublicationV1)
            for item in publications
        ):
            raise ValueError("BEA PIO manifest publications are incomplete")
        for indexed, publication in zip(
            self.release_index.releases, publications, strict=True
        ):
            observation = publication.observation
            if (
                observation.index_entry_id != indexed.entry_id
                or observation.release_date != indexed.release_date
                or observation.artifact_uri != indexed.artifact_uri
            ):
                raise ValueError("BEA PIO publication differs from its index")
        values = tuple(
            measure
            for publication in publications
            for measure in publication.measures
        )
        if not values or len(values) > MAX_BEA_PIO_MEASURE_OCCURRENCES:
            raise ValueError("BEA PIO measure corpus is invalid")
        normalized = {item.normalized_sha256 for item in values}
        if len(normalized) != len(values):
            raise ValueError("BEA PIO manifest repeats normalized content")
        initial_values = tuple(
            item
            for item in values
            if item.value.release_stage is EconomicReleaseStage.INITIAL
        )
        initial_by_measure = {
            key: tuple(
                item for item in initial_values if item.value.measure_key == key
            )
            for key in BEA_PIO_MEASURE_KEYS
        }
        for key in (BEA_PIO_PERSONAL_INCOME, BEA_PIO_CURRENT_DOLLAR_PCE):
            periods = tuple(
                item.value.reference_period for item in initial_by_measure[key]
            )
            if (
                not periods
                or periods[0] != "1999-12"
                or periods[-1] != "2026-07"
            ):
                raise ValueError(
                    "BEA PIO current-dollar coverage endpoints differ"
                )
            for previous, current in pairwise(periods):
                if current != _next_month(previous):
                    raise ValueError(
                        "BEA PIO current-dollar coverage has a gap"
                    )
        for key in (BEA_PIO_HEADLINE_PRICE, BEA_PIO_CORE_PRICE):
            periods = tuple(
                item.value.reference_period for item in initial_by_measure[key]
            )
            if (
                not periods
                or periods[0] != "2000-06"
                or periods[-1] != "2026-07"
            ):
                raise ValueError("BEA PIO price coverage endpoints differ")
            for previous, current in pairwise(periods):
                if current != _next_month(previous):
                    raise ValueError("BEA PIO price coverage has a gap")
        artifacts = {
            (
                publication.observation.artifact_uri,
                publication.observation.content_sha256,
                publication.observation.content_length,
            )
            for publication in publications
        }
        artifacts.add(
            (
                predecessor_uri,
                predecessor_sha,
                self.predecessor_content_length,
            )
        )
        for publication in publications:
            observation = publication.observation
            if observation.support_uri is not None:
                artifacts.add(
                    (
                        observation.support_uri,
                        cast(str, observation.support_content_sha256),
                        cast(int, observation.support_content_length),
                    )
                )
        if len({item[1] for item in artifacts}) != len(artifacts):
            raise ValueError("BEA PIO manifest repeats a raw artifact hash")
        revision_counts = {
            key: sum(
                item.previous_was_revised
                for item in values
                if item.value.measure_key == key
            )
            for key in BEA_PIO_MEASURE_KEYS
        }
        era_counts = {
            era: sum(
                publication.observation.source_era == era
                for publication in publications
            )
            for era in (
                "legacy-html",
                "preformatted-html",
                "semantic-table-html",
                "streamlined-html",
                "pdf-only",
                "xlsx-backed-html",
            )
        }
        expected_counts = {
            "raw_artifact_count": len(artifacts),
            "total_content_bytes": sum(item[2] for item in artifacts),
            "initial_publication_count": sum(
                publication.observation.values[0].release_stage
                is EconomicReleaseStage.INITIAL
                for publication in publications
            ),
            "revision_publication_count": sum(
                publication.observation.values[0].release_stage
                is EconomicReleaseStage.REVISION
                for publication in publications
            ),
            "measure_occurrence_count": len(values),
            "personal_income_initial_count": len(
                initial_by_measure[BEA_PIO_PERSONAL_INCOME]
            ),
            "current_dollar_pce_initial_count": len(
                initial_by_measure[BEA_PIO_CURRENT_DOLLAR_PCE]
            ),
            "headline_price_initial_count": len(
                initial_by_measure[BEA_PIO_HEADLINE_PRICE]
            ),
            "core_price_initial_count": len(
                initial_by_measure[BEA_PIO_CORE_PRICE]
            ),
            "comparable_revision_count": sum(
                item.previous_was_revised for item in values
            ),
            "personal_income_revision_count": revision_counts[
                BEA_PIO_PERSONAL_INCOME
            ],
            "current_dollar_pce_revision_count": revision_counts[
                BEA_PIO_CURRENT_DOLLAR_PCE
            ],
            "headline_price_revision_count": revision_counts[
                BEA_PIO_HEADLINE_PRICE
            ],
            "core_price_revision_count": revision_counts[BEA_PIO_CORE_PRICE],
            "derived_price_occurrence_count": sum(
                item.value.value_basis
                is BeaPioValueBasis.DERIVED_FROM_INDEX_LEVELS
                for item in values
            ),
            "historical_update_notice_count": sum(
                publication.observation.historical_update_notice
                for publication in publications
            ),
            "shutdown_publication_count": sum(
                indexed.publication_kind
                in {
                    "shutdown-split",
                    "shutdown-catch-up",
                    "shutdown-combined",
                    "revision-only-update",
                }
                for indexed in self.release_index.releases
            ),
            "support_artifact_count": sum(
                publication.observation.support_uri is not None
                for publication in publications
            ),
            "exact_minute_count": sum(
                publication.observation.time_precision
                is EconomicTimePrecision.EXACT_MINUTE
                for publication in publications
            ),
            "date_only_count": sum(
                publication.observation.time_precision
                is EconomicTimePrecision.DATE_ONLY
                for publication in publications
            ),
            "zone_mismatch_count": sum(
                publication.observation.zone_consistent is False
                for publication in publications
            ),
            "archive_date_mismatch_count": sum(
                item.release_date != item.archive_published_date
                for item in self.release_index.releases
            ),
            "legacy_html_count": era_counts["legacy-html"],
            "preformatted_html_count": era_counts["preformatted-html"],
            "semantic_table_html_count": era_counts["semantic-table-html"],
            "streamlined_html_count": era_counts["streamlined-html"],
            "pdf_only_count": era_counts["pdf-only"],
            "xlsx_backed_html_count": era_counts["xlsx-backed-html"],
        }
        for name, expected in expected_counts.items():
            maximum = (
                MAX_BEA_PIO_RELEASE_BYTES * (MAX_BEA_PIO_RELEASES + 32)
                if name == "total_content_bytes"
                else MAX_BEA_PIO_MEASURE_OCCURRENCES * 100
            )
            supplied = _positive_int(getattr(self, name), name, maximum)
            if supplied != expected:
                raise ValueError(f"BEA PIO {name} differs from entries")
        object.__setattr__(self, "registry_id", registry)
        object.__setattr__(self, "profile_id", profile)
        object.__setattr__(self, "predecessor_uri", predecessor_uri)
        object.__setattr__(self, "predecessor_content_sha256", predecessor_sha)
        object.__setattr__(self, "publications", publications)
        expected_id = _stable_id(
            "bea-pio-archive-manifest", self.identity_payload()
        )
        if self.manifest_id and self.manifest_id != expected_id:
            raise ValueError("BEA PIO archive-manifest identity differs")
        object.__setattr__(self, "manifest_id", expected_id)

    @property
    def by_publication_uri(self) -> Mapping[str, BeaPioArchivePublicationV1]:
        return MappingProxyType(
            {item.observation.artifact_uri: item for item in self.publications}
        )

    def identity_payload(self) -> dict[str, JSONValue]:
        return {
            "schema_version": self.schema_version,
            "registry_id": self.registry_id,
            "profile_id": self.profile_id,
            "release_index": self.release_index.to_dict(),
            "predecessor_uri": self.predecessor_uri,
            "predecessor_content_sha256": self.predecessor_content_sha256,
            "predecessor_content_length": self.predecessor_content_length,
            "publications": [item.to_dict() for item in self.publications],
            "raw_artifact_count": self.raw_artifact_count,
            "total_content_bytes": self.total_content_bytes,
            "initial_publication_count": self.initial_publication_count,
            "revision_publication_count": self.revision_publication_count,
            "measure_occurrence_count": self.measure_occurrence_count,
            "personal_income_initial_count": self.personal_income_initial_count,
            "current_dollar_pce_initial_count": (
                self.current_dollar_pce_initial_count
            ),
            "headline_price_initial_count": self.headline_price_initial_count,
            "core_price_initial_count": self.core_price_initial_count,
            "comparable_revision_count": self.comparable_revision_count,
            "personal_income_revision_count": (
                self.personal_income_revision_count
            ),
            "current_dollar_pce_revision_count": (
                self.current_dollar_pce_revision_count
            ),
            "headline_price_revision_count": self.headline_price_revision_count,
            "core_price_revision_count": self.core_price_revision_count,
            "derived_price_occurrence_count": (
                self.derived_price_occurrence_count
            ),
            "historical_update_notice_count": (
                self.historical_update_notice_count
            ),
            "shutdown_publication_count": self.shutdown_publication_count,
            "support_artifact_count": self.support_artifact_count,
            "exact_minute_count": self.exact_minute_count,
            "date_only_count": self.date_only_count,
            "zone_mismatch_count": self.zone_mismatch_count,
            "archive_date_mismatch_count": self.archive_date_mismatch_count,
            "legacy_html_count": self.legacy_html_count,
            "preformatted_html_count": self.preformatted_html_count,
            "semantic_table_html_count": self.semantic_table_html_count,
            "streamlined_html_count": self.streamlined_html_count,
            "pdf_only_count": self.pdf_only_count,
            "xlsx_backed_html_count": self.xlsx_backed_html_count,
        }

    def to_dict(self) -> dict[str, JSONValue]:
        return {**self.identity_payload(), "manifest_id": self.manifest_id}

    def to_json(self) -> str:
        return str(canonical_contract_json(self.to_dict()))

    @classmethod
    def from_dict(cls, data: Mapping[str, Any]) -> BeaPioArchiveManifestV1:
        return cls(
            registry_id=str(data.get("registry_id", "")),
            profile_id=str(data.get("profile_id", "")),
            release_index=BeaPioReleaseIndexV1.from_dict(
                _mapping(data.get("release_index"), "BEA PIO release index")
            ),
            predecessor_uri=str(data.get("predecessor_uri", "")),
            predecessor_content_sha256=str(
                data.get("predecessor_content_sha256", "")
            ),
            predecessor_content_length=cast(
                int, data.get("predecessor_content_length")
            ),
            publications=tuple(
                BeaPioArchivePublicationV1.from_dict(
                    _mapping(item, "BEA PIO publication")
                )
                for item in _sequence(data.get("publications"), "publications")
            ),
            raw_artifact_count=cast(int, data.get("raw_artifact_count")),
            total_content_bytes=cast(int, data.get("total_content_bytes")),
            initial_publication_count=cast(
                int, data.get("initial_publication_count")
            ),
            revision_publication_count=cast(
                int, data.get("revision_publication_count")
            ),
            measure_occurrence_count=cast(
                int, data.get("measure_occurrence_count")
            ),
            personal_income_initial_count=cast(
                int, data.get("personal_income_initial_count")
            ),
            current_dollar_pce_initial_count=cast(
                int, data.get("current_dollar_pce_initial_count")
            ),
            headline_price_initial_count=cast(
                int, data.get("headline_price_initial_count")
            ),
            core_price_initial_count=cast(
                int, data.get("core_price_initial_count")
            ),
            comparable_revision_count=cast(
                int, data.get("comparable_revision_count")
            ),
            personal_income_revision_count=cast(
                int, data.get("personal_income_revision_count")
            ),
            current_dollar_pce_revision_count=cast(
                int, data.get("current_dollar_pce_revision_count")
            ),
            headline_price_revision_count=cast(
                int, data.get("headline_price_revision_count")
            ),
            core_price_revision_count=cast(
                int, data.get("core_price_revision_count")
            ),
            derived_price_occurrence_count=cast(
                int, data.get("derived_price_occurrence_count")
            ),
            historical_update_notice_count=cast(
                int, data.get("historical_update_notice_count")
            ),
            shutdown_publication_count=cast(
                int, data.get("shutdown_publication_count")
            ),
            support_artifact_count=cast(
                int, data.get("support_artifact_count")
            ),
            exact_minute_count=cast(int, data.get("exact_minute_count")),
            date_only_count=cast(int, data.get("date_only_count")),
            zone_mismatch_count=cast(int, data.get("zone_mismatch_count")),
            archive_date_mismatch_count=cast(
                int, data.get("archive_date_mismatch_count")
            ),
            legacy_html_count=cast(int, data.get("legacy_html_count")),
            preformatted_html_count=cast(
                int, data.get("preformatted_html_count")
            ),
            semantic_table_html_count=cast(
                int, data.get("semantic_table_html_count")
            ),
            streamlined_html_count=cast(
                int, data.get("streamlined_html_count")
            ),
            pdf_only_count=cast(int, data.get("pdf_only_count")),
            xlsx_backed_html_count=cast(
                int, data.get("xlsx_backed_html_count")
            ),
            manifest_id=str(data.get("manifest_id", "")),
            schema_version=str(data.get("schema_version", "")),
        )

    @classmethod
    def from_json(cls, payload: str) -> BeaPioArchiveManifestV1:
        return cls.from_dict(_mapping(json.loads(payload), "BEA PIO manifest"))


def build_bea_pio_index_requests(
    registry: OfficialSourceRegistryV1,
) -> tuple[OfficialSourceRequestV1, ...]:
    """Build the exact 19-page official PIO archive inventory request set."""
    if not isinstance(registry, OfficialSourceRegistryV1):
        raise TypeError("BEA PIO index requests require a v1 registry")
    source = registry.source(BEA_PIO_SOURCE_KEY)
    if OfficialSourceFormat.HTML not in source.formats:
        raise ValueError("BEA PIO source does not declare HTML")
    return tuple(
        OfficialSourceRequestV1(
            source_key=source.source_key,
            source_id=source.source_id,
            method=OfficialRequestMethod.GET,
            uri=_archive_page_uri(page),
            source_format=OfficialSourceFormat.HTML,
            parser_id=source.parser_id,
            parser_version=source.parser_version,
            page_number=page,
        )
        for page in range(BEA_PIO_INDEX_PAGE_COUNT)
    )


def _validate_index_snapshots(
    snapshots: Sequence[OfficialRawSnapshotV1],
) -> tuple[OfficialRawSnapshotV1, ...]:
    values = tuple(snapshots)
    if len(values) != BEA_PIO_INDEX_PAGE_COUNT or any(
        not isinstance(item, OfficialRawSnapshotV1) for item in values
    ):
        raise ValueError("BEA PIO index snapshots are invalid")
    ordered = tuple(
        sorted(values, key=lambda item: cast(int, item.request.page_number))
    )
    if tuple(item.request.page_number for item in ordered) != tuple(
        range(BEA_PIO_INDEX_PAGE_COUNT)
    ):
        raise ValueError("BEA PIO index snapshots are not contiguous")
    for page, snapshot in enumerate(ordered):
        if (
            snapshot.request.source_key != BEA_PIO_SOURCE_KEY
            or snapshot.request.source_format is not OfficialSourceFormat.HTML
            or snapshot.request.uri != _archive_page_uri(page)
            or not snapshot.content
            or len(snapshot.content) > MAX_BEA_PIO_INDEX_PAGE_BYTES
        ):
            raise ValueError("BEA PIO index snapshot scope differs")
    return ordered


def parse_bea_pio_release_index(
    snapshots: Sequence[OfficialRawSnapshotV1],
    *,
    as_of_date: str,
) -> BeaPioReleaseIndexV1:
    """Parse all archive pages and add the bounded current-release supplement."""
    ordered = _validate_index_snapshots(snapshots)
    as_of = _iso_date(as_of_date, "as_of_date")
    pages: list[BeaPioReleaseIndexPageV1] = []
    releases: list[BeaPioReleaseIndexEntryV1] = []
    for snapshot in ordered:
        page = cast(int, snapshot.request.page_number)
        pages.append(
            BeaPioReleaseIndexPageV1(
                page_number=page,
                source_uri=snapshot.request.uri,
                content_sha256=snapshot.content_sha256,
                content_length=len(snapshot.content),
            )
        )
        markup = _decode_html(snapshot.content)
        for row in _ARCHIVE_ROW_RE.finditer(markup):
            title_match = _ARCHIVE_TITLE_RE.search(row.group("body"))
            time_match = _ARCHIVE_TIME_RE.search(row.group("body"))
            if title_match is None or time_match is None:
                continue
            title = _visible_html(title_match.group("title"))
            if not title.lower().startswith("personal income"):
                continue
            stamp = html.unescape(time_match.group("stamp")).strip()
            try:
                parsed_stamp = datetime.fromisoformat(stamp)
            except ValueError as exc:
                raise ValueError("BEA PIO archive timestamp changed") from exc
            published = parsed_stamp.date().isoformat()
            if published < US_BACKFILL_START_DATE or published > as_of:
                continue
            artifact = urljoin(
                "https://www.bea.gov",
                html.unescape(title_match.group("uri")),
            )
            artifact = re.sub(
                r"^https://www\.bea\.gov/index\.(?:php|ph%70)/news/",
                "https://www.bea.gov/news/",
                artifact,
                flags=re.IGNORECASE,
            )
            released = BEA_PIO_ARCHIVE_DATE_CORRECTIONS.get(artifact, published)
            releases.append(
                BeaPioReleaseIndexEntryV1(
                    primary_reference_period=_primary_period(artifact, title),
                    release_stage=_release_stage(artifact),
                    release_date=released,
                    archive_published_date=published,
                    archive_timestamp=stamp,
                    artifact_uri=artifact,
                    title=title,
                    publication_kind=_publication_kind(artifact),
                    archive_listed=True,
                )
            )
    if as_of >= BEA_PIO_SUPPLEMENTAL_DATE and not any(
        item.artifact_uri == BEA_PIO_SUPPLEMENTAL_URI for item in releases
    ):
        releases.append(
            BeaPioReleaseIndexEntryV1(
                primary_reference_period="2026-07",
                release_stage=EconomicReleaseStage.INITIAL,
                release_date=BEA_PIO_SUPPLEMENTAL_DATE,
                archive_published_date=BEA_PIO_SUPPLEMENTAL_DATE,
                archive_timestamp="2026-08-26T08:30:00-04:00",
                artifact_uri=BEA_PIO_SUPPLEMENTAL_URI,
                title=BEA_PIO_SUPPLEMENTAL_TITLE,
                publication_kind="regular",
                archive_listed=False,
            )
        )
    return BeaPioReleaseIndexV1(
        pages=tuple(pages),
        as_of_date=as_of,
        releases=tuple(sorted(releases, key=lambda item: item.release_date)),
    )


def _release_request(
    registry: OfficialSourceRegistryV1,
    uri: str,
    release_date: str,
    source_format: OfficialSourceFormat,
) -> OfficialSourceRequestV1:
    source = registry.source(BEA_PIO_SOURCE_KEY)
    if source_format not in source.formats:
        raise ValueError("BEA PIO source does not declare an artifact format")
    return OfficialSourceRequestV1(
        source_key=source.source_key,
        source_id=source.source_id,
        method=OfficialRequestMethod.GET,
        uri=uri,
        source_format=source_format,
        parser_id=source.parser_id,
        parser_version=source.parser_version,
        window_start=release_date,
        window_end=release_date,
    )


def build_bea_pio_archive_requests(
    registry: OfficialSourceRegistryV1,
    release_index: BeaPioReleaseIndexV1,
) -> tuple[OfficialSourceRequestV1, ...]:
    """Build exact requests for releases, supplements, and predecessor data."""
    if not isinstance(registry, OfficialSourceRegistryV1):
        raise TypeError("BEA PIO requests require a v1 registry")
    if not isinstance(release_index, BeaPioReleaseIndexV1):
        raise TypeError("BEA PIO requests require a v1 release index")
    result = [
        _release_request(
            registry,
            item.artifact_uri,
            item.release_date,
            OfficialSourceFormat.HTML,
        )
        for item in release_index.releases
    ]
    for item in release_index.releases:
        support = BEA_PIO_SUPPORT_ARTIFACTS.get(item.artifact_uri)
        if support is not None:
            result.append(
                _release_request(
                    registry,
                    support[0],
                    item.release_date,
                    support[1],
                )
            )
    result.append(
        _release_request(
            registry,
            BEA_PIO_PREDECESSOR_URI,
            BEA_PIO_PREDECESSOR_DATE,
            OfficialSourceFormat.PDF,
        )
    )
    if len({item.uri for item in result}) != len(result):
        raise ValueError("BEA PIO request set repeats an artifact")
    return tuple(result)


def _validate_release_snapshot(
    snapshot: OfficialRawSnapshotV1,
    *,
    uri: str,
    release_date: str,
    source_format: OfficialSourceFormat,
) -> None:
    if not isinstance(snapshot, OfficialRawSnapshotV1):
        raise TypeError("BEA PIO parser requires a v1 snapshot")
    if (
        snapshot.request.source_key != BEA_PIO_SOURCE_KEY
        or snapshot.request.uri != uri
        or snapshot.request.source_format is not source_format
        or snapshot.request.window_start != release_date
        or snapshot.request.window_end != release_date
        or not snapshot.content
        or len(snapshot.content) > MAX_BEA_PIO_RELEASE_BYTES
    ):
        raise ValueError("BEA PIO release snapshot scope differs")


def _lexical_number(value: object) -> tuple[str, float]:
    text = str(value).strip()
    if text.startswith("-."):
        text = "-0" + text[1:]
    elif text.startswith("."):
        text = "0" + text
    try:
        number = float(text)
    except ValueError as exc:
        raise ValueError("BEA PIO table value is not numeric") from exc
    if not math.isfinite(number):
        raise ValueError("BEA PIO table value is not finite")
    lexical = f"{number:.1f}"
    return lexical, float(lexical)


def _row_tokens(text: str, patterns: Sequence[str]) -> tuple[str, ...]:
    token = r"(?:\.\.\.|-?(?:\d+\.\d+|\.\d+|\d+))"
    for pattern in patterns:
        match = re.search(
            rf"(?:{pattern})\s*[:;,]?\s*"
            rf"(?:(?:\.(?!\d))+\s*)?"
            rf"(?P<values>{token}(?:\s+{token}){{0,31}})",
            text,
            re.IGNORECASE,
        )
        if match is not None:
            return tuple(_TOKEN_RE.findall(match.group("values")))
    return ()


def _map_tokens(
    tokens: Sequence[str], ending_period: str
) -> dict[str, tuple[str, float]]:
    values = tuple(tokens)
    if not values:
        return {}
    periods = _periods_ending(ending_period, len(values))
    result: dict[str, tuple[str, float]] = {}
    for period, lexical in zip(periods, values, strict=True):
        if lexical == "...":
            continue
        result[period] = _lexical_number(lexical)
    return result


def _first_position(text: str, markers: Sequence[str]) -> int | None:
    positions = [
        found
        for marker in markers
        if (found := text.lower().find(marker.lower())) >= 0
    ]
    return min(positions) if positions else None


def _current_dollar_maps(
    text: str, ending_period: str
) -> dict[str, dict[str, tuple[str, float]]]:
    start = _first_position(
        text,
        (
            "percent change from preceding month",
            "percent change from preceding period",
            "percent change from jan. to feb.",
            "percent change from august to september",
        ),
    )
    section = text if start is None else text[start:]
    income = _row_tokens(
        section,
        (
            r"Personal income\s*,\s*current dollars",
            r"Personal income\s*:\s*Current dollars",
            r"Current-dollar personal income",
        ),
    )
    pce = _row_tokens(
        section,
        (
            r"Personal consumption expenditures(?:\s*\(PCE\))?\s*:\s*Current dollars",
            r"Current-dollar personal consumption expenditures(?:\s*\(PCE\))?",
            r"Current-dollar PCE",
        ),
    )
    if not income or not pce:
        raise ValueError("BEA PIO release omits current-dollar summary values")
    return {
        BEA_PIO_PERSONAL_INCOME: _map_tokens(income, ending_period),
        BEA_PIO_CURRENT_DOLLAR_PCE: _map_tokens(pce, ending_period),
    }


def _derived_rate_map(
    levels: Mapping[str, tuple[str, float]],
) -> dict[str, tuple[str, float]]:
    ordered = tuple(sorted(levels.items()))
    result: dict[str, tuple[str, float]] = {}
    for previous, current in pairwise(ordered):
        if current[0] != _next_month(previous[0]):
            continue
        rate = round(((current[1][1] / previous[1][1]) - 1.0) * 100.0, 1)
        result[current[0]] = (f"{rate:.1f}", rate)
    return result


_NARRATIVE_CHANGE_RE = re.compile(
    r"(?:(?P<direction>increased|decreased)\s+"
    r"(?P<less>less than\s+)?(?P<value>\d+(?:\.\d+)?)\s+percent|"
    r"(?:(?:was|were)\s+)?(?:essentially\s+)?unchanged)",
    re.IGNORECASE,
)


def _narrative_rate_tokens(
    text: str, patterns: Sequence[str]
) -> tuple[str, ...]:
    for pattern in patterns:
        for label in re.finditer(pattern, text, re.IGNORECASE):
            tail = text[label.end() : label.end() + 300]
            matches = tuple(_NARRATIVE_CHANGE_RE.finditer(tail))
            if not matches or matches[0].start() > 80:
                continue
            values: list[str] = []
            for match in matches[:2]:
                if match.group("direction") is None or match.group("less"):
                    values.append("0.0")
                    continue
                magnitude = float(cast(str, match.group("value")))
                if match.group("direction").lower() == "decreased":
                    magnitude = -magnitude
                values.append(f"{magnitude:.1f}")
            if len(values) == 1 and re.search(
                r"same (?:increase|decrease)", tail, re.IGNORECASE
            ):
                values.append(values[0])
            if len(values) > 1:
                values[0], values[1] = values[1], values[0]
            return tuple(values)
    return ()


def _price_maps(text: str, ending_period: str) -> tuple[
    dict[str, dict[str, tuple[str, float]]],
    BeaPioValueBasis,
    str,
]:
    summary_start = text.lower().find("price indexes:")
    if summary_start >= 0:
        section = text[summary_start : summary_start + 2500]
        headline = _row_tokens(section, (r"\bPCE\b",))
        core = _row_tokens(section, (r"PCE\s*,\s*excluding food and energy",))
        if headline and core:
            return (
                {
                    BEA_PIO_HEADLINE_PRICE: _map_tokens(
                        headline, ending_period
                    ),
                    BEA_PIO_CORE_PRICE: _map_tokens(core, ending_period),
                },
                BeaPioValueBasis.SOURCE_PUBLISHED,
                "html:summary-price-rows",
            )
    direct_start = text.lower().rfind(
        "percent change from preceding period in price indexes"
    )
    if direct_start >= 0:
        section = text[direct_start : direct_start + 4000]
        headline = _row_tokens(
            section,
            (
                r"\bPCE\b(?!\s+(?:less|excluding))",
                r"Personal consumption expenditures(?:\s*\(PCE\))?",
            ),
        )
        core = _row_tokens(
            section,
            (
                (
                    r"(?:Personal consumption expenditures|PCE)\s+"
                    r"(?:less|excluding)\s+food and energy"
                ),
            ),
        )
        if headline and core:
            return (
                {
                    BEA_PIO_HEADLINE_PRICE: _map_tokens(
                        headline, ending_period
                    ),
                    BEA_PIO_CORE_PRICE: _map_tokens(core, ending_period),
                },
                BeaPioValueBasis.SOURCE_PUBLISHED,
                "html:detailed-price-rate-rows",
            )
    modern_start = text.lower().find("personal income and related measures")
    if modern_start >= 0:
        section = text[modern_start : modern_start + 5000]
        modern_headline = _row_tokens(
            section,
            (r"PCE price index(?!\s*,?\s*excluding)",),
        )
        modern_core = _row_tokens(
            section,
            (r"PCE price index\s*,?\s*excluding food and energy",),
        )
        if modern_headline and modern_core:
            return (
                {
                    BEA_PIO_HEADLINE_PRICE: _map_tokens(
                        modern_headline, ending_period
                    ),
                    BEA_PIO_CORE_PRICE: _map_tokens(modern_core, ending_period),
                },
                BeaPioValueBasis.SOURCE_PUBLISHED,
                "html:streamlined-summary-table",
            )
    if ending_period >= "2002-02":
        narrative_headline = _narrative_rate_tokens(
            text,
            (
                r"price index for PCE(?!\s*,?\s*excluding)",
                r"PCE price index(?!\s*,\s*excluding)",
            ),
        )
        narrative_core = _narrative_rate_tokens(
            text,
            (
                r"PCE price index\s*,?\s*excluding food and energy",
                r"price index for PCE\s*,?\s*excluding food and energy",
                r"prices?\s*,\s*excluding food and energy",
            ),
        )
        if narrative_headline and narrative_core:
            return (
                {
                    BEA_PIO_HEADLINE_PRICE: _map_tokens(
                        narrative_headline, ending_period
                    ),
                    BEA_PIO_CORE_PRICE: _map_tokens(
                        narrative_core, ending_period
                    ),
                },
                BeaPioValueBasis.SOURCE_PUBLISHED,
                "html:narrative-price-change",
            )
    level_start = text.lower().rfind("chain-type price indexes")
    if level_start >= 0:
        section = text[level_start : level_start + 8000]
        headline_levels = _row_tokens(
            section,
            (r"Personal consumption expenditures(?:\s*\(PCE\))?",),
        )
        core_levels = _row_tokens(
            section,
            (
                (
                    r"(?:Personal consumption expenditures|PCE)\s+"
                    r"(?:less|excluding)\s+food and energy"
                ),
            ),
        )
        if headline_levels and core_levels:
            return (
                {
                    BEA_PIO_HEADLINE_PRICE: _derived_rate_map(
                        _map_tokens(headline_levels, ending_period)
                    ),
                    BEA_PIO_CORE_PRICE: _derived_rate_map(
                        _map_tokens(core_levels, ending_period)
                    ),
                },
                BeaPioValueBasis.DERIVED_FROM_INDEX_LEVELS,
                "derived:html:price-index-level-rows",
            )
    return ({}, BeaPioValueBasis.SOURCE_PUBLISHED, "html:price-unavailable")


def _xlsx_periods(sheet: Any) -> tuple[str, ...]:
    years = tuple(sheet.cell(4, column).value for column in range(3, 11))
    months = tuple(sheet.cell(5, column).value for column in range(3, 11))
    result = []
    for year, month in zip(years, months, strict=True):
        if not isinstance(year, int) or not isinstance(month, str):
            raise TypeError("BEA PIO workbook period header changed")
        cleaned = re.sub(r"\s+[pr]$", "", month.strip(), flags=re.IGNORECASE)
        cleaned = cleaned.rstrip(".").lower()
        if cleaned not in _MONTH_NUMBERS:
            raise ValueError("BEA PIO workbook month header changed")
        result.append(f"{year:04d}-{_MONTH_NUMBERS[cleaned]:02d}")
    return tuple(result)


def _xlsx_row(
    sheet: Any,
    label: str,
    *,
    after_marker: str | None = None,
) -> dict[str, tuple[str, float]]:
    periods = _xlsx_periods(sheet)
    marker_seen = after_marker is None
    for row in range(1, sheet.max_row + 1):
        row_values = tuple(
            sheet.cell(row, column).value for column in range(1, 11)
        )
        if after_marker is not None and any(
            isinstance(item, str)
            and item.strip().lower().startswith(after_marker.lower())
            for item in row_values
        ):
            marker_seen = True
            continue
        if not marker_seen or str(row_values[1]).strip() != label:
            continue
        result: dict[str, tuple[str, float]] = {}
        for period, value in zip(periods, row_values[2:], strict=True):
            if not isinstance(value, (int, float)) or isinstance(value, bool):
                raise TypeError("BEA PIO workbook measure row is not numeric")
            result[period] = _lexical_number(value)
        return result
    raise ValueError(f"BEA PIO workbook omits {label}")


def _xlsx_maps(
    content: bytes,
) -> dict[str, dict[str, tuple[str, float]]]:
    try:
        workbook = load_workbook(
            BytesIO(content), read_only=True, data_only=True
        )
        income_sheet = workbook["Table 3"]
        price_sheet = workbook["Table 5"]
        return {
            BEA_PIO_PERSONAL_INCOME: _xlsx_row(income_sheet, "Personal income"),
            BEA_PIO_CURRENT_DOLLAR_PCE: _xlsx_row(
                income_sheet, "Personal consumption expenditures"
            ),
            BEA_PIO_HEADLINE_PRICE: _xlsx_row(
                price_sheet,
                "Personal consumption expenditures (PCE)",
                after_marker="Percent change from preceding period",
            ),
            BEA_PIO_CORE_PRICE: _xlsx_row(
                price_sheet,
                "PCE excluding food and energy",
                after_marker="Percent change from preceding period",
            ),
        }
    except Exception as exc:
        raise ValueError("BEA PIO workbook cannot be parsed") from exc


def _initial_periods(
    index_entry: BeaPioReleaseIndexEntryV1, measure_key: str
) -> tuple[str, ...]:
    uri = index_entry.artifact_uri
    if index_entry.release_stage is EconomicReleaseStage.REVISION:
        return ("2025-07", "2025-08", "2025-09")
    if uri.endswith(
        "personal-income-and-outlays-december-2018-personal-income-january-2019"
    ):
        if measure_key == BEA_PIO_PERSONAL_INCOME:
            return ("2018-12", "2019-01")
        return ("2018-12",)
    if uri.endswith(
        "personal-income-february-2019-personal-outlays-january-2019"
    ):
        if measure_key == BEA_PIO_PERSONAL_INCOME:
            return ("2019-02",)
        return ("2019-01",)
    if uri.endswith("personal-income-and-outlays-march-2019"):
        if measure_key == BEA_PIO_PERSONAL_INCOME:
            return ("2019-03",)
        return ("2019-02", "2019-03")
    if uri.endswith("personal-income-and-outlays-october-and-november-2025"):
        return ("2025-10", "2025-11")
    return (index_entry.primary_reference_period,)


def _source_era(
    markup: str,
    index_entry: BeaPioReleaseIndexEntryV1,
    support_format: OfficialSourceFormat | None,
) -> str:
    if index_entry.release_stage is EconomicReleaseStage.REVISION:
        return "revision-only-xlsx"
    if support_format is OfficialSourceFormat.XLSX:
        return "xlsx-backed-html"
    if support_format is OfficialSourceFormat.PDF and not re.search(
        r"(?i)personal income (?:increased|decreased)", _visible_html(markup)
    ):
        return "pdf-only"
    body = re.search(
        r"(?is)<div\b[^>]*class=[\"'][^\"']*field--name-body[^\"']*"
        r"[\"'][^>]*>(?P<body>.*?)(?:</div>\s*</div>|</article>)",
        markup,
    )
    body_markup = body.group("body") if body is not None else markup
    if re.search(r"(?is)<pre\b", body_markup):
        return "preformatted-html"
    if (
        "personal income and related measures"
        in _visible_html(body_markup).lower()
    ):
        if index_entry.release_date >= "2025-03-28":
            return "streamlined-html"
        return "semantic-table-html"
    if re.search(r"(?is)<table\b", body_markup):
        return "semantic-table-html"
    return "legacy-html"


def _release_time(
    visible_text: str,
    index_entry: BeaPioReleaseIndexEntryV1,
    *,
    support_text: str | None,
    support_format: OfficialSourceFormat | None,
) -> tuple[
    str,
    EconomicTimePrecision,
    str,
    str | None,
    bool | None,
    str,
]:
    match = _RELEASE_TIME_RE.search(visible_text)
    locator = "html:visible-release-header"
    if match is None and support_text is not None:
        match = _RELEASE_TIME_RE.search(support_text)
        locator = (
            "supporting-pdf:page-1:release-header"
            if support_format is OfficialSourceFormat.PDF
            else "supporting-text:release-header"
        )
    if match is None:
        if index_entry.artifact_uri.endswith(
            "personal-income-and-outlays-data-update-september-2025"
        ):
            stamp = datetime.fromisoformat(index_entry.archive_timestamp)
            local = stamp.astimezone(ZoneInfo("America/New_York"))
            if (local.hour, local.minute) != (8, 30):
                raise ValueError("BEA PIO data-update archive time changed")
            zone = cast(str, local.tzname())
            return (
                f"{index_entry.release_date}T08:30:00",
                EconomicTimePrecision.EXACT_MINUTE,
                index_entry.archive_timestamp,
                zone,
                True,
                "archive-index:datetime",
            )
        if index_entry.release_date == "2000-02-28":
            return (
                index_entry.release_date,
                EconomicTimePrecision.DATE_ONLY,
                index_entry.archive_timestamp,
                None,
                None,
                "archive-index:date-only",
            )
        raise ValueError("BEA PIO release omits its exact release time")
    parsed_date = _long_date(match.group("date"))
    if parsed_date.isoformat() != index_entry.release_date:
        raise ValueError("BEA PIO release header date differs from index")
    hour = int(match.group("hour"))
    minute = int(match.group("minute"))
    if (hour, minute) not in {(8, 30), (10, 0)}:
        raise ValueError("BEA PIO release time is outside observed bounds")
    zone = match.group("zone").upper()
    aware = datetime.combine(
        parsed_date,
        time(hour, minute),
        tzinfo=ZoneInfo("America/New_York"),
    )
    return (
        f"{index_entry.release_date}T{hour:02d}:{minute:02d}:00",
        EconomicTimePrecision.EXACT_MINUTE,
        " ".join(match.group(0).split()),
        zone,
        aware.tzname() == zone,
        locator,
    )


def _value_source(
    snapshot: OfficialRawSnapshotV1,
    support: OfficialRawSnapshotV1 | None,
    support_format: OfficialSourceFormat | None,
) -> OfficialRawSnapshotV1:
    if support_format in {OfficialSourceFormat.PDF, OfficialSourceFormat.XLSX}:
        if support is None:
            raise ValueError("BEA PIO value support is unavailable")
        return support
    return snapshot


def parse_bea_pio_release(
    snapshot: OfficialRawSnapshotV1,
    index_entry: BeaPioReleaseIndexEntryV1,
    *,
    support_snapshot: OfficialRawSnapshotV1 | None = None,
) -> BeaPioReleaseObservationV1:
    """Parse every newly issued/revised value from one official publication."""
    if not isinstance(index_entry, BeaPioReleaseIndexEntryV1):
        raise TypeError("BEA PIO parser requires a v1 index entry")
    _validate_release_snapshot(
        snapshot,
        uri=index_entry.artifact_uri,
        release_date=index_entry.release_date,
        source_format=OfficialSourceFormat.HTML,
    )
    declared_support = BEA_PIO_SUPPORT_ARTIFACTS.get(index_entry.artifact_uri)
    support_uri: str | None = None
    support_format: OfficialSourceFormat | None = None
    support_sha: str | None = None
    support_length: int | None = None
    support_text: str | None = None
    if declared_support is not None:
        if support_snapshot is None:
            raise ValueError(
                "BEA PIO release omits its required support artifact"
            )
        support_uri, support_format = declared_support
        _validate_release_snapshot(
            support_snapshot,
            uri=support_uri,
            release_date=index_entry.release_date,
            source_format=support_format,
        )
        support_sha = support_snapshot.content_sha256
        support_length = len(support_snapshot.content)
        if support_format is OfficialSourceFormat.PDF:
            support_text = _pdf_text(support_snapshot.content)
        elif support_format is OfficialSourceFormat.TEXT:
            support_text = " ".join(
                support_snapshot.content.decode("utf-8-sig", "replace").split()
            )
    elif support_snapshot is not None:
        raise ValueError(
            "BEA PIO release received an unexpected support artifact"
        )
    markup = _decode_html(snapshot.content)
    visible = _visible_html(markup)
    released = _release_time(
        visible,
        index_entry,
        support_text=support_text,
        support_format=support_format,
    )
    if support_format is OfficialSourceFormat.XLSX:
        assert support_snapshot is not None
        maps = _xlsx_maps(support_snapshot.content)
        value_basis = BeaPioValueBasis.SOURCE_PUBLISHED
        value_locator = "supporting-xlsx:table-3-and-table-5"
    else:
        source_text = (
            support_text
            if support_format is OfficialSourceFormat.PDF
            else visible
        )
        assert source_text is not None
        maps = _current_dollar_maps(
            source_text, index_entry.primary_reference_period
        )
        price_maps, value_basis, value_locator = _price_maps(
            source_text, index_entry.primary_reference_period
        )
        maps.update(price_maps)
    value_snapshot = _value_source(snapshot, support_snapshot, support_format)
    values: list[BeaPioMeasureValueV1] = []
    for measure_key in BEA_PIO_MEASURE_KEYS:
        measure_map = maps.get(measure_key, {})
        for reference in _initial_periods(index_entry, measure_key):
            if reference not in measure_map:
                if (
                    measure_key in {BEA_PIO_HEADLINE_PRICE, BEA_PIO_CORE_PRICE}
                    and reference < "2000-06"
                ):
                    continue
                raise ValueError(
                    f"BEA PIO release omits {measure_key} for {reference}"
                )
            lexical, actual = measure_map[reference]
            previous_reference = _previous_month(reference)
            reported = measure_map.get(previous_reference)
            is_revision = (
                index_entry.release_stage is EconomicReleaseStage.REVISION
            )
            basis = (
                value_basis
                if measure_key in {BEA_PIO_HEADLINE_PRICE, BEA_PIO_CORE_PRICE}
                else BeaPioValueBasis.SOURCE_PUBLISHED
            )
            values.append(
                BeaPioMeasureValueV1(
                    reference_period=reference,
                    measure_key=measure_key,
                    release_stage=index_entry.release_stage,
                    actual_value=actual,
                    actual_lexical=lexical,
                    value_basis=basis,
                    value_locator=(
                        value_locator
                        if measure_key
                        in {
                            BEA_PIO_HEADLINE_PRICE,
                            BEA_PIO_CORE_PRICE,
                        }
                        else (
                            "supporting-xlsx:table-3"
                            if support_format is OfficialSourceFormat.XLSX
                            else "html:current-dollar-summary-row"
                        )
                    ),
                    value_artifact_uri=value_snapshot.request.uri,
                    value_content_sha256=value_snapshot.content_sha256,
                    value_content_length=len(value_snapshot.content),
                    reported_previous_value=(
                        None if is_revision or reported is None else reported[1]
                    ),
                    reported_previous_lexical=(
                        None if is_revision or reported is None else reported[0]
                    ),
                    previous_locator=(
                        None
                        if is_revision or reported is None
                        else (
                            value_locator
                            if measure_key
                            in {BEA_PIO_HEADLINE_PRICE, BEA_PIO_CORE_PRICE}
                            else (
                                "supporting-xlsx:table-3"
                                if support_format is OfficialSourceFormat.XLSX
                                else "html:current-dollar-summary-row"
                            )
                        )
                    ),
                )
            )
    title_lower = index_entry.title.lower()
    return BeaPioReleaseObservationV1(
        index_entry_id=index_entry.entry_id,
        release_date=index_entry.release_date,
        artifact_uri=index_entry.artifact_uri,
        content_sha256=snapshot.content_sha256,
        content_length=len(snapshot.content),
        released_lexical=released[0],
        time_precision=released[1],
        release_header_lexical=released[2],
        reported_zone=released[3],
        zone_consistent=released[4],
        release_time_locator=released[5],
        source_era=_source_era(markup, index_entry, support_format),
        historical_update_notice=any(
            marker in title_lower for marker in _UPDATE_TITLE_MARKERS
        ),
        support_uri=support_uri,
        support_format=support_format,
        support_content_sha256=support_sha,
        support_content_length=support_length,
        values=tuple(
            sorted(
                values,
                key=lambda item: (item.reference_period, item.measure_key),
            )
        ),
    )


def _normalized_sha256(
    value: BeaPioMeasureValueV1,
    comparison_period: str,
    comparison_basis: str,
    previous_lexical: str,
    revised_lexical: str,
) -> str:
    payload: dict[str, JSONValue] = {
        "measure_key": value.measure_key,
        "reference_period": value.reference_period,
        "comparison_reference_period": comparison_period,
        "release_stage": value.release_stage.value,
        "value_basis": value.value_basis.value,
        "comparison_basis": comparison_basis,
        "actual_lexical": value.actual_lexical,
        "previous_as_known_lexical": previous_lexical,
        "revised_previous_lexical": revised_lexical,
    }
    return hashlib.sha256(
        str(canonical_contract_json(payload)).encode("utf-8")
    ).hexdigest()


@dataclass(frozen=True, slots=True)
class _KnownValue:
    lexical: str
    value: float
    basis: BeaPioValueBasis
    artifact_uri: str
    content_sha256: str
    content_length: int


def _known_from_value(value: BeaPioMeasureValueV1) -> _KnownValue:
    return _KnownValue(
        lexical=value.actual_lexical,
        value=value.actual_value,
        basis=value.value_basis,
        artifact_uri=value.value_artifact_uri,
        content_sha256=value.value_content_sha256,
        content_length=value.value_content_length,
    )


def _archive_measure(
    value: BeaPioMeasureValueV1,
    before: Mapping[tuple[str, str], _KnownValue],
) -> BeaPioArchiveMeasureV1:
    if value.release_stage is EconomicReleaseStage.REVISION:
        comparison_period = value.reference_period
        previous = before.get((value.measure_key, comparison_period))
        if previous is None:
            raise ValueError("BEA PIO revision has no prior published value")
        revised_lexical = value.actual_lexical
        revised = value.actual_value
        comparison_basis = "revision-publication"
        comparable = True
    else:
        comparison_period = _previous_month(value.reference_period)
        previous = before.get((value.measure_key, comparison_period))
        if previous is None:
            if value.reported_previous_value is None:
                raise ValueError(
                    "BEA PIO initial value has no comparison evidence"
                )
            previous = _KnownValue(
                lexical=cast(str, value.reported_previous_lexical),
                value=value.reported_previous_value,
                basis=value.value_basis,
                artifact_uri=value.value_artifact_uri,
                content_sha256=value.value_content_sha256,
                content_length=value.value_content_length,
            )
            comparison_basis = "same-publication"
            comparable = False
        elif value.reported_previous_value is None:
            comparison_basis = "prior-not-restated"
            comparable = False
        elif previous.basis is not value.value_basis:
            comparison_basis = "basis-transition"
            comparable = False
        else:
            comparison_basis = "prior-publication"
            comparable = True
        if value.reported_previous_value is None:
            revised_lexical = previous.lexical
            revised = previous.value
        else:
            revised_lexical = cast(str, value.reported_previous_lexical)
            revised = value.reported_previous_value
    normalized = _normalized_sha256(
        value,
        comparison_period,
        comparison_basis,
        previous.lexical,
        revised_lexical,
    )
    return BeaPioArchiveMeasureV1(
        value=value,
        comparison_reference_period=comparison_period,
        comparison_basis=comparison_basis,
        previous_artifact_uri=previous.artifact_uri,
        previous_content_sha256=previous.content_sha256,
        previous_content_length=previous.content_length,
        previous_as_known_value=previous.value,
        previous_as_known_lexical=previous.lexical,
        revised_previous_value=revised,
        revised_previous_lexical=revised_lexical,
        revision_comparable=comparable,
        normalized_sha256=normalized,
    )


def _state_from_xlsx(
    snapshot: OfficialRawSnapshotV1,
) -> dict[tuple[str, str], _KnownValue]:
    maps = _xlsx_maps(snapshot.content)
    return {
        (measure_key, period): _KnownValue(
            lexical=lexical,
            value=value,
            basis=BeaPioValueBasis.SOURCE_PUBLISHED,
            artifact_uri=snapshot.request.uri,
            content_sha256=snapshot.content_sha256,
            content_length=len(snapshot.content),
        )
        for measure_key, periods in maps.items()
        for period, (lexical, value) in periods.items()
    }


def build_bea_pio_archive_manifest(
    registry: OfficialSourceRegistryV1,
    profile: UnitedStatesBackfillProfileV1,
    release_index: BeaPioReleaseIndexV1,
    snapshots_by_uri: Mapping[str, OfficialRawSnapshotV1],
) -> BeaPioArchiveManifestV1:
    """Build the complete PIO manifest from an exact retained corpus."""
    if not isinstance(registry, OfficialSourceRegistryV1):
        raise TypeError("BEA PIO manifest requires a v1 registry")
    if not isinstance(profile, UnitedStatesBackfillProfileV1):
        raise TypeError("BEA PIO manifest requires a v1 U.S. profile")
    if not isinstance(release_index, BeaPioReleaseIndexV1):
        raise TypeError("BEA PIO manifest requires a v1 release index")
    program = profile.by_key[BEA_PIO_PROGRAM_KEY]
    if program.source_key != BEA_PIO_SOURCE_KEY:
        raise ValueError("BEA PIO profile source differs")
    expected_uris = {
        *(item.artifact_uri for item in release_index.releases),
        *(item[0] for item in BEA_PIO_SUPPORT_ARTIFACTS.values()),
        BEA_PIO_PREDECESSOR_URI,
    }
    if set(snapshots_by_uri) != expected_uris:
        raise ValueError("BEA PIO retained corpus URI set differs")
    predecessor = snapshots_by_uri[BEA_PIO_PREDECESSOR_URI]
    _validate_release_snapshot(
        predecessor,
        uri=BEA_PIO_PREDECESSOR_URI,
        release_date=BEA_PIO_PREDECESSOR_DATE,
        source_format=OfficialSourceFormat.PDF,
    )
    predecessor_maps = _current_dollar_maps(
        _pdf_text(predecessor.content),
        BEA_PIO_PREDECESSOR_REFERENCE_PERIOD,
    )
    known: dict[tuple[str, str], _KnownValue] = {}
    for measure_key in (
        BEA_PIO_PERSONAL_INCOME,
        BEA_PIO_CURRENT_DOLLAR_PCE,
    ):
        lexical, predecessor_value = predecessor_maps[measure_key][
            BEA_PIO_PREDECESSOR_REFERENCE_PERIOD
        ]
        known[(measure_key, BEA_PIO_PREDECESSOR_REFERENCE_PERIOD)] = (
            _KnownValue(
                lexical=lexical,
                value=predecessor_value,
                basis=BeaPioValueBasis.SOURCE_PUBLISHED,
                artifact_uri=BEA_PIO_PREDECESSOR_URI,
                content_sha256=predecessor.content_sha256,
                content_length=len(predecessor.content),
            )
        )
    publications: list[BeaPioArchivePublicationV1] = []
    for index_entry in release_index.releases:
        support_spec = BEA_PIO_SUPPORT_ARTIFACTS.get(index_entry.artifact_uri)
        support = (
            None if support_spec is None else snapshots_by_uri[support_spec[0]]
        )
        observation = parse_bea_pio_release(
            snapshots_by_uri[index_entry.artifact_uri],
            index_entry,
            support_snapshot=support,
        )
        before = dict(known)
        measures = tuple(
            _archive_measure(value, before) for value in observation.values
        )
        publications.append(
            BeaPioArchivePublicationV1(
                observation=observation,
                measures=measures,
            )
        )
        for measure_value in observation.values:
            known[
                (measure_value.measure_key, measure_value.reference_period)
            ] = _known_from_value(measure_value)
        if support is not None and (
            support.request.source_format is OfficialSourceFormat.XLSX
        ):
            known.update(_state_from_xlsx(support))
    values = tuple(
        measure
        for publication in publications
        for measure in publication.measures
    )
    initial_values = tuple(
        item
        for item in values
        if item.value.release_stage is EconomicReleaseStage.INITIAL
    )
    artifacts = {
        (
            publication.observation.artifact_uri,
            publication.observation.content_sha256,
            publication.observation.content_length,
        )
        for publication in publications
    }
    artifacts.add(
        (
            BEA_PIO_PREDECESSOR_URI,
            predecessor.content_sha256,
            len(predecessor.content),
        )
    )
    for publication in publications:
        observation = publication.observation
        if observation.support_uri is not None:
            artifacts.add(
                (
                    observation.support_uri,
                    cast(str, observation.support_content_sha256),
                    cast(int, observation.support_content_length),
                )
            )
    revision_counts = {
        key: sum(
            item.previous_was_revised
            for item in values
            if item.value.measure_key == key
        )
        for key in BEA_PIO_MEASURE_KEYS
    }
    era_counts = {
        era: sum(
            publication.observation.source_era == era
            for publication in publications
        )
        for era in (
            "legacy-html",
            "preformatted-html",
            "semantic-table-html",
            "streamlined-html",
            "pdf-only",
            "xlsx-backed-html",
        )
    }
    initial_counts = {
        key: sum(item.value.measure_key == key for item in initial_values)
        for key in BEA_PIO_MEASURE_KEYS
    }
    return BeaPioArchiveManifestV1(
        registry_id=registry.registry_id,
        profile_id=profile.profile_id,
        release_index=release_index,
        predecessor_uri=BEA_PIO_PREDECESSOR_URI,
        predecessor_content_sha256=predecessor.content_sha256,
        predecessor_content_length=len(predecessor.content),
        publications=tuple(publications),
        raw_artifact_count=len(artifacts),
        total_content_bytes=sum(item[2] for item in artifacts),
        initial_publication_count=sum(
            publication.observation.values[0].release_stage
            is EconomicReleaseStage.INITIAL
            for publication in publications
        ),
        revision_publication_count=sum(
            publication.observation.values[0].release_stage
            is EconomicReleaseStage.REVISION
            for publication in publications
        ),
        measure_occurrence_count=len(values),
        personal_income_initial_count=initial_counts[BEA_PIO_PERSONAL_INCOME],
        current_dollar_pce_initial_count=initial_counts[
            BEA_PIO_CURRENT_DOLLAR_PCE
        ],
        headline_price_initial_count=initial_counts[BEA_PIO_HEADLINE_PRICE],
        core_price_initial_count=initial_counts[BEA_PIO_CORE_PRICE],
        comparable_revision_count=sum(
            item.previous_was_revised for item in values
        ),
        personal_income_revision_count=revision_counts[BEA_PIO_PERSONAL_INCOME],
        current_dollar_pce_revision_count=revision_counts[
            BEA_PIO_CURRENT_DOLLAR_PCE
        ],
        headline_price_revision_count=revision_counts[BEA_PIO_HEADLINE_PRICE],
        core_price_revision_count=revision_counts[BEA_PIO_CORE_PRICE],
        derived_price_occurrence_count=sum(
            item.value.value_basis is BeaPioValueBasis.DERIVED_FROM_INDEX_LEVELS
            for item in values
        ),
        historical_update_notice_count=sum(
            publication.observation.historical_update_notice
            for publication in publications
        ),
        shutdown_publication_count=sum(
            item.publication_kind
            in {
                "shutdown-split",
                "shutdown-catch-up",
                "shutdown-combined",
                "revision-only-update",
            }
            for item in release_index.releases
        ),
        support_artifact_count=sum(
            publication.observation.support_uri is not None
            for publication in publications
        ),
        exact_minute_count=sum(
            publication.observation.time_precision
            is EconomicTimePrecision.EXACT_MINUTE
            for publication in publications
        ),
        date_only_count=sum(
            publication.observation.time_precision
            is EconomicTimePrecision.DATE_ONLY
            for publication in publications
        ),
        zone_mismatch_count=sum(
            publication.observation.zone_consistent is False
            for publication in publications
        ),
        archive_date_mismatch_count=sum(
            item.release_date != item.archive_published_date
            for item in release_index.releases
        ),
        legacy_html_count=era_counts["legacy-html"],
        preformatted_html_count=era_counts["preformatted-html"],
        semantic_table_html_count=era_counts["semantic-table-html"],
        streamlined_html_count=era_counts["streamlined-html"],
        pdf_only_count=era_counts["pdf-only"],
        xlsx_backed_html_count=era_counts["xlsx-backed-html"],
    )


def replay_bea_pio_archive(
    registry: OfficialSourceRegistryV1,
    profile: UnitedStatesBackfillProfileV1,
    index_snapshots: Sequence[OfficialRawSnapshotV1],
    snapshots_by_uri: Mapping[str, OfficialRawSnapshotV1],
    expected: BeaPioArchiveManifestV1,
) -> BeaPioArchiveManifestV1:
    """Recompute and compare the complete PIO archive receipt."""
    if not isinstance(expected, BeaPioArchiveManifestV1):
        raise TypeError("BEA PIO replay requires a v1 expected manifest")
    release_index = parse_bea_pio_release_index(
        index_snapshots,
        as_of_date=expected.release_index.as_of_date,
    )
    rebuilt = build_bea_pio_archive_manifest(
        registry, profile, release_index, snapshots_by_uri
    )
    if rebuilt != expected:
        raise ValueError("BEA PIO retained-corpus replay differs")
    return rebuilt


def bea_pio_coverage_from_manifest(
    manifest: BeaPioArchiveManifestV1,
) -> UnitedStatesProgramCoverageV1:
    """Derive the complete U.S. PIO program coverage from a manifest."""
    if not isinstance(manifest, BeaPioArchiveManifestV1):
        raise TypeError("BEA PIO coverage requires a v1 manifest")
    return UnitedStatesProgramCoverageV1(
        program_key=BEA_PIO_PROGRAM_KEY,
        window_start_date=US_BACKFILL_START_DATE,
        window_end_date=manifest.release_index.as_of_date,
        expected_occurrence_count=len(manifest.publications),
        schedule_count=len(manifest.publications),
        initial_actual_count=len(manifest.publications),
        previous_as_known_count=len(manifest.publications),
        revision_count=manifest.comparable_revision_count,
        exact_minute_count=manifest.exact_minute_count,
        forecast_count=0,
        artifact_sha256s=tuple(
            sorted(
                {
                    manifest.predecessor_content_sha256,
                    *(
                        publication.observation.content_sha256
                        for publication in manifest.publications
                    ),
                    *(
                        publication.observation.support_content_sha256
                        for publication in manifest.publications
                        if publication.observation.support_content_sha256
                        is not None
                    ),
                }
            )
        ),
        gap_reasons=(),
        notes=(
            "Each official publication remains one schedule occurrence while its four measure values retain independent revision lineage.",
            "The February 2000 archive exposes only date precision; all other publications retain exact local release minutes.",
            "Headline and core PCE changes from June 2000 through January 2002 are derived only from contemporaneous source-published index levels.",
            "No event-level historical consensus is projected from a monthly survey series.",
        ),
    )


def packaged_bea_pio_archive_manifest_path() -> Path:
    """Return the packaged PIO manifest path."""
    return Path(__file__).with_name("assets") / "us_pio_archive_v1.json"


def packaged_bea_pio_index_path() -> Path:
    """Return the packaged PIO index-envelope path."""
    return Path(__file__).with_name("assets") / "us_pio_release_index_v1.json"


def load_packaged_bea_pio_archive_manifest() -> BeaPioArchiveManifestV1:
    """Load and validate the packaged PIO archive receipt."""
    return BeaPioArchiveManifestV1.from_json(
        packaged_bea_pio_archive_manifest_path().read_text(encoding="utf-8")
    )


def load_packaged_bea_pio_index_snapshots() -> (
    tuple[OfficialRawSnapshotV1, ...]
):
    """Load exact archive pages from the packaged base64 JSON envelope."""
    payload = _mapping(
        json.loads(packaged_bea_pio_index_path().read_text(encoding="utf-8")),
        "BEA PIO index envelope",
    )
    if payload.get("schema_version") != BEA_PIO_INDEX_ENVELOPE_SCHEMA_VERSION:
        raise ValueError("unsupported BEA PIO index-envelope schema")
    captured_at_ns = cast(int, payload.get("captured_at_ns"))
    _positive_int(captured_at_ns, "captured_at_ns", 2**63 - 1)
    registry = load_packaged_official_source_registry()
    requests = build_bea_pio_index_requests(registry)
    snapshots: list[OfficialRawSnapshotV1] = []
    for request, item in zip(
        requests,
        _sequence(payload.get("pages"), "pages"),
        strict=True,
    ):
        page = _mapping(item, "BEA PIO index-envelope page")
        if (
            cast(int, page.get("page_number")) != request.page_number
            or str(page.get("uri", "")) != request.uri
        ):
            raise ValueError("BEA PIO index-envelope page identity differs")
        try:
            content = base64.b64decode(
                str(page.get("content_base64", "")), validate=True
            )
        except (TypeError, ValueError) as exc:
            raise ValueError(
                "BEA PIO index envelope has invalid base64"
            ) from exc
        snapshots.append(
            OfficialRawSnapshotV1(
                request=request,
                retrieved_at_ns=captured_at_ns,
                completed_at_ns=captured_at_ns + 1,
                status_code=200,
                resolved_uri=request.uri,
                response_headers={"Content-Type": "text/html; charset=UTF-8"},
                content=content,
                content_type="text/html",
            )
        )
    return _validate_index_snapshots(snapshots)


def load_packaged_bea_pio_index() -> BeaPioReleaseIndexV1:
    """Parse the exact packaged PIO archive-page evidence."""
    manifest = load_packaged_bea_pio_archive_manifest()
    return parse_bea_pio_release_index(
        load_packaged_bea_pio_index_snapshots(),
        as_of_date=manifest.release_index.as_of_date,
    )


__all__ = [
    "BEA_PIO_ARCHIVE_DATE_CORRECTIONS",
    "BEA_PIO_ARCHIVE_MANIFEST_SCHEMA_VERSION",
    "BEA_PIO_ARCHIVE_MEASURE_SCHEMA_VERSION",
    "BEA_PIO_ARCHIVE_PUBLICATION_SCHEMA_VERSION",
    "BEA_PIO_CORE_PRICE",
    "BEA_PIO_CURRENT_DOLLAR_PCE",
    "BEA_PIO_HEADLINE_PRICE",
    "BEA_PIO_INDEX_ENTRY_SCHEMA_VERSION",
    "BEA_PIO_INDEX_ENVELOPE_SCHEMA_VERSION",
    "BEA_PIO_INDEX_PAGE_COUNT",
    "BEA_PIO_INDEX_PAGE_SCHEMA_VERSION",
    "BEA_PIO_INDEX_SCHEMA_VERSION",
    "BEA_PIO_INDEX_URI",
    "BEA_PIO_MEASURE_KEYS",
    "BEA_PIO_MEASURE_VALUE_SCHEMA_VERSION",
    "BEA_PIO_PERSONAL_INCOME",
    "BEA_PIO_PREDECESSOR_DATE",
    "BEA_PIO_PREDECESSOR_REFERENCE_PERIOD",
    "BEA_PIO_PREDECESSOR_URI",
    "BEA_PIO_PROGRAM_KEY",
    "BEA_PIO_RELEASE_OBSERVATION_SCHEMA_VERSION",
    "BEA_PIO_SOURCE_KEY",
    "BEA_PIO_SUPPLEMENTAL_DATE",
    "BEA_PIO_SUPPLEMENTAL_TITLE",
    "BEA_PIO_SUPPLEMENTAL_URI",
    "BEA_PIO_SUPPORT_ARTIFACTS",
    "BeaPioArchiveManifestV1",
    "BeaPioArchiveMeasureV1",
    "BeaPioArchivePublicationV1",
    "BeaPioMeasureValueV1",
    "BeaPioReleaseIndexEntryV1",
    "BeaPioReleaseIndexPageV1",
    "BeaPioReleaseIndexV1",
    "BeaPioReleaseObservationV1",
    "BeaPioValueBasis",
    "bea_pio_coverage_from_manifest",
    "build_bea_pio_archive_manifest",
    "build_bea_pio_archive_requests",
    "build_bea_pio_index_requests",
    "load_packaged_bea_pio_archive_manifest",
    "load_packaged_bea_pio_index",
    "load_packaged_bea_pio_index_snapshots",
    "packaged_bea_pio_archive_manifest_path",
    "packaged_bea_pio_index_path",
    "parse_bea_pio_release",
    "parse_bea_pio_release_index",
    "replay_bea_pio_archive",
]
