"""Occurrence-specific Census Advance Monthly Retail Sales evidence.

The official MARTS archive mixes scanned and born-digital PDFs, optional XLS
and XLSX table companions, a SIC-to-NAICS boundary, annual benchmark/sample
revisions, and two government-shutdown delays.  This module retains those
facts while reconstructing total, ex-autos, ex-gasoline, and retail-control
monthly changes from the adjusted levels published in each occurrence.
"""

from __future__ import annotations

import base64
import calendar
import hashlib
import json
import math
import re
import warnings
from collections.abc import Mapping, Sequence
from dataclasses import dataclass
from datetime import date, datetime, time
from decimal import ROUND_HALF_UP, Decimal
from enum import Enum
from io import BytesIO
from pathlib import Path
from types import MappingProxyType
from typing import Any, cast
from urllib.parse import urlparse
from zipfile import ZIP_DEFLATED, ZipFile
from zoneinfo import ZoneInfo

import xlrd
from openpyxl import load_workbook
from pypdf import PdfReader

from histdatacom.market_context.contracts import canonical_contract_json
from histdatacom.market_context.official_sources import (
    OfficialRawSnapshotV1,
    OfficialRequestMethod,
    OfficialSourceFormat,
    OfficialSourceRegistryV1,
    OfficialSourceRequestV1,
    load_packaged_official_source_registry,
)
from histdatacom.market_context.us_backfill import (
    US_BACKFILL_START_DATE,
    UnitedStatesBackfillProfileV1,
    UnitedStatesProgramCoverageV1,
)
from histdatacom.runtime_contracts import JSONValue

CENSUS_RETAIL_INDEX_ENTRY_SCHEMA_VERSION = (
    "histdatacom.census-retail-index-entry.v1"
)
CENSUS_RETAIL_INDEX_SCHEMA_VERSION = "histdatacom.census-retail-index.v1"
CENSUS_RETAIL_BENCHMARK_MEASURE_SCHEMA_VERSION = (
    "histdatacom.census-retail-benchmark-measure.v1"
)
CENSUS_RETAIL_BENCHMARK_ENTRY_SCHEMA_VERSION = (
    "histdatacom.census-retail-benchmark-entry.v1"
)
CENSUS_RETAIL_BENCHMARK_INDEX_SCHEMA_VERSION = (
    "histdatacom.census-retail-benchmark-index.v1"
)
CENSUS_RETAIL_OCR_ENTRY_SCHEMA_VERSION = (
    "histdatacom.census-retail-ocr-entry.v1"
)
CENSUS_RETAIL_OCR_CORPUS_SCHEMA_VERSION = (
    "histdatacom.census-retail-ocr-corpus.v1"
)
CENSUS_RETAIL_MEASURE_SCHEMA_VERSION = "histdatacom.census-retail-measure.v1"
CENSUS_RETAIL_PUBLICATION_SCHEMA_VERSION = (
    "histdatacom.census-retail-publication.v1"
)
CENSUS_RETAIL_ARCHIVE_MANIFEST_SCHEMA_VERSION = (
    "histdatacom.census-retail-archive-manifest.v1"
)
CENSUS_RETAIL_INDEX_ENVELOPE_SCHEMA_VERSION = (
    "histdatacom.census-retail-index-envelope.v1"
)
CENSUS_RETAIL_BENCHMARK_INDEX_ENVELOPE_SCHEMA_VERSION = (
    "histdatacom.census-retail-benchmark-index-envelope.v1"
)

CENSUS_RETAIL_SOURCE_KEY = "us.census.retail-sales"
CENSUS_RETAIL_PROGRAM_KEY = CENSUS_RETAIL_SOURCE_KEY
CENSUS_RETAIL_INDEX_URI = (
    "https://www.census.gov/retail/marts/historic_releases.html"
)
CENSUS_RETAIL_BENCHMARK_INDEX_URI = (
    "https://www.census.gov/retail/mrts/historic_releases.html"
)
CENSUS_RETAIL_CURRENT_PDF_URI = (
    "https://www.census.gov/retail/marts/www/marts_current.pdf"
)
CENSUS_RETAIL_CURRENT_XLSX_URI = (
    "https://www.census.gov/retail/marts/www/marts_current.xlsx"
)
CENSUS_RETAIL_CURRENT_REFERENCE_PERIOD = "2026-07"
CENSUS_RETAIL_CURRENT_RELEASE_DATE = "2026-08-14"
CENSUS_RETAIL_PREDECESSOR_URI = (
    "https://www2.census.gov/retail/releases/historical/marts/adv9911.pdf"
)
CENSUS_RETAIL_PREDECESSOR_REFERENCE_PERIOD = "1999-11"
CENSUS_RETAIL_FIRST_REFERENCE_PERIOD = "1999-12"
CENSUS_RETAIL_NAICS_BOUNDARY = "2001-05"
CENSUS_RETAIL_DIRECT_EX_GAS_BOUNDARY = "2018-05"


def _benchmark_artifact_uri(report_year: int) -> str:
    if not 2000 <= report_year <= 2025:
        raise ValueError("Census retail benchmark report year is invalid")
    if report_year <= 2007:
        return (
            "https://www2.census.gov/retail/releases/benchmark/annpub"
            f"{(report_year - 2001) % 100:02d}.pdf"
        )
    suffix = "xlsx" if report_year >= 2023 else "xls"
    return (
        "https://www.census.gov/retail/mrts/www/benchmark/"
        f"{report_year}/excel/benchsales{report_year % 100:02d}.{suffix}"
    )


CENSUS_RETAIL_BENCHMARK_URIS = MappingProxyType(
    {year: _benchmark_artifact_uri(year) for year in range(2000, 2026)}
)
CENSUS_RETAIL_BENCHMARK_PERIOD_ENDS = MappingProxyType(
    {
        2000: "1999-12",
        2001: "2000-12",
        2002: "2002-03",
        2003: "2003-03",
        2004: "2004-02",
        2005: "2005-02",
        2006: "2006-02",
        2007: "2007-02",
        2008: "2008-03",
        2009: "2009-03",
        2010: "2010-03",
        2011: "2011-03",
        2012: "2012-03",
        2013: "2013-04",
        2014: "2014-03",
        2015: "2015-03",
        2016: "2016-03",
        2017: "2017-03",
        2018: "2018-04",
        2019: "2019-05",
        2020: "2020-03",
        2021: "2021-03",
        2022: "2022-03",
        2023: "2023-03",
        2024: "2024-03",
        2025: "2025-03",
    }
)
# The 2023 workbook's terminal header says "Mar. 2022(a)" on its 2023 sheet;
# its title, values, surrounding columns, and archive placement bind it to 2023.
# Raw bytes remain unchanged and only this exact source-authored typo is accepted.
_BENCHMARK_TERMINAL_HEADER_TYPOS = MappingProxyType({2023: "mar 2022 a"})

CENSUS_RETAIL_TOTAL = "retail-and-food-services-total"
CENSUS_RETAIL_EX_AUTOS = "retail-and-food-services-excluding-autos"
CENSUS_RETAIL_EX_GAS = "retail-and-food-services-excluding-gasoline"
CENSUS_RETAIL_CONTROL_GROUP = "retail-control-group"
CENSUS_RETAIL_MEASURE_KEYS = (
    CENSUS_RETAIL_TOTAL,
    CENSUS_RETAIL_EX_AUTOS,
    CENSUS_RETAIL_EX_GAS,
    CENSUS_RETAIL_CONTROL_GROUP,
)

CENSUS_RETAIL_SUPPORT_LINK_CORRECTIONS = MappingProxyType(
    {
        (
            "https://www2.census.gov/retail/releases/historical/marts/"
            f"adv20{month:02d}.xls"
        ): (
            "https://www2.census.gov/retail/releases/historical/marts/"
            f"rs20{month:02d}.xls"
        )
        for month in range(7, 13)
    }
)

# Page-aware Tesseract output is hash-bound before these narrowly reviewed
# cell corrections are applied.  A changed official PDF therefore fails
# closed instead of inheriting an obsolete OCR correction.
_OCR_LEVEL_OVERRIDES = MappingProxyType(
    {
        (
            "f4b70a61bbe58d6e78095155f5ec5a957f99f4980c0a1beb63056b96ce7c6ac3",
            "food",
        ): (25_478, 25_326, 25_369),
        (
            "c7126d216ea7cb59edd9e8780c0d84d46b7b84059ae2c76eb1ffa48ab97ce454",
            "auto",
        ): (72_807, 73_185, 72_647),
        (
            "0c0c53aded49cba04da14b3f58681f9af682f5276534438741504b32abd02f41",
            "gas",
        ): (28_039, 27_773, 26_386),
        (
            "9ec0e9827080759e7422f13e1f2dd57b78661992546bc5ba7a8611ff1cb65e43",
            "gas",
        ): (29_963, 30_439, 29_791),
        (
            "6ae7aa09f1850a67f6f7321397b5ba79367f6639286328518388bf0191c54bd1",
            "gas",
        ): (34_152, 32_728, 31_471),
        (
            "91f3094ba2e823ebb7bda9754f9d65d044d8c6455c5a9244746d309f2bfb6545",
            "food",
        ): (25_841, 25_577, 25_698),
        (
            "0c0c53aded49cba04da14b3f58681f9af682f5276534438741504b32abd02f41",
            "food",
        ): (33_325, 33_265, 32_803),
        (
            "3fd1505c5bbcbdca30636eb5abdeb287617942f7e0c5aa9f6b58a7931612840e",
            "gas",
        ): (18_689, 18_642, 18_549),
    }
)

MAX_CENSUS_RETAIL_RELEASES = 512
MAX_CENSUS_RETAIL_MEASURES = 4096
MAX_CENSUS_RETAIL_INDEX_BYTES = 4 * 1024 * 1024
MAX_CENSUS_RETAIL_ARTIFACT_BYTES = 32 * 1024 * 1024
MAX_CENSUS_RETAIL_OCR_TEXT_BYTES = 2 * 1024 * 1024

_SHA256_RE = re.compile(r"^[0-9a-f]{64}$")
_YEAR_MONTH_RE = re.compile(r"^(?P<year>\d{4})-(?P<month>0[1-9]|1[0-2])$")
_ARCHIVE_PDF_RE = re.compile(
    r"https://www2\.census\.gov/retail/releases/historical/marts/"
    r"adv(?P<year>\d{2})(?P<month>\d{2})\.pdf$",
    re.IGNORECASE,
)
_ARCHIVE_SUPPORT_RE = re.compile(
    r"https://www2\.census\.gov/retail/releases/historical/marts/"
    r"(?P<prefix>rs|adv)(?P<year>\d{2})(?P<month>\d{2})\."
    r"(?P<suffix>xls|xlsx)$",
    re.IGNORECASE,
)
_HREF_RE = re.compile(
    r"\bhref=[\"'](?P<uri>https?://[^\"']+)[\"']", re.IGNORECASE
)
_THOUSANDS_RE = re.compile(
    r"(?<![\d.])(?:\d{1,3}(?:[,.]\d{3})+|\d{4,7})(?![\d.])"
)
_SPACED_THOUSANDS_RE = re.compile(
    r"(?<!\d)(?P<head>\d(?:\s*\d){0,2})\s*,\s*" r"(?P<tail>\d\s*\d\s*\d)"
)
_MONTH_NUMBERS = {
    name.lower(): number
    for number, name in enumerate(calendar.month_name)
    if name
}
_LONG_DATE_FRAGMENT = r"[A-Z]+\s+\d{1,2}\s*,?\s*\d{4}"
_WEEKDAY_FRAGMENT = r"(?:MONDAY|TUESDAY|WEDNESDAY|THURSDAY|FRIDAY)"
_TIME_THEN_DATE_RE = re.compile(
    r"8\s*:\s*30\s*(?:A\.?\s*M\.?|AM)\s*"
    r"(?P<zone>E\s*(?:D\s*T|S\s*T|T))\s*,?\s*"
    rf"(?:{_WEEKDAY_FRAGMENT}\s*,?\s*)?"
    rf"(?P<date>{_LONG_DATE_FRAGMENT})",
    re.IGNORECASE,
)
_DATE_THEN_TIME_RE = re.compile(
    rf"(?:{_WEEKDAY_FRAGMENT}\s*,?\s*)?"
    rf"(?P<date>{_LONG_DATE_FRAGMENT})\s*,?\s*"
    r"(?:,?\s*AT\s*)8\s*:\s*30\s*(?:A\.?\s*M\.?|AM)\s*"
    r"(?P<zone>E\s*(?:D\s*T|S\s*T|T))",
    re.IGNORECASE,
)
_ROW_PATTERNS = MappingProxyType(
    {
        "exauto": re.compile(
            r"total\s*[^\n]{0,10}\s*excl[^\n]{0,60}"
            r"(?:auto dealers|motor vehic(?:le|ie)\s*&?\s*parts)",
            re.IGNORECASE,
        ),
        "auto": re.compile(
            r"(?:automotive dealers|motor vehicle\s*&?\s*parts dealers)",
            re.IGNORECASE,
        ),
        "building": re.compile(
            r"building\s+mat[^\n]{0,150}"
            r"(?:mobile home dealers|suppl(?:y|ies) dealers)",
            re.IGNORECASE,
        ),
        "gas": re.compile(r"gasoline\s+(?:service\s+)?stations", re.IGNORECASE),
        "food": re.compile(
            r"(?:eating and (?:drinking|crinking)\s+places|"
            r"food services\s*&?\s*drinking\s+places)",
            re.IGNORECASE,
        ),
    }
)
_LEVEL_RANGES = MappingProxyType(
    {
        "exauto": (100_000, 800_000),
        "auto": (30_000, 200_000),
        "building": (5_000, 100_000),
        "gas": (5_000, 100_000),
        "food": (10_000, 150_000),
    }
)
_REVISION_NOTICE_RE = re.compile(
    r"notice of (?:sample )?revision|"
    r"monthly retail sales estimates were revised",
    re.IGNORECASE,
)
_BENCHMARK_PDF_ROW_PATTERNS = MappingProxyType(
    {
        "total": re.compile(
            r"R\s*e\s*t\s*a\s*i\s*l\s*a\s*n\s*d\s*"
            r"f\s*o\s*o\s*d\s*s\s*e\s*r\s*v\s*i\s*c\s*e\s*s\s*"
            r"s\s*a\s*l\s*e\s*s\s*,?\s*t\s*o\s*t\s*a\s*l",
            re.IGNORECASE,
        ),
        "exauto": re.compile(
            r"Total \(excl\. motor vehicle and parts dealers\)", re.IGNORECASE
        ),
        "auto": re.compile(
            r"^\s*441\s+Motor vehicle and parts dealers", re.IGNORECASE
        ),
        "building": re.compile(
            r"^\s*444\s+Building mat[^\n]{0,100}supplies dealers", re.IGNORECASE
        ),
        "gas": re.compile(r"^\s*447\s+Gasoline stations", re.IGNORECASE),
        "food": re.compile(
            r"^\s*722\s+Food services and drinking places", re.IGNORECASE
        ),
    }
)


class CensusRetailLevelBasis(str, Enum):
    """How the adjusted level underlying a percentage was obtained."""

    SOURCE_PUBLISHED = "source-published-adjusted-level"
    DERIVED_AGGREGATE = "derived-aggregate-adjusted-level"

    @classmethod
    def from_value(
        cls, value: str | CensusRetailLevelBasis
    ) -> CensusRetailLevelBasis:
        if isinstance(value, cls):
            return value
        try:
            return cls(str(value))
        except ValueError as exc:
            raise ValueError("unsupported Census retail level basis") from exc


class CensusRetailBenchmarkDisposition(str, Enum):
    """Whether an annual report can replace the next release's prior state."""

    RETAINED_NON_OVERLAPPING = "retained-non-overlapping"
    PRIOR_STATE_OVERRIDE = "prior-state-override"

    @classmethod
    def from_value(
        cls, value: str | CensusRetailBenchmarkDisposition
    ) -> CensusRetailBenchmarkDisposition:
        if isinstance(value, cls):
            return value
        try:
            return cls(str(value))
        except ValueError as exc:
            raise ValueError(
                "unsupported Census retail benchmark disposition"
            ) from exc


def _required_text(value: object, name: str, *, maximum: int = 8192) -> str:
    result = str(value).strip()
    if not result or len(result.encode("utf-8")) > maximum:
        raise ValueError(f"{name} is invalid")
    return result


def _optional_text(value: object | None, name: str) -> str | None:
    if value is None:
        return None
    return _required_text(value, name)


def _iso_date(value: object, name: str) -> str:
    result = _required_text(value, name)
    try:
        parsed = date.fromisoformat(result)
    except ValueError as exc:
        raise ValueError(f"{name} must be an ISO date") from exc
    if parsed.isoformat() != result:
        raise ValueError(f"{name} must be a canonical ISO date")
    return result


def _year_month(value: object, name: str) -> str:
    result = _required_text(value, name)
    if _YEAR_MONTH_RE.fullmatch(result) is None:
        raise ValueError(f"{name} must be a canonical year-month")
    return result


def _previous_month(value: str) -> str:
    match = _YEAR_MONTH_RE.fullmatch(_year_month(value, "reference_period"))
    assert match is not None
    year = int(match.group("year"))
    month = int(match.group("month"))
    if month == 1:
        return f"{year - 1:04d}-12"
    return f"{year:04d}-{month - 1:02d}"


def _next_month(value: str) -> str:
    match = _YEAR_MONTH_RE.fullmatch(_year_month(value, "reference_period"))
    assert match is not None
    year = int(match.group("year"))
    month = int(match.group("month"))
    if month == 12:
        return f"{year + 1:04d}-01"
    return f"{year:04d}-{month + 1:02d}"


def _https_uri(value: object, name: str) -> str:
    result = _required_text(value, name)
    parsed = urlparse(result)
    if (
        parsed.scheme != "https"
        or parsed.hostname not in {"www.census.gov", "www2.census.gov"}
        or parsed.username is not None
        or parsed.password is not None
        or parsed.fragment
    ):
        raise ValueError(f"{name} is not an official Census HTTPS URI")
    return result


def _sha256(value: object, name: str) -> str:
    result = _required_text(value, name)
    if _SHA256_RE.fullmatch(result) is None:
        raise ValueError(f"{name} must be a lowercase SHA-256")
    return result


def _positive_int(value: object, name: str, maximum: int) -> int:
    if (
        not isinstance(value, int)
        or isinstance(value, bool)
        or not 1 <= value <= maximum
    ):
        raise ValueError(f"{name} is outside its bound")
    return value


def _nonnegative_int(value: object, name: str, maximum: int) -> int:
    if (
        not isinstance(value, int)
        or isinstance(value, bool)
        or not 0 <= value <= maximum
    ):
        raise ValueError(f"{name} is outside its bound")
    return value


def _finite(value: object, name: str) -> float:
    if not isinstance(value, (int, float)) or isinstance(value, bool):
        raise TypeError(f"{name} must be numeric")
    result = float(value)
    if not math.isfinite(result):
        raise ValueError(f"{name} must be finite")
    return result


def _mapping(value: object, name: str) -> Mapping[str, Any]:
    if not isinstance(value, Mapping):
        raise TypeError(f"{name} must be a mapping")
    return cast(Mapping[str, Any], value)


def _sequence(value: object, name: str) -> Sequence[Any]:
    if not isinstance(value, Sequence) or isinstance(value, (str, bytes)):
        raise TypeError(f"{name} must be a sequence")
    return value


def _stable_id(prefix: str, payload: Mapping[str, JSONValue]) -> str:
    digest = hashlib.sha256(
        canonical_contract_json(payload).encode("utf-8")
    ).hexdigest()
    return f"{prefix}:sha256:{digest}"


def _period_from_suffix(year: str, month: str) -> str:
    short_year = int(year)
    full_year = 1900 + short_year if short_year >= 90 else 2000 + short_year
    month_number = int(month)
    if not 1 <= month_number <= 12:
        raise ValueError("Census retail artifact month is invalid")
    return f"{full_year:04d}-{month_number:02d}"


def _archive_pdf_uri(reference_period: str) -> str:
    match = _YEAR_MONTH_RE.fullmatch(_year_month(reference_period, "period"))
    assert match is not None
    return (
        "https://www2.census.gov/retail/releases/historical/marts/adv"
        f"{int(match.group('year')) % 100:02d}{match.group('month')}.pdf"
    )


def _support_format(uri: str) -> OfficialSourceFormat:
    suffix = Path(urlparse(uri).path).suffix.lower()
    if suffix == ".xls":
        return OfficialSourceFormat.XLS
    if suffix == ".xlsx":
        return OfficialSourceFormat.XLSX
    raise ValueError("Census retail support URI has an unsupported suffix")


@dataclass(frozen=True, slots=True)
class CensusRetailReleaseIndexEntryV1:
    """One PDF publication and its optional official table companion."""

    reference_period: str
    artifact_uri: str
    support_uri: str | None
    support_format: OfficialSourceFormat | None
    archive_listed: bool
    support_link_corrected: bool
    entry_id: str = ""
    schema_version: str = CENSUS_RETAIL_INDEX_ENTRY_SCHEMA_VERSION

    def __post_init__(self) -> None:
        if self.schema_version != CENSUS_RETAIL_INDEX_ENTRY_SCHEMA_VERSION:
            raise ValueError("unsupported Census retail index-entry schema")
        period = _year_month(self.reference_period, "reference_period")
        artifact = _https_uri(self.artifact_uri, "artifact_uri")
        support = _optional_text(self.support_uri, "support_uri")
        support_format = self.support_format
        if support is None:
            if support_format is not None or self.support_link_corrected:
                raise ValueError("Census retail support metadata is incomplete")
        else:
            support = _https_uri(support, "support_uri")
            if support_format is None:
                raise ValueError("Census retail support format is missing")
            support_format = OfficialSourceFormat.from_value(support_format)
            if support_format is not _support_format(support):
                raise ValueError(
                    "Census retail support format differs from URI"
                )
        expected_artifact = (
            CENSUS_RETAIL_CURRENT_PDF_URI
            if not self.archive_listed
            else _archive_pdf_uri(period)
        )
        if artifact != expected_artifact:
            raise ValueError("Census retail artifact URI differs from period")
        if not isinstance(self.archive_listed, bool) or not isinstance(
            self.support_link_corrected, bool
        ):
            raise TypeError("Census retail index flags must be boolean")
        object.__setattr__(self, "reference_period", period)
        object.__setattr__(self, "artifact_uri", artifact)
        object.__setattr__(self, "support_uri", support)
        object.__setattr__(self, "support_format", support_format)
        expected = _stable_id(
            "census-retail-index-entry", self.identity_payload()
        )
        if self.entry_id and self.entry_id != expected:
            raise ValueError("Census retail index-entry identity differs")
        object.__setattr__(self, "entry_id", expected)

    def identity_payload(self) -> dict[str, JSONValue]:
        return {
            "schema_version": self.schema_version,
            "reference_period": self.reference_period,
            "artifact_uri": self.artifact_uri,
            "support_uri": self.support_uri,
            "support_format": (
                None
                if self.support_format is None
                else self.support_format.value
            ),
            "archive_listed": self.archive_listed,
            "support_link_corrected": self.support_link_corrected,
        }

    def to_dict(self) -> dict[str, JSONValue]:
        return {**self.identity_payload(), "entry_id": self.entry_id}

    @classmethod
    def from_dict(
        cls, data: Mapping[str, Any]
    ) -> CensusRetailReleaseIndexEntryV1:
        raw_format = data.get("support_format")
        return cls(
            reference_period=str(data.get("reference_period", "")),
            artifact_uri=str(data.get("artifact_uri", "")),
            support_uri=(
                None
                if data.get("support_uri") is None
                else str(data.get("support_uri"))
            ),
            support_format=(
                None
                if raw_format is None
                else OfficialSourceFormat.from_value(str(raw_format))
            ),
            archive_listed=cast(bool, data.get("archive_listed")),
            support_link_corrected=cast(
                bool, data.get("support_link_corrected")
            ),
            entry_id=str(data.get("entry_id", "")),
            schema_version=str(data.get("schema_version", "")),
        )


@dataclass(frozen=True, slots=True)
class CensusRetailReleaseIndexV1:
    """Hash-bound official archive enumeration as observed on one date."""

    source_uri: str
    content_sha256: str
    content_length: int
    as_of_date: str
    releases: tuple[CensusRetailReleaseIndexEntryV1, ...]
    index_id: str = ""
    schema_version: str = CENSUS_RETAIL_INDEX_SCHEMA_VERSION

    def __post_init__(self) -> None:
        if self.schema_version != CENSUS_RETAIL_INDEX_SCHEMA_VERSION:
            raise ValueError("unsupported Census retail index schema")
        source_uri = _https_uri(self.source_uri, "source_uri")
        if source_uri != CENSUS_RETAIL_INDEX_URI:
            raise ValueError("Census retail index source URI differs")
        content_sha = _sha256(self.content_sha256, "content_sha256")
        _positive_int(
            self.content_length,
            "content_length",
            MAX_CENSUS_RETAIL_INDEX_BYTES,
        )
        as_of = _iso_date(self.as_of_date, "as_of_date")
        releases = tuple(self.releases)
        if (
            not releases
            or len(releases) > MAX_CENSUS_RETAIL_RELEASES
            or any(
                not isinstance(item, CensusRetailReleaseIndexEntryV1)
                for item in releases
            )
        ):
            raise TypeError("Census retail index releases are invalid")
        periods = tuple(item.reference_period for item in releases)
        if periods != tuple(sorted(set(periods))):
            raise ValueError("Census retail periods must be unique and sorted")
        expected_periods: list[str] = []
        period = CENSUS_RETAIL_FIRST_REFERENCE_PERIOD
        while period <= periods[-1]:
            expected_periods.append(period)
            period = _next_month(period)
        if periods != tuple(expected_periods):
            raise ValueError("Census retail index has a reference-period gap")
        if releases[-1].archive_listed:
            raise ValueError("Census retail current supplement is missing")
        if CENSUS_RETAIL_CURRENT_RELEASE_DATE > as_of:
            raise ValueError(
                "Census retail supplement postdates as-of boundary"
            )
        object.__setattr__(self, "source_uri", source_uri)
        object.__setattr__(self, "content_sha256", content_sha)
        object.__setattr__(self, "as_of_date", as_of)
        object.__setattr__(self, "releases", releases)
        expected = _stable_id("census-retail-index", self.identity_payload())
        if self.index_id and self.index_id != expected:
            raise ValueError("Census retail index identity differs")
        object.__setattr__(self, "index_id", expected)

    @property
    def by_period(self) -> Mapping[str, CensusRetailReleaseIndexEntryV1]:
        return MappingProxyType(
            {item.reference_period: item for item in self.releases}
        )

    def identity_payload(self) -> dict[str, JSONValue]:
        return {
            "schema_version": self.schema_version,
            "source_uri": self.source_uri,
            "content_sha256": self.content_sha256,
            "content_length": self.content_length,
            "as_of_date": self.as_of_date,
            "releases": [item.to_dict() for item in self.releases],
        }

    def to_dict(self) -> dict[str, JSONValue]:
        return {**self.identity_payload(), "index_id": self.index_id}

    @classmethod
    def from_dict(cls, data: Mapping[str, Any]) -> CensusRetailReleaseIndexV1:
        return cls(
            source_uri=str(data.get("source_uri", "")),
            content_sha256=str(data.get("content_sha256", "")),
            content_length=cast(int, data.get("content_length")),
            as_of_date=str(data.get("as_of_date", "")),
            releases=tuple(
                CensusRetailReleaseIndexEntryV1.from_dict(
                    _mapping(item, "Census retail index entry")
                )
                for item in _sequence(data.get("releases"), "releases")
            ),
            index_id=str(data.get("index_id", "")),
            schema_version=str(data.get("schema_version", "")),
        )


@dataclass(frozen=True, slots=True)
class CensusRetailBenchmarkMeasureV1:
    """One annual-revision prior-month rate reconstructed from adjusted levels."""

    measure_key: str
    previous_level_millions: int
    current_level_millions: int
    value: float
    lexical: str
    measure_id: str = ""
    schema_version: str = CENSUS_RETAIL_BENCHMARK_MEASURE_SCHEMA_VERSION

    def __post_init__(self) -> None:
        if (
            self.schema_version
            != CENSUS_RETAIL_BENCHMARK_MEASURE_SCHEMA_VERSION
        ):
            raise ValueError(
                "unsupported Census retail benchmark-measure schema"
            )
        key = _required_text(self.measure_key, "measure_key")
        if key not in CENSUS_RETAIL_MEASURE_KEYS:
            raise ValueError("Census retail benchmark measure key is invalid")
        previous = _positive_int(
            self.previous_level_millions,
            "previous_level_millions",
            10_000_000,
        )
        current = _positive_int(
            self.current_level_millions,
            "current_level_millions",
            10_000_000,
        )
        lexical, numeric = _percent_change(current, previous)
        if self.lexical != lexical or not math.isclose(
            self.value, numeric, rel_tol=0.0, abs_tol=1e-15
        ):
            raise ValueError("Census retail benchmark rate differs from levels")
        object.__setattr__(self, "measure_key", key)
        expected = _stable_id(
            "census-retail-benchmark-measure", self.identity_payload()
        )
        if self.measure_id and self.measure_id != expected:
            raise ValueError("Census retail benchmark-measure identity differs")
        object.__setattr__(self, "measure_id", expected)

    def identity_payload(self) -> dict[str, JSONValue]:
        return {
            "schema_version": self.schema_version,
            "measure_key": self.measure_key,
            "previous_level_millions": self.previous_level_millions,
            "current_level_millions": self.current_level_millions,
            "value": self.value,
            "lexical": self.lexical,
        }

    def to_dict(self) -> dict[str, JSONValue]:
        return {**self.identity_payload(), "measure_id": self.measure_id}

    @classmethod
    def from_dict(
        cls, data: Mapping[str, Any]
    ) -> CensusRetailBenchmarkMeasureV1:
        return cls(
            measure_key=str(data.get("measure_key", "")),
            previous_level_millions=cast(
                int, data.get("previous_level_millions")
            ),
            current_level_millions=cast(
                int, data.get("current_level_millions")
            ),
            value=cast(float, data.get("value")),
            lexical=str(data.get("lexical", "")),
            measure_id=str(data.get("measure_id", "")),
            schema_version=str(data.get("schema_version", "")),
        )


@dataclass(frozen=True, slots=True)
class CensusRetailBenchmarkEntryV1:
    """One official annual benchmark artifact and its usable terminal state."""

    report_year: int
    report_period_end: str
    applies_to_reference_period: str | None
    disposition: CensusRetailBenchmarkDisposition
    artifact_uri: str
    source_format: OfficialSourceFormat
    content_sha256: str
    content_length: int
    measures: tuple[CensusRetailBenchmarkMeasureV1, ...]
    entry_id: str = ""
    schema_version: str = CENSUS_RETAIL_BENCHMARK_ENTRY_SCHEMA_VERSION

    def __post_init__(self) -> None:
        if self.schema_version != CENSUS_RETAIL_BENCHMARK_ENTRY_SCHEMA_VERSION:
            raise ValueError("unsupported Census retail benchmark-entry schema")
        if (
            isinstance(self.report_year, bool)
            or not isinstance(self.report_year, int)
            or self.report_year not in CENSUS_RETAIL_BENCHMARK_URIS
        ):
            raise ValueError("Census retail benchmark report year is invalid")
        period_end = _year_month(self.report_period_end, "report_period_end")
        if period_end != CENSUS_RETAIL_BENCHMARK_PERIOD_ENDS[self.report_year]:
            raise ValueError("Census retail benchmark period end differs")
        disposition = CensusRetailBenchmarkDisposition.from_value(
            self.disposition
        )
        applies_to = _optional_text(
            self.applies_to_reference_period, "applies_to_reference_period"
        )
        if applies_to is not None:
            applies_to = _year_month(applies_to, "applies_to_reference_period")
        expected_applies_to = (
            None if self.report_year < 2002 else _next_month(period_end)
        )
        expected_disposition = (
            CensusRetailBenchmarkDisposition.RETAINED_NON_OVERLAPPING
            if expected_applies_to is None
            else CensusRetailBenchmarkDisposition.PRIOR_STATE_OVERRIDE
        )
        if (
            applies_to != expected_applies_to
            or disposition is not expected_disposition
        ):
            raise ValueError("Census retail benchmark application differs")
        uri = _https_uri(self.artifact_uri, "artifact_uri")
        if uri != CENSUS_RETAIL_BENCHMARK_URIS[self.report_year]:
            raise ValueError("Census retail benchmark artifact URI differs")
        source_format = OfficialSourceFormat.from_value(self.source_format)
        expected_format = (
            OfficialSourceFormat.PDF
            if self.report_year <= 2007
            else (
                OfficialSourceFormat.XLSX
                if self.report_year >= 2023
                else OfficialSourceFormat.XLS
            )
        )
        if source_format is not expected_format:
            raise ValueError("Census retail benchmark source format differs")
        content_sha = _sha256(self.content_sha256, "content_sha256")
        _positive_int(
            self.content_length,
            "content_length",
            MAX_CENSUS_RETAIL_ARTIFACT_BYTES,
        )
        measures = tuple(self.measures)
        if self.report_year < 2002:
            if measures:
                raise ValueError(
                    "non-overlapping Census retail benchmark has measures"
                )
        elif (
            len(measures) != len(CENSUS_RETAIL_MEASURE_KEYS)
            or any(
                not isinstance(item, CensusRetailBenchmarkMeasureV1)
                for item in measures
            )
            or tuple(item.measure_key for item in measures)
            != CENSUS_RETAIL_MEASURE_KEYS
        ):
            raise ValueError("Census retail benchmark measures differ")
        object.__setattr__(self, "report_period_end", period_end)
        object.__setattr__(self, "applies_to_reference_period", applies_to)
        object.__setattr__(self, "disposition", disposition)
        object.__setattr__(self, "artifact_uri", uri)
        object.__setattr__(self, "source_format", source_format)
        object.__setattr__(self, "content_sha256", content_sha)
        object.__setattr__(self, "measures", measures)
        expected = _stable_id(
            "census-retail-benchmark-entry", self.identity_payload()
        )
        if self.entry_id and self.entry_id != expected:
            raise ValueError("Census retail benchmark-entry identity differs")
        object.__setattr__(self, "entry_id", expected)

    @property
    def by_measure(self) -> Mapping[str, CensusRetailBenchmarkMeasureV1]:
        return MappingProxyType(
            {item.measure_key: item for item in self.measures}
        )

    def identity_payload(self) -> dict[str, JSONValue]:
        return {
            "schema_version": self.schema_version,
            "report_year": self.report_year,
            "report_period_end": self.report_period_end,
            "applies_to_reference_period": self.applies_to_reference_period,
            "disposition": self.disposition.value,
            "artifact_uri": self.artifact_uri,
            "source_format": self.source_format.value,
            "content_sha256": self.content_sha256,
            "content_length": self.content_length,
            "measures": [item.to_dict() for item in self.measures],
        }

    def to_dict(self) -> dict[str, JSONValue]:
        return {**self.identity_payload(), "entry_id": self.entry_id}

    @classmethod
    def from_dict(cls, data: Mapping[str, Any]) -> CensusRetailBenchmarkEntryV1:
        return cls(
            report_year=cast(int, data.get("report_year")),
            report_period_end=str(data.get("report_period_end", "")),
            applies_to_reference_period=(
                None
                if data.get("applies_to_reference_period") is None
                else str(data.get("applies_to_reference_period"))
            ),
            disposition=CensusRetailBenchmarkDisposition.from_value(
                str(data.get("disposition", ""))
            ),
            artifact_uri=str(data.get("artifact_uri", "")),
            source_format=OfficialSourceFormat.from_value(
                str(data.get("source_format", ""))
            ),
            content_sha256=str(data.get("content_sha256", "")),
            content_length=cast(int, data.get("content_length")),
            measures=tuple(
                CensusRetailBenchmarkMeasureV1.from_dict(
                    _mapping(item, "Census retail benchmark measure")
                )
                for item in _sequence(data.get("measures"), "measures")
            ),
            entry_id=str(data.get("entry_id", "")),
            schema_version=str(data.get("schema_version", "")),
        )


@dataclass(frozen=True, slots=True)
class CensusRetailBenchmarkIndexV1:
    """Hash-bound annual-revision index plus all retained benchmark artifacts."""

    source_uri: str
    content_sha256: str
    content_length: int
    as_of_date: str
    entries: tuple[CensusRetailBenchmarkEntryV1, ...]
    index_id: str = ""
    schema_version: str = CENSUS_RETAIL_BENCHMARK_INDEX_SCHEMA_VERSION

    def __post_init__(self) -> None:
        if self.schema_version != CENSUS_RETAIL_BENCHMARK_INDEX_SCHEMA_VERSION:
            raise ValueError("unsupported Census retail benchmark-index schema")
        source_uri = _https_uri(self.source_uri, "source_uri")
        if source_uri != CENSUS_RETAIL_BENCHMARK_INDEX_URI:
            raise ValueError("Census retail benchmark-index URI differs")
        content_sha = _sha256(self.content_sha256, "content_sha256")
        _positive_int(
            self.content_length, "content_length", MAX_CENSUS_RETAIL_INDEX_BYTES
        )
        as_of = _iso_date(self.as_of_date, "as_of_date")
        entries = tuple(self.entries)
        if any(
            not isinstance(item, CensusRetailBenchmarkEntryV1)
            for item in entries
        ) or tuple(item.report_year for item in entries) != tuple(
            range(2000, 2026)
        ):
            raise ValueError("Census retail benchmark entries differ")
        object.__setattr__(self, "source_uri", source_uri)
        object.__setattr__(self, "content_sha256", content_sha)
        object.__setattr__(self, "as_of_date", as_of)
        object.__setattr__(self, "entries", entries)
        expected = _stable_id(
            "census-retail-benchmark-index", self.identity_payload()
        )
        if self.index_id and self.index_id != expected:
            raise ValueError("Census retail benchmark-index identity differs")
        object.__setattr__(self, "index_id", expected)

    @property
    def by_application_period(
        self,
    ) -> Mapping[str, CensusRetailBenchmarkEntryV1]:
        return MappingProxyType(
            {
                item.applies_to_reference_period: item
                for item in self.entries
                if item.applies_to_reference_period is not None
            }
        )

    def identity_payload(self) -> dict[str, JSONValue]:
        return {
            "schema_version": self.schema_version,
            "source_uri": self.source_uri,
            "content_sha256": self.content_sha256,
            "content_length": self.content_length,
            "as_of_date": self.as_of_date,
            "entries": [item.to_dict() for item in self.entries],
        }

    def to_dict(self) -> dict[str, JSONValue]:
        return {**self.identity_payload(), "index_id": self.index_id}

    @classmethod
    def from_dict(cls, data: Mapping[str, Any]) -> CensusRetailBenchmarkIndexV1:
        return cls(
            source_uri=str(data.get("source_uri", "")),
            content_sha256=str(data.get("content_sha256", "")),
            content_length=cast(int, data.get("content_length")),
            as_of_date=str(data.get("as_of_date", "")),
            entries=tuple(
                CensusRetailBenchmarkEntryV1.from_dict(
                    _mapping(item, "Census retail benchmark entry")
                )
                for item in _sequence(data.get("entries"), "entries")
            ),
            index_id=str(data.get("index_id", "")),
            schema_version=str(data.get("schema_version", "")),
        )


@dataclass(frozen=True, slots=True)
class CensusRetailOcrEntryV1:
    """Reviewed page-aware OCR bound to one image-only official PDF."""

    artifact_uri: str
    content_sha256: str
    page_texts: tuple[str, ...]
    entry_id: str = ""
    schema_version: str = CENSUS_RETAIL_OCR_ENTRY_SCHEMA_VERSION

    def __post_init__(self) -> None:
        if self.schema_version != CENSUS_RETAIL_OCR_ENTRY_SCHEMA_VERSION:
            raise ValueError("unsupported Census retail OCR-entry schema")
        uri = _https_uri(self.artifact_uri, "artifact_uri")
        sha = _sha256(self.content_sha256, "content_sha256")
        pages = tuple(
            _required_text(
                item,
                "page_texts",
                maximum=MAX_CENSUS_RETAIL_OCR_TEXT_BYTES,
            )
            for item in self.page_texts
        )
        if len(pages) != 3:
            raise ValueError("Census retail OCR requires pages 1 through 3")
        if sum(len(item.encode("utf-8")) for item in pages) > (
            MAX_CENSUS_RETAIL_OCR_TEXT_BYTES
        ):
            raise ValueError("Census retail OCR text exceeds its bound")
        object.__setattr__(self, "artifact_uri", uri)
        object.__setattr__(self, "content_sha256", sha)
        object.__setattr__(self, "page_texts", pages)
        expected = _stable_id(
            "census-retail-ocr-entry", self.identity_payload()
        )
        if self.entry_id and self.entry_id != expected:
            raise ValueError("Census retail OCR-entry identity differs")
        object.__setattr__(self, "entry_id", expected)

    def identity_payload(self) -> dict[str, JSONValue]:
        return {
            "schema_version": self.schema_version,
            "artifact_uri": self.artifact_uri,
            "content_sha256": self.content_sha256,
            "page_texts": list(self.page_texts),
        }

    def to_dict(self) -> dict[str, JSONValue]:
        return {**self.identity_payload(), "entry_id": self.entry_id}

    @classmethod
    def from_dict(cls, data: Mapping[str, Any]) -> CensusRetailOcrEntryV1:
        return cls(
            artifact_uri=str(data.get("artifact_uri", "")),
            content_sha256=str(data.get("content_sha256", "")),
            page_texts=tuple(
                str(item)
                for item in _sequence(data.get("page_texts"), "page_texts")
            ),
            entry_id=str(data.get("entry_id", "")),
            schema_version=str(data.get("schema_version", "")),
        )


@dataclass(frozen=True, slots=True)
class CensusRetailOcrCorpusV1:
    """Deterministic reviewed OCR configuration and evidence corpus."""

    engine: str
    engine_version: str
    render_dpi: int
    page_modes: tuple[int, ...]
    entries: tuple[CensusRetailOcrEntryV1, ...]
    corpus_id: str = ""
    schema_version: str = CENSUS_RETAIL_OCR_CORPUS_SCHEMA_VERSION

    def __post_init__(self) -> None:
        if self.schema_version != CENSUS_RETAIL_OCR_CORPUS_SCHEMA_VERSION:
            raise ValueError("unsupported Census retail OCR-corpus schema")
        engine = _required_text(self.engine, "engine")
        version = _required_text(self.engine_version, "engine_version")
        if engine != "tesseract":
            raise ValueError("Census retail OCR engine differs")
        if self.render_dpi != 300 or self.page_modes != (3, 6, 6):
            raise ValueError("Census retail OCR settings differ")
        entries = tuple(self.entries)
        if (
            not entries
            or len(entries) > MAX_CENSUS_RETAIL_RELEASES
            or any(
                not isinstance(item, CensusRetailOcrEntryV1) for item in entries
            )
        ):
            raise TypeError("Census retail OCR entries are invalid")
        uris = tuple(item.artifact_uri for item in entries)
        if uris != tuple(sorted(set(uris))):
            raise ValueError(
                "Census retail OCR entries must be unique and sorted"
            )
        if len({item.content_sha256 for item in entries}) != len(entries):
            raise ValueError("Census retail OCR repeats source content")
        object.__setattr__(self, "engine", engine)
        object.__setattr__(self, "engine_version", version)
        object.__setattr__(self, "entries", entries)
        expected = _stable_id(
            "census-retail-ocr-corpus", self.identity_payload()
        )
        if self.corpus_id and self.corpus_id != expected:
            raise ValueError("Census retail OCR-corpus identity differs")
        object.__setattr__(self, "corpus_id", expected)

    @property
    def by_uri(self) -> Mapping[str, CensusRetailOcrEntryV1]:
        return MappingProxyType(
            {item.artifact_uri: item for item in self.entries}
        )

    def identity_payload(self) -> dict[str, JSONValue]:
        return {
            "schema_version": self.schema_version,
            "engine": self.engine,
            "engine_version": self.engine_version,
            "render_dpi": self.render_dpi,
            "page_modes": list(self.page_modes),
            "entries": [item.to_dict() for item in self.entries],
        }

    def to_dict(self) -> dict[str, JSONValue]:
        return {**self.identity_payload(), "corpus_id": self.corpus_id}

    def to_json(self) -> str:
        return str(canonical_contract_json(self.to_dict()))

    @classmethod
    def from_dict(cls, data: Mapping[str, Any]) -> CensusRetailOcrCorpusV1:
        return cls(
            engine=str(data.get("engine", "")),
            engine_version=str(data.get("engine_version", "")),
            render_dpi=cast(int, data.get("render_dpi")),
            page_modes=tuple(
                cast(int, item)
                for item in _sequence(data.get("page_modes"), "page_modes")
            ),
            entries=tuple(
                CensusRetailOcrEntryV1.from_dict(
                    _mapping(item, "Census retail OCR entry")
                )
                for item in _sequence(data.get("entries"), "entries")
            ),
            corpus_id=str(data.get("corpus_id", "")),
            schema_version=str(data.get("schema_version", "")),
        )

    @classmethod
    def from_json(cls, text: str) -> CensusRetailOcrCorpusV1:
        try:
            payload = json.loads(text)
        except json.JSONDecodeError as exc:
            raise ValueError(
                "Census retail OCR corpus is invalid JSON"
            ) from exc
        return cls.from_dict(_mapping(payload, "Census retail OCR corpus"))


def _percentage_lexical(value: object, name: str) -> tuple[str, float]:
    text = _required_text(value, name)
    if re.fullmatch(r"-?(?:0|[1-9]\d*)\.\d", text) is None:
        raise ValueError(f"{name} must have one decimal place")
    numeric = float(text)
    if text == "-0.0":
        raise ValueError(f"{name} must not use negative zero")
    return text, numeric


def _canonical_percentage(value: Decimal) -> tuple[str, float]:
    rounded = value.quantize(Decimal("0.1"), rounding=ROUND_HALF_UP)
    if rounded == 0:
        rounded = Decimal("0.0")
    lexical = format(rounded, ".1f")
    return lexical, float(rounded)


@dataclass(frozen=True, slots=True)
class CensusRetailMeasureV1:
    """One occurrence-specific monthly percentage and revision lineage."""

    measure_key: str
    reference_period: str
    actual_value: float
    actual_lexical: str
    current_level_millions: int
    previous_level_millions: int
    level_basis: CensusRetailLevelBasis
    previous_as_known_value: float
    previous_as_known_lexical: str
    revised_previous_value: float
    revised_previous_lexical: str
    revision_comparable: bool
    comparison_basis: str
    previous_artifact_uri: str
    measure_id: str = ""
    schema_version: str = CENSUS_RETAIL_MEASURE_SCHEMA_VERSION

    def __post_init__(self) -> None:
        if self.schema_version != CENSUS_RETAIL_MEASURE_SCHEMA_VERSION:
            raise ValueError("unsupported Census retail measure schema")
        key = _required_text(self.measure_key, "measure_key")
        if key not in CENSUS_RETAIL_MEASURE_KEYS:
            raise ValueError("unsupported Census retail measure key")
        period = _year_month(self.reference_period, "reference_period")
        basis = CensusRetailLevelBasis.from_value(self.level_basis)
        for name in (
            "actual_value",
            "previous_as_known_value",
            "revised_previous_value",
        ):
            object.__setattr__(self, name, _finite(getattr(self, name), name))
        for lexical_name, numeric_name in (
            ("actual_lexical", "actual_value"),
            ("previous_as_known_lexical", "previous_as_known_value"),
            ("revised_previous_lexical", "revised_previous_value"),
        ):
            lexical, numeric = _percentage_lexical(
                getattr(self, lexical_name), lexical_name
            )
            if not math.isclose(
                numeric,
                getattr(self, numeric_name),
                rel_tol=0.0,
                abs_tol=1e-15,
            ):
                raise ValueError(f"{lexical_name} differs from {numeric_name}")
            object.__setattr__(self, lexical_name, lexical)
        for name in ("current_level_millions", "previous_level_millions"):
            _positive_int(
                getattr(self, name), name, MAX_CENSUS_RETAIL_ARTIFACT_BYTES
            )
        if not isinstance(self.revision_comparable, bool):
            raise TypeError("revision_comparable must be boolean")
        comparison = _required_text(self.comparison_basis, "comparison_basis")
        expected_comparison = (
            "sic-to-naics-boundary"
            if period == CENSUS_RETAIL_NAICS_BOUNDARY
            else "same-classification"
        )
        if comparison != expected_comparison:
            raise ValueError("Census retail comparison basis differs")
        if self.revision_comparable != (comparison == "same-classification"):
            raise ValueError("Census retail comparison flag differs")
        previous_uri = _https_uri(
            self.previous_artifact_uri, "previous_artifact_uri"
        )
        object.__setattr__(self, "measure_key", key)
        object.__setattr__(self, "reference_period", period)
        object.__setattr__(self, "level_basis", basis)
        object.__setattr__(self, "comparison_basis", comparison)
        object.__setattr__(self, "previous_artifact_uri", previous_uri)
        expected = _stable_id("census-retail-measure", self.identity_payload())
        if self.measure_id and self.measure_id != expected:
            raise ValueError("Census retail measure identity differs")
        object.__setattr__(self, "measure_id", expected)

    @property
    def previous_was_revised(self) -> bool:
        return self.revision_comparable and not math.isclose(
            self.previous_as_known_value,
            self.revised_previous_value,
            rel_tol=0.0,
            abs_tol=1e-15,
        )

    @property
    def revision_value(self) -> float | None:
        if not self.revision_comparable:
            return None
        value = self.revised_previous_value - self.previous_as_known_value
        return 0.0 if math.isclose(value, 0.0, abs_tol=1e-15) else value

    def identity_payload(self) -> dict[str, JSONValue]:
        return {
            "schema_version": self.schema_version,
            "measure_key": self.measure_key,
            "reference_period": self.reference_period,
            "actual_value": self.actual_value,
            "actual_lexical": self.actual_lexical,
            "current_level_millions": self.current_level_millions,
            "previous_level_millions": self.previous_level_millions,
            "level_basis": self.level_basis.value,
            "previous_as_known_value": self.previous_as_known_value,
            "previous_as_known_lexical": self.previous_as_known_lexical,
            "revised_previous_value": self.revised_previous_value,
            "revised_previous_lexical": self.revised_previous_lexical,
            "revision_comparable": self.revision_comparable,
            "comparison_basis": self.comparison_basis,
            "previous_artifact_uri": self.previous_artifact_uri,
        }

    def to_dict(self) -> dict[str, JSONValue]:
        return {**self.identity_payload(), "measure_id": self.measure_id}

    @classmethod
    def from_dict(cls, data: Mapping[str, Any]) -> CensusRetailMeasureV1:
        return cls(
            measure_key=str(data.get("measure_key", "")),
            reference_period=str(data.get("reference_period", "")),
            actual_value=cast(float, data.get("actual_value")),
            actual_lexical=str(data.get("actual_lexical", "")),
            current_level_millions=cast(
                int, data.get("current_level_millions")
            ),
            previous_level_millions=cast(
                int, data.get("previous_level_millions")
            ),
            level_basis=CensusRetailLevelBasis.from_value(
                str(data.get("level_basis", ""))
            ),
            previous_as_known_value=cast(
                float, data.get("previous_as_known_value")
            ),
            previous_as_known_lexical=str(
                data.get("previous_as_known_lexical", "")
            ),
            revised_previous_value=cast(
                float, data.get("revised_previous_value")
            ),
            revised_previous_lexical=str(
                data.get("revised_previous_lexical", "")
            ),
            revision_comparable=cast(bool, data.get("revision_comparable")),
            comparison_basis=str(data.get("comparison_basis", "")),
            previous_artifact_uri=str(data.get("previous_artifact_uri", "")),
            measure_id=str(data.get("measure_id", "")),
            schema_version=str(data.get("schema_version", "")),
        )


@dataclass(frozen=True, slots=True)
class CensusRetailPublicationV1:
    """One real PDF occurrence with optional exact spreadsheet evidence."""

    reference_period: str
    release_date: str
    released_lexical: str
    released_at_ns: int
    reported_zone: str
    zone_consistent: bool
    artifact_uri: str
    content_sha256: str
    content_length: int
    support_uri: str | None
    support_format: OfficialSourceFormat | None
    support_content_sha256: str | None
    support_content_length: int | None
    parser_era: str
    ocr_entry_id: str | None
    benchmark_revision_notice: bool
    shutdown_delayed: bool
    reported_rate_cross_check_count: int
    measures: tuple[CensusRetailMeasureV1, ...]
    publication_id: str = ""
    schema_version: str = CENSUS_RETAIL_PUBLICATION_SCHEMA_VERSION

    def __post_init__(self) -> None:
        if self.schema_version != CENSUS_RETAIL_PUBLICATION_SCHEMA_VERSION:
            raise ValueError("unsupported Census retail publication schema")
        period = _year_month(self.reference_period, "reference_period")
        release_date = _iso_date(self.release_date, "release_date")
        released = _required_text(self.released_lexical, "released_lexical")
        if released != f"{release_date}T08:30:00":
            raise ValueError("Census retail release time differs from date")
        expected_ns = int(
            datetime.combine(
                date.fromisoformat(release_date),
                time(8, 30),
                ZoneInfo("America/New_York"),
            ).timestamp()
            * 1_000_000_000
        )
        if self.released_at_ns != expected_ns:
            raise ValueError("Census retail normalized release time differs")
        reported_zone = _required_text(self.reported_zone, "reported_zone")
        if reported_zone not in {"ET", "EST", "EDT"}:
            raise ValueError("Census retail reported zone is invalid")
        if (
            not isinstance(self.zone_consistent, bool)
            or not self.zone_consistent
        ):
            raise ValueError("Census retail reported zone is inconsistent")
        artifact_uri = _https_uri(self.artifact_uri, "artifact_uri")
        content_sha = _sha256(self.content_sha256, "content_sha256")
        _positive_int(
            self.content_length,
            "content_length",
            MAX_CENSUS_RETAIL_ARTIFACT_BYTES,
        )
        support_uri = _optional_text(self.support_uri, "support_uri")
        support_format = self.support_format
        support_sha = self.support_content_sha256
        support_length = self.support_content_length
        if support_uri is None:
            if any(
                item is not None
                for item in (support_format, support_sha, support_length)
            ):
                raise ValueError("Census retail support evidence is incomplete")
        else:
            support_uri = _https_uri(support_uri, "support_uri")
            if (
                support_format is None
                or support_sha is None
                or support_length is None
            ):
                raise ValueError("Census retail support evidence is incomplete")
            support_format = OfficialSourceFormat.from_value(support_format)
            if support_format is not _support_format(support_uri):
                raise ValueError("Census retail support format differs")
            support_sha = _sha256(support_sha, "support_content_sha256")
            _positive_int(
                support_length,
                "support_content_length",
                MAX_CENSUS_RETAIL_ARTIFACT_BYTES,
            )
        era = _required_text(self.parser_era, "parser_era")
        if era not in {
            "ocr-pdf-table",
            "native-pdf-table",
            "xls-table",
            "xlsx-table",
        }:
            raise ValueError("Census retail parser era is invalid")
        ocr_id = _optional_text(self.ocr_entry_id, "ocr_entry_id")
        if (era == "ocr-pdf-table") != (ocr_id is not None):
            raise ValueError(
                "Census retail OCR identity differs from parser era"
            )
        if ocr_id is not None and not ocr_id.startswith(
            "census-retail-ocr-entry:sha256:"
        ):
            raise ValueError("Census retail OCR identity is invalid")
        for name in ("benchmark_revision_notice", "shutdown_delayed"):
            if not isinstance(getattr(self, name), bool):
                raise TypeError(f"{name} must be boolean")
        cross_checks = _nonnegative_int(
            self.reported_rate_cross_check_count,
            "reported_rate_cross_check_count",
            8,
        )
        if (support_uri is None and cross_checks != 0) or (
            support_uri is not None and cross_checks not in {4, 6}
        ):
            raise ValueError(
                "Census retail reported-rate cross-check count differs"
            )
        measures = tuple(self.measures)
        if (
            len(measures) != len(CENSUS_RETAIL_MEASURE_KEYS)
            or any(
                not isinstance(item, CensusRetailMeasureV1) for item in measures
            )
            or tuple(item.measure_key for item in measures)
            != CENSUS_RETAIL_MEASURE_KEYS
            or any(item.reference_period != period for item in measures)
        ):
            raise ValueError("Census retail publication measures differ")
        object.__setattr__(self, "reference_period", period)
        object.__setattr__(self, "release_date", release_date)
        object.__setattr__(self, "reported_zone", reported_zone)
        object.__setattr__(self, "artifact_uri", artifact_uri)
        object.__setattr__(self, "content_sha256", content_sha)
        object.__setattr__(self, "support_uri", support_uri)
        object.__setattr__(self, "support_format", support_format)
        object.__setattr__(self, "support_content_sha256", support_sha)
        object.__setattr__(self, "parser_era", era)
        object.__setattr__(self, "ocr_entry_id", ocr_id)
        object.__setattr__(self, "measures", measures)
        expected = _stable_id(
            "census-retail-publication", self.identity_payload()
        )
        if self.publication_id and self.publication_id != expected:
            raise ValueError("Census retail publication identity differs")
        object.__setattr__(self, "publication_id", expected)

    @property
    def by_measure(self) -> Mapping[str, CensusRetailMeasureV1]:
        return MappingProxyType(
            {item.measure_key: item for item in self.measures}
        )

    def identity_payload(self) -> dict[str, JSONValue]:
        return {
            "schema_version": self.schema_version,
            "reference_period": self.reference_period,
            "release_date": self.release_date,
            "released_lexical": self.released_lexical,
            "released_at_ns": self.released_at_ns,
            "reported_zone": self.reported_zone,
            "zone_consistent": self.zone_consistent,
            "artifact_uri": self.artifact_uri,
            "content_sha256": self.content_sha256,
            "content_length": self.content_length,
            "support_uri": self.support_uri,
            "support_format": (
                None
                if self.support_format is None
                else self.support_format.value
            ),
            "support_content_sha256": self.support_content_sha256,
            "support_content_length": self.support_content_length,
            "parser_era": self.parser_era,
            "ocr_entry_id": self.ocr_entry_id,
            "benchmark_revision_notice": self.benchmark_revision_notice,
            "shutdown_delayed": self.shutdown_delayed,
            "reported_rate_cross_check_count": self.reported_rate_cross_check_count,
            "measures": [item.to_dict() for item in self.measures],
        }

    def to_dict(self) -> dict[str, JSONValue]:
        return {
            **self.identity_payload(),
            "publication_id": self.publication_id,
        }

    @classmethod
    def from_dict(cls, data: Mapping[str, Any]) -> CensusRetailPublicationV1:
        raw_format = data.get("support_format")
        return cls(
            reference_period=str(data.get("reference_period", "")),
            release_date=str(data.get("release_date", "")),
            released_lexical=str(data.get("released_lexical", "")),
            released_at_ns=cast(int, data.get("released_at_ns")),
            reported_zone=str(data.get("reported_zone", "")),
            zone_consistent=cast(bool, data.get("zone_consistent")),
            artifact_uri=str(data.get("artifact_uri", "")),
            content_sha256=str(data.get("content_sha256", "")),
            content_length=cast(int, data.get("content_length")),
            support_uri=(
                None
                if data.get("support_uri") is None
                else str(data.get("support_uri"))
            ),
            support_format=(
                None
                if raw_format is None
                else OfficialSourceFormat.from_value(str(raw_format))
            ),
            support_content_sha256=(
                None
                if data.get("support_content_sha256") is None
                else str(data.get("support_content_sha256"))
            ),
            support_content_length=cast(
                int | None, data.get("support_content_length")
            ),
            parser_era=str(data.get("parser_era", "")),
            ocr_entry_id=(
                None
                if data.get("ocr_entry_id") is None
                else str(data.get("ocr_entry_id"))
            ),
            benchmark_revision_notice=cast(
                bool, data.get("benchmark_revision_notice")
            ),
            shutdown_delayed=cast(bool, data.get("shutdown_delayed")),
            reported_rate_cross_check_count=cast(
                int, data.get("reported_rate_cross_check_count")
            ),
            measures=tuple(
                CensusRetailMeasureV1.from_dict(
                    _mapping(item, "Census retail measure")
                )
                for item in _sequence(data.get("measures"), "measures")
            ),
            publication_id=str(data.get("publication_id", "")),
            schema_version=str(data.get("schema_version", "")),
        )


@dataclass(frozen=True, slots=True)
class CensusRetailArchiveManifestV1:
    """Compact deterministic closure receipt for the complete MARTS corpus."""

    registry_id: str
    profile_id: str
    release_index: CensusRetailReleaseIndexV1
    benchmark_index: CensusRetailBenchmarkIndexV1
    ocr_corpus_id: str
    predecessor_content_sha256: str
    predecessor_content_length: int
    publications: tuple[CensusRetailPublicationV1, ...]
    raw_artifact_count: int
    total_content_bytes: int
    measure_occurrence_count: int
    comparable_revision_count: int
    derived_aggregate_occurrence_count: int
    support_artifact_count: int
    corrected_support_link_count: int
    ocr_publication_count: int
    benchmark_revision_notice_count: int
    benchmark_state_override_count: int
    shutdown_delayed_count: int
    reported_rate_cross_check_count: int
    exact_minute_count: int
    manifest_id: str = ""
    schema_version: str = CENSUS_RETAIL_ARCHIVE_MANIFEST_SCHEMA_VERSION

    def __post_init__(self) -> None:
        if self.schema_version != CENSUS_RETAIL_ARCHIVE_MANIFEST_SCHEMA_VERSION:
            raise ValueError(
                "unsupported Census retail archive-manifest schema"
            )
        registry = _required_text(self.registry_id, "registry_id")
        profile = _required_text(self.profile_id, "profile_id")
        if not registry.startswith("official-source-registry:sha256:"):
            raise ValueError("Census retail registry identity is invalid")
        if not profile.startswith("us-backfill-profile:sha256:"):
            raise ValueError("Census retail profile identity is invalid")
        if not isinstance(self.release_index, CensusRetailReleaseIndexV1):
            raise TypeError(
                "Census retail manifest requires a v1 release index"
            )
        if not isinstance(self.benchmark_index, CensusRetailBenchmarkIndexV1):
            raise TypeError(
                "Census retail manifest requires a v1 benchmark index"
            )
        if self.benchmark_index.as_of_date != self.release_index.as_of_date:
            raise ValueError("Census retail index observation dates differ")
        ocr_id = _required_text(self.ocr_corpus_id, "ocr_corpus_id")
        if not ocr_id.startswith("census-retail-ocr-corpus:sha256:"):
            raise ValueError("Census retail OCR-corpus identity is invalid")
        predecessor_sha = _sha256(
            self.predecessor_content_sha256, "predecessor_content_sha256"
        )
        _positive_int(
            self.predecessor_content_length,
            "predecessor_content_length",
            MAX_CENSUS_RETAIL_ARTIFACT_BYTES,
        )
        publications = tuple(self.publications)
        if (
            not publications
            or len(publications) > MAX_CENSUS_RETAIL_RELEASES
            or any(
                not isinstance(item, CensusRetailPublicationV1)
                for item in publications
            )
        ):
            raise TypeError("Census retail manifest publications are invalid")
        periods = tuple(item.reference_period for item in publications)
        if periods != tuple(
            item.reference_period for item in self.release_index.releases
        ):
            raise ValueError(
                "Census retail manifest differs from release index"
            )
        dates = tuple(item.release_date for item in publications)
        if dates != tuple(sorted(set(dates))):
            raise ValueError(
                "Census retail release dates must be unique and sorted"
            )
        if tuple(item.artifact_uri for item in publications) != tuple(
            item.artifact_uri for item in self.release_index.releases
        ):
            raise ValueError("Census retail publication URIs differ from index")
        expected_override_uris = {
            item.artifact_uri
            for item in self.benchmark_index.entries
            if item.disposition
            is CensusRetailBenchmarkDisposition.PRIOR_STATE_OVERRIDE
        }
        used_override_uris = {
            measure.previous_artifact_uri
            for item in publications
            for measure in item.measures
            if measure.previous_artifact_uri in expected_override_uris
        }
        if used_override_uris != expected_override_uris:
            raise ValueError("Census retail benchmark state overrides differ")
        expected_counts = {
            "raw_artifact_count": 3
            + len(publications)
            + sum(item.support_uri is not None for item in publications)
            + len(self.benchmark_index.entries),
            "total_content_bytes": self.release_index.content_length
            + self.benchmark_index.content_length
            + self.predecessor_content_length
            + sum(item.content_length for item in self.benchmark_index.entries)
            + sum(
                item.content_length + (item.support_content_length or 0)
                for item in publications
            ),
            "measure_occurrence_count": sum(
                len(item.measures) for item in publications
            ),
            "comparable_revision_count": sum(
                measure.previous_was_revised
                for item in publications
                for measure in item.measures
            ),
            "derived_aggregate_occurrence_count": sum(
                measure.level_basis is CensusRetailLevelBasis.DERIVED_AGGREGATE
                for item in publications
                for measure in item.measures
            ),
            "support_artifact_count": sum(
                item.support_uri is not None for item in publications
            ),
            "corrected_support_link_count": sum(
                item.support_link_corrected
                for item in self.release_index.releases
            ),
            "ocr_publication_count": sum(
                item.ocr_entry_id is not None for item in publications
            ),
            "benchmark_revision_notice_count": sum(
                item.benchmark_revision_notice for item in publications
            ),
            "benchmark_state_override_count": len(used_override_uris),
            "shutdown_delayed_count": sum(
                item.shutdown_delayed for item in publications
            ),
            "reported_rate_cross_check_count": sum(
                item.reported_rate_cross_check_count for item in publications
            ),
            "exact_minute_count": len(publications),
        }
        for name, expected_value in expected_counts.items():
            _nonnegative_int(getattr(self, name), name, 1_000_000_000)
            if getattr(self, name) != expected_value:
                raise ValueError(f"Census retail {name} differs")
        object.__setattr__(self, "registry_id", registry)
        object.__setattr__(self, "profile_id", profile)
        object.__setattr__(self, "ocr_corpus_id", ocr_id)
        object.__setattr__(self, "predecessor_content_sha256", predecessor_sha)
        object.__setattr__(self, "publications", publications)
        expected = _stable_id(
            "census-retail-archive-manifest", self.identity_payload()
        )
        if self.manifest_id and self.manifest_id != expected:
            raise ValueError("Census retail archive identity differs")
        object.__setattr__(self, "manifest_id", expected)

    @property
    def by_period(self) -> Mapping[str, CensusRetailPublicationV1]:
        return MappingProxyType(
            {item.reference_period: item for item in self.publications}
        )

    def identity_payload(self) -> dict[str, JSONValue]:
        return {
            "schema_version": self.schema_version,
            "registry_id": self.registry_id,
            "profile_id": self.profile_id,
            "release_index": self.release_index.to_dict(),
            "benchmark_index": self.benchmark_index.to_dict(),
            "ocr_corpus_id": self.ocr_corpus_id,
            "predecessor_content_sha256": self.predecessor_content_sha256,
            "predecessor_content_length": self.predecessor_content_length,
            "publications": [item.to_dict() for item in self.publications],
            "raw_artifact_count": self.raw_artifact_count,
            "total_content_bytes": self.total_content_bytes,
            "measure_occurrence_count": self.measure_occurrence_count,
            "comparable_revision_count": self.comparable_revision_count,
            "derived_aggregate_occurrence_count": self.derived_aggregate_occurrence_count,
            "support_artifact_count": self.support_artifact_count,
            "corrected_support_link_count": self.corrected_support_link_count,
            "ocr_publication_count": self.ocr_publication_count,
            "benchmark_revision_notice_count": self.benchmark_revision_notice_count,
            "benchmark_state_override_count": self.benchmark_state_override_count,
            "shutdown_delayed_count": self.shutdown_delayed_count,
            "reported_rate_cross_check_count": self.reported_rate_cross_check_count,
            "exact_minute_count": self.exact_minute_count,
        }

    def to_dict(self) -> dict[str, JSONValue]:
        return {**self.identity_payload(), "manifest_id": self.manifest_id}

    def to_json(self) -> str:
        return str(canonical_contract_json(self.to_dict()))

    @classmethod
    def from_dict(
        cls, data: Mapping[str, Any]
    ) -> CensusRetailArchiveManifestV1:
        return cls(
            registry_id=str(data.get("registry_id", "")),
            profile_id=str(data.get("profile_id", "")),
            release_index=CensusRetailReleaseIndexV1.from_dict(
                _mapping(data.get("release_index"), "release_index")
            ),
            benchmark_index=CensusRetailBenchmarkIndexV1.from_dict(
                _mapping(data.get("benchmark_index"), "benchmark_index")
            ),
            ocr_corpus_id=str(data.get("ocr_corpus_id", "")),
            predecessor_content_sha256=str(
                data.get("predecessor_content_sha256", "")
            ),
            predecessor_content_length=cast(
                int, data.get("predecessor_content_length")
            ),
            publications=tuple(
                CensusRetailPublicationV1.from_dict(
                    _mapping(item, "Census retail publication")
                )
                for item in _sequence(data.get("publications"), "publications")
            ),
            raw_artifact_count=cast(int, data.get("raw_artifact_count")),
            total_content_bytes=cast(int, data.get("total_content_bytes")),
            measure_occurrence_count=cast(
                int, data.get("measure_occurrence_count")
            ),
            comparable_revision_count=cast(
                int, data.get("comparable_revision_count")
            ),
            derived_aggregate_occurrence_count=cast(
                int, data.get("derived_aggregate_occurrence_count")
            ),
            support_artifact_count=cast(
                int, data.get("support_artifact_count")
            ),
            corrected_support_link_count=cast(
                int, data.get("corrected_support_link_count")
            ),
            ocr_publication_count=cast(int, data.get("ocr_publication_count")),
            benchmark_revision_notice_count=cast(
                int, data.get("benchmark_revision_notice_count")
            ),
            benchmark_state_override_count=cast(
                int, data.get("benchmark_state_override_count")
            ),
            shutdown_delayed_count=cast(
                int, data.get("shutdown_delayed_count")
            ),
            reported_rate_cross_check_count=cast(
                int, data.get("reported_rate_cross_check_count")
            ),
            exact_minute_count=cast(int, data.get("exact_minute_count")),
            manifest_id=str(data.get("manifest_id", "")),
            schema_version=str(data.get("schema_version", "")),
        )

    @classmethod
    def from_json(cls, text: str) -> CensusRetailArchiveManifestV1:
        try:
            payload = json.loads(text)
        except json.JSONDecodeError as exc:
            raise ValueError("Census retail manifest is invalid JSON") from exc
        return cls.from_dict(_mapping(payload, "Census retail manifest"))


def build_census_retail_index_request(
    registry: OfficialSourceRegistryV1,
    *,
    as_of_date: str,
) -> OfficialSourceRequestV1:
    """Plan the exact official MARTS historical-release index request."""
    if not isinstance(registry, OfficialSourceRegistryV1):
        raise TypeError("Census retail index request requires a v1 registry")
    as_of = _iso_date(as_of_date, "as_of_date")
    source = registry.source(CENSUS_RETAIL_SOURCE_KEY)
    return OfficialSourceRequestV1(
        source_key=source.source_key,
        source_id=source.source_id,
        method=OfficialRequestMethod.GET,
        uri=CENSUS_RETAIL_INDEX_URI,
        source_format=OfficialSourceFormat.HTML,
        parser_id=source.parser_id,
        parser_version=source.parser_version,
        window_start=US_BACKFILL_START_DATE,
        window_end=as_of,
        page_number=1,
    )


def build_census_retail_benchmark_index_request(
    registry: OfficialSourceRegistryV1,
    *,
    as_of_date: str,
) -> OfficialSourceRequestV1:
    """Plan the exact official annual-revision archive-index request."""
    if not isinstance(registry, OfficialSourceRegistryV1):
        raise TypeError(
            "Census retail benchmark request requires a v1 registry"
        )
    as_of = _iso_date(as_of_date, "as_of_date")
    source = registry.source(CENSUS_RETAIL_SOURCE_KEY)
    return OfficialSourceRequestV1(
        source_key=source.source_key,
        source_id=source.source_id,
        method=OfficialRequestMethod.GET,
        uri=CENSUS_RETAIL_BENCHMARK_INDEX_URI,
        source_format=OfficialSourceFormat.HTML,
        parser_id=source.parser_id,
        parser_version=source.parser_version,
        window_start=US_BACKFILL_START_DATE,
        window_end=as_of,
        page_number=1,
    )


def parse_census_retail_release_index(
    snapshot: OfficialRawSnapshotV1,
    *,
    as_of_date: str,
) -> CensusRetailReleaseIndexV1:
    """Parse the official index and append the bounded current supplement."""
    if not isinstance(snapshot, OfficialRawSnapshotV1):
        raise TypeError("Census retail index parser requires a v1 snapshot")
    if (
        snapshot.request.source_key != CENSUS_RETAIL_SOURCE_KEY
        or snapshot.request.source_format is not OfficialSourceFormat.HTML
        or snapshot.request.uri != CENSUS_RETAIL_INDEX_URI
    ):
        raise ValueError(
            "Census retail index parser received a different source"
        )
    as_of = _iso_date(as_of_date, "as_of_date")
    try:
        text = snapshot.content.decode("utf-8-sig")
    except UnicodeDecodeError as exc:
        raise ValueError("Census retail index is not UTF-8") from exc
    pdfs: dict[str, str] = {}
    supports: dict[str, tuple[str, bool]] = {}
    for match in _HREF_RE.finditer(text):
        observed_uri = match.group("uri").replace("&amp;", "&")
        pdf_match = _ARCHIVE_PDF_RE.fullmatch(observed_uri)
        if pdf_match is not None:
            period = _period_from_suffix(
                pdf_match.group("year"), pdf_match.group("month")
            )
            pdfs.setdefault(period, observed_uri)
            continue
        corrected_uri = CENSUS_RETAIL_SUPPORT_LINK_CORRECTIONS.get(
            observed_uri, observed_uri
        )
        support_match = _ARCHIVE_SUPPORT_RE.fullmatch(corrected_uri)
        if support_match is None:
            continue
        period = _period_from_suffix(
            support_match.group("year"), support_match.group("month")
        )
        candidate = (
            corrected_uri,
            corrected_uri != observed_uri,
        )
        existing = supports.get(period)
        if existing is not None and existing != candidate:
            raise ValueError("Census retail index repeats a support period")
        supports[period] = candidate
    archive_periods = tuple(
        period
        for period in sorted(pdfs)
        if CENSUS_RETAIL_FIRST_REFERENCE_PERIOD
        <= period
        < CENSUS_RETAIL_CURRENT_REFERENCE_PERIOD
    )
    if not archive_periods or archive_periods[-1] != "2026-06":
        raise ValueError("Census retail archive boundary differs")
    releases: list[CensusRetailReleaseIndexEntryV1] = []
    for period in archive_periods:
        support = supports.get(period)
        releases.append(
            CensusRetailReleaseIndexEntryV1(
                reference_period=period,
                artifact_uri=pdfs[period],
                support_uri=None if support is None else support[0],
                support_format=(
                    None if support is None else _support_format(support[0])
                ),
                archive_listed=True,
                support_link_corrected=(
                    False if support is None else support[1]
                ),
            )
        )
    if CENSUS_RETAIL_CURRENT_RELEASE_DATE > as_of:
        raise ValueError("Census retail current release postdates as-of date")
    releases.append(
        CensusRetailReleaseIndexEntryV1(
            reference_period=CENSUS_RETAIL_CURRENT_REFERENCE_PERIOD,
            artifact_uri=CENSUS_RETAIL_CURRENT_PDF_URI,
            support_uri=CENSUS_RETAIL_CURRENT_XLSX_URI,
            support_format=OfficialSourceFormat.XLSX,
            archive_listed=False,
            support_link_corrected=False,
        )
    )
    return CensusRetailReleaseIndexV1(
        source_uri=CENSUS_RETAIL_INDEX_URI,
        content_sha256=snapshot.content_sha256,
        content_length=len(snapshot.content),
        as_of_date=as_of,
        releases=tuple(releases),
    )


def _release_request(
    registry: OfficialSourceRegistryV1,
    *,
    uri: str,
    source_format: OfficialSourceFormat,
    reference_period: str,
) -> OfficialSourceRequestV1:
    source = registry.source(CENSUS_RETAIL_SOURCE_KEY)
    period = _year_month(reference_period, "reference_period")
    return OfficialSourceRequestV1(
        source_key=source.source_key,
        source_id=source.source_id,
        method=OfficialRequestMethod.GET,
        uri=uri,
        source_format=source_format,
        parser_id=source.parser_id,
        parser_version=source.parser_version,
        window_start=f"{period}-01",
        window_end=f"{period}-01",
    )


def build_census_retail_benchmark_requests(
    registry: OfficialSourceRegistryV1,
    index_snapshot: OfficialRawSnapshotV1,
) -> tuple[OfficialSourceRequestV1, ...]:
    """Verify the annual index and plan all in-scope benchmark artifacts."""
    if not isinstance(registry, OfficialSourceRegistryV1):
        raise TypeError(
            "Census retail benchmark requests require a v1 registry"
        )
    _validate_snapshot(
        index_snapshot,
        uri=CENSUS_RETAIL_BENCHMARK_INDEX_URI,
        source_format=OfficialSourceFormat.HTML,
    )
    try:
        text = index_snapshot.content.decode("utf-8-sig")
    except UnicodeDecodeError as exc:
        raise ValueError("Census retail benchmark index is not UTF-8") from exc
    observed = {
        match.group("uri")
        .replace("&amp;", "&")
        .replace("http://", "https://", 1)
        for match in _HREF_RE.finditer(text)
    }
    expected = set(CENSUS_RETAIL_BENCHMARK_URIS.values())
    if not expected.issubset(observed):
        missing = sorted(expected - observed)
        raise ValueError(
            f"Census retail benchmark index is incomplete: {missing!r}"
        )
    requests: list[OfficialSourceRequestV1] = []
    for report_year, uri in CENSUS_RETAIL_BENCHMARK_URIS.items():
        source_format = (
            OfficialSourceFormat.PDF
            if report_year <= 2007
            else (
                OfficialSourceFormat.XLSX
                if report_year >= 2023
                else OfficialSourceFormat.XLS
            )
        )
        requests.append(
            _release_request(
                registry,
                uri=uri,
                source_format=source_format,
                reference_period=CENSUS_RETAIL_BENCHMARK_PERIOD_ENDS[
                    report_year
                ],
            )
        )
    return tuple(requests)


def build_census_retail_archive_requests(
    registry: OfficialSourceRegistryV1,
    release_index: CensusRetailReleaseIndexV1,
) -> tuple[OfficialSourceRequestV1, ...]:
    """Plan predecessor, PDF, and table-companion acquisition."""
    if not isinstance(registry, OfficialSourceRegistryV1):
        raise TypeError("Census retail requests require a v1 registry")
    if not isinstance(release_index, CensusRetailReleaseIndexV1):
        raise TypeError("Census retail requests require a v1 release index")
    requests = [
        _release_request(
            registry,
            uri=CENSUS_RETAIL_PREDECESSOR_URI,
            source_format=OfficialSourceFormat.PDF,
            reference_period=CENSUS_RETAIL_PREDECESSOR_REFERENCE_PERIOD,
        )
    ]
    for release in release_index.releases:
        requests.append(
            _release_request(
                registry,
                uri=release.artifact_uri,
                source_format=OfficialSourceFormat.PDF,
                reference_period=release.reference_period,
            )
        )
        if release.support_uri is not None:
            assert release.support_format is not None
            requests.append(
                _release_request(
                    registry,
                    uri=release.support_uri,
                    source_format=release.support_format,
                    reference_period=release.reference_period,
                )
            )
    if len({item.uri for item in requests}) != len(requests):
        raise ValueError("Census retail request plan repeats a URI")
    return tuple(requests)


def _validate_snapshot(
    snapshot: OfficialRawSnapshotV1,
    *,
    uri: str,
    source_format: OfficialSourceFormat,
) -> None:
    if not isinstance(snapshot, OfficialRawSnapshotV1):
        raise TypeError("Census retail parser requires v1 snapshots")
    if (
        snapshot.request.source_key != CENSUS_RETAIL_SOURCE_KEY
        or snapshot.request.uri != uri
        or snapshot.request.source_format is not source_format
        or snapshot.status_code != 200
        or not snapshot.content
    ):
        raise ValueError("Census retail snapshot identity differs")
    if len(snapshot.content) > MAX_CENSUS_RETAIL_ARTIFACT_BYTES:
        raise ValueError("Census retail artifact exceeds its byte bound")


def _pdf_text(content: bytes) -> str:
    try:
        reader = PdfReader(BytesIO(content))
        return "\f".join(page.extract_text() or "" for page in reader.pages)
    except Exception as exc:
        raise ValueError(
            "Census retail artifact is not a readable PDF"
        ) from exc


def _release_time(text: str) -> tuple[str, str, int, bool]:
    flattened = " ".join(text.split())
    match = _TIME_THEN_DATE_RE.search(flattened)
    if match is None:
        match = _DATE_THEN_TIME_RE.search(flattened)
    if match is None:
        raise ValueError("Census retail release header time is unavailable")
    date_text = re.sub(r",\s*", ", ", match.group("date"))
    date_match = re.fullmatch(
        r"(?P<month>[A-Za-z]+)\s+(?P<day>\d{1,2}),\s+(?P<year>\d{4})",
        date_text,
    )
    if date_match is None:
        raise ValueError("Census retail release header date is invalid")
    month_number = _MONTH_NUMBERS.get(date_match.group("month").lower())
    if month_number is None:
        raise ValueError("Census retail release header month is invalid")
    try:
        parsed_date = date(
            int(date_match.group("year")),
            month_number,
            int(date_match.group("day")),
        )
    except ValueError as exc:
        raise ValueError(
            "Census retail release header date is invalid"
        ) from exc
    reported_zone = re.sub(r"\s+", "", match.group("zone")).upper()
    local = datetime.combine(
        parsed_date, time(8, 30), ZoneInfo("America/New_York")
    )
    utc_offset = local.dst()
    expected_zone = (
        "EDT"
        if utc_offset is not None and utc_offset.total_seconds()
        else "EST"
    )
    zone_consistent = reported_zone == "ET" or reported_zone == expected_zone
    return (
        parsed_date.isoformat(),
        reported_zone,
        int(local.timestamp() * 1_000_000_000),
        zone_consistent,
    )


def _int_tokens(value: str) -> tuple[int, ...]:
    return tuple(
        int(item.replace(",", "").replace(".", ""))
        for item in _THOUSANDS_RE.findall(value)
    )


def _benchmark_int_tokens(value: str) -> tuple[int, ...]:
    spaced = tuple(
        int(
            re.sub(r"\s+", "", match.group("head"))
            + re.sub(r"\s+", "", match.group("tail"))
        )
        for match in _SPACED_THOUSANDS_RE.finditer(value)
    )
    return spaced if spaced else _int_tokens(value)


def _pdf_component_levels(
    text: str,
    *,
    content_sha256: str,
) -> Mapping[str, tuple[int, int, int]]:
    table_one = re.search(r"Table\s*1\.", text, re.IGNORECASE)
    if table_one is None:
        raise ValueError("Census retail PDF Table 1 is unavailable")
    table_two = re.search(
        r"Table\s*2\.", text[table_one.end() :], re.IGNORECASE
    )
    table_text = (
        text[table_one.start() : table_one.end() + table_two.start()]
        if table_two is not None
        else text[table_one.start() :]
    )
    lines = table_text.splitlines()
    levels: dict[str, tuple[int, int, int]] = {}
    for key, pattern in _ROW_PATTERNS.items():
        override = _OCR_LEVEL_OVERRIDES.get((content_sha256, key))
        if override is not None:
            levels[key] = override
            continue
        candidates: list[tuple[int, ...]] = []
        for position in range(len(lines)):
            segment = " ".join(lines[position : position + 2])
            match = pattern.search(segment)
            if match is None:
                continue
            tokens = _int_tokens(segment[match.end() :])
            if len(tokens) >= 9:
                candidates.append(tokens)
        if not candidates:
            raise ValueError(f"Census retail PDF {key} row is unavailable")
        tokens = candidates[0]
        if len(tokens) >= 11:
            selected = cast(tuple[int, int, int], tuple(tokens[6:9]))
        elif len(tokens) == 10:
            selected = cast(tuple[int, int, int], tuple(tokens[5:8]))
        else:
            raise ValueError(f"Census retail PDF {key} row is incomplete")
        lower, upper = _LEVEL_RANGES[key]
        if any(not lower <= item <= upper for item in selected):
            raise ValueError(f"Census retail PDF {key} level is implausible")
        levels[key] = selected
    return MappingProxyType(levels)


def _normalized_cell(value: object) -> str:
    return re.sub(r"[^a-z0-9&]+", " ", str(value or "").lower()).strip()


def _cell_code(value: object) -> str:
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value or "").strip()


def _row_level_triple(row: Sequence[object]) -> tuple[int, int, int] | None:
    values = tuple(row[9:12])
    if len(values) != 3 or any(
        not isinstance(item, (int, float)) or isinstance(item, bool)
        for item in values
    ):
        return None
    first, second, third = values
    assert isinstance(first, (int, float))
    assert isinstance(second, (int, float))
    assert isinstance(third, (int, float))
    return int(first), int(second), int(third)


def _scrub_strict_ooxml(content: bytes) -> bytes:
    """Normalize two source-authored Strict OOXML external-link defects."""
    source = ZipFile(BytesIO(content))
    output = BytesIO()
    with source, ZipFile(output, "w", ZIP_DEFLATED) as target:
        for member in source.infolist():
            data = source.read(member.filename)
            if member.filename.endswith((".xml", ".rels")):
                data = data.replace(
                    b"http://purl.oclc.org/ooxml/spreadsheetml/main",
                    b"http://schemas.openxmlformats.org/spreadsheetml/2006/main",
                )
                data = data.replace(
                    b"http://purl.oclc.org/ooxml/officeDocument/relationships",
                    b"http://schemas.openxmlformats.org/officeDocument/2006/relationships",
                )
                data = data.replace(
                    b"http://purl.oclc.org/ooxml/package/relationships",
                    b"http://schemas.openxmlformats.org/package/2006/relationships",
                )
            if member.filename == "xl/workbook.xml":
                data = re.sub(
                    rb"<externalReferences>.*?</externalReferences>",
                    b"",
                    data,
                    flags=re.DOTALL,
                )
            target.writestr(member, data)
    return output.getvalue()


def _spreadsheet_rows(
    content: bytes,
    source_format: OfficialSourceFormat,
) -> Mapping[str, tuple[tuple[object, ...], ...]]:
    sheet_names = ("Table 1.", "Table 2.")
    if source_format is OfficialSourceFormat.XLS:
        try:
            workbook = xlrd.open_workbook(file_contents=content)
            return MappingProxyType(
                {
                    name: tuple(
                        tuple(workbook.sheet_by_name(name).row_values(position))
                        for position in range(
                            workbook.sheet_by_name(name).nrows
                        )
                    )
                    for name in sheet_names
                }
            )
        except (xlrd.XLRDError, IndexError, ValueError) as exc:
            raise ValueError("Census retail XLS tables are unreadable") from exc
    if source_format is not OfficialSourceFormat.XLSX:
        raise ValueError("Census retail spreadsheet format is unsupported")
    try:
        try:
            workbook = load_workbook(
                BytesIO(content),
                data_only=True,
                read_only=True,
                keep_links=False,
            )
            if any(name not in workbook.sheetnames for name in sheet_names):
                workbook.close()
                raise TypeError("Strict OOXML sheet namespace is unsupported")
        except TypeError:
            with warnings.catch_warnings():
                warnings.simplefilter("ignore")
                workbook = load_workbook(
                    BytesIO(_scrub_strict_ooxml(content)),
                    data_only=True,
                    read_only=True,
                    keep_links=False,
                )
        result = {
            name: tuple(
                tuple(row) for row in workbook[name].iter_rows(values_only=True)
            )
            for name in sheet_names
        }
        workbook.close()
        return MappingProxyType(result)
    except (KeyError, OSError, ValueError) as exc:
        raise ValueError("Census retail XLSX tables are unreadable") from exc


def _percent_change(current: int, previous: int) -> tuple[str, float]:
    if current <= 0 or previous <= 0:
        raise ValueError("Census retail levels must be positive")
    return _canonical_percentage(
        Decimal(100) * (Decimal(current) / Decimal(previous) - Decimal(1))
    )


def _spreadsheet_component_levels(
    content: bytes,
    source_format: OfficialSourceFormat,
) -> tuple[Mapping[str, tuple[int, int, int]], int]:
    sheets = _spreadsheet_rows(content, source_format)
    table_one = sheets["Table 1."]
    raw: dict[str, tuple[int, int, int]] = {}
    for position, row in enumerate(table_one):
        label = _normalized_cell(row[1] if len(row) > 1 else None)
        code = _cell_code(row[0] if row else None)
        values = _row_level_triple(row)
        if (
            label.startswith("total")
            and "excl" not in label
            and values is not None
            and "total" not in raw
        ):
            raw["total"] = values
        if (
            label.startswith("total excl motor vehicle")
            and "gasoline" not in label
            and values is not None
        ):
            raw["exauto"] = values
        if label.startswith("total excl gasoline") and values is not None:
            raw["exgas"] = values
        if code == "441" and "motor vehicle" in label and values is not None:
            raw["auto"] = values
        if code == "444" and label.startswith("building material"):
            for candidate in table_one[position : position + 3]:
                candidate_values = _row_level_triple(candidate)
                if candidate_values is not None:
                    raw["building"] = candidate_values
                    break
        if code == "447" and "gasoline" in label and values is not None:
            raw["gas"] = values
        if code == "722" and "food services" in label and values is not None:
            raw["food"] = values
    required = {"total", "exauto", "auto", "building", "gas", "food"}
    if not required.issubset(raw):
        raise ValueError("Census retail spreadsheet lacks required level rows")
    if raw["total"] != tuple(
        left + right
        for left, right in zip(raw["exauto"], raw["auto"], strict=True)
    ):
        raise ValueError("Census retail total accounting identity differs")
    if "exgas" in raw and raw["exgas"] != tuple(
        left - right
        for left, right in zip(raw["total"], raw["gas"], strict=True)
    ):
        raise ValueError("Census retail ex-gas accounting identity differs")

    reported: dict[str, tuple[float, float]] = {}
    for row in sheets["Table 2."]:
        label = _normalized_cell(row[1] if len(row) > 1 else None)
        pair = (
            (float(row[2]), float(row[4]))
            if len(row) > 4
            and isinstance(row[2], (int, float))
            and not isinstance(row[2], bool)
            and isinstance(row[4], (int, float))
            and not isinstance(row[4], bool)
            else None
        )
        if (
            label.startswith("total")
            and "excl" not in label
            and pair is not None
            and "total" not in reported
        ):
            reported["total"] = pair
        if (
            label.startswith("total excl motor vehicle")
            and "gasoline" not in label
            and pair is not None
        ):
            reported["exauto"] = pair
        if label.startswith("total excl gasoline") and pair is not None:
            reported["exgas"] = pair
    expected_reported = {"total", "exauto"}
    if "exgas" in raw:
        expected_reported.add("exgas")
    if set(reported) != expected_reported:
        raise ValueError("Census retail reported-rate rows differ")
    cross_checks = 0
    for key, pair in reported.items():
        levels = raw[key]
        calculated = (
            _percent_change(levels[0], levels[1])[1],
            _percent_change(levels[1], levels[2])[1],
        )
        if any(
            not math.isclose(actual, expected, rel_tol=0.0, abs_tol=1e-15)
            for actual, expected in zip(pair, calculated, strict=True)
        ):
            raise ValueError("Census retail reported rate differs from levels")
        cross_checks += 2
    return MappingProxyType(raw), cross_checks


def _benchmark_pdf_component_levels(
    content: bytes,
) -> Mapping[str, tuple[int, int]]:
    text = _pdf_text(content)
    table_one = re.search(r"Table\s*1a\.", text, re.IGNORECASE)
    if table_one is None:
        raise ValueError("Census retail benchmark PDF Table 1a is unavailable")
    table_two = re.search(
        r"Table\s*1b\.", text[table_one.end() :], re.IGNORECASE
    )
    if table_two is None:
        raise ValueError(
            "Census retail benchmark PDF Table 1b boundary is unavailable"
        )
    table_text = text[table_one.start() : table_one.end() + table_two.start()]
    raw: dict[str, tuple[int, int]] = {}
    for key, pattern in _BENCHMARK_PDF_ROW_PATTERNS.items():
        for line in table_text.splitlines():
            match = pattern.search(line)
            if match is None:
                continue
            tokens = _benchmark_int_tokens(line[match.end() :])
            if len(tokens) >= 2:
                raw[key] = cast(tuple[int, int], tuple(tokens[-2:]))
                break
    required = {"total", "exauto", "auto", "building", "gas", "food"}
    if set(raw) != required:
        raise ValueError("Census retail benchmark PDF lacks required rows")
    if raw["total"] != tuple(
        left + right
        for left, right in zip(raw["exauto"], raw["auto"], strict=True)
    ):
        raise ValueError("Census retail benchmark PDF total identity differs")
    return MappingProxyType(raw)


def _benchmark_spreadsheet_rows(
    content: bytes,
    source_format: OfficialSourceFormat,
    report_year: int,
) -> tuple[tuple[object, ...], ...]:
    sheet_name = str(report_year)
    if source_format is OfficialSourceFormat.XLS:
        try:
            workbook = xlrd.open_workbook(file_contents=content)
            sheet = workbook.sheet_by_name(sheet_name)
            return tuple(
                tuple(sheet.row_values(position))
                for position in range(sheet.nrows)
            )
        except (xlrd.XLRDError, IndexError, ValueError) as exc:
            raise ValueError(
                "Census retail benchmark XLS is unreadable"
            ) from exc
    if source_format is not OfficialSourceFormat.XLSX:
        raise ValueError(
            "Census retail benchmark spreadsheet format is invalid"
        )
    try:
        workbook = load_workbook(
            BytesIO(content), data_only=True, read_only=True, keep_links=False
        )
        rows = tuple(
            tuple(row)
            for row in workbook[sheet_name].iter_rows(values_only=True)
        )
        workbook.close()
        return rows
    except (KeyError, OSError, ValueError) as exc:
        raise ValueError("Census retail benchmark XLSX is unreadable") from exc


def _benchmark_row_pair(row: Sequence[object]) -> tuple[int, int] | None:
    values = tuple(
        int(item)
        for item in row[2:]
        if isinstance(item, (int, float)) and not isinstance(item, bool)
    )
    if len(values) < 2:
        return None
    return cast(tuple[int, int], values[-2:])


def _benchmark_spreadsheet_component_levels(
    content: bytes,
    source_format: OfficialSourceFormat,
    report_year: int,
) -> Mapping[str, tuple[int, int]]:
    rows = _benchmark_spreadsheet_rows(content, source_format, report_year)
    if len(rows) < 75:
        raise ValueError("Census retail benchmark spreadsheet is incomplete")
    expected_period = CENSUS_RETAIL_BENCHMARK_PERIOD_ENDS[report_year]
    expected_year, expected_month = map(int, expected_period.split("-"))
    final_headers = tuple(
        item for item in rows[4][2:] if str(item or "").strip()
    )
    if not final_headers:
        raise ValueError(
            "Census retail benchmark month headers are unavailable"
        )
    final_header = _normalized_cell(final_headers[-1])
    expected_month_name = calendar.month_abbr[expected_month].lower()
    header_differs = (
        expected_month_name not in final_header
        or str(expected_year) not in final_header
    )
    if header_differs and final_header != _BENCHMARK_TERMINAL_HEADER_TYPOS.get(
        report_year
    ):
        raise ValueError("Census retail benchmark terminal period differs")
    adjusted_position = next(
        (
            position
            for position, row in enumerate(rows)
            if any(
                _normalized_cell(item).startswith("adjusted") for item in row
            )
        ),
        None,
    )
    if adjusted_position is None:
        raise ValueError(
            "Census retail benchmark adjusted table is unavailable"
        )
    raw: dict[str, tuple[int, int]] = {}
    for row in rows[adjusted_position + 1 :]:
        label = _normalized_cell(row[1] if len(row) > 1 else None)
        code = _cell_code(row[0] if row else None)
        values = _benchmark_row_pair(row)
        if values is None:
            continue
        if label == "retail and food services sales total":
            raw["total"] = values
        elif (
            label.startswith(
                "retail sales and food services excl motor vehicle and parts"
            )
            and "gasoline" not in label
        ):
            raw["exauto"] = values
        elif label.startswith("retail sales and food services excl gasoline"):
            raw["exgas"] = values
        elif code == "441" and "motor vehicle" in label:
            raw["auto"] = values
        elif code == "444" and label.startswith("building mat"):
            raw["building"] = values
        elif code == "447" and "gasoline" in label:
            raw["gas"] = values
        elif code == "722" and "food services" in label:
            raw["food"] = values
    required = {"total", "exauto", "auto", "building", "gas", "food"}
    if not required.issubset(raw):
        raise ValueError(
            "Census retail benchmark spreadsheet lacks required rows"
        )
    if raw["total"] != tuple(
        left + right
        for left, right in zip(raw["exauto"], raw["auto"], strict=True)
    ):
        raise ValueError(
            "Census retail benchmark spreadsheet total identity differs"
        )
    if "exgas" in raw and raw["exgas"] != tuple(
        left - right
        for left, right in zip(raw["total"], raw["gas"], strict=True)
    ):
        raise ValueError(
            "Census retail benchmark spreadsheet ex-gas identity differs"
        )
    return MappingProxyType(raw)


def _benchmark_measures(
    raw: Mapping[str, tuple[int, int]],
) -> tuple[CensusRetailBenchmarkMeasureV1, ...]:
    total = raw["total"]
    auto = raw["auto"]
    building = raw["building"]
    gas = raw["gas"]
    food = raw["food"]
    level_maps = {
        CENSUS_RETAIL_TOTAL: total,
        CENSUS_RETAIL_EX_AUTOS: raw["exauto"],
        CENSUS_RETAIL_EX_GAS: raw.get(
            "exgas",
            cast(
                tuple[int, int],
                tuple(
                    left - right for left, right in zip(total, gas, strict=True)
                ),
            ),
        ),
        CENSUS_RETAIL_CONTROL_GROUP: cast(
            tuple[int, int],
            tuple(
                total_value
                - auto_value
                - gas_value
                - building_value
                - food_value
                for (
                    total_value,
                    auto_value,
                    gas_value,
                    building_value,
                    food_value,
                ) in zip(total, auto, gas, building, food, strict=True)
            ),
        ),
    }
    result: list[CensusRetailBenchmarkMeasureV1] = []
    for measure_key in CENSUS_RETAIL_MEASURE_KEYS:
        previous, current = level_maps[measure_key]
        lexical, value = _percent_change(current, previous)
        result.append(
            CensusRetailBenchmarkMeasureV1(
                measure_key=measure_key,
                previous_level_millions=previous,
                current_level_millions=current,
                value=value,
                lexical=lexical,
            )
        )
    return tuple(result)


def parse_census_retail_benchmark_index(
    registry: OfficialSourceRegistryV1,
    index_snapshot: OfficialRawSnapshotV1,
    snapshots_by_uri: Mapping[str, OfficialRawSnapshotV1],
    *,
    as_of_date: str,
) -> CensusRetailBenchmarkIndexV1:
    """Parse annual artifacts and retain their latest pre-release rate states."""
    if not isinstance(registry, OfficialSourceRegistryV1):
        raise TypeError("Census retail benchmark parser requires a v1 registry")
    as_of = _iso_date(as_of_date, "as_of_date")
    requests = build_census_retail_benchmark_requests(registry, index_snapshot)
    if set(snapshots_by_uri) != {item.uri for item in requests}:
        raise ValueError("Census retail benchmark artifact set differs")
    entries: list[CensusRetailBenchmarkEntryV1] = []
    for report_year, request in zip(
        CENSUS_RETAIL_BENCHMARK_URIS, requests, strict=True
    ):
        snapshot = snapshots_by_uri[request.uri]
        _validate_snapshot(
            snapshot, uri=request.uri, source_format=request.source_format
        )
        measures: tuple[CensusRetailBenchmarkMeasureV1, ...] = ()
        if report_year >= 2002:
            raw = (
                _benchmark_pdf_component_levels(snapshot.content)
                if request.source_format is OfficialSourceFormat.PDF
                else _benchmark_spreadsheet_component_levels(
                    snapshot.content, request.source_format, report_year
                )
            )
            measures = _benchmark_measures(raw)
        period_end = CENSUS_RETAIL_BENCHMARK_PERIOD_ENDS[report_year]
        entries.append(
            CensusRetailBenchmarkEntryV1(
                report_year=report_year,
                report_period_end=period_end,
                applies_to_reference_period=(
                    None if report_year < 2002 else _next_month(period_end)
                ),
                disposition=(
                    CensusRetailBenchmarkDisposition.RETAINED_NON_OVERLAPPING
                    if report_year < 2002
                    else CensusRetailBenchmarkDisposition.PRIOR_STATE_OVERRIDE
                ),
                artifact_uri=request.uri,
                source_format=request.source_format,
                content_sha256=snapshot.content_sha256,
                content_length=len(snapshot.content),
                measures=measures,
            )
        )
    return CensusRetailBenchmarkIndexV1(
        source_uri=CENSUS_RETAIL_BENCHMARK_INDEX_URI,
        content_sha256=index_snapshot.content_sha256,
        content_length=len(index_snapshot.content),
        as_of_date=as_of,
        entries=tuple(entries),
    )


@dataclass(frozen=True, slots=True)
class _ParsedRelease:
    release_date: str
    released_at_ns: int
    reported_zone: str
    zone_consistent: bool
    parser_era: str
    ocr_entry_id: str | None
    benchmark_revision_notice: bool
    shutdown_delayed: bool
    reported_rate_cross_check_count: int
    level_maps: Mapping[str, tuple[int, int, int]]
    level_bases: Mapping[str, CensusRetailLevelBasis]


def _release_is_delayed(reference_period: str, release_date: str) -> bool:
    expected_release_month = _next_month(reference_period)
    return release_date[:7] > expected_release_month


def _measure_level_maps(
    components: Mapping[str, tuple[int, int, int]],
    *,
    reference_period: str,
) -> tuple[
    Mapping[str, tuple[int, int, int]],
    Mapping[str, CensusRetailLevelBasis],
]:
    exauto = components["exauto"]
    auto = components["auto"]
    gas = components["gas"]
    building = components["building"]
    food = components["food"]
    total = cast(
        tuple[int, int, int],
        tuple(left + right for left, right in zip(exauto, auto, strict=True)),
    )
    if "total" in components and components["total"] != total:
        raise ValueError("Census retail reconstructed total differs")
    derived_exgas = cast(
        tuple[int, int, int],
        tuple(left - right for left, right in zip(total, gas, strict=True)),
    )
    if "exgas" in components:
        if components["exgas"] != derived_exgas:
            raise ValueError("Census retail reconstructed ex-gas differs")
        exgas = components["exgas"]
        exgas_basis = CensusRetailLevelBasis.SOURCE_PUBLISHED
    else:
        if reference_period >= CENSUS_RETAIL_DIRECT_EX_GAS_BOUNDARY:
            raise ValueError("Census retail direct ex-gas level is missing")
        exgas = derived_exgas
        exgas_basis = CensusRetailLevelBasis.DERIVED_AGGREGATE
    control = cast(
        tuple[int, int, int],
        tuple(
            value - auto_value - gas_value - building_value - food_value
            for value, auto_value, gas_value, building_value, food_value in zip(
                total, auto, gas, building, food, strict=True
            )
        ),
    )
    if any(item <= 0 for item in (*total, *exgas, *control)):
        raise ValueError("Census retail derived level is not positive")
    return (
        MappingProxyType(
            {
                CENSUS_RETAIL_TOTAL: total,
                CENSUS_RETAIL_EX_AUTOS: exauto,
                CENSUS_RETAIL_EX_GAS: exgas,
                CENSUS_RETAIL_CONTROL_GROUP: control,
            }
        ),
        MappingProxyType(
            {
                CENSUS_RETAIL_TOTAL: CensusRetailLevelBasis.SOURCE_PUBLISHED,
                CENSUS_RETAIL_EX_AUTOS: CensusRetailLevelBasis.SOURCE_PUBLISHED,
                CENSUS_RETAIL_EX_GAS: exgas_basis,
                CENSUS_RETAIL_CONTROL_GROUP: CensusRetailLevelBasis.DERIVED_AGGREGATE,
            }
        ),
    )


def _parse_release(
    *,
    reference_period: str,
    pdf_snapshot: OfficialRawSnapshotV1,
    support_snapshot: OfficialRawSnapshotV1 | None,
    ocr_corpus: CensusRetailOcrCorpusV1,
) -> _ParsedRelease:
    _validate_snapshot(
        pdf_snapshot,
        uri=pdf_snapshot.request.uri,
        source_format=OfficialSourceFormat.PDF,
    )
    native_text = _pdf_text(pdf_snapshot.content)
    ocr_entry = ocr_corpus.by_uri.get(pdf_snapshot.request.uri)
    if ocr_entry is not None:
        if ocr_entry.content_sha256 != pdf_snapshot.content_sha256:
            raise ValueError("Census retail OCR source hash differs")
        text = "\f".join(ocr_entry.page_texts)
        ocr_entry_id: str | None = ocr_entry.entry_id
    else:
        if len(native_text.strip()) < 1000:
            raise ValueError("Census retail image-only PDF lacks OCR evidence")
        text = native_text
        ocr_entry_id = None
    release_date, reported_zone, released_at_ns, zone_consistent = (
        _release_time(text)
    )
    if not zone_consistent:
        raise ValueError("Census retail release reports an inconsistent zone")
    if support_snapshot is None:
        components = _pdf_component_levels(
            text, content_sha256=pdf_snapshot.content_sha256
        )
        parser_era = (
            "ocr-pdf-table" if ocr_entry is not None else "native-pdf-table"
        )
        cross_checks = 0
    else:
        _validate_snapshot(
            support_snapshot,
            uri=support_snapshot.request.uri,
            source_format=support_snapshot.request.source_format,
        )
        components, cross_checks = _spreadsheet_component_levels(
            support_snapshot.content, support_snapshot.request.source_format
        )
        parser_era = (
            "xls-table"
            if support_snapshot.request.source_format
            is OfficialSourceFormat.XLS
            else "xlsx-table"
        )
        if ocr_entry is not None:
            raise ValueError(
                "Census retail spreadsheet release unexpectedly uses OCR"
            )
    level_maps, level_bases = _measure_level_maps(
        components, reference_period=reference_period
    )
    return _ParsedRelease(
        release_date=release_date,
        released_at_ns=released_at_ns,
        reported_zone=reported_zone,
        zone_consistent=zone_consistent,
        parser_era=parser_era,
        ocr_entry_id=ocr_entry_id,
        benchmark_revision_notice=bool(_REVISION_NOTICE_RE.search(text)),
        shutdown_delayed=_release_is_delayed(reference_period, release_date),
        reported_rate_cross_check_count=cross_checks,
        level_maps=level_maps,
        level_bases=level_bases,
    )


def build_census_retail_archive_manifest(
    registry: OfficialSourceRegistryV1,
    profile: UnitedStatesBackfillProfileV1,
    release_index: CensusRetailReleaseIndexV1,
    benchmark_index: CensusRetailBenchmarkIndexV1,
    snapshots_by_uri: Mapping[str, OfficialRawSnapshotV1],
    ocr_corpus: CensusRetailOcrCorpusV1,
) -> CensusRetailArchiveManifestV1:
    """Reconstruct every occurrence and its previous-as-known lineage."""
    if not isinstance(registry, OfficialSourceRegistryV1):
        raise TypeError("Census retail manifest requires a v1 registry")
    if not isinstance(profile, UnitedStatesBackfillProfileV1):
        raise TypeError("Census retail manifest requires a v1 profile")
    if not isinstance(release_index, CensusRetailReleaseIndexV1):
        raise TypeError("Census retail manifest requires a v1 index")
    if not isinstance(benchmark_index, CensusRetailBenchmarkIndexV1):
        raise TypeError("Census retail manifest requires a v1 benchmark index")
    if benchmark_index.as_of_date != release_index.as_of_date:
        raise ValueError("Census retail manifest index dates differ")
    if not isinstance(ocr_corpus, CensusRetailOcrCorpusV1):
        raise TypeError("Census retail manifest requires a v1 OCR corpus")
    source = registry.source(CENSUS_RETAIL_SOURCE_KEY)
    program = profile.by_key[CENSUS_RETAIL_PROGRAM_KEY]
    if (
        profile.registry_id != registry.registry_id
        or program.source_key != source.source_key
        or program.archive_uri != CENSUS_RETAIL_INDEX_URI
    ):
        raise ValueError("Census retail registry and profile differ")
    expected_requests = build_census_retail_archive_requests(
        registry, release_index
    )
    expected_uris = {item.uri for item in expected_requests}
    if set(snapshots_by_uri) != expected_uris:
        raise ValueError("Census retail retained artifact set differs")
    for request in expected_requests:
        snapshot = snapshots_by_uri[request.uri]
        _validate_snapshot(
            snapshot, uri=request.uri, source_format=request.source_format
        )

    predecessor_snapshot = snapshots_by_uri[CENSUS_RETAIL_PREDECESSOR_URI]
    predecessor = _parse_release(
        reference_period=CENSUS_RETAIL_PREDECESSOR_REFERENCE_PERIOD,
        pdf_snapshot=predecessor_snapshot,
        support_snapshot=None,
        ocr_corpus=ocr_corpus,
    )
    known: dict[str, tuple[str, float, str]] = {}
    for measure_key in CENSUS_RETAIL_MEASURE_KEYS:
        levels = predecessor.level_maps[measure_key]
        lexical, numeric = _percent_change(levels[0], levels[1])
        known[measure_key] = (
            lexical,
            numeric,
            CENSUS_RETAIL_PREDECESSOR_URI,
        )

    publications: list[CensusRetailPublicationV1] = []
    used_ocr_uris = {CENSUS_RETAIL_PREDECESSOR_URI}
    previous_release_date = predecessor.release_date
    for index_entry in release_index.releases:
        benchmark = benchmark_index.by_application_period.get(
            index_entry.reference_period
        )
        if benchmark is not None:
            for measure_key in CENSUS_RETAIL_MEASURE_KEYS:
                benchmark_measure = benchmark.by_measure[measure_key]
                known[measure_key] = (
                    benchmark_measure.lexical,
                    benchmark_measure.value,
                    benchmark.artifact_uri,
                )
        pdf_snapshot = snapshots_by_uri[index_entry.artifact_uri]
        support_snapshot = (
            None
            if index_entry.support_uri is None
            else snapshots_by_uri[index_entry.support_uri]
        )
        parsed = _parse_release(
            reference_period=index_entry.reference_period,
            pdf_snapshot=pdf_snapshot,
            support_snapshot=support_snapshot,
            ocr_corpus=ocr_corpus,
        )
        if parsed.release_date <= previous_release_date:
            raise ValueError(
                "Census retail release dates are not increasing: "
                f"{index_entry.reference_period}={parsed.release_date} after "
                f"{previous_release_date}"
            )
        previous_release_date = parsed.release_date
        if (
            index_entry.reference_period
            == CENSUS_RETAIL_CURRENT_REFERENCE_PERIOD
            and parsed.release_date != CENSUS_RETAIL_CURRENT_RELEASE_DATE
        ):
            raise ValueError("Census retail current release date differs")
        if parsed.ocr_entry_id is not None:
            used_ocr_uris.add(index_entry.artifact_uri)
        measures: list[CensusRetailMeasureV1] = []
        comparison_basis = (
            "sic-to-naics-boundary"
            if index_entry.reference_period == CENSUS_RETAIL_NAICS_BOUNDARY
            else "same-classification"
        )
        for measure_key in CENSUS_RETAIL_MEASURE_KEYS:
            levels = parsed.level_maps[measure_key]
            actual_lexical, actual = _percent_change(levels[0], levels[1])
            revised_lexical, revised = _percent_change(levels[1], levels[2])
            previous_lexical, previous, previous_uri = known[measure_key]
            measures.append(
                CensusRetailMeasureV1(
                    measure_key=measure_key,
                    reference_period=index_entry.reference_period,
                    actual_value=actual,
                    actual_lexical=actual_lexical,
                    current_level_millions=levels[0],
                    previous_level_millions=levels[1],
                    level_basis=parsed.level_bases[measure_key],
                    previous_as_known_value=previous,
                    previous_as_known_lexical=previous_lexical,
                    revised_previous_value=revised,
                    revised_previous_lexical=revised_lexical,
                    revision_comparable=(
                        comparison_basis == "same-classification"
                    ),
                    comparison_basis=comparison_basis,
                    previous_artifact_uri=previous_uri,
                )
            )
            known[measure_key] = (
                actual_lexical,
                actual,
                index_entry.artifact_uri,
            )
        publications.append(
            CensusRetailPublicationV1(
                reference_period=index_entry.reference_period,
                release_date=parsed.release_date,
                released_lexical=f"{parsed.release_date}T08:30:00",
                released_at_ns=parsed.released_at_ns,
                reported_zone=parsed.reported_zone,
                zone_consistent=parsed.zone_consistent,
                artifact_uri=index_entry.artifact_uri,
                content_sha256=pdf_snapshot.content_sha256,
                content_length=len(pdf_snapshot.content),
                support_uri=index_entry.support_uri,
                support_format=index_entry.support_format,
                support_content_sha256=(
                    None
                    if support_snapshot is None
                    else support_snapshot.content_sha256
                ),
                support_content_length=(
                    None
                    if support_snapshot is None
                    else len(support_snapshot.content)
                ),
                parser_era=parsed.parser_era,
                ocr_entry_id=parsed.ocr_entry_id,
                benchmark_revision_notice=parsed.benchmark_revision_notice,
                shutdown_delayed=parsed.shutdown_delayed,
                reported_rate_cross_check_count=(
                    parsed.reported_rate_cross_check_count
                ),
                measures=tuple(measures),
            )
        )
    if used_ocr_uris != set(ocr_corpus.by_uri):
        raise ValueError(
            "Census retail OCR corpus has unused or missing evidence"
        )
    publication_tuple = tuple(publications)
    raw_count = (
        3
        + len(publication_tuple)
        + sum(item.support_uri is not None for item in publication_tuple)
        + len(benchmark_index.entries)
    )
    total_bytes = (
        release_index.content_length
        + benchmark_index.content_length
        + len(predecessor_snapshot.content)
        + sum(item.content_length for item in benchmark_index.entries)
        + sum(
            item.content_length + (item.support_content_length or 0)
            for item in publication_tuple
        )
    )
    measure_count = sum(len(item.measures) for item in publication_tuple)
    revision_count = sum(
        measure.previous_was_revised
        for item in publication_tuple
        for measure in item.measures
    )
    derived_count = sum(
        measure.level_basis is CensusRetailLevelBasis.DERIVED_AGGREGATE
        for item in publication_tuple
        for measure in item.measures
    )
    support_count = sum(
        item.support_uri is not None for item in publication_tuple
    )
    return CensusRetailArchiveManifestV1(
        registry_id=registry.registry_id,
        profile_id=profile.profile_id,
        release_index=release_index,
        benchmark_index=benchmark_index,
        ocr_corpus_id=ocr_corpus.corpus_id,
        predecessor_content_sha256=predecessor_snapshot.content_sha256,
        predecessor_content_length=len(predecessor_snapshot.content),
        publications=publication_tuple,
        raw_artifact_count=raw_count,
        total_content_bytes=total_bytes,
        measure_occurrence_count=measure_count,
        comparable_revision_count=revision_count,
        derived_aggregate_occurrence_count=derived_count,
        support_artifact_count=support_count,
        corrected_support_link_count=sum(
            item.support_link_corrected for item in release_index.releases
        ),
        ocr_publication_count=sum(
            item.ocr_entry_id is not None for item in publication_tuple
        ),
        benchmark_revision_notice_count=sum(
            item.benchmark_revision_notice for item in publication_tuple
        ),
        benchmark_state_override_count=len(
            benchmark_index.by_application_period
        ),
        shutdown_delayed_count=sum(
            item.shutdown_delayed for item in publication_tuple
        ),
        reported_rate_cross_check_count=sum(
            item.reported_rate_cross_check_count for item in publication_tuple
        ),
        exact_minute_count=len(publication_tuple),
    )


def replay_census_retail_archive(
    registry: OfficialSourceRegistryV1,
    profile: UnitedStatesBackfillProfileV1,
    index_snapshot: OfficialRawSnapshotV1,
    benchmark_index_snapshot: OfficialRawSnapshotV1,
    benchmark_snapshots_by_uri: Mapping[str, OfficialRawSnapshotV1],
    snapshots_by_uri: Mapping[str, OfficialRawSnapshotV1],
    ocr_corpus: CensusRetailOcrCorpusV1,
    expected: CensusRetailArchiveManifestV1,
) -> CensusRetailArchiveManifestV1:
    """Recompute and compare the complete retained MARTS corpus."""
    if not isinstance(expected, CensusRetailArchiveManifestV1):
        raise TypeError("Census retail replay requires a v1 expected manifest")
    release_index = parse_census_retail_release_index(
        index_snapshot, as_of_date=expected.release_index.as_of_date
    )
    benchmark_index = parse_census_retail_benchmark_index(
        registry,
        benchmark_index_snapshot,
        benchmark_snapshots_by_uri,
        as_of_date=expected.release_index.as_of_date,
    )
    rebuilt = build_census_retail_archive_manifest(
        registry,
        profile,
        release_index,
        benchmark_index,
        snapshots_by_uri,
        ocr_corpus,
    )
    if rebuilt != expected:
        raise ValueError("Census retail retained-corpus replay differs")
    return rebuilt


def census_retail_coverage_from_manifest(
    manifest: CensusRetailArchiveManifestV1,
) -> UnitedStatesProgramCoverageV1:
    """Derive complete U.S. retail-sales coverage from the manifest."""
    if not isinstance(manifest, CensusRetailArchiveManifestV1):
        raise TypeError("Census retail coverage requires a v1 manifest")
    return UnitedStatesProgramCoverageV1(
        program_key=CENSUS_RETAIL_PROGRAM_KEY,
        window_start_date=US_BACKFILL_START_DATE,
        window_end_date=manifest.release_index.as_of_date,
        expected_occurrence_count=len(manifest.publications),
        schedule_count=len(manifest.publications),
        initial_actual_count=len(manifest.publications),
        previous_as_known_count=len(manifest.publications),
        revision_count=manifest.comparable_revision_count,
        exact_minute_count=manifest.exact_minute_count,
        forecast_count=0,
        artifact_sha256s=tuple(
            sorted(
                {
                    manifest.release_index.content_sha256,
                    manifest.benchmark_index.content_sha256,
                    manifest.predecessor_content_sha256,
                    *(
                        item.content_sha256
                        for item in manifest.benchmark_index.entries
                    ),
                    *(item.content_sha256 for item in manifest.publications),
                    *(
                        item.support_content_sha256
                        for item in manifest.publications
                        if item.support_content_sha256 is not None
                    ),
                }
            )
        ),
        gap_reasons=(),
        notes=(
            "Each official PDF is one schedule occurrence; its four measures retain independent revision lineage.",
            "The May 2001 SIC-to-NAICS comparison is retained but explicitly noncomparable.",
            "Ex-gasoline levels before May 2018 and all retail-control levels are derived only from contemporaneous published components.",
            "Twenty-four annual revisions replace the prior-month state before the next advance release; the 2000 and 2001 reports are retained but do not overlap that boundary.",
            "No monthly event-level historical consensus is manufactured from another survey.",
        ),
    )


def packaged_census_retail_archive_manifest_path() -> Path:
    """Return the packaged Census retail manifest path."""
    return (
        Path(__file__).with_name("assets") / "us_retail_sales_archive_v1.json"
    )


def packaged_census_retail_index_path() -> Path:
    """Return the packaged exact archive-index envelope path."""
    return (
        Path(__file__).with_name("assets")
        / "us_retail_sales_release_index_v1.json"
    )


def packaged_census_retail_benchmark_index_path() -> Path:
    """Return the packaged exact annual-revision index envelope path."""
    return (
        Path(__file__).with_name("assets")
        / "us_retail_sales_benchmark_index_v1.json"
    )


def packaged_census_retail_ocr_path() -> Path:
    """Return the packaged reviewed OCR evidence path."""
    return Path(__file__).with_name("assets") / "us_retail_sales_ocr_v1.json"


def load_packaged_census_retail_archive_manifest() -> (
    CensusRetailArchiveManifestV1
):
    """Load and validate the packaged MARTS archive receipt."""
    return CensusRetailArchiveManifestV1.from_json(
        packaged_census_retail_archive_manifest_path().read_text(
            encoding="utf-8"
        )
    )


def load_packaged_census_retail_ocr_corpus() -> CensusRetailOcrCorpusV1:
    """Load and validate reviewed page-aware OCR evidence."""
    return CensusRetailOcrCorpusV1.from_json(
        packaged_census_retail_ocr_path().read_text(encoding="utf-8")
    )


def load_packaged_census_retail_index_snapshot() -> OfficialRawSnapshotV1:
    """Load the exact official archive page from its base64 envelope."""
    try:
        payload = _mapping(
            json.loads(
                packaged_census_retail_index_path().read_text(encoding="utf-8")
            ),
            "Census retail index envelope",
        )
    except json.JSONDecodeError as exc:
        raise ValueError(
            "Census retail index envelope is invalid JSON"
        ) from exc
    if (
        payload.get("schema_version")
        != CENSUS_RETAIL_INDEX_ENVELOPE_SCHEMA_VERSION
    ):
        raise ValueError("unsupported Census retail index-envelope schema")
    as_of = _iso_date(payload.get("as_of_date"), "as_of_date")
    captured_at_ns = cast(int, payload.get("captured_at_ns"))
    _positive_int(captured_at_ns, "captured_at_ns", 2**63 - 1)
    uri = str(payload.get("uri", ""))
    if uri != CENSUS_RETAIL_INDEX_URI:
        raise ValueError("Census retail index-envelope URI differs")
    try:
        content = base64.b64decode(
            str(payload.get("content_base64", "")), validate=True
        )
    except (TypeError, ValueError) as exc:
        raise ValueError(
            "Census retail index envelope has invalid base64"
        ) from exc
    registry = load_packaged_official_source_registry()
    request = build_census_retail_index_request(registry, as_of_date=as_of)
    snapshot = OfficialRawSnapshotV1(
        request=request,
        retrieved_at_ns=captured_at_ns,
        completed_at_ns=captured_at_ns,
        status_code=200,
        resolved_uri=uri,
        response_headers={"Content-Type": "text/html"},
        content=content,
        content_type="text/html",
    )
    expected_sha = _sha256(payload.get("content_sha256"), "content_sha256")
    if snapshot.content_sha256 != expected_sha:
        raise ValueError("Census retail index-envelope content differs")
    return snapshot


def load_packaged_census_retail_index() -> CensusRetailReleaseIndexV1:
    """Replay the packaged exact archive page into its typed index."""
    snapshot = load_packaged_census_retail_index_snapshot()
    assert snapshot.request.window_end is not None
    return parse_census_retail_release_index(
        snapshot, as_of_date=snapshot.request.window_end
    )


def load_packaged_census_retail_benchmark_index_snapshot() -> (
    OfficialRawSnapshotV1
):
    """Load the exact official annual-revision archive page from its envelope."""
    try:
        payload = _mapping(
            json.loads(
                packaged_census_retail_benchmark_index_path().read_text(
                    encoding="utf-8"
                )
            ),
            "Census retail benchmark-index envelope",
        )
    except json.JSONDecodeError as exc:
        raise ValueError(
            "Census retail benchmark-index envelope is invalid JSON"
        ) from exc
    if (
        payload.get("schema_version")
        != CENSUS_RETAIL_BENCHMARK_INDEX_ENVELOPE_SCHEMA_VERSION
    ):
        raise ValueError(
            "unsupported Census retail benchmark-index envelope schema"
        )
    as_of = _iso_date(payload.get("as_of_date"), "as_of_date")
    captured_at_ns = cast(int, payload.get("captured_at_ns"))
    _positive_int(captured_at_ns, "captured_at_ns", 2**63 - 1)
    uri = str(payload.get("uri", ""))
    if uri != CENSUS_RETAIL_BENCHMARK_INDEX_URI:
        raise ValueError("Census retail benchmark-index envelope URI differs")
    try:
        content = base64.b64decode(
            str(payload.get("content_base64", "")), validate=True
        )
    except (TypeError, ValueError) as exc:
        raise ValueError(
            "Census retail benchmark-index envelope has invalid base64"
        ) from exc
    registry = load_packaged_official_source_registry()
    request = build_census_retail_benchmark_index_request(
        registry, as_of_date=as_of
    )
    snapshot = OfficialRawSnapshotV1(
        request=request,
        retrieved_at_ns=captured_at_ns,
        completed_at_ns=captured_at_ns,
        status_code=200,
        resolved_uri=uri,
        response_headers={"Content-Type": "text/html"},
        content=content,
        content_type="text/html",
    )
    expected_sha = _sha256(payload.get("content_sha256"), "content_sha256")
    if snapshot.content_sha256 != expected_sha:
        raise ValueError(
            "Census retail benchmark-index envelope content differs"
        )
    return snapshot


def load_packaged_census_retail_benchmark_index() -> (
    CensusRetailBenchmarkIndexV1
):
    """Load the typed annual-revision index and verify its exact HTML envelope."""
    manifest = load_packaged_census_retail_archive_manifest()
    snapshot = load_packaged_census_retail_benchmark_index_snapshot()
    if (
        manifest.benchmark_index.content_sha256 != snapshot.content_sha256
        or manifest.benchmark_index.content_length != len(snapshot.content)
        or manifest.benchmark_index.as_of_date != snapshot.request.window_end
    ):
        raise ValueError("Census retail packaged benchmark index differs")
    return manifest.benchmark_index
